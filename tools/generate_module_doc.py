# -*- coding: utf-8 -*-
"""Generate MyLunarCore module relationship Word document."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from pathlib import Path


def set_run_font(run, name="微软雅黑", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=16 if level == 1 else 14 if level == 2 else 12, bold=True)
    return h


def add_para(doc, text, bold=False, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Inches(0.25)
    run = p.add_run(text)
    set_run_font(run, bold=bold)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.25)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                set_run_font(run, bold=True, size=10)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()
    return table


def add_code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    run = p.add_run(text)
    set_run_font(run, name="Consolas", size=9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p


def main():
    out = Path(__file__).resolve().parent.parent / "docs" / "MyLunarCore-Module-Relations.docx"
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    # Default style
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(11)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("MyLunarCore 各模块联系说明")
    set_run_font(tr, size=22, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("项目架构与模块依赖关系文档")
    set_run_font(sr, size=12, color=RGBColor(0x66, 0x66, 0x66))

    doc.add_paragraph()

    add_para(doc, "项目采用 Spring Boot 单体游戏服为主架构，另有一套 Spring Cloud Alibaba 微服务脚手架（microservices/），二者目前相对独立，主业务逻辑集中在根目录的 MyLunarCore 模块中。")

    # Section 1
    add_heading(doc, "一、整体分层关系", 1)
    add_para(doc, "架构自上而下分为：客户端 → 入口层（net / admin）→ 玩法业务层 → 基础设施层（common / config）→ 数据层（repo / model / protocol / MySQL / 静态资源）。")
    add_para(doc, "分层结构示意：", bold=True)
    add_code_block(doc, """客户端（游戏客户端 TCP/KCP / 管理后台 HTTP）
    ↓
入口层：net（Netty/KCP 编解码与分发）、admin（HTTP + Spring Security）
    ↓
玩法业务层：player、battle、scene、challenge、gacha、item、rogue
    ↓
基础设施层：common（主循环/日志/热更/事件）、config（配置与线程池）
    ↓
数据层：repo（JDBC 仓储）、model（领域实体）、protocol（Protobuf）、MySQL、CSV/JSON 静态资源""")

    # Section 2
    add_heading(doc, "二、主模块职责与依赖", 1)
    add_para(doc, "各模块职责与依赖关系如下表：")

    add_table(doc,
        ["模块", "职责", "主要依赖", "被谁依赖"],
        [
            ["config", "LunarCoreProperties 配置绑定、Netty 业务线程池、GameConfigurationManager 桥接", "Spring Boot", "net、common、player"],
            ["model", "玩家、抽卡、肉鸽、挑战等领域实体/值对象", "无（纯 POJO）", "repo、player、各玩法模块"],
            ["repo", "MySQL JDBC 访问（玩家数据、战斗配置、抽卡保底等）", "model、JdbcTemplate", "所有需要持久化的业务模块"],
            ["protocol", "由 .proto 生成的 Protobuf 消息类", "—", "net、各 *NettyService"],
            ["net", "TCP/KCP 服务器、帧编解码、@PacketCmd 命令路由", "config、common、各玩法 Handler", "客户端连接入口"],
            ["player", "登录/登出、会话、异步加载、统一同步、Tick 注册", "repo、model、scene、common、net", "net（会话包）、common（OnlinePlayer）"],
            ["battle", "战斗运行时、波次、技能/Buff、战斗协议", "repo、common（事件）、net", "net → BattlePacketHandlers"],
            ["scene", "场景上下文、NPC/怪物/道具实体", "repo、model", "net、player（OnlinePlayer）"],
            ["challenge", "挑战关卡运行时索引", "repo", "net → ChallengePacketHandlers"],
            ["gacha", "抽卡逻辑、卡池配置、保底", "repo、gacha 配置、item（发奖）", "net、common（热更）"],
            ["item", "背包、使用/强化/装备等", "repo", "net、gacha"],
            ["rogue", "肉鸽玩法运行时", "repo", "net → RoguePacketHandlers"],
            ["common", "游戏主循环、日志、监控、热更新、静态资源、领域事件", "config、gacha 等", "全局基础设施"],
            ["admin", "后台 RBAC、登录、HTTP API", "repo、Spring Security", "运维/管理端"],
        ])

    add_para(doc, "模块划分说明（摘自 package-info.java）：")
    add_bullet(doc, "config — 配置与 Spring 装配")
    add_bullet(doc, "model — 玩家/玩法领域实体与值对象")
    add_bullet(doc, "repo — 数据库访问")
    add_bullet(doc, "net — Netty/KCP 协议、编解码与命令分发")
    add_bullet(doc, "player — 会话、登录与数据同步")
    add_bullet(doc, "battle|scene|rogue|gacha|challenge|item — 各玩法业务与 Netty 服务")
    add_bullet(doc, "common — 日志、监控、热更新、主循环、活动等基础设施")
    add_bullet(doc, "admin — 后台 RBAC 与 HTTP 接口")

    # Section 3
    add_heading(doc, "三、核心调用链", 1)

    add_heading(doc, "3.1 启动与装配", 2)
    add_code_block(doc, """MyLunarCoreApplication (Spring Boot)
    ├── config: LunarCoreProperties、NettyGameBusinessConfiguration
    ├── net: GameNettyServer / GameKcpServer 监听端口
    ├── net: PacketCommandRegistry 扫描 @PacketCmd 注册路由
    ├── common: GameServer 启动 Timer 主循环
    └── common: HotReloadCoordinator、StaticResourceRegistry 加载配置""")

    add_heading(doc, "3.2 客户端请求路径（以战斗为例）", 2)
    add_code_block(doc, """客户端 TCP/KCP
  → LunarFrameDecoder（解码）
  → GameServerChannelHandler
  → GamePacketDispatcher（按 cmdId 分发）
  → BattlePacketHandlers（@PacketCmd 入口，Protobuf 反序列化）
  → BattleNettyService（业务逻辑）
      → BattleManager / BattleContext（内存战局）
      → BattleRepository 等 repo（持久化）
      → GameEventPublisher（发布 BattleStartedEvent / BattleEndedEvent）
  → 响应 GamePacket 写回客户端""")

    add_para(doc, "网络层与业务层通过注解路由解耦：PacketCommandRegistry 在启动时扫描 Spring 容器中所有 @PacketCmd 方法，运行时 O(1) 查找，无需手写 switch。")

    add_heading(doc, "3.3 玩家模块作为「枢纽」", 2)
    add_bullet(doc, "登录：PlayerSessionService → PlayerDataRepository 加载数据 → GameSessionManager 建会话 → PlayerTickRegistry 注册 OnlinePlayer")
    add_bullet(doc, "同步：PlayerSyncCoordinator 聚合多个 Syncable 实现（PlayerCoreSyncable、AvatarSyncable、LineupSyncable、UnlockSyncable 等），组装统一下行包")
    add_bullet(doc, "场景：OnlinePlayer 依赖 SceneManager，登录后绑定场景状态")
    add_bullet(doc, "Channel 绑定：各 *NettyService 通过 Channel 属性 playerUid 关联当前玩家")

    add_heading(doc, "3.4 玩法模块之间的横向联系", 2)
    add_table(doc,
        ["关系", "说明"],
        [
            ["gacha → item", "抽卡发奖调用 ItemRepository 写入背包"],
            ["battle → common", "战斗开始/结束通过 GameEventPublisher 发布事件，BattleEventLoggingListener 等监听"],
            ["scene ↔ battle", "场景内遇敌可触发战斗；BattleSceneFactory 从场景/波次配置创建战局"],
            ["player → scene", "玩家坐标、plane/floor 由场景模块管理"],
            ["common 热更", "HotReloadCoordinator 统一触发 hotfix、活动排期、抽卡配置、CSV 静态表重载"],
        ])

    add_heading(doc, "3.5 数据访问模式", 2)
    add_code_block(doc, """各 *NettyService / *Manager
    → repo（JdbcTemplate）
        → model（Entity / PlayerData 聚合）
            → MySQL（lunarcore_merged.sql）""")

    add_para(doc, "静态配置有两条路径：")
    add_bullet(doc, "数据库：GameDataRepository 读 game_data 表 JSON")
    add_bullet(doc, "文件：StaticResourceRegistry + CSV（如 items_config.csv），由 @ResourceType 标注")

    # Section 4
    add_heading(doc, "四、协议与模块映射", 1)
    add_para(doc, "7 个 .proto 文件对应 7 个玩法/系统：")
    add_table(doc,
        ["Proto 文件", "对应 Java 模块", "PacketHandlers"],
        [
            ["player_session.proto", "player", "PlayerSessionPacketHandlers"],
            ["battle_system.proto", "battle", "BattlePacketHandlers"],
            ["scene_system.proto", "scene", "ScenePacketHandlers"],
            ["challenge_system.proto", "challenge", "ChallengePacketHandlers"],
            ["gacha_system.proto", "gacha", "GachaPacketHandlers"],
            ["item_system.proto", "item", "ItemPacketHandlers"],
            ["rogue_system.proto", "rogue", "RoguePacketHandlers"],
        ])

    # Section 5
    add_heading(doc, "五、common 基础设施的内部协作", 1)
    add_bullet(doc, "GameServer：定时驱动在线玩家 Tick 与活动日切")
    add_bullet(doc, "PlayerTickRegistry：维护 OnlinePlayer 注册表，供主循环遍历")
    add_bullet(doc, "HotReloadCoordinator：运维 /reload 的统一入口")
    add_bullet(doc, "GameEventPublisher：Spring 事件总线，解耦战斗等业务与监听方（如日志、统计）")
    add_bullet(doc, "GameTrafficMetrics：网络包流量统计")
    add_bullet(doc, "StaticResourceRegistry：CSV 等静态表格资源的延迟加载与热失效")

    add_para(doc, "协作关系示意：")
    add_code_block(doc, """GameServer 主循环
    → PlayerTickRegistry → OnlinePlayer（每帧 Tick）
    → ActivityScheduleService（活动日切）

GamePacketDispatcher → GameTrafficMetrics（流量统计）

BattleNettyService → GameEventPublisher → BattleEventLoggingListener 等

HotReloadCoordinator → HotfixDataService / ActivityScheduleService / GachaConfigService / StaticResourceRegistry""")

    # Section 6
    add_heading(doc, "六、admin 模块（独立 HTTP 通道）", 1)
    add_para(doc, "与游戏 Netty 通道并行，走 Spring MVC + Security：")
    add_code_block(doc, """HTTP 请求 → AdminSecurityConfig（鉴权）
         → AdminLoginController / AdminRbacDemoController
         → AdminRbacRepository（数据库）""")
    add_para(doc, "管理后台不参与游戏协议，只负责运维、RBAC 演示等。")

    # Section 7
    add_heading(doc, "七、microservices 子项目（迁移脚手架）", 1)
    add_para(doc, "位于 microservices/，与主单体尚未深度集成，是未来拆分方向。")
    add_code_block(doc, """HTTP 客户端 → api-gateway
    → JWT 校验 / 限流 / 粘性负载均衡
    → auth-service（JWT 签发与认证）
    → player-service（玩家 REST API 雏形）
    各服务 → common-api（ApiResponse 等统一响应）
    各服务 → Nacos（注册与配置中心）""")

    add_table(doc,
        ["子模块", "作用"],
        [
            ["common-api", "统一 API 响应结构 ApiResponse"],
            ["api-gateway", "Spring Cloud Gateway、JWT 过滤、限流、用户粘性负载均衡"],
            ["auth-service", "JWT 签发与认证"],
            ["player-service", "玩家 REST API 雏形（PlayerController）"],
        ])

    add_para(doc, "当前主游戏逻辑（Netty、战斗、场景等）仍在根目录单体 MyLunarCore 中；微服务部分主要是 Spring Cloud Alibaba + Nacos 的骨架代码。")

    # Section 8
    add_heading(doc, "八、总结：依赖方向规则", 1)
    add_bullet(doc, "单向分层：net → 玩法业务 → repo → model，避免 repo 反向依赖 net。")
    add_bullet(doc, "player 是中心：会话、uid、Channel 绑定、数据同步，各玩法通过 uid 关联玩家。")
    add_bullet(doc, "net 与业务解耦：*PacketHandlers 只做协议转换，*NettyService 承载业务。")
    add_bullet(doc, "common 横切：日志、主循环、热更、事件、监控，被各层按需引用。")
    add_bullet(doc, "config 贯穿：端口、线程数、TLS、游戏循环开关等全局参数。")
    add_bullet(doc, "admin 与 game 双通道：HTTP 管理 vs TCP/KCP 游戏，共用数据库但入口分离。")

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer.add_run("文档生成自 MyLunarCore 项目源码分析")
    set_run_font(fr, size=9, color=RGBColor(0x99, 0x99, 0x99))

    doc.save(str(out))
    print(f"Generated: {out}")


if __name__ == "__main__":
    main()
