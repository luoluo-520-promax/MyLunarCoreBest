# -*- coding: utf-8 -*-
"""将 AI 助手「沉浸增强六维」同步进桌面原 bak 总结报告，覆盖保存原文件（不另存新文档）。

幂等标记：正文出现「7.28 AI 沉浸增强六维」即视为已写入。
"""

import os
from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DOC_PATH = Path.home() / "Desktop" / "MyLunarCore项目各方面详细总结报告.bak.docx"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def replace_paragraph_text(paragraph, new_text, bold=False):
    style = paragraph.style
    for r in list(paragraph.runs):
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        set_run_font(paragraph.runs[0], bold=bold)
    else:
        run = paragraph.add_run(new_text)
        set_run_font(run, bold=bold)
    paragraph.style = style


def insert_paragraph_after(paragraph, text, style_name="List Bullet"):
    new_para = paragraph._parent.add_paragraph(text, style=style_name)
    paragraph._p.addnext(new_para._p)
    for r in new_para.runs:
        set_run_font(
            r,
            size=13 if style_name.startswith("Heading") else 11,
            bold=style_name.startswith("Heading"),
        )
        if style_name.startswith("Heading"):
            r.font.color.rgb = RGBColor(25, 75, 140)
    return new_para


def find_para(doc, exact=None, startswith=None, contains=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if exact is not None and t == exact:
            return p
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    return None


def set_cell_text(cell, text):
    if not cell.paragraphs:
        cell.text = text
        return
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r.text = ""
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0])
    else:
        run = p.add_run(text)
        set_run_font(run)
    for extra in cell.paragraphs[1:]:
        extra.clear()


def add_table_row(table, values):
    row = table.add_row()
    for i, value in enumerate(values):
        if i < len(row.cells):
            set_cell_text(row.cells[i], value)
    return row


def table_has(table, needle):
    for row in table.rows:
        for cell in row.cells:
            if needle in cell.text:
                return True
    return False


def find_table_by_header(doc, *needles):
    for table in doc.tables:
        if not table.rows:
            continue
        header = " ".join(c.text.strip() for c in table.rows[0].cells)
        if all(n in header for n in needles):
            return table
    return None


def patch_table_cell_contains(table, contains, append=None, col=0):
    for row in table.rows:
        if contains in row.cells[col].text:
            if append:
                target = row.cells[-1] if append[0] is None else row.cells[append[0]]
                cur = target.text.strip()
                extra = append[1]
                if extra not in cur:
                    set_cell_text(target, cur.rstrip("。") + extra)
            return True
    return False


def safe_save(doc: Document, path: Path) -> None:
    """原地覆盖保存；若被 Word 占用则写入 .pending.docx 并提示合并。"""
    tmp = path.with_name(path.stem + "._patch_tmp.docx")
    pending = path.with_name(path.stem + ".pending.docx")
    doc.save(str(tmp))
    try:
        os.replace(str(tmp), str(path))
        if pending.exists():
            pending.unlink(missing_ok=True)
        print(f"已更新原文档: {path}")
    except OSError as err:
        if pending.exists():
            pending.unlink()
        os.replace(str(tmp), str(pending))
        raise SystemExit(
            f"原文件被占用，无法直接覆盖: {path}\n"
            f"已写入待合并文件: {pending}\n"
            f"请关闭 Word 后执行:\n"
            f'  Move-Item -Force "{pending}" "{path}"\n'
            f"原因: {err}"
        ) from err


def main():
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    doc = Document(str(DOC_PATH))
    if find_para(doc, contains="7.28 AI 沉浸增强六维") is not None:
        print(f"Already patched (7.28 present): {DOC_PATH}")
        return

    today = f"{date.today().year}年{date.today().month:02d}月{date.today().day:02d}日"
    n = 0

    # ---------- 报告头 ----------
    p = find_para(doc, startswith="报告生成日期")
    if p is not None and "沉浸增强六维" not in p.text:
        replace_paragraph_text(
            p,
            f"报告生成日期：{today}（在 AI 助手增强 CmdId 1098–1104 基础上，已同步「沉浸增强六维」："
            "箱庭 VLM 时序帧对比 + 多步空间锚点链（1100 anchor_chains）、"
            "战斗状态向量 + 微观干预 lock_next_skill_id（211/1101）、"
            "UID 长期记忆 + 治愈系 ProactiveCoach、"
            "三级漏斗本地/缓存/远程 + estimated_wait_ms、"
            "TEAM_CARD/SIM_BLESSING 结构化渲染、"
            "装备快照 + 1046 expected_dmg_increase_percent 养成优先级）",
        )
        n += 1

    # ---------- 玩家一天 ----------
    p = find_para(doc, startswith="为了让非技术同学建立整体画面")
    if p is not None and "空间锚点" not in p.text:
        text = p.text
        if "战斗里可一键执行 AI 建议的 Auto 策略" in text:
            text = text.replace(
                "战斗里可一键执行 AI 建议的 Auto 策略",
                "战斗里 AI 可锁定下一动技能并在执行后自动恢复 Auto；截图解惑会给出 1→2→3 步空间锚点 AR 指引",
            )
        if "长按助手可上传屏幕截图" in text and "配队海报" not in text:
            text = text.replace(
                "长按助手可上传屏幕截图做机关/敌人解惑（AR 高亮）",
                "长按助手可上传屏幕截图做机关/敌人解惑（时序帧对比 + 多步锚点链）；"
                "配队/模拟宇宙问题可返回 TEAM_CARD 等结构化 JSON 本地渲染",
            )
        replace_paragraph_text(p, text)
        n += 1

    # ---------- 7.26 补一句指向 7.28 ----------
    p = find_para(doc, startswith="测试：AssistEnhancementBusinessFlowTest")
    if p is not None and "AssistImmersionIntegrationTest" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；沉浸增强六维见 7.28 与第十章 10.7，回归 AssistImmersionIntegrationTest。",
        )
        n += 1

    # ---------- 新增 7.28（插在 7.27 之后） ----------
    anchor = find_para(doc, startswith="测试入口：ExperienceImmersionV3FlowsTest")
    if anchor is None:
        anchor = find_para(doc, contains="7.27 战斗手感")
    if anchor is None:
        anchor = find_para(doc, contains="WalletWalService")
    if anchor is not None:
        h = insert_paragraph_after(anchor, "7.28 AI 沉浸增强六维（协议字段扩展）", "Heading 2")
        p1 = insert_paragraph_after(
            h,
            "在 7.26 截图 VLM / 一键 Auto / 人设 TTS 之上，进一步补齐「看得懂动态机关、打得准下一动、"
            "记得住玩家情绪、答得快且可预期、渲染得像游戏内 UI、养谁算得清」六条链路。"
            "（1）箱庭空间智能：VisualFrameTemporalAnalyzer 对 2 秒窗口 JPEG 做帧差，识别浮台/位移；"
            "SpatialPuzzleHintBuilder 生成 AssistSpatialAnchorChain（sequence_step + world_x/y/z + direction_hint），"
            "经 AssistVisualHintScNotify（1100）下发 anchor_chains 与带 sequence_step 的 overlays。"
            "（2）动态战斗策略：BattleStateVectorSerializer 序列化战技点/终结技充能/BUFF 预警；"
            "HeuristicBattleAssistPolicy 输出 battle_state_summary；SuggestedAutoOverride 扩展 "
            "lock_next_skill_id / lock_caster_entity_id / auto_revert_after_action；"
            "ApplyAiSuggestion（1101）写入 BattleContext 微观干预队列，BattleAutoService 执行后 consumeRevertAutoAfterMicro。"
            "（3）长期记忆与情感：AssistLongTermMemoryService（Redis/内存，TTL 可配）按 UID 记录 battle_fail/gacha_fail 等；"
            "AssistGachaMemoryHook 在抽卡后写入；AssistProactiveCoachService 读取 shouldUseHealingPersona 切换治愈话术。"
            "（4）三级漏斗：AssistIntentFunnelService tier1 本地规则 <100ms、tier2 Jaccard≥0.85 缓存、tier3 远程 LLM；"
            "AskAiAssist / AiHintNotify / AssistVisualHint 携带 estimated_wait_ms。"
            "（5）AIGC 结构化渲染：AssistStructuredRenderService 输出 render_type + render_payload_json"
            "（TEAM_CARD / MATERIAL_TREE / SIM_BLESSING），客户端本地渲染竖版海报/材料树/祝福优先级。"
            "（6）养成决策：EquipmentAffixSnapshotService 注入 AssistPlayerProfileService；"
            "AssistCultivationAdvisorService 对比两名角色；CalculateUpgradeMaterialsScRsp（1046）"
            "返回 expected_dmg_increase_percent / recommended_priority_avatar_id / cultivation_reason。",
            "Normal",
        )
        insert_paragraph_after(
            p1,
            "实现入口：assist/visual/*、assist/memory/*、assist/funnel/*、assist/render/*、"
            "assist/cultivation/*、equipment/EquipmentAffixSnapshotService、"
            "assist_system.proto / character_system.proto 字段扩展；"
            "配置 lunarcore.ai-assist.intent-funnel-enabled、funnel-similar-cache-threshold、long-term-memory-ttl-days；"
            "测试 AssistImmersionIntegrationTest + AssistImmersionEnhancementsTest + AssistEnhancementBusinessFlowTest。",
            "List Bullet",
        )
        n += 1

    # ---------- 第十章引言 ----------
    p = find_para(doc, startswith="AI 模块的定位是「游戏内向导」")
    if p is not None and "三级漏斗" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "。沉浸增强另含：截图时序对比与多步空间锚点、战斗微观干预与状态向量、"
            "玩家长期记忆与治愈人设、三级漏斗 estimated_wait、结构化 TEAM_CARD 渲染、"
            "装备快照驱动的养成优先级与 1046 伤害预期字段。",
        )
        n += 1

    # ---------- 10.1 新能力子弹 ----------
    if find_para(doc, startswith="箱庭空间锚点链（时序 VLM）") is None:
        anchor = find_para(doc, startswith="语音输入：客户端 ASR")
        if anchor is None:
            anchor = find_para(doc, startswith="角色人设与 TTS")
        if anchor is not None:
            b1 = insert_paragraph_after(
                anchor,
                "箱庭空间锚点链（时序 VLM）：连续 UploadScreenshot 在 2 秒窗口做帧差；"
                "机关类问句生成 3 步 AssistSpatialAnchorChain（①花坛 ②石柱 ③浮台），"
                "1100 推送 anchor_chains + overlays.sequence_step + estimated_wait_ms。",
                "List Bullet",
            )
            b2 = insert_paragraph_after(
                b1,
                "战斗微观干预：AskBattleHint 返回 battle_state_summary；suggested_auto_override 可带 "
                "lock_next_skill_id；ApplyAiSuggestion 排队下一动技能，可选 auto_revert_after_action 执行后恢复 Auto。",
                "List Bullet",
            )
            b3 = insert_paragraph_after(
                b2,
                "长期记忆与情感陪伴：AssistLongTermMemoryService 记录连败/抽卡歪池等 salience 事件；"
                "promptBlock 注入多轮上下文；ProactiveCoach 对「连败愤怒」等标签启用治愈系 persona。",
                "List Bullet",
            )
            b4 = insert_paragraph_after(
                b3,
                "三级漏斗延迟优化：tier1 本地「怎么选/附近/对话」<100ms；tier2 高相似答案缓存 Jaccard≥0.85；"
                "tier3 远程 LLM 按场景返回 estimated_wait_ms（战斗/配队/养成不同档位）。",
                "List Bullet",
            )
            b5 = insert_paragraph_after(
                b4,
                "结构化渲染（AIGC JSON）：配队问句 → TEAM_CARD；模拟宇宙 → SIM_BLESSING；"
                "养成路线 → MATERIAL_TREE；AskAiAssistScRsp / AiHintNotify 带 render_type + render_payload_json。",
                "List Bullet",
            )
            insert_paragraph_after(
                b5,
                "养成优先级预览：1046 CalculateUpgradeMaterials 在材料清单之外附带 "
                "expected_dmg_increase_percent、recommended_priority_avatar_id、cultivation_reason"
                "（结合背包遗器词条与深渊弱点环境）。",
                "List Bullet",
            )
            n += 1

    # ---------- 10.2 类名 ----------
    p = find_para(doc, startswith="主工程：AiAssistApplicationService")
    if p is not None and "AssistIntentFunnelService" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；沉浸增强：VisualFrameTemporalAnalyzer、SpatialPuzzleHintBuilder、"
            "BattleStateVectorSerializer、AssistLongTermMemoryService、AssistGachaMemoryHook、"
            "AssistIntentFunnelService、AssistStructuredRenderService、AssistCultivationAdvisorService、"
            "EquipmentAffixSnapshotService。",
        )
        n += 1

    p = find_para(doc, startswith="多轮/画像/主动推送")
    if p is not None and "equipmentSnapshot" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；AssistPlayerProfileService 画像含 equipmentSnapshot；"
            "AssistLongTermMemoryService 供 promptBlock / 治愈人设判定。",
        )
        n += 1

    # ---------- 10.3 协议 ----------
    p = find_para(doc, startswith="协议扩展：AskAiAssist 支持 session_id")
    if p is not None and "anchor_chains" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "。AssistVisualHintScNotify 扩展 anchor_chains、estimated_wait_ms；"
            "SuggestedAutoOverride 扩展 lock_next_skill_id / auto_revert_after_action / battle_state_summary；"
            "AskAiAssistScRsp / AiHintNotify / AskBattleHintScRsp 扩展 estimated_wait_ms、render_type、render_payload_json；"
            "CalculateUpgradeMaterialsScRsp 扩展 expected_dmg_increase_percent 等养成预览字段。",
        )
        n += 1

    # ---------- 10.7 专节 ----------
    if find_para(doc, contains="10.7 AI 沉浸增强六维业务流程") is None:
        anchor = find_para(doc, startswith="实现入口：assist_system.proto、CmdIds 1098")
        if anchor is None:
            anchor = find_para(doc, contains="10.6 AI 助手增强七项")
        if anchor is not None:
            h = insert_paragraph_after(anchor, "10.7 AI 沉浸增强六维业务流程", "Heading 2")
            p1 = insert_paragraph_after(
                h,
                "（1）箱庭解谜：UploadScreenshot → 帧差检测动态机关 → 匹配 POI puzzleHint → "
                "生成 3 步 anchor_chains → 1100 AR 叠加。（2）战斗副官：AskBattleHint 读状态向量 → "
                "建议 lock 下一技能 → ApplyAiSuggestion → Auto 执行 → 可选恢复策略。"
                "（3）情感记忆：连败/歪池 → record → 下次 ask/coach promptBlock 引用 → 治愈 persona。"
                "（4）漏斗：tryTier1Local / tryTier2Cache 命中则秒回；否则异步 remote 并带 estimated_wait_ms。"
                "（5）渲染：tryRender 命中 TEAM_CARD 等 → 协议下发 JSON → 客户端本地海报。"
                "（6）养成：1046 计算材料同时 advise 两名角色 → 返回预期伤害提升百分比与理由。",
                "Normal",
            )
            insert_paragraph_after(
                p1,
                "测试：AssistImmersionIntegrationTest（17 条端到端）；"
                "配置 intent-funnel-enabled=true、funnel-similar-cache-threshold=0.85、long-term-memory-ttl-days=30。",
                "List Bullet",
            )
            n += 1

    # ---------- 7.2 养成计划补 AI 预览 ----------
    p = find_para(doc, startswith="养成计划 CalculateUpgradeMaterials")
    if p is None:
        p = find_para(doc, contains="CalculateUpgradeMaterials（1046")
    if p is not None and "expected_dmg" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "。1046 响应另含 AI 养成优先级预览：expected_dmg_increase_percent、"
            "recommended_priority_avatar_id、cultivation_reason（AssistCultivationAdvisorService + 装备快照）。",
        )
        n += 1

    # ---------- 结语 / 成熟度 ----------
    p = find_para(doc, startswith="Admin 热更/导入/审计")
    if p is not None and "沉浸增强" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；AI 沉浸增强：空间锚点链、战斗微观干预、长期记忆、三级漏斗、结构化渲染、1046 养成预览。",
        )
        n += 1

    p = find_para(doc, contains="AI 助手增强（1098–1104")
    if p is not None and "沉浸增强六维" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。") + "；沉浸增强六维（锚点链/微观干预/长期记忆/漏斗/结构化渲染/养成预览）。",
        )
        n += 1

    p = find_para(doc, startswith="继续以模块化单体打磨")
    if p is not None and "anchor_chains" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；客户端对齐 1100 anchor_chains、战斗 lock_next_skill 气泡、"
            "estimated_wait_ms 排队 UI、TEAM_CARD 本地渲染、1046 养成伤害预期条。",
        )
        n += 1

    p = find_para(doc, startswith="MyLunarCore 已经不是「只有空壳的演示仓库」")
    if p is not None and "三级漏斗" not in p.text:
        replace_paragraph_text(
            p,
            p.text.replace(
                "屏幕 VLM 解惑、一键 Auto、角色人设 TTS、语音输入与公告灌库 RAG）。",
                "屏幕 VLM 解惑（含多步空间锚点）、一键 Auto 与微观干预、角色人设 TTS、"
                "三级漏斗与 structured render、长期记忆陪伴、1046 养成预览、语音输入与公告灌库 RAG）。",
            ),
        )
        n += 1

    # ---------- 表格 ----------
    t7 = find_table_by_header(doc, "能力面", "代表类")
    if t7 is not None and not table_has(t7, "AssistIntentFunnelService"):
        add_table_row(
            t7,
            [
                "AI 沉浸增强六维",
                "VisualFrameTemporalAnalyzer / SpatialPuzzleHintBuilder / "
                "AssistLongTermMemoryService / AssistIntentFunnelService / "
                "AssistStructuredRenderService / AssistCultivationAdvisorService",
                "箱庭锚点链、战斗微观干预、长期记忆、三级漏斗、TEAM_CARD 渲染、1046 养成预览；"
                "协议字段 anchor_chains / estimated_wait_ms / render_type",
            ],
        )
        n += 1

    t15 = find_table_by_header(doc, "配置", "含义")
    if t15 is not None:
        rows = [
            ("lunarcore.ai-assist.intent-funnel-enabled", "是否启用三级漏斗（本地/缓存/远程，默认 true）"),
            (
                "lunarcore.ai-assist.funnel-similar-cache-threshold",
                "二级缓存 Jaccard 相似度阈值（默认 0.85）",
            ),
            ("lunarcore.ai-assist.long-term-memory-ttl-days", "玩家长期记忆 TTL 天数（默认 30）"),
        ]
        for key, meaning in rows:
            if not table_has(t15, key):
                add_table_row(t15, [key, meaning])
        n += 1

    t17 = find_table_by_header(doc, "术语", "人话解释")
    if t17 is not None:
        terms = [
            ("AssistSpatialAnchorChain", "解谜时服务端下发的 1→2→3 步世界坐标与方位描述，客户端按 sequence_step 画 AR 路线"),
            ("lock_next_skill_id", "AI 锁定「下一动」释放哪个技能，执行完可自动恢复 Auto 策略"),
            ("estimated_wait_ms", "异步远程大模型预计等待毫秒数，客户端可展示排队/加载条"),
            ("TEAM_CARD", "配队问题的结构化 JSON，客户端本地渲染竖版阵容海报而非纯文本"),
            ("expected_dmg_increase_percent", "1046 养成计划里 AI 对比两名角色后给出的预期伤害提升百分比"),
        ]
        for term, meaning in terms:
            if not table_has(t17, term):
                add_table_row(t17, [term, meaning])
        n += 1

    t18 = find_table_by_header(doc, "别人常问", "你可以这样答")
    if t18 is not None and not table_has(t18, "空间锚点"):
        add_table_row(
            t18,
            [
                "AI 能不能指引多步机关？养谁更划算？",
                "可以：截图解惑会下发 anchor_chains 多步 AR 锚点；1046 养成计划附带 expected_dmg 对比；"
                "配队问题可返回 TEAM_CARD JSON 本地渲染。",
            ],
        )
        n += 1

    t3 = find_table_by_header(doc, "包名", "普通人理解")
    if t3 is not None:
        patch_table_cell_contains(
            t3,
            "assist",
            append=(2, "；沉浸增强：锚点链、微观干预、长期记忆、三级漏斗、结构化渲染、养成预览"),
        )
        n += 1

    t6 = find_table_by_header(doc, "文件/目录", "用途")
    if t6 is not None and not table_has(t6, "EquipmentAffixPool"):
        add_table_row(
            t6,
            [
                "EquipmentAffixPool.json / 背包 subAffixesJson",
                "装备词条快照：AssistCultivationAdvisorService 对比养成优先级与 1046 expected_dmg",
            ],
        )
        n += 1

    safe_save(doc, DOC_PATH)
    print(f"（命中补丁约 {n} 处）")


if __name__ == "__main__":
    main()
