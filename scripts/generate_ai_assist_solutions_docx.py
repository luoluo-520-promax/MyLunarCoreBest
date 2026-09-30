# -*- coding: utf-8 -*-
"""Generate MyLunarCore AI player-assist solutions Word document."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DESKTOP = Path.home() / "Desktop"
# Explicit Unicode avoids Windows console/source encoding mangling the filename.
OUTPUT = DESKTOP / ("MyLunarCore_AI" + "\u8f85\u52a9\u529f\u80fd\u65b9\u6848" + ".docx")


def set_doc_fonts(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15


def add_title(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(25, 75, 140)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_subtitle(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(100, 100, 100)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_heading(doc: Document, text: str, level: int = 1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        if level == 1:
            run.font.color.rgb = RGBColor(25, 75, 140)


def add_body(doc: Document, text: str):
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_bullet(doc: Document, text: str, level: int = 0):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.5 * (level + 1))
    for run in p.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_code(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 30, 30)


def add_table(doc: Document, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.name = "微软雅黑"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                run.font.size = Pt(10)
    for r_idx, row in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = val
            for p in row_cells[c_idx].paragraphs:
                for run in p.runs:
                    run.font.name = "微软雅黑"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                    run.font.size = Pt(10)
    return table


def build():
    doc = Document()
    set_doc_fonts(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    add_title(doc, "MyLunarCore AI 辅助功能接入方案")
    add_subtitle(doc, f"面向玩家游玩体验增强 · 方案对比与落地建议 · {date.today().isoformat()}")
    add_body(
        doc,
        "本文档结合当前 MyLunarCore 游戏服（Spring Boot + Netty/KCP + Protobuf，模块化单体为主，"
        "含 battle / gacha / quest / hall / character / economy / matchmaking / activity / item / scene）"
        "的现有架构，给出若干可落地的 AI 辅助玩法方案，帮助玩家更好理解系统、规划养成与完成目标。",
    )

    # ---------- 1 ----------
    add_heading(doc, "一、目标与原则", 1)
    add_heading(doc, "1.1 产品目标", 2)
    add_bullet(doc, "降低新手门槛：任务、场景、养成路径可解释、可引导。")
    add_bullet(doc, "提升中后期效率：阵容/天赋/商店/活动优先级给出可执行建议。")
    add_bullet(doc, "增强战斗体验：可选战术提示、复盘、智能托管（需防作弊与公平性约束）。")
    add_bullet(doc, "不破坏核心玩法乐趣：建议为主、代操为辅；关键消费决策需透明可关闭。")

    add_heading(doc, "1.2 技术原则", 2)
    add_bullet(doc, "AI 不得直接写权威游戏状态（货币、抽卡结果、战斗结算）；只读上下文 + 建议输出。")
    add_bullet(doc, "与 Netty 主游戏循环解耦：推理走异步/旁路服务，避免阻塞 tick 与战斗帧。")
    add_bullet(doc, "优先复用现有 ApplicationService / NettyService 作为上下文采集点，不侵入结算逻辑。")
    add_bullet(doc, "配置热更（活动、卡池、商店）后，知识库/规则库需可同步刷新。")
    add_bullet(doc, "合规：提示词与日志脱敏；禁止泄露其他玩家隐私；抽卡/付费建议需合规文案。")

    # ---------- 2 ----------
    add_heading(doc, "二、现有系统与 AI 切入点", 1)
    add_body(doc, "下表列出玩家侧系统与建议的服务侧接入点（均已存在于仓库）：")
    add_table(
        doc,
        ["系统", "玩家价值", "建议接入点", "典型 AI 能力"],
        [
            ["战斗 Battle", "通关与操作", "BattleNettyService / BattleManager", "技能建议、复盘、智能托管"],
            ["抽卡 Gacha", "资源决策", "GachaNettyService / GachaDrawEngine", "卡池解读、保底进度说明"],
            ["任务 Quest", "目标导航", "QuestNettyService / QuestProgress*", "下一步提示、路线规划"],
            ["大厅 Hall", "社交沟通", "HallNettyService / ChatService", "聊天助手、邮件摘要"],
            ["角色 Character", "养成规划", "CharacterNettyService / Talent*", "Build、天赋、突破建议"],
            ["经济 Shop", "资源分配", "EconomyNettyService / Shop*", "性价比推荐、消费规划"],
            ["匹配 Match", "组队体验", "MatchNettyService / MatchmakingService", "阵容匹配、队列建议"],
            ["活动 Activity", "限时目标", "ActivityNettyService / ActivityQuery*", "今日优先事项"],
            ["道具 Item", "背包管理", "ItemNettyService / ItemApplicationService", "强化/分解/对比建议"],
            ["场景 Scene", "世界探索", "SceneNettyService / SceneManager", "导航、NPC/任务点提示"],
        ],
    )

    # ---------- 3 ----------
    add_heading(doc, "三、方案总览对比", 1)
    add_table(
        doc,
        ["方案", "核心思路", "实现成本", "效果上限", "运维成本", "推荐阶段"],
        [
            ["方案A 规则教练", "配置+规则引擎", "低", "中", "低", "第1期 MVP"],
            ["方案B LLM+RAG 助手", "大模型+游戏知识库", "中", "高", "中", "第2期体验"],
            ["方案C AI 旁路微服务", "独立 ai-service", "中高", "高", "中高", "与微服务演进同步"],
            ["方案D 领域小模型", "战斗/匹配专用模型", "高", "很高（垂直）", "高", "有数据后"],
            ["方案E 客户端本地助手", "端侧推理/混合", "中", "中高", "中", "移动端成熟后"],
        ],
    )

    # ---------- 方案 A ----------
    add_heading(doc, "四、方案 A：规则 / 启发式「游戏教练」（推荐先做）", 1)
    add_heading(doc, "4.1 思路", 2)
    add_body(
        doc,
        "不依赖大模型。用策划配置表 + 确定性规则，根据玩家当前进度（任务、等级、体力类资源、活动开启状态）"
        "生成「下一步建议」。适合快速上线、结果可预期、易测试、零 Token 成本。",
    )

    add_heading(doc, "4.2 架构", 2)
    add_code(
        doc,
        "Client  --Netty/Proto-->  PlayerAssistNettyService\n"
        "                              |\n"
        "                              v\n"
        "                     PlayerCoachApplicationService\n"
        "                     (读取 Quest/Activity/Character/Item 快照)\n"
        "                              |\n"
        "                              v\n"
        "                     CoachRuleEngine + CoachTips.json",
    )

    add_heading(doc, "4.3 落地步骤", 2)
    add_bullet(doc, "新增协议：AskCoachHintReq / AskCoachHintRsp（CmdIds 新增区间），或复用大厅聊天式查询。")
    add_bullet(doc, "新增 PlayerCoachApplicationService：通过 PlayerContextResolver 取只读 PlayerData。")
    add_bullet(doc, "规则示例：未完成主线 → 推任务；活动即将结束且可领奖 → 推活动；突破材料齐备 → 推养成。")
    add_bullet(doc, "配置放 data/CoachTips.json，走现有 ConfigFileService / 热更管线。")
    add_bullet(doc, "单测：给定 PlayerData fixture，断言 tip 优先级与文案。")

    add_heading(doc, "4.4 优缺点", 2)
    add_bullet(doc, "优点：可控、便宜、延迟低、易合规；与现有热更体系契合。")
    add_bullet(doc, "缺点：表达能力弱，难处理开放式提问与复杂阵容推理。")

    # ---------- 方案 B ----------
    add_heading(doc, "五、方案 B：LLM + RAG「智能玩伴」", 1)
    add_heading(doc, "5.1 思路", 2)
    add_body(
        doc,
        "玩家用自然语言提问（「今天该刷什么」「这角色怎么加点」「保底还有多少」）。"
        "服务端组装玩家上下文 + 检索游戏配置知识库（活动/任务/卡池/商店说明），调用大模型生成回答。"
        "模型只做解释与建议，不执行扣费/抽卡/改战斗结果。",
    )

    add_heading(doc, "5.2 架构", 2)
    add_code(
        doc,
        "Client  -->  AiAssistNettyService (限流)\n"
        "               |\n"
        "               +--> ContextBuilder (PlayerAggregate 只读投影)\n"
        "               +--> GameKnowledgeIndexer (ActivityConfigs / QuestConfigs / Banners / Shop)\n"
        "               |\n"
        "               v\n"
        "          AiAssistGateway (HTTP) --> 云端/私有 LLM\n"
        "               |\n"
        "               v\n"
        "          SafetyFilter + 审计日志 --> Rsp 下发",
    )

    add_heading(doc, "5.3 落地步骤", 2)
    add_bullet(doc, "知识库：将 data/*.json 与活动说明切块入库（向量库可用 Redis/本地 FAISS/轻量 pgvector）。")
    add_bullet(doc, "热更后触发 VersionHotReload / ActivityConfigHotReload 钩子，异步重建索引。")
    add_bullet(doc, "ContextBuilder 字段白名单：等级、主线进度、背包关键材料、卡池 pity、今日活动状态等。")
    add_bullet(doc, "限流：复用 ConnectionPacketRateLimiter / GameLoginRateLimiter 模式，按 uid 配额。")
    add_bullet(doc, "提示词模板：角色设定为「官方游戏助手」；禁止编造未上架内容；不确定则引导打开对应界面。")
    add_bullet(doc, "可与 Hall ChatService 结合：@助手 触发；邮件摘要走 MailApplicationService 只读。")

    add_heading(doc, "5.4 场景示例", 2)
    add_bullet(doc, "任务：根据 QuestProgress 解释当前目标与推荐场景入口。")
    add_bullet(doc, "抽卡：解释当前 Banner 机制与 pity，不做「必出」承诺。")
    add_bullet(doc, "养成：结合 AttributeCalculator / Talent 配置给出加点方向。")
    add_bullet(doc, "活动：ActivityScheduleService 输出「今日 ROI 排序」。")

    add_heading(doc, "5.5 优缺点", 2)
    add_bullet(doc, "优点：体验强、覆盖开放问答；与配置驱动内容天然契合。")
    add_bullet(doc, "缺点：成本与延迟、幻觉风险；需内容安全与付费话术审核。")

    # ---------- 方案 C ----------
    add_heading(doc, "六、方案 C：独立 AI 旁路微服务（契合 microservices 演进）", 1)
    add_heading(doc, "6.1 思路", 2)
    add_body(
        doc,
        "仓库已有 microservices 脚手架（api-gateway / auth-service / player-service）。"
        "将 AI 作为 ops/体验旁路服务 ai-assist-service，游戏核心仍保持模块化单体；"
        "游戏服通过 HTTP/Feign 或消息队列异步请求建议，结果用 Push/Notify 回客户端。"
        "符合现有「Admin/Config、Leaderboard 可拆；Battle/Player 不轻易拆」的判断。",
    )

    add_heading(doc, "6.2 架构", 2)
    add_code(
        doc,
        "GameServer (Netty) --异步--> ai-assist-service\n"
        "                         |-- LLM Provider Adapter\n"
        "                         |-- Rule Coach Engine\n"
        "                         |-- RAG Index\n"
        "                         |-- Audit / Quota\n"
        "Client <-- UpdateNotifyBroadcaster / 专用 AiHintNotify",
    )

    add_heading(doc, "6.3 落地步骤", 2)
    add_bullet(doc, "在 microservices 下新增 services/ai-assist-service，经 api-gateway 暴露管理与健康检查。")
    add_bullet(doc, "游戏服增加 AiAssistClient（类似 HttpCenterRoutingClient 模式），失败降级到规则教练。")
    add_bullet(doc, "鉴权：内部服务 token；玩家请求必须带本服校验过的 uid/session。")
    add_bullet(doc, "观测：接入现有 GameTrafficMetrics / ServerMetricsMonitor 维度（ai_latency、ai_error、quota）。")

    add_heading(doc, "6.4 优缺点", 2)
    add_bullet(doc, "优点：弹性扩缩容、模型升级不重启游戏服；故障隔离好。")
    add_bullet(doc, "缺点：多一跳网络与部署复杂度；需定义稳定的上下文 DTO 契约。")

    # ---------- 方案 D ----------
    add_heading(doc, "七、方案 D：领域专用小模型（战斗 / 匹配 / 养成）", 1)
    add_heading(doc, "7.1 思路", 2)
    add_body(
        doc,
        "对延迟与确定性要求高的场景，不用通用大模型，而训练/部署轻量策略模型："
        "战斗技能选择、匹配分池、阵容推荐。输入为结构化特征，输出为离散动作或分数。",
    )

    add_heading(doc, "7.2 与现有战斗系统结合", 2)
    add_bullet(doc, "BattleManager / BattleContext：在「玩家请求提示」或「自动战斗」分支调用 PolicyModel。")
    add_bullet(doc, "离线用历史战报或自对弈模拟生成训练集；在线推理用 ONNX Runtime / 纯 Java 规则+轻量 NN。")
    add_bullet(doc, "公平性：PVP 必须同源模型版本；禁止偷看对手隐藏信息；可开关「智能托管」。")
    add_bullet(doc, "MatchmakingService：用 embedding 或评分模型做角色互补匹配（在 MatchQueue 入队评分）。")

    add_heading(doc, "7.3 落地步骤", 2)
    add_bullet(doc, "先做可解释启发式 auto-battle，再替换为模型；保持接口 BattleAssistPolicy 稳定。")
    add_bullet(doc, "建立 BattleEndedEvent 落库特征，形成数据飞轮。")
    add_bullet(doc, "灰度：按账号白名单 / 难度副本开启。")

    add_heading(doc, "7.4 优缺点", 2)
    add_bullet(doc, "优点：低延迟、可嵌入战斗循环；垂直效果强。")
    add_bullet(doc, "缺点：需要数据和算法人力；版本迭代与反作弊成本高。")

    # ---------- 方案 E ----------
    add_heading(doc, "八、方案 E：客户端本地 / 混合助手", 1)
    add_heading(doc, "8.1 思路", 2)
    add_body(
        doc,
        "端上缓存静态攻略与轻量模型，常见问题本地回答；复杂问题或需实时玩家权威数据时再请求服务端。"
        "适合弱网、降本与隐私敏感场景。",
    )
    add_heading(doc, "8.2 落地要点", 2)
    add_bullet(doc, "服务端提供「攻略包」版本号（可挂 VersionNettyService / 热更通知）。")
    add_bullet(doc, "权威数值仍以服务端 PlayerData 为准；客户端仅展示建议。")
    add_bullet(doc, "与方案 B/C 组合：Local First + Cloud Fallback。")

    # ---------- 组合 ----------
    add_heading(doc, "九、推荐落地路线（组合策略）", 1)
    add_body(doc, "建议按「先规则、后 LLM、再垂直模型、微服务化」推进，降低一次上线风险：")
    add_table(
        doc,
        ["阶段", "内容", "周期建议", "成功标准"],
        [
            ["P0", "方案A：今日建议 + 任务下一步 + 活动提醒", "1–2 周", "点击率、任务完成率提升"],
            ["P1", "方案B：大厅自然语言助手（限量）+ RAG 配置库", "3–6 周", "有用率反馈、投诉可控"],
            ["P2", "方案C：抽离 ai-assist-service，游戏服仅客户端", "与微服务节奏对齐", "故障不影响战斗"],
            ["P3", "方案D：副本智能托管 / 阵容推荐小模型", "有数据后 1–2 季", "通关时长、胜率、留存"],
            ["可选", "方案E：端侧攻略包", "客户端排期", "弱网可用、Token 下降"],
        ],
    )

    # ---------- 协议与代码骨架 ----------
    add_heading(doc, "十、建议的最小代码骨架（与仓库风格对齐）", 1)
    add_body(doc, "包路径建议：cn.itcast.demo.mylunarcore.assist（与 character/economy/hall 同级）。")
    add_bullet(doc, "AiAssistProperties：开关、模型 endpoint、超时、每分钟配额、是否允许战斗提示。")
    add_bullet(doc, "PlayerCoachApplicationService：方案 A 实现。")
    add_bullet(doc, "LlmAssistGateway + RagKnowledgeService：方案 B 实现。")
    add_bullet(doc, "AssistNettyService + AssistPacketHandlers：协议入口。")
    add_bullet(doc, "UpdateNotifyBroadcaster 可推送 AiHintNotify（异步结果）。")
    add_body(doc, "伪代码示意：")
    add_code(
        doc,
        "@Service\n"
        "public class PlayerCoachApplicationService {\n"
        "  public CoachHint suggest(long uid) {\n"
        "    PlayerData pd = playerContextResolver.require(uid); // 只读\n"
        "    return coachRuleEngine.evaluate(pd, activityQueryService.listActive());\n"
        "  }\n"
        "}\n"
        "\n"
        "// Netty 层：限流后调用 ApplicationService，禁止在 Handler 内调 LLM 阻塞 EventLoop",
    )

    # ---------- 风险 ----------
    add_heading(doc, "十一、风险与对策", 1)
    add_table(
        doc,
        ["风险", "对策"],
        [
            ["LLM 幻觉导致误导消耗", "RAG 强约束 + 「非配置内容不建议」+ 引用配置 ID"],
            ["阻塞 Netty EventLoop", "业务线程池/旁路服务；超时降级规则教练"],
            ["成本失控", "按 uid 日配额、缓存相同问题、端侧优先"],
            ["PVP 不公平 / 外挂质疑", "仅 PVE 开启智能操作；PVP 禁用或双方同源"],
            ["隐私与日志合规", "上下文脱敏、审计、保留期限"],
            ["与抽卡/付费暗示冲突", "禁止诱导充值话术；概率说明走官方文案模板"],
            ["热更后知识过期", "挂钩 HotReloadCoordinator / ActivityConfigHotReloadService"],
        ],
    )

    # ---------- 结论 ----------
    add_heading(doc, "十二、结论", 1)
    add_body(
        doc,
        "对 MyLunarCore 而言，最优路径不是一上来接大模型，而是："
        "用方案 A 快速建立「可解释的下一步建议」产品形态与协议面；"
        "再用方案 B 提升开放问答体验；"
        "在微服务脚手架成熟时用方案 C 隔离成本与故障；"
        "待战斗与匹配数据充足后，用方案 D 做垂直增强；"
        "客户端能力允许时叠加方案 E 降本。",
    )
    add_body(
        doc,
        "无论采用哪条路径，请坚持：AI 只读权威状态、异步推理、可降级、可热更、可审计。"
        "这样既能帮助玩家更好游玩，又不会破坏现有战斗结算、经济与抽卡的一致性边界。",
    )

    add_subtitle(doc, "—— 文档结束 ——")
    doc.save(OUTPUT)
    print(f"Wrote: {OUTPUT}")


if __name__ == "__main__":
    build()
