# -*- coding: utf-8 -*-
"""Generate MyLunarCore character skin design + shop listing solutions Word document."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DESKTOP = Path.home() / "Desktop"
OUTPUT = DESKTOP / ("MyLunarCore_" + "\u89d2\u8272\u76ae\u80a4\u8bbe\u8ba1\u4e0e\u5546\u5e97\u4e0a\u67b6\u65b9\u6848" + ".docx")


def set_doc_fonts(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15


def add_title(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(25, 75, 140)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_subtitle(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(100, 100, 100)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_heading(doc: Document, text: str, level: int = 1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        if level == 1:
            run.font.color.rgb = RGBColor(25, 75, 140)


def add_body(doc: Document, text: str):
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_bullet(doc: Document, text: str, level: int = 0):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.5 * (level + 1))
    for run in p.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_code(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 30, 30)


def add_table(doc: Document, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.name = "微软雅黑"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                run.font.size = Pt(10)
    for r_idx, row in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = val
            for p in row_cells[c_idx].paragraphs:
                for run in p.runs:
                    run.font.name = "微软雅黑"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                    run.font.size = Pt(9)
    doc.add_paragraph()


def build():
    doc = Document()
    set_doc_fonts(doc)
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    add_title(doc, "MyLunarCore 角色皮肤设计与商店上架方案")
    add_subtitle(
        doc,
        "皮肤配置 · 拥有与穿戴 · 商城/IAP 上架  |  " + date.today().isoformat(),
    )
    add_body(
        doc,
        "本文基于当前 MyLunarCore 实现给出可落地的角色皮肤方案：角色侧已有 AvatarEntity / "
        "AvatarRepository（成长、编队），商城侧已有 ShopConfigs.json、游戏币购买"
        "（ShopApplicationService）与 IAP 发货（IapOrderService / IapGrantService，"
        "rewards 支持 CURRENCY / ITEM / DAILY_CLAIM）。现状缺少「皮肤」领域模型与穿戴协议，"
        "本文补齐设计、配置、发货与上架流程。",
    )

    # ---------- 1 ----------
    add_heading(doc, "1. 目标与范围", 1)
    add_bullet(doc, "为目标角色设计可售皮肤（外观替换，不改变战斗数值）。")
    add_bullet(doc, "玩家可购买、拥有、穿戴/卸下皮肤；默认皮肤永久免费。")
    add_bullet(doc, "商店同时支持：游戏币兑换皮肤、晶石兑换皮肤、真钱 IAP 直购皮肤/皮肤礼包。")
    add_bullet(doc, "与现有商城限购、售卖时间窗、热更配置发布链路对齐。")
    add_body(doc, "非目标（本期不做）：皮肤带来属性加成、皮肤染色自由组合、跨角色通用时装。")

    # ---------- 2 ----------
    add_heading(doc, "2. 现状与缺口", 1)
    add_heading(doc, "2.1 已具备", 2)
    add_bullet(doc, "角色实例：avatar 表 / AvatarEntity（avatarId、level、exp、promotion、rank）。")
    add_bullet(doc, "道具背包：GameItemEntity（含 equipAvatarId，可用于“皮肤道具绑定到角色”的过渡方案）。")
    add_bullet(doc, "游戏币商店：shopId=1 兑换商店，BuyShopItem → 扣币 + addSimpleItem。")
    add_bullet(
        doc,
        "氪金商城：shopId=2，productCategory=DIRECT_TOPUP / DISCOUNT_PACK / PAID_PACK，"
        "payType=IAP，CreateOrder → 验单 → Grant。",
    )
    add_bullet(doc, "配置热更：ConfigImport / HotReload 可发布 ShopConfigs.json。")

    add_heading(doc, "2.2 缺口", 2)
    add_bullet(doc, "无 SkinConfig / 皮肤归属表；AvatarEntity 无 equippedSkinId 字段。")
    add_bullet(doc, "RewardSpec 无 SKIN 类型；IAP 发货只认货币与普通道具。")
    add_bullet(doc, "无「穿戴皮肤」Netty 协议；战斗/场景同步未带皮肤外观 id。")
    add_bullet(doc, "ShopConfigs 尚无 SKIN / SKIN_PACK 品类与皮肤页签。")

    # ---------- 3 ----------
    add_heading(doc, "3. 总体方案（推荐）", 1)
    add_body(
        doc,
        "采用「配置驱动 + 皮肤道具发货 + 角色穿戴状态」三段式：策划在 SkinConfigs.json "
        "定义皮肤；商城以 ITEM（皮肤道具）或新增 SKIN 奖励发货；玩家背包拥有后对指定 "
        "avatarId 执行 EquipSkin，写入角色当前外观。数值战斗不读皮肤，仅客户端表现与"
        "场景同步使用 equippedSkinId。",
    )
    add_heading(doc, "3.1 为什么推荐「皮肤道具」而非仅 entitlement 标记", 2)
    add_bullet(doc, "复用现有背包、邮件附件、活动奖励、IAP rewards.type=ITEM，改造成本最低。")
    add_bullet(doc, "皮肤道具 itemId 与 skinId 一一映射，运营可走礼包、兑换、活动多通道投放。")
    add_bullet(doc, "穿戴层仍以「已拥有 skinId」校验，避免重复穿戴未拥有皮肤。")
    add_heading(doc, "3.2 数据流", 2)
    add_code(
        doc,
        "配置: SkinConfigs.json + Items(皮肤道具) + ShopConfigs(上架)\n"
        "购买: 游戏币 BuyShopItem 或 IAP CreateOrder/Confirm\n"
        "发货: IapGrantService / ShopApplicationService → 背包获得皮肤道具(type=SKIN)\n"
        "穿戴: EquipSkin(avatarId, skinId) → 校验拥有 → avatar.equippedSkinId\n"
        "同步: GetAvatar / 编队 / 场景外观包带上 skinId",
    )

    # ---------- 4 ----------
    add_heading(doc, "4. 皮肤配置设计", 1)
    add_body(doc, "新增 data/SkinConfigs.json（热更友好），建议字段如下：")
    add_table(
        doc,
        ["字段", "类型", "说明"],
        [
            ["skinId", "int", "皮肤唯一 ID，建议 8xxxxxx 段"],
            ["avatarId", "int", "绑定角色模板 ID；仅该角色可穿戴"],
            ["itemId", "int", "对应背包道具 ID，购买/发货用"],
            ["name", "string", "展示名"],
            ["rarity", "int", "稀有度 3/4/5，影响商城角标"],
            ["isDefault", "bool", "默认皮肤；免费且创角自动拥有"],
            ["resourceKey", "string", "客户端资源包/预制体 key"],
            ["previewIcon", "string", "商城/衣柜图标"],
            ["tags", "string[]", "如 LIMITED / FESTIVAL / COLLAB"],
            ["obtainTips", "string", "获取途径文案"],
            ["enabled", "bool", "总开关；false 不下发"],
        ],
    )
    add_code(
        doc,
        """{
  "schemaVersion": 1,
  "skins": [
    {
      "skinId": 8001001,
      "avatarId": 1001,
      "itemId": 8101001,
      "name": "默认外观",
      "rarity": 3,
      "isDefault": true,
      "resourceKey": "avatar/1001/skin_default",
      "previewIcon": "ui/skin/8001001",
      "tags": ["DEFAULT"],
      "enabled": true
    },
    {
      "skinId": 8001002,
      "avatarId": 1001,
      "itemId": 8101002,
      "name": "星海礼赞",
      "rarity": 5,
      "isDefault": false,
      "resourceKey": "avatar/1001/skin_star_ocean",
      "previewIcon": "ui/skin/8001002",
      "tags": ["LIMITED", "FESTIVAL"],
      "obtainTips": "皮肤商店 / 限时礼包",
      "enabled": true
    }
  ]
}""",
    )

    add_heading(doc, "4.1 道具表约定", 2)
    add_bullet(doc, "皮肤道具 type 建议固定为 SKIN（或数值 type=7，与现有 type=3 普通道具区分）。")
    add_bullet(doc, "count 恒为 1；重复购买按「已拥有」拦截或转为碎片（二期）。")
    add_bullet(doc, "不可堆叠、不可交易、不可分解（服务端校验）。")

    add_heading(doc, "4.2 角色穿戴状态", 2)
    add_body(doc, "AvatarEntity / avatar 表增加字段：")
    add_bullet(doc, "equippedSkinId int：当前穿戴皮肤；0 或空表示使用该角色默认皮肤。")
    add_body(doc, "可选冗余表 player_skin_owned(player_id, skin_id, obtained_at, source) 便于查询衣柜；MVP 可用「背包是否持有对应 itemId」判定拥有。")

    # ---------- 5 ----------
    add_heading(doc, "5. 商店上架方案", 1)
    add_heading(doc, "5.1 商品品类扩展", 2)
    add_table(
        doc,
        ["品类", "productCategory", "payType", "玩家感知", "限购建议"],
        [
            [
                "皮肤兑换",
                "SKIN_SHOP",
                "CURRENCY",
                "皮肤商店页：晶石/活动币兑换",
                "终身限购 1；已拥有置灰",
            ],
            [
                "皮肤直购",
                "SKIN",
                "IAP",
                "真钱直购单皮肤",
                "终身限购 1",
            ],
            [
                "皮肤礼包",
                "SKIN_PACK",
                "IAP 或 CURRENCY",
                "皮肤+货币/材料组合；可打折",
                "限时窗 + 终身/活动期限购",
            ],
        ],
    )
    add_body(
        doc,
        "在 ShopConfigs 增加 shopId=3「皮肤商店」，tabs=[\"SKIN\", \"SKIN_PACK\"]；"
        "也可把皮肤商品挂到现有氪金商城 shopId=2 的新页签，两种皆可，推荐独立商店便于运营筛选。",
    )

    add_heading(doc, "5.2 游戏币上架示例", 2)
    add_code(
        doc,
        """{
  "shopId": 3,
  "shopName": "皮肤商店",
  "tabs": ["SKIN", "SKIN_PACK"],
  "items": [
    {
      "shopItemId": 3001,
      "productCategory": "SKIN_SHOP",
      "payType": "CURRENCY",
      "itemId": 8101002,
      "itemCount": 1,
      "currencyId": 2,
      "price": 1680,
      "displayName": "星海礼赞",
      "dailyLimit": 0,
      "totalLimit": 1,
      "limitPeriod": "LIFETIME",
      "sortOrder": 10,
      "enabled": true
    }
  ]
}""",
    )

    add_heading(doc, "5.3 IAP 直购 / 礼包示例", 2)
    add_code(
        doc,
        """{
  "shopItemId": 3101,
  "productCategory": "SKIN",
  "payType": "IAP",
  "skuId": "com.mylunarcore.skin.8001002",
  "displayName": "星海礼赞",
  "priceCents": 6800,
  "currencyCode": "CNY",
  "rewards": [
    {"type": "ITEM", "itemId": 8101002, "count": 1}
  ],
  "totalLimit": 1,
  "limitPeriod": "LIFETIME",
  "sortOrder": 10,
  "enabled": true
},
{
  "shopItemId": 3201,
  "productCategory": "SKIN_PACK",
  "payType": "IAP",
  "skuId": "com.mylunarcore.skinpack.star_ocean.68",
  "displayName": "星海礼赞礼包",
  "priceCents": 6800,
  "originalPriceCents": 12800,
  "discountRate": 53,
  "saleStartAt": "2026-08-01T00:00:00+08:00",
  "saleEndAt": "2026-08-31T23:59:59+08:00",
  "packKind": "SKIN_PACK",
  "rewards": [
    {"type": "ITEM", "itemId": 8101002, "count": 1},
    {"type": "CURRENCY", "currencyId": 2, "amount": 300},
    {"type": "ITEM", "itemId": 101, "count": 5}
  ],
  "totalLimit": 1,
  "limitPeriod": "LIFETIME",
  "sortOrder": 20,
  "enabled": true
}""",
    )

    add_heading(doc, "5.4 发货与幂等", 2)
    add_bullet(doc, "游戏币路径：沿用 ShopApplicationService.buy；购买前增加「已拥有皮肤则失败」。")
    add_bullet(doc, "IAP 路径：沿用 IapGrantService，rewards 发 ITEM；order_id 幂等不变。")
    add_bullet(doc, "可选增强：RewardSpec 增加 type=SKIN + skinId，Grant 时写 owned 表并同步塞道具。")
    add_bullet(doc, "重复发货：若已拥有，记审计日志并跳过道具，避免多份皮肤道具；礼包内其它奖励仍正常发。")

    add_heading(doc, "5.5 定价与运营建议", 2)
    add_table(
        doc,
        ["档位", "建议定价", "适用", "备注"],
        [
            ["普通皮肤", "680~1280 晶石 或 ¥12~¥30", "常驻店", "可用活动币部分兑换"],
            ["限定皮肤", "¥68 / 1680 晶石", "节日/版本", "强依赖 saleStart/End"],
            ["联动皮肤", "¥98~¥128 + 礼包", "IP 联动", "礼包含材料拉高客单"],
            ["默认皮肤", "免费", "创角赠送", "不可上架付费"],
        ],
    )

    # ---------- 6 ----------
    add_heading(doc, "6. 服务端模块落地", 1)
    add_heading(doc, "6.1 新增类（建议包路径）", 2)
    add_table(
        doc,
        ["类", "职责"],
        [
            ["skin.SkinConfigRepository", "加载/热更 SkinConfigs.json"],
            ["skin.SkinOwnershipService", "拥有判定、默认皮肤授予、重复购买拦截"],
            ["skin.SkinEquipApplicationService", "穿戴/卸下校验与写库"],
            ["skin.SkinNettyService", "GetSkinWardrobe / EquipSkin 协议"],
            ["character 扩展", "创角时授予默认皮肤；Avatar 快照带 equippedSkinId"],
        ],
    )

    add_heading(doc, "6.2 关键协议（建议）", 2)
    add_code(
        doc,
        "GetSkinWardrobeReq { avatarId }\n"
        "GetSkinWardrobeRsp { skins: [{skinId, owned, equipped, name, rarity}] }\n"
        "EquipSkinReq { avatarId, skinId }  // skinId=0 还原默认\n"
        "EquipSkinRsp { retcode, avatarId, equippedSkinId }\n"
        "商店列表已有 ShopItem：展示时按 itemId 反查 SkinConfig 填预览与角色绑定信息",
    )

    add_heading(doc, "6.3 与现有代码的衔接点", 2)
    add_bullet(doc, "ShopApplicationService.buy：皮肤商品增加 ownership 前置校验。")
    add_bullet(doc, "IapGrantService.grantRewards：ITEM 发货后若 item 为皮肤类型，触发 SkinOwnershipService.markOwned。")
    add_bullet(doc, "EconomyNettyService 拉商店列表：对 SKIN* 品类附加 avatarId/skinId/owned 标记。")
    add_bullet(doc, "战斗/场景同步：EntityState 或角色外观字段增加 skinId（仅表现）。")
    add_bullet(doc, "配置发布：SkinConfigs.json 纳入 ConfigImport / HotReload，与 ShopConfigs 同批上架。")

    # ---------- 7 ----------
    add_heading(doc, "7. 客户端配合要点", 1)
    add_bullet(doc, "资源按 resourceKey 分包下载；未下载完时商城可预览静态图，穿戴提示「资源下载中」。")
    add_bullet(doc, "衣柜 UI：按角色筛选；未拥有显示获取途径（商店跳转 shopItemId）。")
    add_bullet(doc, "编队/战斗准备界面展示当前皮肤；切换后本地乐观更新，等 EquipSkinRsp 确认。")
    add_bullet(doc, "IAP：皮肤直购走现有 CreateOrder/Confirm，SKU 与渠道后台预先录入。")

    # ---------- 8 ----------
    add_heading(doc, "8. 上架操作清单（策划/运营）", 1)
    add_body(doc, "单皮肤上架标准流程：")
    add_bullet(doc, "1）美术交付 resourceKey 资源包与 previewIcon。")
    add_bullet(doc, "2）在 SkinConfigs.json 登记 skinId / avatarId / itemId / rarity / tags。")
    add_bullet(doc, "3）在道具表登记皮肤道具（type=SKIN，不可交易）。")
    add_bullet(doc, "4）在 ShopConfigs 增加商品：选 CURRENCY 或 IAP；设 totalLimit=1、limitPeriod=LIFETIME。")
    add_bullet(doc, "5）若 IAP：渠道后台创建 skuId，与配置一致。")
    add_bullet(doc, "6）配置导入 → 沙盒验购买/穿戴/重复购买 → 正式热更启用 enabled=true。")
    add_bullet(doc, "7）限时皮肤到期：saleEndAt 到期自动下架展示；已购玩家仍可穿戴。")

    # ---------- 9 ----------
    add_heading(doc, "9. 分期实施计划", 1)
    add_table(
        doc,
        ["阶段", "内容", "产出"],
        [
            [
                "P0 MVP",
                "SkinConfigs + 拥有判定 + EquipSkin + 晶石商店上架 1~2 个皮肤",
                "可买可穿，战斗不同步也可先本地表现",
            ],
            [
                "P1",
                "IAP 皮肤直购/礼包、商店 owned 标记、场景/战斗外观同步",
                "完整付费闭环",
            ],
            [
                "P2",
                "活动投放、邮件发皮肤、重复购买转碎片、衣柜筛选与推荐",
                "运营自动化",
            ],
        ],
    )

    # ---------- 10 ----------
    add_heading(doc, "10. 验收标准", 1)
    add_bullet(doc, "未拥有角色不能穿戴该角色皮肤；未拥有皮肤 EquipSkin 失败。")
    add_bullet(doc, "默认皮肤创角后自动拥有且可穿戴。")
    add_bullet(doc, "游戏币与 IAP 购买成功后衣柜 owned=true；重复购买被拒绝或幂等跳过。")
    add_bullet(doc, "限时商品窗口外不可下单；已购玩家窗口外仍可穿戴。")
    add_bullet(doc, "IAP 断线重试 Confirm 不重复发放第二件皮肤。")
    add_bullet(doc, "热更关闭 enabled=false 后新玩家不可见；老玩家已有皮肤仍可用。")

    # ---------- 11 ----------
    add_heading(doc, "11. 风险与对策", 1)
    add_table(
        doc,
        ["风险", "对策"],
        [
            ["皮肤与道具 ID 不一致导致发货错皮", "SkinConfig 加载时强校验 skinId↔itemId↔avatarId 唯一映射"],
            ["IAP 与游戏币双通道重复获得", "统一走 SkinOwnershipService；已拥有则拒单/跳过道具"],
            ["客户端资源未就绪穿戴黑模", "Equip 允许但客户端降级默认皮 + 提示下载"],
            ["限时下架误删配置导致无法穿戴", "下架只改商店 enabled/时间窗，SkinConfigs 保留历史皮肤"],
            ["外观被当成数值付费", "明确皮肤零属性；审核与公告文案禁止战力暗示"],
        ],
    )

    # ---------- 12 ----------
    add_heading(doc, "12. 结论", 1)
    add_body(
        doc,
        "推荐以 SkinConfigs 定义皮肤、以现有商城 ITEM/IAP 发货获得皮肤道具、以 Avatar.equippedSkinId "
        "表达穿戴状态。该方案最大程度复用 ShopApplicationService 与 IapGrantService，"
        "可在 P0 快速上架晶石皮肤，P1 接入真钱直购与礼包，满足「角色设计皮肤并在商店上架」的产品目标。",
    )

    doc.save(OUTPUT)
    print("Wrote:", OUTPUT)


if __name__ == "__main__":
    build()
