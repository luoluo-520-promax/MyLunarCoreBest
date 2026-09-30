# -*- coding: utf-8 -*-
"""就地修订桌面 MyLunarCore 详细总结报告.bak.docx：写入 P11 体验/社交缺口补齐（CmdId 1200–1247）。"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DESKTOP = Path.home() / "Desktop"
DOC_PATH = DESKTOP / (
    "MyLunarCore"
    + "\u9879\u76ee\u5404\u65b9\u9762\u8be6\u7ec6\u603b\u7ed3\u62a5\u544a"
    + ".bak.docx"
)


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_text(paragraph, text: str, *, bold=False, size=Pt(11), color=None):
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def insert_paragraph_after(paragraph, text: str, style: str | None = None):
    new_p = deepcopy(paragraph._p)
    paragraph._p.addnext(new_p)
    # rebuild as empty then fill
    from docx.text.paragraph import Paragraph

    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    set_paragraph_text(new_para, text)
    return new_para


def find_para(doc: Document, predicate):
    for i, p in enumerate(doc.paragraphs):
        if predicate(p):
            return i, p
    return None, None


def patch_cover_and_toc(doc: Document):
    # 封面副标题（含…）
    for p in doc.paragraphs:
        t = p.text or ""
        if t.startswith("（含技术架构、业务能力"):
            if "P11" not in t and "1200" not in t:
                set_paragraph_text(
                    p,
                    t.rstrip("。")
                    + "；以及 P11 体验/社交缺口补齐：一键领奖、区域探索度、签名/自定义状态、"
                    "举报屏蔽、回流活动、战斗录像分享、每日挑战提醒、自定义表情、月卡补签、"
                    "好友公会模糊搜索、荣誉称号、组队一键邀请、活动结束提醒、保底计数与消费导出"
                    "（CmdId 1200–1247）。",
                )
            break

    for p in doc.paragraphs:
        t = p.text or ""
        if t.startswith("报告生成日期："):
            set_paragraph_text(
                p,
                "报告生成日期：2026年09月04日（在 CmdId 1130–1143 与环境生态假死感优化基础上，"
                "已同步落地 P11 体验/社交缺口补齐：qol_social_system.proto、QolSocialPacketHandlers、"
                "QolLoginHookService 登录聚合推送；配置 ClaimAllDailyRewardsConfig / ExplorationConfigs / "
                "DailyReminderConfig / ReturnPlayerActivityConfig / PlayerStatusPresets；"
                "迁移 migration_qol_social_gap_fill.sql；回归 QolSocialGapFillFlowsTest）。",
            )
            break

    for p in doc.paragraphs:
        t = p.text or ""
        if "第七章已新增 7.31" in t or "第七章已新增 7.32" in t or "第七章已新增 7.33" in t:
            if "7.34" not in t:
                set_paragraph_text(
                    p,
                    t
                    + " 第七章已新增 7.34「P11 体验/社交缺口补齐（CmdId 1200–1247）」。",
                )
            break

    for p in doc.paragraphs:
        t = p.text or ""
        if t.startswith("报告依据："):
            if "1200–1247" not in t and "P11" not in t:
                set_paragraph_text(
                    p,
                    t.rstrip("。")
                    + "；本轮再把 P11 体验/社交缺口补齐写入第七章 7.34 与第八、九章相关段落"
                    "（CmdId 1200–1247，测试 QolSocialGapFillFlowsTest）。",
                )
            break


def patch_chapter7_social_gacha(doc: Document):
    for p in doc.paragraphs:
        t = p.text or ""
        if t.startswith("大厅模块把社交能力收在一起"):
            if "1205" not in t and "签名" not in t:
                set_paragraph_text(
                    p,
                    t
                    + " P11 起：PlayerProfileService 支持签名/心情/自定义状态（Cmd 1205–1208），"
                    "FriendOnlineNotify 扩展 custom_status / status_message / equipped_title；"
                    "PlayerReportBlockService 提供举报与双向屏蔽（1209–1213），私聊屏蔽返回 retcode=5；"
                    "FriendSearchService / GuildSearchService 支持名称前缀与 UID 后缀模糊搜索及最近联系人（1228–1233）；"
                    "TitleService 将成就转化为可装备称号（1234–1237）。",
                )
            break

    for p in doc.paragraphs:
        t = p.text or ""
        if "Party 是大世界组队雏形" in t and "inviteMany" not in t:
            set_paragraph_text(
                p,
                t
                + " P11：PartyService.inviteMany 支持一键邀请好友 UID 列表与公会成员批量邀请"
                "（Cmd 1238–1242，含 5 秒冷却与 PartyInviteScNotify）。",
            )
            break

    for p in doc.paragraphs:
        t = p.text or ""
        # 抽卡小节第一段通常含保底/卡池
        if ("抽卡" in t or "Gacha" in t) and "pity" in t.lower() and "1244" not in t:
            # too risky; skip if not exact
            pass

    # 找 7.6 抽卡正文：较短定位
    idx, p = find_para(
        doc,
        lambda x: (x.style and x.style.name == "Heading 2" and "7.6" in (x.text or "")),
    )
    if p is not None and idx + 1 < len(doc.paragraphs):
        body = doc.paragraphs[idx + 1]
        bt = body.text or ""
        if "1244" not in bt:
            set_paragraph_text(
                body,
                bt
                + " P11：GetGachaGuaranteeInfo（1244–1245）下发各卡池 pity5/pity4/硬保底剩余与 UP 软保底；"
                "ExportTransactionHistory（1246–1247）导出抽卡/消费 CSV 经邮件下发。",
            )


def insert_section_734(doc: Document):
    # 若已存在则跳过
    for p in doc.paragraphs:
        if p.style and p.style.name == "Heading 2" and "7.34" in (p.text or ""):
            return False

    _, anchor = find_para(
        doc,
        lambda x: x.style
        and x.style.name == "Heading 2"
        and "7.33" in (x.text or ""),
    )
    if anchor is None:
        # 回退：插在「八、网络」之前
        _, anchor = find_para(
            doc,
            lambda x: x.style
            and x.style.name == "Heading 1"
            and (x.text or "").startswith("八、"),
        )
        if anchor is None:
            raise RuntimeError("找不到插入锚点（7.33 或 第八章）")
        # 插在第八章标题之前：找到 7.33 段落后的最后一个内容段较难，直接在第八章前插入
        # 使用第八章标题的前一个段落作为 after 点
        prev = None
        for p in doc.paragraphs:
            if p is anchor:
                break
            prev = p
        if prev is None:
            raise RuntimeError("无法定位第八章前一段")
        anchor = prev

    # 找到 7.33 小节末尾：从 7.33 heading 往后直到下一个 Heading 1/2
    start_i, start_p = find_para(
        doc,
        lambda x: x.style
        and x.style.name == "Heading 2"
        and "7.33" in (x.text or ""),
    )
    if start_i is None:
        # insert before chapter 8
        ch8_i, ch8 = find_para(
            doc,
            lambda x: x.style
            and x.style.name == "Heading 1"
            and (x.text or "").startswith("八、"),
        )
        insert_after = doc.paragraphs[ch8_i - 1]
    else:
        insert_after = start_p
        for j in range(start_i + 1, len(doc.paragraphs)):
            st = doc.paragraphs[j].style.name if doc.paragraphs[j].style else ""
            txt = doc.paragraphs[j].text or ""
            if st.startswith("Heading") and (
                txt.startswith("八、") or (st == "Heading 2" and txt[:3] in ("7.3", "8."))
            ):
                break
            insert_after = doc.paragraphs[j]

    heading = insert_paragraph_after(
        insert_after,
        "7.34 P11 体验/社交缺口补齐（CmdId 1200–1247）",
        style="Heading 2",
    )
    # Heading style font
    for r in heading.runs:
        set_run_font(r, size=Pt(13), bold=True, color=RGBColor(25, 75, 140))

    body1 = insert_paragraph_after(
        heading,
        "在 7.31–7.33（CmdId 1130–1143）之后，本轮按「日常收菜、箱庭探索、社交表达、社区健康、"
        "回流召回、精彩时刻传播、挑战次数提醒、自制表情、月卡补签、找人找公会、荣誉展示、"
        "组队效率、活动临期提醒、抽卡透明度」十四条产品缺口，一次性补齐服务端框架。"
        "协议文件 qol_social_system.proto（外层类 QolSocialSystemProto）；入口 QolSocialPacketHandlers；"
        "登录后由 QolLoginHookService 聚合推送回流/每日次数/活动临期/月卡补发。"
        "号段占用 1200–1247，成对规则仍为 CS_REQ=N、SC_RSP=N+1；推送类单独占号。",
    )

    body2 = insert_paragraph_after(
        body1,
        "（1）一键领取 ClaimAllDailyRewards（1200–1201）：ClaimAllDailyRewardsService 按 "
        "ClaimAllDailyRewardsConfig.json 开关聚合每日任务/战令免费·付费轨/邮件（可排除高价值）/"
        "成就/签到，失败项跳过并返回模块摘要。（2）区域探索度：ExplorationService 按 Plane/Floor "
        "记录宝箱/解密/观景点；GetExplorationInfo（1202–1203）下发进度%；SceneInteractHandler "
        "在拾取等交互时 recordCollect 并推送 ExplorationUpdateScNotify（1204）。"
        "（3）签名与自定义状态：PlayerProfileService + SetPlayerStatus/GetPlayerProfileDetail（1205–1208）；"
        "预设模板 PlayerStatusPresets.json；FriendOnlineNotify 扩展字段广播好友。"
        "（4）举报/屏蔽：ReportPlayer/BlockPlayer（1209–1212）与 ReportResultScNotify（1213）；"
        "后台 /api/admin/social/reports*；屏蔽后双方私聊过滤。"
        "（5）回流：ReturnCheckService 按离线天数阈值激活 RETURN_PLAYER，推送 ReturnActivityScNotify（1214）。"
        "（6）战斗录像分享：BattleReplayShareService 生成短 TTL 分享码（Redis/本地），"
        "Share/GetSharedReplay（1215–1218）与 GuildReplayScNotify（1219）。"
        "（7）每日挑战提醒：DailyReminderService + DailyReminderConfig.json，登录推送 "
        "DailyChallengeReminderScNotify（1220），偏好 SetDailyReminderPref（1221–1222）。"
        "（8）自定义表情：UploadCustomEmote（1223–1224），ID≥100000，后台审核启用/封禁。"
        "（9）月卡/补签：MonthlyCardService.grantTodayIfNeeded + MonthlyCardScNotify（1225）；"
        "MakeupSignIn / ActivityTemplateService.makeupSignIn（COMPENSATE，1226–1227）。"
        "（10）搜索：SearchPlayers/SearchGuilds/GetRecentPlayers（1228–1233）。"
        "（11）称号：成就领取联动 TitleService，GetTitleList/EquipTitle（1234–1237）。"
        "（12）组队一键邀请：InviteFriendsToParty/InviteGuildMembers（1238–1242）。"
        "（13）活动临期：ActivityReminderService 登录 24h 提醒与定时 1h 紧急推送（1243）。"
        "（14）保底与导出：GetGachaGuaranteeInfo（1244–1245）、ExportTransactionHistory（1246–1247）。",
    )

    bullet1 = insert_paragraph_after(
        body2,
        "实现入口：qol/ClaimAllDailyRewardsService、DailyReminderService、QolLoginHookService、"
        "QolSocialNettyService；exploration/ExplorationService；scene/SceneInteractHandler；"
        "profile/PlayerProfileService、TitleService；social/PlayerReportBlockService、"
        "CustomEmoteService、FriendSearchService、GuildSearchService、FriendOnlineStatusService；"
        "activity/ReturnCheckService、ActivityReminderService、ActivityTemplateService；"
        "battle/BattleReplayShareService；gacha/GachaGuaranteeQueryService、"
        "TransactionHistoryExportService；party/PartyService.inviteMany；"
        "admin/ReportAdminController；net/QolSocialPacketHandlers、CmdIds 1200–1247；"
        "proto/qol_social_system.proto；迁移 db/migration_qol_social_gap_fill.sql；"
        "配置 data/ClaimAllDailyRewardsConfig.json、ExplorationConfigs.json、"
        "DailyReminderConfig.json、ReturnPlayerActivityConfig.json、PlayerStatusPresets.json。",
        style="List Bullet",
    )
    bullet2 = insert_paragraph_after(
        bullet1,
        "测试入口：QolSocialGapFillFlowsTest（协议门禁 + 14 项业务流 + 登录钩子 + 跨模块 E2E）；"
        "ClaimAllDailyRewardsServiceTest；CmdIdUniquenessTest / CmdIdProtoCompletenessTest"
        "（已纳入 QolSocialPacketHandlers）。说明见 docs/feature-gap-fill.md「P11」节。",
        style="List Bullet",
    )
    insert_paragraph_after(
        bullet2,
        "客户端联调要点：对接 1200–1247；登录后消费 ReturnActivity / DailyChallengeReminder / "
        "ActivityEndingSoon / MonthlyCard 等 Notify；一键领奖展示各模块摘要；探索进度条消费 "
        "ExplorationUpdate；好友列表展示 custom_status 与称号；举报/屏蔽入口与结果回执；"
        "抽卡页展示保底剩余；支持分享码回放与公会频道 GuildReplay。",
        style="List Bullet",
    )
    return True


def patch_protocol_and_data(doc: Document):
    for p in doc.paragraphs:
        t = p.text or ""
        if t.startswith("src/main/proto 下按系统拆分") and "qol_social" not in t:
            set_paragraph_text(
                p,
                t
                + " P11 新增 qol_social_system.proto（体验/社交缺口，CmdId 1200–1247），"
                "由 QolSocialPacketHandlers 统一注册。",
            )
            break

    for p in doc.paragraphs:
        t = p.text or ""
        if t.startswith("很多玩法数值不以硬编码") and "ClaimAllDailyRewardsConfig" not in t:
            set_paragraph_text(
                p,
                t
                + " P11 配置：ClaimAllDailyRewardsConfig.json（一键领取模块开关）、"
                "ExplorationConfigs.json（区域收集品）、DailyReminderConfig.json、"
                "ReturnPlayerActivityConfig.json、PlayerStatusPresets.json。",
            )
            break

    for p in doc.paragraphs:
        t = p.text or ""
        if t.startswith("MyLunarCore 已经不是") and "P11" not in t and "一键领奖" not in t:
            set_paragraph_text(
                p,
                t.rstrip("。")
                + "；并已补齐 P11 日常一键领奖、探索收集、社交状态/举报屏蔽/称号、"
                "回流与活动临期提醒、录像分享、自定义表情、月卡补签、模糊搜索与组队一键邀请、"
                "保底透明与消费导出等体验与社交能力（CmdId 1200–1247）。",
            )
            break


def main():
    if not DOC_PATH.is_file():
        raise SystemExit(f"找不到文档: {DOC_PATH}")
    doc = Document(str(DOC_PATH))
    patch_cover_and_toc(doc)
    patch_chapter7_social_gacha(doc)
    inserted = insert_section_734(doc)
    patch_protocol_and_data(doc)
    doc.save(str(DOC_PATH))
    print(f"Updated in-place: {DOC_PATH}")
    print(f"Section 7.34 inserted: {inserted}")


if __name__ == "__main__":
    main()
