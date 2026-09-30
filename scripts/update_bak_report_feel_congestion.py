# -*- coding: utf-8 -*-
"""
将「战斗手感 / 服务器拥堵 / 运维防雪崩」本轮新增能力同步进桌面原 bak 总结报告，
覆盖保存原文件（不另存新文档）。

目标文件：~/Desktop/MyLunarCore项目各方面详细总结报告.bak.docx
"""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

DOC_PATH = Path.home() / "Desktop" / "MyLunarCore项目各方面详细总结报告.bak.docx"


def set_run_font(run, size=11, bold=False):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold


def replace_paragraph_text(paragraph, new_text, bold=False):
    style = paragraph.style
    for r in list(paragraph.runs):
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        set_run_font(paragraph.runs[0], bold=bold)
    else:
        run = paragraph.add_run(new_text)
        set_run_font(run, bold=bold)
    paragraph.style = style


def insert_paragraph_after(paragraph, text, style_name="List Bullet"):
    new_para = paragraph._parent.add_paragraph(text, style=style_name)
    paragraph._p.addnext(new_para._p)
    for r in new_para.runs:
        set_run_font(r)
    return new_para


def find_para(doc, exact=None, startswith=None, contains=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if exact is not None and t == exact:
            return p
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    return None


def set_cell_text(cell, text):
    if not cell.paragraphs:
        cell.text = text
        return
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r.text = ""
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0])
    else:
        run = p.add_run(text)
        set_run_font(run)
    for extra in cell.paragraphs[1:]:
        extra.clear()


def add_table_row(table, values):
    row = table.add_row()
    for i, value in enumerate(values):
        if i < len(row.cells):
            set_cell_text(row.cells[i], value)
    return row


def table_has(table, needle):
    for row in table.rows:
        for cell in row.cells:
            if needle in cell.text:
                return True
    return False


def find_table_by_header(doc, *needles):
    for table in doc.tables:
        if not table.rows:
            continue
        header = " ".join(c.text.strip() for c in table.rows[0].cells)
        if all(n in header for n in needles):
            return table
    return None


def main():
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    doc = Document(str(DOC_PATH))
    n = 0

    # ---------- 报告头 ----------
    p = find_para(doc, startswith="报告生成日期")
    if p and "1105" not in p.text:
        replace_paragraph_text(
            p,
            "报告生成日期：2026年08月29日（在 AI 助手增强 CmdId 1098–1104 基础上，已同步"
            "「战斗手感 / 服务器拥堵 / 运维防雪崩」三项深度优化："
            "服务端权威 Hit-stop（BattleFxScNotify.hit_stop_frames）、预输入缓冲与取消层级、"
            "断线回放 BattleReplayDeltaScNotify（1105）；"
            "Scene/Battle 独立时钟池 + ThrottleScNotify（1106）、钱包 WAL 本地预扣异步落盘、"
            "匹配对象池与 100ms 批次成房、KCP fastack/RTT 分段自适应；"
            "ScenePreload 预加载风暴限流、VisualAssist VLM 全异步且队列>100 繁忙降级）",
        )
        n += 1

    p = find_para(doc, startswith="报告依据：当前仓库 README")
    if p and "FeelCongestionAvalancheFlowTest" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；本轮再将战斗手感微观节奏、拥堵确定性与防雪崩能力写入 7.4/7.7/7.11/7.18/7.27、"
            "第八/九/十一章等，回归 FeelCongestionAvalancheFlowTest 等约 140 条相关用例。",
        )
        n += 1

    # ---------- 一句话定位 / GameServer ----------
    p = find_para(doc, startswith="主工程启动类是 MyLunarCoreApplication")
    if p and "独立线程池" not in p.text:
        replace_paragraph_text(
            p,
            "主工程启动类是 MyLunarCoreApplication。游戏侧由 GameServer 驱动全局/场景/战斗时钟；"
            "生产默认开启 isolatedTickPools：Scene-IO 与 Battle-CPU 使用独立 ScheduledExecutor，"
            "BattleTick 积压超过约 50ms 时下发 ThrottleScNotify（1106）将场景 NPC 刷新由约 5Hz 降至 2Hz。"
            "网络侧由 GameNettyServer（TCP）与 GameKcpServer（KCP）监听同一业务端口（默认 9000）；"
            "管理侧由嵌入式 Web 容器提供 8080 端口的后台接口。",
        )
        n += 1

    # ---------- 7.4 回合制战斗 ----------
    p = find_para(doc, startswith="场景中触发遭遇后进入战斗实例")
    if p and "hit_stop_frames" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "。"
            "打击感方面：BattleFxComposer + HitStopComposer 在 BattleFxScNotify（213）下发 "
            "hit_stop_frames / hit_stop_total_ms / expected_hit_stop_end_ms，"
            "轻击/重击/战技/终结技冻结时长由服务端权威同源下发，避免组队「各冻各的」。"
            "BattleDeterministicValidator.validateHitStopEnd 校验客户端 client_hit_stop_end_ms，"
            "防止外挂缩短顿帧变相提速。"
            "手动模式由 PlayerInputBufferService 维护长度 2~3 的预输入队列，"
            "取消层级 闪避>终结技>战技>普攻，在 ActionWindow（顿帧结束后）弹出执行，减少「点了没反应」。",
        )
        n += 1

    # 在 7.4 相关列表后补子弹（若有明确列表锚点）
    if find_para(doc, startswith="预输入缓冲：PlayerInputBufferService") is None:
        anchor = find_para(doc, startswith="场景中触发遭遇后进入战斗实例")
        if anchor is not None:
            b1 = insert_paragraph_after(
                anchor,
                "Hit-stop：HitStopComposer → BattleFxScNotify.hit_stop_frames；"
                "FightActionCsReq.client_hit_stop_end_ms 由 BattleDeterministicValidator 校验。",
                "List Bullet",
            )
            insert_paragraph_after(
                b1,
                "预输入缓冲：PlayerInputBufferService（队列≤3，取消层级闪避>终结技>战技>普攻）；"
                "断线回放：BattleSnapshotService 最近 10 条 BattleDeltaRecord → "
                "BattleReplayDeltaScNotify（CmdId 1105）。",
                "List Bullet",
            )
            n += 1

    # ---------- 7.7 钱包 ----------
    p = find_para(doc, startswith="商店购买会变更钱包并推送货币变化")
    if p and "WalletWalService" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "。"
            "高并发下默认启用 WalletWalService：扣币/加币先改 ConcurrentHashMap 内存余额（亚毫秒返回），"
            "写 WalletDeltaLog 本地 WAL；后台约每 200ms 或积压 50 条批量合并后 "
            "FOR UPDATE + data_version CAS 落 MySQL，显著减少热点行锁竞争。"
            "负资产追回等特殊路径仍走同步事务。开关 lunarcore.wallet-wal.enabled。",
        )
        n += 1

    # ---------- 7.11 匹配 ----------
    p = find_para(doc, startswith="Party 是大世界组队雏形")
    if p and "batchMatchTick" not in p.text and "100ms" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "。"
            "万人秒排场景：MatchObjectPools（Apache Commons Pool2）复用 MatchSession/Room 壳对象，"
            "降低 YGC；MatchmakingService.batchMatchTick 每约 100ms、每批最多约 500 人统一成房，"
            "入队本身不再触发全量匹配计算。",
        )
        n += 1

    p = find_para(doc, startswith="匹配：MatchNettyService")
    if p and "MatchObjectPools" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；对象池 MatchObjectPools；批次 Tick batchMatchTick（interval≈100ms，batchSize≈500）。",
        )
        n += 1

    # ---------- 7.18 快照 ----------
    p = find_para(doc, startswith="回合战不能「客户端报多少伤害服务器就信多少」")
    if p and "BattleReplayDeltaScNotify" not in p.text:
        replace_paragraph_text(
            p,
            "回合战不能「客户端报多少伤害服务器就信多少」。"
            "BattleDeterministicValidator 与 BattleDamageFormula 用服务端公式做确定性校验，"
            "并校验顿帧结束时间戳防缩短 Hit-stop；BattleAuditService 留下审计线索。"
            "BattleSnapshotService 除最终态外保留最近约 10 条 BattleDeltaRecord"
            "（行动ID、目标、伤害、暴击、击杀、结算后 HP）；"
            "断线重连 / GetBattleInfo 时推送 BattleReplayDeltaScNotify（1105），"
            "客户端按时间戳强制回放 3~5 个关键动作，保持「我没掉线」的沉浸感。",
        )
        n += 1

    # ---------- 切图预加载风暴 ----------
    p = find_para(doc, startswith="切图加载：PlayerLoadingStateService")
    if p and "风暴" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；预加载风暴限流：10 秒内请求超过 3 次则 storm_throttled=true、返回 BLACK 遮罩并拒绝额外计算"
            "（ScenePreloadService.isPreloadStorm），防止异常重连拖死 Zone 加载队列。",
        )
        n += 1

    p = find_para(doc, startswith="预加载协议：ScenePreloadCsReq")
    if p and "storm_throttled" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；风暴限流时 ScenePreloadScRsp.retcode=5 / ScenePreloadPushScNotify.storm_throttled。",
        )
        n += 1

    # ---------- 新增 7.27（幂等） ----------
    if find_para(doc, contains="7.27 战斗手感与拥堵防雪崩") is None:
        anchor = find_para(doc, startswith="测试：AssistEnhancementBusinessFlowTest")
        if anchor is None:
            anchor = find_para(doc, exact="7.26 AI 助手增强（CmdId 1098–1104）")
        if anchor is None:
            anchor = find_para(doc, contains="7.26 AI 助手增强")
        if anchor is not None:
            # 找到 7.26 节后一段再插入
            base = find_para(doc, startswith="测试：AssistEnhancementBusinessFlowTest") or anchor
            h = insert_paragraph_after(
                base,
                "7.27 战斗手感与拥堵防雪崩（CmdId 1105–1106）",
                "Heading 2",
            )
            p1 = insert_paragraph_after(
                h,
                "本轮补齐三类「二游高水准」缺口。"
                "（1）战斗手感：服务端权威 Hit-stop 数组同源下发；预输入缓冲与取消层级；"
                "断线增量回放 1105。"
                "（2）拥堵确定性：Scene/Battle 时钟隔离；积压节流 1106；钱包 WAL；"
                "匹配对象池+批次；KCP 按 RTT<50 / >200 切换 fastack 与 nodelay 画像（默认 adaptive）。"
                "（3）运维防雪崩：预加载风暴限流；VLM 解码/推理全进 assistInferenceExecutor，"
                "队列长度>100 直接「服务器繁忙，稍后重试」，严禁堵 Netty Worker。",
                "Normal",
            )
            insert_paragraph_after(
                p1,
                "实现入口：HitStopComposer、PlayerInputBufferService、BattleSnapshotService、"
                "TickSchedulerConfiguration、BattleSceneThrottleService、WalletWalService、"
                "MatchObjectPools、AdaptiveKcpRetransmitAlgo、ScenePreloadService.isPreloadStorm、"
                "VisualAssistService；协议 battle_system.proto / scene_system.proto；"
                "CmdIds 1105–1106；测试 FeelCongestionAvalancheFlowTest（约 14 条）+ 相关回归约 140 条。",
                "List Bullet",
            )
            n += 1

    # ---------- 第八章 KCP ----------
    p = find_para(doc, startswith="KCP 拥塞与重传：KcpRetransmitAlgo")
    if p and "fastack" not in p.text.lower() and "fastAck" not in p.text and "RTT" not in p.text[20:]:
        replace_paragraph_text(
            p,
            "KCP 拥塞与重传：默认 AdaptiveKcpRetransmitAlgo（algo=adaptive）+ KcpRttMonitor；"
            "启用 fastack：RTT<50ms 强网降重传倍数提吞吐，RTT>200ms 弱网强制 nodelay 短间隔，"
            "牺牲少量可靠性换战斗指令实时性；见 KcpCongestionProfile。",
        )
        n += 1
    elif p and "KcpCongestionProfile" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；默认 adaptive + fastack，KcpCongestionProfile 按 RTT 分强网/常态/弱网。",
        )
        n += 1

    # ---------- 第九章 钱包 ----------
    p = find_para(doc, startswith="账号、角色、背包、钱包等需要长期保存")
    if p and "WalletWalService" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "。"
            "钱包热点写路径可叠加 WalletWalService：先内存预扣再异步批量 CAS 落库，"
            "WAL 文件落在 data/wal/；与 wallet_ledger / data_version 乐观锁并存。",
        )
        n += 1

    # ---------- 10.1 / VLM 异步 ----------
    p = find_para(doc, startswith="屏幕级视觉理解（VLM）")
    if p and "队列>100" not in p.text and "队列＞100" not in p.text and "繁忙" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；analyze/解码已下沉 assistInferenceExecutor，禁止堵 Netty Worker；"
            "队列长度>100 时 UploadScreenshotScRsp.retcode=5「服务器繁忙，稍后重试」。",
        )
        n += 1

    # ---------- 观测 / Zone ----------
    p = find_para(doc, startswith="Zone 默认人数上限约 300")
    if p and "ThrottleScNotify" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；Battle 时钟积压可触发 ThrottleScNotify（1106）降场景刷新率保出手流畅。",
        )
        n += 1

    # ---------- 成熟度已完成清单 ----------
    p = find_para(doc, startswith="移动反作弊、抽卡扣费流水")
    if p and "Hit-stop" not in p.text and "hit_stop" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；权威 Hit-stop + 预输入缓冲 + 战斗增量回放（1105）；"
            "Scene/Battle 隔离池与节流（1106）；钱包 WAL；匹配对象池/批次；"
            "KCP adaptive/fastack；预加载风暴限流；VLM 异步繁忙降级。",
        )
        n += 1

    p = find_para(doc, startswith="Admin 热更/导入/审计")
    if p and "1105" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；手感/拥堵/防雪崩：BattleReplayDelta（1105）、ThrottleScNotify（1106）、"
            "WalletWal、MatchObjectPools、HitStop/InputBuffer。",
        )
        n += 1

    # ---------- 已知限制：钱包多节点 ----------
    p = find_para(doc, startswith="钱包多节点依赖 Redis 分布式锁")
    if p and "WAL" not in p.text:
        replace_paragraph_text(
            p,
            "钱包多节点依赖 Redis 分布式锁；单节点洪峰已可用 WalletWalService 内存预扣+批量落盘减压。"
            "跨节点仍须锁+最终一致策略，仅 DB 事务在跨节点并发下仍有双扣风险窗口。",
        )
        n += 1

    # ---------- 中期路线 ----------
    p = find_para(doc, startswith="箱庭回合制战斗虽是 CPU 密集")
    if p and "独立线程池" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "。"
            "单体 JVM 内已先做 Scene/Battle 时钟隔离与战斗积压时场景节流，"
            "避免 Zone AOI 挤占 BattleTick；匹配仍优先作为可拆洪峰能力。",
        )
        n += 1

    # ---------- 配置摘录表（若有） ----------
    table = find_table_by_header(doc, "配置项") or find_table_by_header(doc, "配置键")
    if table is not None:
        if not table_has(table, "wallet-wal"):
            add_table_row(
                table,
                [
                    "lunarcore.wallet-wal.enabled",
                    "true",
                    "钱包本地预扣 + WAL 异步落盘",
                ],
            )
            n += 1
        if not table_has(table, "isolated-tick"):
            add_table_row(
                table,
                [
                    "lunarcore.game-loop.isolated-tick-pools",
                    "true",
                    "Scene/Battle 独立调度线程池",
                ],
            )
            n += 1
        if not table_has(table, "retransmit.algo") and not table_has(table, "kcp.retransmit"):
            add_table_row(
                table,
                [
                    "lunarcore.kcp.retransmit.algo",
                    "adaptive",
                    "KCP 自适应重传 + fastack（强/弱网画像）",
                ],
            )
            n += 1

    # ---------- 结语 ----------
    p = find_para(doc, startswith="MyLunarCore 已经不是「只有空壳的演示仓库」")
    if p and "Hit-stop" not in p.text and "1105" not in p.text:
        # 在战斗相关描述处追加一句
        text = p.text
        if "回合战（含 Auto" in text and "Hit-stop" not in text:
            text = text.replace(
                "回合战（含 Auto 策略权重与 target_focus、抢占式手动大招、倍速出手、打击感 FX/弱网降档与设备触觉波形）",
                "回合战（含 Auto 策略权重与 target_focus、抢占式手动大招、倍速出手、"
                "权威 Hit-stop/预输入缓冲、打击感 FX/弱网降档与设备触觉波形、断线增量回放 1105）",
            )
        if "1105" not in text:
            text = text.rstrip("。") + (
                "。拥堵侧已具备 Scene/Battle 隔离池与节流 1106、钱包 WAL、匹配对象池/批次、"
                "KCP adaptive/fastack；运维侧预加载风暴限流与 VLM 异步繁忙降级已落地。"
            )
        replace_paragraph_text(p, text)
        n += 1

    p = find_para(doc, startswith="本报告基于仓库当前代码与文档快照整理")
    if p and "1105" not in p.text:
        replace_paragraph_text(
            p,
            "本报告基于仓库当前代码与文档快照整理，2026年08月29日 已在本文档内同步"
            "战斗手感/拥堵防雪崩能力（CmdId 1105–1106；此前 AI 增强为 1098–1104、"
            "体验优化三期 1078–1097、二期 1070–1076、沉浸 1046–1069、闭环 1020–1045、"
            "PROTOCOL_WIRE_VERSION=3），可用于汇报、培训与决策讨论。"
            "若后续架构或配置有重大变更，建议修订对应章节后刷新目录域，以便目录页码与正文一致。",
        )
        n += 1

    # ---------- 玩家一天串场补一句 ----------
    p = find_para(doc, startswith="为了让非技术同学建立整体画面")
    if p and "顿帧" not in p.text and "Hit-stop" not in p.text:
        text = p.text
        if "战斗里可一键执行 AI 建议" in text:
            text = text.replace(
                "战斗里可一键执行 AI 建议的 Auto 策略",
                "战斗里可一键执行 AI 建议的 Auto 策略，普攻后摇中狂点战技由服务端预输入缓冲消化，"
                "多人共战 Hit-stop 时长由服务器同源下发",
            )
        elif "触发遭遇进入战斗" in text or "进入战斗" in text:
            text = text.rstrip("。") + (
                "；战斗中服务端下发权威顿帧与预输入缓冲，掉线重连可快速回放最近行动增量"
            )
        replace_paragraph_text(p, text)
        n += 1

    doc.save(str(DOC_PATH))
    print(f"已覆盖更新原文档: {DOC_PATH}")
    print(f"本轮变更点数: {n}")


if __name__ == "__main__":
    main()
