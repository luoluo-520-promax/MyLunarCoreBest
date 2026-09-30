# -*- coding: utf-8 -*-
"""将 AI 助手七项增强能力同步进桌面原 bak 总结报告，覆盖保存原文件（不另存新文档）。"""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

DOC_PATH = Path.home() / "Desktop" / "MyLunarCore项目各方面详细总结报告.bak.docx"


def set_run_font(run, size=11, bold=False):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold


def replace_paragraph_text(paragraph, new_text, bold=False):
    style = paragraph.style
    for r in list(paragraph.runs):
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        set_run_font(paragraph.runs[0], bold=bold)
    else:
        run = paragraph.add_run(new_text)
        set_run_font(run, bold=bold)
    paragraph.style = style


def insert_paragraph_after(paragraph, text, style_name="List Bullet"):
    new_para = paragraph._parent.add_paragraph(text, style=style_name)
    paragraph._p.addnext(new_para._p)
    for r in new_para.runs:
        set_run_font(r)
    return new_para


def find_para(doc, exact=None, startswith=None, contains=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if exact is not None and t == exact:
            return p
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    return None


def set_cell_text(cell, text):
    if not cell.paragraphs:
        cell.text = text
        return
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r.text = ""
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0])
    else:
        run = p.add_run(text)
        set_run_font(run)
    for extra in cell.paragraphs[1:]:
        extra.clear()


def add_table_row(table, values):
    row = table.add_row()
    for i, value in enumerate(values):
        if i < len(row.cells):
            set_cell_text(row.cells[i], value)
    return row


def table_has(table, needle):
    for row in table.rows:
        for cell in row.cells:
            if needle in cell.text:
                return True
    return False


def find_table_by_header(doc, *needles):
    for table in doc.tables:
        if not table.rows:
            continue
        header = " ".join(c.text.strip() for c in table.rows[0].cells)
        if all(n in header for n in needles):
            return table
    return None


def patch_table_cell_contains(table, contains, new_text=None, col=0, append=None):
    for row in table.rows:
        if contains in row.cells[col].text:
            if new_text is not None:
                set_cell_text(row.cells[col], new_text)
            if append:
                target = row.cells[-1] if append[0] is None else row.cells[append[0]]
                cur = target.text.strip()
                extra = append[1]
                if extra not in cur:
                    set_cell_text(target, cur.rstrip("。") + extra)
            return True
    return False


def main():
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    doc = Document(str(DOC_PATH))
    n = 0

    # ---------- 报告头：协议号段扩展 ----------
    p = find_para(doc, startswith="报告生成日期")
    if p and "1098" not in p.text:
        replace_paragraph_text(
            p,
            "报告生成日期：2026年08月29日（在体验优化三期 CmdId 1078–1097 基础上，已同步 AI 助手增强七项："
            "屏幕 VLM 截图解惑 UploadScreenshot/AssistVisualHint（1098–1100）、"
            "战斗建议一键落地 ApplyAiSuggestion（1101–1102）、"
            "角色人设 AssistPersona + TTS 音频旁路、主动教练事件驱动（BOSS 门口徘徊/打扰系数）、"
            "语音输入 input_type=VOICE 配额减免、UI 深链 AssistDeepLink（1103）、"
            "公告灌库 RAG ingest + 新鲜语料加权；协议号段扩展至 1098–1104）",
        )
        n += 1

    # ---------- 玩家一天 ----------
    p = find_para(doc, startswith="为了让非技术同学建立整体画面")
    if p and "屏幕解惑" not in p.text:
        text = p.text
        text = text.replace(
            "迷路时问 AI 助手可直接在地图画荧光引路线或自动寻路到材料门前（也可站内 WebView 打开攻略）",
            "迷路时问 AI 助手可直接在地图画荧光引路线/自动寻路到材料门前，或「去合成台/强化光锥」后 DeepLink 自动打开二级界面；"
            "战斗里可一键执行 AI 建议的 Auto 策略；长按助手可上传屏幕截图做机关/敌人解惑（AR 高亮）；"
            "手机可语音提问（配额减免）；回复可带角色人设语气与 TTS 声线；也可站内 WebView 打开攻略",
        )
        replace_paragraph_text(p, text)
        n += 1

    # ---------- 7.25 正文补一句 AI 增强入口 ----------
    p = find_para(doc, startswith="在体验优化二期（1070–1076）之后")
    if p and "1098" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "。后续 AI 助手增强另占 CmdId 1098–1104（截图 VLM、一键采纳 Auto、UI 深链、语音埋点），详见第十章 10.6。",
        )
        n += 1

    # ---------- 新增 7.26（幂等） ----------
    if find_para(doc, contains="7.26 AI 助手增强") is None:
        anchor = find_para(doc, startswith="移动：EnvInteractDetector 微交互池推送 1078")
        if anchor is None:
            anchor = find_para(doc, exact="7.25 体验优化三期（CmdId 1078–1097）")
        if anchor is not None:
            # 找到 7.25 小节末尾附近：触觉那一段
            haptic = find_para(doc, startswith="移动：EnvInteractDetector")
            base = haptic or anchor
            h = insert_paragraph_after(base, "7.26 AI 助手增强（CmdId 1098–1104）", "Heading 2")
            p1 = insert_paragraph_after(
                h,
                "针对「看不见屏幕、建议落不了地、口吻工业、推送死板、打字成本高、导航停在门口、攻略滞后」七个短板，"
                "服务端补齐协议与业务编排（客户端需对接对应 Cmd）。"
                "UploadScreenshot（1098–1099）上传 ≤50KB JPEG，VisualAssistService 结合 Plane/Floor/POI 做启发式/可配远程 VLM，"
                "推送 AssistVisualHintScNotify（1100）半透明箭头/高亮圈。"
                "AskBattleHint / AskAiAssist 带回 suggested_auto_override；ApplyAiSuggestion（1101–1102）直接更新 BattleAutoService。"
                "AssistPersonaConfig + AssistTtsService 经 AiHintNotify.audio_chunk 下发角色声线。"
                "AssistProactiveCoachService 监听连败与场景徘徊，interruptScore 低打扰时才弹。"
                "AskAiAssist.input_type=VOICE 减免 daily 配额并回 AssistVoiceUsageScNotify（1104）。"
                "IntentExecuteHandler.UI_INTERACTION + AssistDeepLinkScNotify（1103）实现到达后开合成/强化/抽卡等二级界面。"
                "Admin POST /api/admin/assist/rag/ingest 公告灌库；RAG 新鲜窗口内语料默认 +30% 权重。",
                "Normal",
            )
            insert_paragraph_after(
                p1,
                "测试：AssistEnhancementBusinessFlowTest（12 条业务流程）+ AssistEnhancementFeaturesTest；"
                "相关回归约 85 条（含 Assist/Battle/CmdId/HotReload/外链攻略）。",
                "List Bullet",
            )
            n += 1

    # ---------- 第十章引言 ----------
    p = find_para(doc, startswith="AI 模块的定位是「游戏内向导」")
    if p and "VLM" not in p.text:
        replace_paragraph_text(
            p,
            "AI 模块的定位是「游戏内向导」，不是替代游戏逻辑的裁判。默认启用规则教练（按配置提示），"
            "本地大模型与远程 sidecar 默认关闭。处理链路大致是：安全分类 → 意图执行（导航/UI 深链）→ "
            "本地规则/FAQ/缓存快路径 → 拼装玩家上下文与多轮历史 → 可选远程/本地 LLM → "
            "人设包装与可选 TTS → 出站清洗、合规声明、缓存、审计、配额与埋点；"
            "攻略意图再附加主流平台外链（B站、抖音、官方网站）；屏幕解惑走 UploadScreenshot→VLM/启发式→VisualHint。"
            "近期已按「对话体验 / 业务联动 / 架构性能 / 多语言运维 / 合规」补强，并落地屏幕视觉、一键 Auto、"
            "角色人设 TTS、事件驱动主动教练、语音输入、UI 深链与公告灌库，下列小节说明能力边界与实现入口。",
        )
        n += 1

    # ---------- 10.1 能力列表 ----------
    p = find_para(doc, startswith="知识检索（RAG）回答设定/玩法问题")
    if p and "公告灌库" not in p.text:
        replace_paragraph_text(
            p,
            "知识检索（RAG）回答设定/玩法问题；配置热更后优先增量刷新语料；"
            "Admin 公告发布可 POST /api/admin/assist/rag/ingest 切块灌库，新鲜语料默认上调约 30% 召回权重",
        )
        n += 1

    p = find_para(doc, startswith="主动教练推送：连败、主线卡关、体力满溢")
    if p and "徘徊" not in p.text:
        replace_paragraph_text(
            p,
            "主动教练推送：连败、主线卡关、体力满溢；场景移动钩子检测 BOSS/目标门口徘徊（默认 ≥15s）并按打扰系数 "
            "interruptScore 决定是否弹出（战斗中/刚移动完降低打扰）",
        )
        n += 1

    p = find_para(doc, startswith="主流平台外链/视频攻略")
    if p and "UI_INTERACTION" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；IntentExecuteHandler 另支持 UI_INTERACTION（合成/强化光锥/抽卡/领取），"
            "先 SceneNavigate 再 AssistDeepLinkScNotify（1103）延迟打开二级界面。",
        )
        n += 1

    # 在外链配置子弹后插入新能力子弹（若尚未存在）
    if find_para(doc, startswith="屏幕级视觉理解（VLM）") is None:
        anchor = find_para(doc, startswith="配置与热更：开关 lunarcore.ai-assist.external-guide-enabled")
        if anchor is None:
            anchor = find_para(doc, startswith="主流平台外链/视频攻略")
        if anchor is not None:
            b1 = insert_paragraph_after(
                anchor,
                "屏幕级视觉理解（VLM）：UploadScreenshotCsReq（1098）上传低分辨率 JPEG（≤50KB）；"
                "VisualAssistService 结合坐标/POI/敌人弱点启发式（可配 vlmEndpoint 接 Qwen-VL 等）；"
                "AssistVisualHintScNotify（1100）在画面叠加箭头/高亮圈，而非只丢外链。",
                "List Bullet",
            )
            b2 = insert_paragraph_after(
                b1,
                "战斗建议一键落地：AskBattleHint/AskAiAssist 下发 suggested_auto_override"
                "（target_focus / skill_priority / ult_reserve，建议 3 秒倒计时气泡）；"
                "ApplyAiSuggestionCsReq（1101）无缝更新 BattleAutoService，无需进二级 Auto 菜单。",
                "List Bullet",
            )
            b3 = insert_paragraph_after(
                b2,
                "角色人设与 TTS：data/AssistPersonaConfig.json 按 UID 绑定/已拥有高练度角色（如停云）；"
                "AssistPersonaService 包装口语；AssistTtsService 生成 16kbps OPUS 占位/可配远程 TTS，"
                "经 AiHintNotify.audio_chunk 下发。",
                "List Bullet",
            )
            b4 = insert_paragraph_after(
                b3,
                "语音输入：客户端 ASR 转写后仍走 AskAiAssist；input_type=VOICE 时减免 daily LLM 配额，"
                "并推送 AssistVoiceUsageScNotify（1104）便于统计语音使用率。",
                "List Bullet",
            )
            n += 1

    # ---------- 类名 / 协议 / 素材 ----------
    p = find_para(doc, startswith="主工程：AiAssistApplicationService")
    if p and "VisualAssistService" not in p.text:
        replace_paragraph_text(
            p,
            "主工程：AiAssistApplicationService、CoachRuleEngine、AssistSafetyFilter、AiAssistClient（含熔断）；"
            "外链攻略：ExternalGuideSearchService、ExternalGuideCatalogRepository、AssistMediaLink；"
            "增强：VisualAssistService、AssistPersonaService/Repository、AssistTtsService、"
            "AiAutoSuggestionFactory、IntentExecuteHandler（含 UI_INTERACTION）、"
            "AssistProactiveCoachService（SceneStay/interruptScore）、RagKnowledgeService.ingestAnnouncement",
        )
        n += 1

    p = find_para(doc, startswith="多轮/画像/主动推送")
    if p and "SceneStay" not in p.text:
        replace_paragraph_text(
            p,
            "多轮/画像/主动推送：AssistConversationHistoryService、AssistPlayerProfileService、"
            "AssistProactiveCoachService（连败 + BOSS 门口徘徊 + 打扰系数）",
        )
        n += 1

    p = find_para(doc, startswith="协议扩展：AskAiAssist 支持 session_id/locale")
    if p and "suggested_auto_override" not in p.text:
        replace_paragraph_text(
            p,
            "协议扩展：AskAiAssist 支持 session_id/locale/input_type（TEXT|VOICE）；"
            "响应与推送带 disclaimer/strategy_version/persona_id/suggested_auto_override；"
            "AskAiAssistScRsp.media_links / AiHintNotify.media_links 下发外链；"
            "AiHintNotify 另可带 audio_chunk（opus）与 hint_type；"
            "新增 UploadScreenshot（1098–1099）、AssistVisualHint（1100）、"
            "ApplyAiSuggestion（1101–1102）、AssistDeepLink（1103）、AssistVoiceUsage（1104）。",
        )
        n += 1

    p = find_para(doc, contains="内容素材来自 GuidePack")
    if p and "AssistPersonaConfig" not in p.text:
        replace_paragraph_text(
            p,
            "入站安全在关键词/正则（AssistSafetyRules）之外增加轻量语义意图分类；出站仍做敏感句清洗。"
            "送入 LLM 前对 UID/手机号/邮箱等做 PII 脱敏；审计用 uidHash。"
            "配额按教练/LLM 分桶，语音输入可减免 daily；支持 VIP 日预算与同问短时限流。"
            "内容素材来自 GuidePack、CoachTips、AssistFeatureContent、AssistPersonaConfig、"
            "ExternalGuideCatalog 与公告灌库块；热更时可增量刷新 RAG 并加载人设配置。"
            "记住：AI 是建议面，真正扣费/开战/发奖仍由游戏权威逻辑决定。",
        )
        n += 1

    # ---------- 10.6 专节（幂等） ----------
    if find_para(doc, contains="10.6 AI 助手增强七项") is None:
        anchor = find_para(doc, startswith="测试入口：ExternalGuideBusinessFlowTest")
        if anchor is None:
            anchor = find_para(doc, startswith="三平台约定：bilibili")
        if anchor is None:
            anchor = find_para(doc, exact="10.5 主流平台外链/视频攻略业务流程")
        if anchor is not None:
            h = insert_paragraph_after(anchor, "10.6 AI 助手增强七项业务流程", "Heading 2")
            p1 = insert_paragraph_after(
                h,
                "（1）屏幕解惑：玩家长按助手/点「屏幕解惑」→ UploadScreenshot → 服务端校验 JPEG → "
                "启发式/远程 VLM → AssistVisualHint 叠加 AR 提示。（2）一键副官：战斗提示或自然语言建议带 "
                "suggested_auto_override → 客户端 3 秒气泡 → ApplyAiSuggestion → BattleAutoService。"
                "（3）人设陪伴：按配置/已拥有角色选 persona → 包装答句 → 可选 TTS 音频块随 AiHintNotify。"
                "（4）时机智能：BattleEnded 连败、Move 门口徘徊 → interruptScore 门槛 → AiHintNotify。"
                "（5）语音：端侧 ASR → AskAiAssist(VOICE) → 配额折扣 + VoiceUsage 埋点。"
                "（6）执行确定：UI_INTERACTION → 引路线 + DeepLink 延迟开界面。"
                "（7）知识鲜度：策划发公告 → Admin rag/ingest → 新鲜块加权召回。",
                "Normal",
            )
            insert_paragraph_after(
                p1,
                "实现入口：assist_system.proto、CmdIds 1098–1104、AssistPacketHandlers、BattlePacketHandlers、"
                "AssistNettyService、VisualAssistService、AssistRagIngestController；"
                "配置 lunarcore.ai-assist.vlm-* / persona-* / tts-* / voice-quota-discount-enabled / "
                "proactive-boss-linger-seconds / proactive-interrupt-score-max / rag-freshness-*；"
                "测试 AssistEnhancementBusinessFlowTest。",
                "List Bullet",
            )
            n += 1

    # ---------- 热更 ----------
    p = find_para(doc, startswith="可以把热更想成")
    if p and "AssistPersonaConfig" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "。AssistPersonaConfig.json 已纳入 HotReloadCoordinator（失败不阻断全量热更）；"
            "公告类语料另可通过 Admin RAG ingest 即时灌库。",
        )
        n += 1

    p = find_para(doc, startswith="核心类：HotReloadCoordinator")
    if p and "AssistPersonaRepository" not in p.text:
        replace_paragraph_text(
            p,
            p.text + "；人设：AssistPersonaRepository；RAG 灌库：AssistRagIngestController",
        )
        n += 1

    # ---------- 成熟度 / 结语 ----------
    p = find_para(doc, startswith="Admin 热更/导入/审计")
    if p and "1098" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；AI 增强：VLM 截图（1098–1100）、一键 Auto（1101–1102）、DeepLink（1103）、"
            "语音配额（1104）、人设 TTS、事件主动教练、公告 RAG 灌库。",
        )
        n += 1

    p = find_para(doc, contains="体验优化三期：微交互/预测移动")
    if p and "1098" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；AI 助手增强（1098–1104：截图 VLM、一键 Auto、UI 深链、语音埋点）。",
        )
        n += 1

    p = find_para(doc, startswith="继续以模块化单体打磨")
    if p and "屏幕解惑" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；客户端对齐屏幕截图上传与 AR 叠加、执行 AI 策略气泡、角色 TTS 播放、麦克风 ASR、"
            "DeepLink 开二级界面。",
        )
        n += 1

    p = find_para(doc, startswith="MyLunarCore 已经不是「只有空壳的演示仓库」")
    if p and "VLM" not in p.text:
        replace_paragraph_text(
            p,
            p.text.replace(
                "以及能力增强后的 AI 助手旁路（站内 WebView 攻略摘要 + IntentExecute 地图引路线/自动寻路）。",
                "以及能力增强后的 AI 助手旁路（站内 WebView 攻略摘要、IntentExecute 引路线/UI 深链、"
                "屏幕 VLM 解惑、一键 Auto、角色人设 TTS、语音输入与公告灌库 RAG）。",
            ),
        )
        n += 1

    # ---------- 表格 ----------
    t6 = find_table_by_header(doc, "文件/目录", "用途")
    if t6:
        if not table_has(t6, "AssistPersonaConfig"):
            add_table_row(
                t6,
                [
                    "AssistPersonaConfig.json",
                    "AI 角色人设：personaId/声线/语气前缀、UID 绑定与已拥有角色匹配",
                ],
            )
        patch_table_cell_contains(
            t6,
            "GuidePack",
            new_text="GuidePack / CoachTips / AssistSafetyRules / AssistFeatureContent / "
            "ExternalGuideCatalog.json / AssistPersonaConfig.json",
        )
        n += 1

    t3 = find_table_by_header(doc, "包名", "普通人理解")
    if t3:
        patch_table_cell_contains(
            t3,
            "assist",
            append=(2, "；VLM 截图、一键 Auto、人设 TTS、语音配额、UI 深链、公告 RAG"),
        )
        n += 1

    t7 = find_table_by_header(doc, "能力面", "代表类")
    if t7 and not table_has(t7, "VisualAssistService"):
        add_table_row(
            t7,
            [
                "AI 助手增强",
                "VisualAssistService / AssistPersonaService / AssistTtsService / "
                "IntentExecuteHandler / AssistProactiveCoachService / RagKnowledgeService",
                "截图 VLM、人设 TTS、UI 深链、事件主动推送、公告灌库与新鲜度加权；"
                "协议 1098–1104",
            ],
        )
        n += 1

    t15 = find_table_by_header(doc, "配置", "含义")
    if t15:
        rows = [
            ("lunarcore.ai-assist.vlm-enabled", "是否启用屏幕截图 VLM/启发式解惑（默认 true）"),
            ("lunarcore.ai-assist.persona-enabled", "是否启用角色人设包装回答（默认 true）"),
            ("lunarcore.ai-assist.tts-enabled", "是否启用 TTS 音频旁路（默认 true）"),
            ("lunarcore.ai-assist.voice-quota-discount-enabled", "语音输入是否减免 daily LLM 配额（默认 true）"),
            ("lunarcore.ai-assist.proactive-boss-linger-seconds", "BOSS 门口徘徊触发主动辅导秒数（默认 15）"),
            ("lunarcore.ai-assist.rag-freshness-boost", "新鲜语料召回上调比例（默认 0.30= +30%）"),
        ]
        for key, meaning in rows:
            if not table_has(t15, key.split(".")[-1] if key.count(".") > 3 else key):
                # 用更稳的整键匹配
                if not table_has(t15, key):
                    add_table_row(t15, [key, meaning])
        n += 1

    t17 = find_table_by_header(doc, "术语", "人话解释")
    if t17:
        terms = [
            ("AssistVisualHint", "把箭头/高亮圈叠在当前游戏画面上，告诉你点哪里，而不是只甩一篇攻略链接"),
            ("suggested_auto_override", "AI 给出的自动战斗参数（集火谁、优先战技还是攒大招），点一下就能生效"),
            ("AssistDeepLink", "导航到地点后，服务端再推一条「打开合成/强化/抽卡界面」的指令"),
            ("interruptScore", "打扰系数：正在激战或刚猛冲时数值高，主动弹窗会被压住，避免挡视线"),
            ("RAG 新鲜度加权", "新版本公告灌进知识库后，回答「打新 BOSS」时更优先用新机制，而不是旧攻略"),
        ]
        for term, meaning in terms:
            if not table_has(t17, term):
                add_table_row(t17, [term, meaning])
        n += 1

    t18 = find_table_by_header(doc, "别人常问", "你可以这样答")
    if t18 and not table_has(t18, "屏幕解惑"):
        add_table_row(
            t18,
            [
                "AI 能不能看玩家屏幕？战斗建议要不要自己去开 Auto？",
                "可以上传截图做屏幕解惑（VLM/启发式+AR 叠加）；战斗建议可一键 Apply 到 Auto。"
                "另外支持角色人设语音、语音提问、门口主动配队提醒，以及导航后自动打开二级界面。",
            ],
        )
        n += 1

    doc.save(str(DOC_PATH))
    print(f"已更新原文档: {DOC_PATH}（命中补丁约 {n} 处）")


if __name__ == "__main__":
    main()
