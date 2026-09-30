# -*- coding: utf-8 -*-
"""Generate MyLunarCore project issues & solutions Word document (current codebase)."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DESKTOP = Path.home() / "Desktop"
# ASCII 文件名，避免 Windows 控制台/脚本编码导致桌面文件名乱码；文档标题仍为中文
OUTPUT = DESKTOP / "MyLunarCore_Project_Issues_And_Solutions.docx"


def set_doc_fonts(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15


def add_title(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(25, 75, 140)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_subtitle(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(100, 100, 100)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_heading(doc: Document, text: str, level: int = 1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        if level == 1:
            run.font.color.rgb = RGBColor(25, 75, 140)


def add_body(doc: Document, text: str):
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_bullet(doc: Document, text: str, level: int = 0):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.5 * (level + 1))
    for run in p.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_code(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 30, 30)


def add_table(doc: Document, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.name = "微软雅黑"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                run.font.size = Pt(10)
    for r_idx, row in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = val
            for p in row_cells[c_idx].paragraphs:
                for run in p.runs:
                    run.font.name = "微软雅黑"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                    run.font.size = Pt(9)
    doc.add_paragraph()


def add_issue(doc: Document, title: str, severity: str, refs: str, problem: str, solution: str):
    add_heading(doc, f"{title}（严重度：{severity}）", 3)
    add_body(doc, f"涉及位置：{refs}")
    add_heading(doc, "问题与风险", 4)
    add_body(doc, problem)
    add_heading(doc, "解决方案", 4)
    for line in solution.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line[0].isdigit() and "）" in line[:4]:
            add_bullet(doc, line)
        else:
            add_body(doc, line)


def build_document() -> Document:
    doc = Document()
    set_doc_fonts(doc)
    today = date.today().strftime("%Y年%m月%d日")

    add_title(doc, "MyLunarCore 项目问题分析与解决方案")
    add_subtitle(doc, "项目路径：c:\\Users\\ASUS\\IdeaProjects\\test\\MyLunarCore")
    add_subtitle(doc, f"评估基准：当前工作区静态代码审查 · 生成日期：{today}")
    doc.add_paragraph()

    # ========== 一、总体结论 ==========
    add_heading(doc, "一、总体结论", 1)
    add_body(
        doc,
        "MyLunarCore 是基于 Spring Boot 4 + Netty/KCP + Protobuf + MySQL 的游戏私服工程。"
        "主工程为模块化单体（游戏核心）；microservices/ 为 Spring Cloud Alibaba 迁移脚手架"
        "（api-gateway、auth-service、player-service、ai-assist-service），尚未替换主业务。",
    )
    add_body(
        doc,
        "相对早期骨架阶段，多项基础设施已明显改善：管理端 Session+CSRF+RBAC+登录锁定；"
        "/internal/** 由 InternalApiAuthFilter 校验；ProductionSecretsValidator 在 prod 拦截弱密钥；"
        "钱包乐观更新与部分经济路径事务化；抽卡历史已落库可查；排行榜支持 Redis ZSET；"
        "跨节点迁移已具备票据重定向雏形；主工程与微服务均对齐 Spring Boot 4.0.4。",
    )
    add_body(
        doc,
        "当前主要矛盾变为：资金与身份面仍有硬缺口（IAP 无真实验签、微服务弱 JWT/明文密码）、"
        "任务触发仍为占位、战斗/场景并发与 Netty 阻塞风险、热更后半段 static/RAG 回滚不完整、"
        "配置路径逃逸与 Schema 校验不足、AI 契约双份维护、反作弊与审计偏「演示级」。"
        "整体判断：适合继续开发与内网联调，不宜在未完成 P0 整改前按生产/商业化标准对外。",
    )

    add_heading(doc, "规模与形态速览", 2)
    add_table(
        doc,
        ["维度", "现状"],
        [
            ["架构形态", "模块化单体为主；microservices/ 为脚手架，未替换主业务"],
            ["网络入口", "TCP/KCP 游戏口 + Admin HTTP:8080；网关 HTTPS 演示口"],
            ["持久化", "MySQL + JDBC；在线态会话缓存 + 周期/异步落盘"],
            ["经济写路径", "钱包事务+乐观锁；抽卡已记历史；IAP 仍为 Mock 验签"],
            ["配置热更", "dry-run + 原子写 + staged reload；static/RAG 回滚缺口"],
            ["中心服", "local/remote；跨节点票据重定向已有，会话迁移未闭环"],
            ["AI 助手", "单体完整能力 + ai-assist-service 旁路；契约/Prompt 双份"],
            ["Boot 版本", "主工程与微服务均 Spring Boot 4.0.4（已对齐）"],
        ],
    )

    add_heading(doc, "风险分布总表", 2)
    add_table(
        doc,
        ["类别", "严重度", "一句话结论"],
        [
            ["IAP 验签", "严重", "Mock 可刷单；关闭 Mock 后无真实渠道，充值全失败"],
            ["微服务默认密钥", "严重", "auth/gateway JWT 与明文演示密码写死在 yml"],
            ["战斗/场景并发", "严重", "可变 HashMap 外泄，锁靠约定，易竞态"],
            ["任务触发引擎", "高", "applyTrigger 空实现，杀怪/NPC 无法推进任务"],
            ["Netty 阻塞", "高", "包处理同步打 JDBC/HTTP/LLM，可能卡住事件循环"],
            ["热更回滚", "高", "失败时 static/RAG 无快照恢复，配置可能撕裂"],
            ["配置路径逃逸", "高", "normalize 未校验仍落在 dataRoot 内"],
            ["Assist 契约双份", "高", "Prompt/DTO 人工同步，易静默不兼容"],
            ["内部 Token/TLS", "高", "dev-internal-token / changeit 默认回落"],
            ["网关鉴权面", "高", "actuator 整段排除；身份头可能被伪造"],
            ["Redis 双轨", "高", "默认关闭时多节点状态与单机内存行为不一致"],
            ["反作弊/审计", "中高", "移速检查偏浅；战斗审计仅写日志"],
            ["经济幂等", "中高", "重试缺稳定幂等键；unsafe 全量覆写钱包 API 仍在"],
            ["测试缺口", "中", "内部鉴权/IAP/停机/并发关键路径缺专项回归"],
            ["文档", "中", "根目录无 README，主工程与脚手架关系易误解"],
        ],
    )

    add_heading(doc, "已改善项（避免重复整改）", 2)
    add_bullet(doc, "AdminSecurityConfig：CSRF + Session + /api/admin/** 鉴权；anyRequest().denyAll()")
    add_bullet(doc, "InternalApiAuthFilter：/internal/** 强制 X-Internal-Token；空 Token 拒绝")
    add_bullet(doc, "ProductionSecretsValidator：prod 禁止弱默认密钥（主工程）")
    add_bullet(doc, "WalletApplicationService：FOR UPDATE + 乐观版本 + 流水；抽卡历史 insertBatch/分页查询已落地")
    add_bullet(doc, "LeaderboardService：可选 Redis ZSET，并持久化昵称快照")
    add_bullet(doc, "PlayerMigrationService：跨节点可签发 sessionTicket 并返回重定向")
    add_bullet(doc, "AiAssistController：internal token 为空时 UNAUTHORIZED，不再静默放行")
    add_bullet(doc, "主工程与 microservices 均对齐 Spring Boot 4.0.4 + Cloud 2025.1.0")
    add_bullet(doc, "ConfigFileService 原子写 .tmp → move；GracefulShutdown 等运维能力已具备雏形")

    # ========== 二、详细问题 ==========
    add_heading(doc, "二、详细问题与解决方案", 1)

    # 2.1 资金与安全
    add_heading(doc, "2.1 资金与身份安全（最高优先级）", 2)
    add_issue(
        doc,
        "1. IAP 生产关闭 Mock 后仍无法真实验签",
        "严重",
        "economy/iap/IapVerifyGateway.java；application.properties（mock-verify=true）；"
        "application-prod.properties（mock-verify=false）",
        "Mock 开启时非空收据即通过；关闭后直接返回 real_channel_not_configured，无 App Store / "
        "Google Play 等真实渠道实现。开发环境可随意发货；生产要么仍用 Mock 被刷单，要么所有充值失败。"
        "默认配置即为 mock-verify=true，漏切 profile 时风险极高。",
        "1）按渠道拆 IapChannelVerifier 接口，分别实现 Apple/Google 验签与幂等落单。\n"
        "2）生产启动时若 mock-verify=false 且无渠道实现则 fail-fast；mock=true 禁止与 prod 同启。\n"
        "3）补单测：mock=false 时任意收据必须失败；伪造收据不可发货。\n"
        "4）发货与订单状态机保持幂等，验签成功才进入发货 commit。",
    )

    add_issue(
        doc,
        "2. auth-service / api-gateway 弱 JWT 密钥与明文演示密码",
        "严重",
        "microservices/services/auth-service/.../application.yml（password: 123456/admin123，"
        "secret: change-me-...）；api-gateway application.yml（jwt-secret 同源弱值）；"
        "AuthJwtProperties / GatewayAuthProperties 默认值同类",
        "演示凭证与 JWT 密钥入库进仓库；登录无哈希。一旦按默认配置部署，账号与 Bearer Token 可伪造，"
        "网关后业务接口形同虚设。微服务侧无对等的 ProductionSecretsValidator。",
        "1）JWT secret / 用户口令一律走环境变量或密钥管理；禁止默认弱密钥启动。\n"
        "2）用户表使用 PasswordEncoder（bcrypt/Argon2）；删除 yml 内硬编码用户或仅限 local profile。\n"
        "3）对齐主工程 ProductionSecretsValidator，在 auth/gateway 增加 prod fail-fast。\n"
        "4）README 明确：微服务为脚手架，不可直接当生产 IdP；已部署需立即轮换密钥。",
    )

    add_issue(
        doc,
        "3. 默认内部 Token / SSL 密码与开发库弱口令",
        "高",
        "application.properties（dev-internal-token、changeit）；ai-assist-service internal-token 默认；"
        "LunarCoreProperties / AiAssistServiceProperties 字段默认值",
        "非 prod 或漏配环境变量时回落弱默认。主工程 prod 有校验，微服务与旁路服务无对等约束。"
        "keystore 密码默认 changeit；单体 TLS 默认关闭，外网暴露时明文传输。",
        "1）微服务与 ai-assist-service 增加 prod 启动校验，禁止默认 token/密钥。\n"
        "2）移除 changeit / dev-internal-token 生产回落；证书落盘于制品外。\n"
        "3）CI 扫描配置中的 changeit、123456、dev-internal-token、change-me。\n"
        "4）对外可达的 HTTP/游戏口启用 TLS，管理端与游戏口分网段。",
    )

    add_issue(
        doc,
        "4. 网关鉴权排除面过宽 + 身份头未剥离",
        "高",
        "api-gateway application.yml（exclude-paths 含 /actuator/**、/ai/fallback）；"
        "GatewayAuthFilter（写入 X-User-Id/Name/Role）",
        "actuator 整段免 JWT；/ai/fallback 未认证。过滤器写入信任身份头前未见删除入站同名头，"
        "下游若信任该头，存在身份伪造风险。discovery.locator 开启也可能暴露未评审路由。",
        "1）仅公开最小健康检查；actuator/gateway 绑私有管理端口并鉴权。\n"
        "2）mutate 前显式 removeHeader 入站 X-User-*，再用网关签发值覆盖。\n"
        "3）生产关闭 discovery locator，或加服务白名单。\n"
        "4）AI 路由独立限流与认证；fallback 仅允许熔断器内部调用。",
    )

    add_issue(
        doc,
        "5. 内部 Token 非恒定时间比较 + 信任域共用",
        "中",
        "InternalApiAuthFilter.java；AiAssistController.assertInternalToken；"
        "application.properties 中 INTERNAL_API_TOKEN 与 AI_ASSIST_INTERNAL_TOKEN 互为回落",
        "String.equals 比较共享密钥存在时序侧信道（利用门槛较高）。"
        "AI 旁路与中心路由/内部管理共用同一类 Token，一处泄露即可横向访问。",
        "1）改用 MessageDigest.isEqual 比较 UTF-8 字节。\n"
        "2）按服务拆分凭证与权限范围；优先 mTLS/短时服务凭证。\n"
        "3）统一封装 InternalTokenVerifier，主工程与旁路共用。",
    )

    # 2.2 并发与运行时
    add_heading(doc, "2.2 运行时并发、网络与数据一致性", 2)
    add_issue(
        doc,
        "6. BattleContext / SceneContext 可变状态外泄，锁靠约定",
        "严重",
        "battle/BattleContext.java；scene/SceneContext.java；BattleManager",
        "实体集合为 HashMap/LinkedHashSet，通过 getter 直接返回可变引用；lock 对象交给调用方自行同步。"
        "包处理、tick、超时清理、断线可能并发访问。同一配置怪 ID 作实体键时，同波次多实例可能互相覆盖。"
        "场景 onTick 遍历 Buff 时包线程可能增删实体，存在 ConcurrentModificationException 与状态撕裂风险。",
        "1）每场战斗/每个 Zone 使用单线程命令队列（Actor）串行化变更。\n"
        "2）对外只暴露不可变快照；运行时实体 ID 与配置怪 ID 分离。\n"
        "3）结束/移除使用原子操作；移动与战斗指令带序号。\n"
        "4）补并发单测：同战斗双线程写入与 tick 交叉。",
    )

    add_issue(
        doc,
        "7. Netty 包处理路径可能同步阻塞",
        "高",
        "net/*PacketHandlers.java → *NettyService → JDBC / HTTP / Redis / LLM",
        "Handler 同步调用业务；业务可触达数据库、中心服 HTTP、文件系统、AI 远程调用。"
        "慢查询或远程超时会占用事件循环，拖垮同 loop 上其他连接的吞吐与延迟。",
        "1）Handler 仅做解码、鉴权上下文、入队；阻塞 IO 进有界业务线程池。\n"
        "2）配置连接/读写超时、队列满拒绝策略与指标（排队时长、拒绝数）。\n"
        "3）CI/压测断言：事件循环任务不得直接调用 JDBC/HTTP。",
    )

    add_issue(
        doc,
        "8. Redis 默认关闭导致多节点行为分叉",
        "高",
        "application.properties（lunarcore.redis.enabled=false）；"
        "MigrationTicketService / OnlinePresenceService / LeaderboardService / ChatService",
        "双节点文档要求共享 Redis，但默认本地内存回落。误开 remote center 却不开 Redis 时，"
        "票据、在线表、排行榜、世界聊天在节点间不一致，表现为「偶现丢榜/票据无效」。",
        "1）部署拓扑显式配置（single-node / multi-node）；multi-node 未开 Redis 则启动失败。\n"
        "2）local fallback 仅限 dev/test profile。\n"
        "3）补 Redis 宕机与节点重启故障演练。",
    )

    add_issue(
        doc,
        "9. 经济幂等与不安全钱包覆写 API",
        "中高",
        "economy/WalletApplicationService；repo/WalletRepository.updateCurrency；"
        "gacha/GachaApplicationService",
        "抽卡已有 @Transactional 扣费+写历史，但调用方若无稳定业务幂等键，超时重试可能重复扣费。"
        "canAfford 无锁读不能当授权边界。WalletRepository.updateCurrency 仍可全量覆写 currency JSON，"
        "被误用会丢失并发更新。",
        "1）所有扣费/发货命令带客户端或网关稳定 requestId，DB 唯一约束防重。\n"
        "2）废弃或严格限制 updateCurrency 全量覆写；仅保留乐观/条件更新。\n"
        "3）崩溃注入测试：扣费后、发货前杀进程，验证可恢复且不重复。",
    )

    # 2.3 核心玩法占位
    add_heading(doc, "2.3 核心玩法占位与防作弊", 2)
    add_issue(
        doc,
        "10. 任务触发引擎未实现",
        "高",
        "repo/QuestProgressRepository.applyTrigger；quest/QuestTriggerEngine.java",
        "applyTrigger 方法体为空，注释写明占位。杀怪/NPC/场景事件经 QuestTriggerEngine 透传后"
        "无法更新 objectives_json，任务进度链路断裂。现有单测仅验证「透传到仓储」，掩盖空实现。",
        "1）与 QuestConfig 目标类型联动：匹配 triggerType/targetId，递增计数并判定完成。\n"
        "2）写库后若玩家在线，经聚合写或明确 sync 推送客户端。\n"
        "3）补集成测试：杀怪 N 次后 objectives 与 status 正确变化；禁止再只测「被调用」。",
    )

    add_issue(
        doc,
        "11. 场景 Tick / Buff 过期为占位",
        "中",
        "scene/SceneContext.onTick",
        "主循环遍历怪物 Buff 但不处理过期与 AI tick，场景内时效 Buff 不会随时间推进。",
        "1）实现 Buff TTL 递减与过期移除，必要时广播状态变化。\n"
        "2）将怪物 AI tick 钩子与 onTick 对齐，补单元测试验证过期行为。",
    )

    add_issue(
        doc,
        "12. 反作弊偏浅，战斗审计仅日志",
        "中高",
        "anticheat/MoveSpeedGuard.java；anticheat/BattleAuditService.java",
        "移速守卫校验有限坐标与水平速度，缺少垂直、导航网格、传送授权、包序号/重放、违规累计与升级处置。"
        "关闭检查时几乎直接接受坐标。BattleAuditService 仅写结构化日志，无落库、不可篡改链路，"
        "无法支撑事后裁决与奖励关联。",
        "1）移动改为服务端权威；校验导航/碰撞与传送 grant。\n"
        "2）包序号单调、违规计分与踢出/软封策略。\n"
        "3）关键战斗结算写入 durable audit（或 outbox），与奖励发放关联。\n"
        "4）伤害与结果以服务端计算为准，客户端仅输入意图。",
    )

    # 2.4 热更与运维
    add_heading(doc, "2.4 热更新、配置与运维", 2)
    add_issue(
        doc,
        "13. 热更回滚未覆盖 static / RAG",
        "高",
        "common/HotReloadCoordinator.reloadAllStaged",
        "成功路径会调用 staticResourceRegistry.reloadAll() 与 ragKnowledgeService.rebuild()；"
        "catch 中仅 restore 活动/抽卡/商店/助手规则等先前列出的仓库快照，无 static/RAG 回滚。"
        "后半段失败时内存配置可能部分新、部分旧。",
        "1）对 static/RAG 也做 before 快照 + restore，或失败时强制从磁盘全量重载恢复。\n"
        "2）将 static/RAG 纳入 staged 阶段顺序，失败尽早中断。\n"
        "3）发布演练：故意让 RAG rebuild 失败，断言运行时配置一致。",
    )

    add_issue(
        doc,
        "14. 配置路径可能逃逸 dataRoot + 校验仅语法层",
        "高",
        "common/ConfigFileService.resolve；ConfigImportService.validateJsonText；"
        "ActivityConfigService 单文件失败仅打日志跳过",
        "resolve 只做 normalize，未断言结果仍位于 dataRoot 之下，相对路径含 ../ 时可写到目录外。"
        "dry-run 多停留在 JSON 语法，缺少字段类型、枚举、引用完整性。单活动文件损坏被跳过会导致"
        "「部分生效」的脏快照；并发 read-modify-write 可能互相覆盖。",
        "1）写入前校验 resolved.startsWith(dataRoot)；拒绝越界路径。\n"
        "2）为 Banners/Activity/AssistSafetyRules 等增加 Schema/业务校验器。\n"
        "3）候选配置全量校验通过后再原子发布快照；导入加版本号/文件锁。\n"
        "4）publish_drill 纳入「故意坏配置/路径逃逸必须失败」。",
    )

    add_issue(
        doc,
        "15. 跨节点会话迁移未完全闭环",
        "中",
        "center/PlayerMigrationService、HttpCenterRoutingClient、LunarCoreProperties.advertiseHost",
        "已支持签发 ticket 并重定向；但缺 host/port 时仍 retcode=3。"
        "advertiseHost 默认 127.0.0.1 易登记错误地址；HTTP 拉 plan 未见强 HTTPS/主机白名单，"
        "内部 Token 可能打向错误 URL；响应 body 进入异常信息有信息泄露风险。",
        "1）强制配置 advertise 可达地址；remote-base-url 仅允许 https + 白名单。\n"
        "2）校验返回的 nodeHost/Port 是否在节点注册表内。\n"
        "3）中期补 MigrationSession：导出→导入→踢旧连接→指引重连与超时回收。\n"
        "4）异常勿把远程 body 直接回传客户端。",
    )

    add_issue(
        doc,
        "16. 根目录缺少项目 README",
        "中",
        "仓库根无 README.md；仅 microservices/README.md",
        "主工程启动、profile、Netty/KCP、热更、内部 Token、与微服务关系文档缺失，新人与运维易误用脚手架当生产。",
        "1）根 README 写清：单体为主、微服务为迁移脚手架、关键配置与发布演练入口。\n"
        "2）链到 microservices/README，并标注「不可单独当生产账号体系」。",
    )

    # 2.5 AI 与微服务
    add_heading(doc, "2.5 AI 助手与微服务演进", 2)
    add_issue(
        doc,
        "17. AI 契约 / Prompt 双份拷贝，主工程未依赖 common-api",
        "高",
        "src/.../assist/AssistPromptTemplates.java；assist/remote/RemoteAssistAskRequest.java；"
        "microservices/common-api/.../AssistAskRequest.java、AssistSchemaVersions",
        "模板与 DTO 靠人工同步；SCHEMA_VERSION 各自维护。一端改契约另一端易静默不兼容或 Prompt 漂移。"
        "旁路已对 schema 不匹配返回 fail，但本地副本仍易漂移。",
        "1）主工程 Maven 依赖 common-api，删除本地 Prompt/DTO 副本。\n"
        "2）共享安全规则与提示词包从 data/*.json 热更，两端同构加载。\n"
        "3）明确唯一权威：旁路为推理面或仅压测旁路，避免能力预期偏差。",
    )

    add_issue(
        doc,
        "18. LLM 可配置端点带来 SSRF / 数据外泄风险",
        "中",
        "application.properties（llm-endpoint）；RemoteLlmGateway；ai-assist-service yml",
        "端点可指向内网/元数据地址；玩家上下文可能发往第三方。"
        "未见 HTTPS 强制与主机白名单、敏感字段脱敏约束。",
        "1）允许列表 + 解析后拒绝私网/链路本地/元数据 IP；强制 HTTPS。\n"
        "2）脱敏账号、会话、凭证；禁止完整 Prompt/Key 入日志。\n"
        "3）出站 egress 控制与请求体大小限制。",
    )

    add_issue(
        doc,
        "19. player-service / auth-service 仍为演示桩",
        "中",
        "microservices/.../PlayerController.java；AuthJwtProperties.users",
        "player-service 仅返回 {uid, service}；auth 为内存用户表，无法对接主工程 MySQL account。"
        "网关可暴露「看起来真实」的玩家接口，易被误认为已拆分玩家域。",
        "1）README 与接口响应明确 scaffold-only / 501。\n"
        "2）真实拆分前对接主库或统一 IdP，删除演示用户。\n"
        "3）能力归属矩阵：钱包/战斗/场景继续单体，先拆 Admin/排行榜/AI。",
    )

    add_issue(
        doc,
        "20. AiAssistClient 同步阻塞与排障信号不足",
        "中",
        "assist/remote/AiAssistClient.java",
        "httpClient.send 同步调用可能占业务线程；Token 配错时易落到 Optional.empty()，"
        "若运维未盯 metrics，排障滞后。prefer-remote-only=true 时远程失败会跳过本地 LLM。",
        "1）远程调用放入专用线程池/异步，避免阻塞 Netty/业务线程。\n"
        "2）remoteEnabled 且 token 为空时启动拒绝；调用失败打结构化告警。\n"
        "3）默认 prefer-remote-only=false；区分「远程优先」与「远程独占」。",
    )

    # 2.6 测试与卫生
    add_heading(doc, "2.6 测试覆盖与代码卫生", 2)
    add_issue(
        doc,
        "21. 关键路径缺专项测试",
        "中",
        "建议补齐：InternalApiAuthFilter、HttpCenterRoutingClient、GracefulShutdownCoordinator、"
        "AiAssistClient 真实 HTTP、IapVerifyGateway（mock=false）、Battle/Scene 并发、Config 路径逃逸",
        "主工程单测数量不少，但安全、资金、并发与停机相关回归仍依赖人工，易在后续改动中回退。",
        "1）优先补：Filter 有/无/错误 Token；IAP mock off；路径逃逸；并发战斗；停机落盘失败。\n"
        "2）将上述用例纳入 CI 必跑集合；定期 OWASP dependency-check。",
    )

    add_issue(
        doc,
        "22. 管理端登录锁定仅进程内存",
        "低",
        "admin/AdminLoginAttemptService",
        "失败计数在 ConcurrentHashMap，多实例 Admin 时锁定不共享，可被分散爆破。",
        "生产多实例改为 Redis/DB；单实例可接受，但需文档标明。",
    )

    add_issue(
        doc,
        "23. Assist 历史编码修复脚本与 __pycache__ 残留",
        "低",
        "scripts/fix_assist_*.py、repair_assist_*.py、scripts/__pycache__/",
        "暗示历史乱码问题；缓存文件不应入库，增加仓库噪声。",
        "1）.gitignore 排除 __pycache__。\n"
        "2）一次性修复脚本移到 scripts/archive/ 并注明勿再执行。",
    )

    add_issue(
        doc,
        "24. MySQL 驱动版本在根 POM 钉死",
        "低",
        "根 pom.xml（mysql-connector 显式 8.0.33）",
        "未跟随 Boot BOM，与 Boot 4 管理版本可能不一致；另有 Netty 4.2 + KCP、JUnit6+TestNG 双栈需在干净 CI 验证。",
        "去掉显式 version，交给 parent BOM；用 enforcer 做依赖收敛；单独验证 KCP 与 Boot 4 兼容。",
    )

    # ========== 三、整改路线 ==========
    add_heading(doc, "三、建议整改路线（按优先级）", 1)

    add_heading(doc, "P0（1–2 周）— 资金、身份与并发止损", 2)
    add_bullet(doc, "接入真实 IAP 渠道验签；prod 无渠道或 mock=true 则 fail-fast")
    add_bullet(doc, "清理/阻断 auth-service 与 api-gateway 默认 JWT 密钥与明文密码")
    add_bullet(doc, "微服务与 ai-assist 对齐 ProductionSecretsValidator；去掉弱 Token/changeit 回落")
    add_bullet(doc, "网关剥离伪造身份头；收紧 actuator/AI 公开面")
    add_bullet(doc, "战斗/场景改为每实例串行命令队列，禁止可变 map 外泄")
    add_bullet(doc, "实现 QuestProgressRepository.applyTrigger 与 QuestConfig 联动")

    add_heading(doc, "P1（2–4 周）— 发布安全与数据正确性", 2)
    add_bullet(doc, "Netty Handler 与阻塞 IO 解耦；有界线程池 + 超时指标")
    add_bullet(doc, "HotReloadCoordinator 补齐 static/RAG 回滚或失败全量恢复")
    add_bullet(doc, "ConfigFileService 路径围栏 + JSON Schema/业务校验 + 全量快照发布")
    add_bullet(doc, "主工程依赖 common-api，删除 Assist 双份契约")
    add_bullet(doc, "经济请求幂等键 + 废弃 unsafe 钱包覆写；multi-node 强制 Redis")
    add_bullet(doc, "补 InternalApi / IAP mock=false / 路径逃逸 / 并发战斗专项测试")

    add_heading(doc, "P2（1–2 月）— 架构演进与防作弊", 2)
    add_bullet(doc, "场景 Buff TTL；服务端权威移动与 durable 战斗审计")
    add_bullet(doc, "Center MigrationSession 闭环；remote HTTPS + 节点白名单")
    add_bullet(doc, "按 README 建议先拆 Admin/排行榜/AI，不拆战斗与钱包")
    add_bullet(doc, "根 README + 端到端故障演练（热更坏配置、钱包并发、AI 旁路降级、Redis 宕机）")

    add_code(
        doc,
        "建议落地顺序（简图）\n"
        "IAP 真实验签 ──► 微服务密钥治理 ──► 任务 applyTrigger\n"
        "      │                │\n"
        "      ▼                ▼\n"
        "战斗串行化 ──► 网关身份头/actuator ──► Netty 解耦阻塞 IO\n"
        "      │\n"
        "      ▼\n"
        "热更 static/RAG 回滚 ──► 配置路径围栏+Schema ──► common-api 统一\n"
        "      │\n"
        "      ▼\n"
        "经济幂等+Redis 强制 ──► 反作弊/审计 ──► 中心服迁移 PoC",
    )

    # ========== 四、验收标准 ==========
    add_heading(doc, "四、整改验收标准（可衡量）", 1)
    add_table(
        doc,
        ["目标", "验收标准"],
        [
            ["IAP 安全", "prod 下 mock=true 无法启动；有渠道时伪造收据不可发货"],
            ["微服务密钥", "默认 change-me / 123456 / admin123 在 prod profile 启动失败"],
            ["任务进度", "杀怪/NPC 触发后 objectives_json 与 status 按配置更新"],
            ["战斗并发", "同 battle 并发写入+tick 无 CME/丢状态；实体 ID 唯一"],
            ["热更一致", "RAG/static 阶段失败后运行时配置与失败前一致"],
            ["配置安全", "../ 路径写入被拒绝；坏 Schema dry-run 失败"],
            ["契约统一", "主工程无本地 AssistPromptTemplates 副本，依赖 common-api"],
            ["网关身份", "伪造入站 X-User-Id 被剥离，下游仅见网关签发值"],
            ["内部面安全", "无 Token 访问 /internal/** 返回 401；比较为恒定时间"],
            ["多节点", "未开 Redis 的 multi-node 配置启动失败"],
            ["停机可靠", "在线 N 人停机，超时内落盘或明确失败 uid；匹配玩家收到取消"],
        ],
    )

    # ========== 五、总结 ==========
    add_heading(doc, "五、总结", 1)
    add_body(
        doc,
        "项目已从「骨架缺失」进入「骨架可用、关键业务与资金/并发面仍有硬缺口」阶段："
        "管理端鉴权、内部 Token、钱包乐观更新、抽卡历史、排行榜 Redis、热更与停机等基础设施明显增强；"
        "Boot 版本双轨问题已消除。但 IAP 真实验签缺失、微服务默认密钥、任务占位、战斗/场景并发、"
        "Netty 阻塞、热更后半段回滚、配置路径围栏与 Assist 双轨仍是上线与扩展的主要拦路虎。",
    )
    add_body(
        doc,
        "建议在继续横向加玩法或推进微服务拆分之前，先完成 P0 的资金验签、密钥治理、任务触发与战斗串行化；"
        "否则 Mock 充值、弱 JWT、空任务进度与偶现战斗态错乱会直接体现在玩家体验与资金风险上。"
        "本文档基于当前工作区静态审查生成，建议结合一次压测与故障演练结果修订。",
    )

    return doc


def main():
    DESKTOP.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(str(OUTPUT))
    print(f"Wrote: {OUTPUT}")


if __name__ == "__main__":
    main()
