# -*- coding: utf-8 -*-
"""Generate MyLunarCore AI assist optimization solutions Word document to Desktop."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DESKTOP = Path.home() / "Desktop"
OUTPUT = DESKTOP / ("MyLunarCore_AI" + "\u8f85\u52a9\u529f\u80fd\u4f18\u5316\u65b9\u6848" + ".docx")


def set_doc_fonts(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15


def add_title(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(25, 75, 140)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_subtitle(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(100, 100, 100)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_heading(doc: Document, text: str, level: int = 1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        if level == 1:
            run.font.color.rgb = RGBColor(25, 75, 140)


def add_body(doc: Document, text: str):
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_bullet(doc: Document, text: str, level: int = 0):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.5 * (level + 1))
    for run in p.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_code(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 30, 30)


def add_table(doc: Document, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.name = "微软雅黑"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                run.font.size = Pt(10)
    for r_idx, row in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = val
            for p in row_cells[c_idx].paragraphs:
                for run in p.runs:
                    run.font.name = "微软雅黑"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                    run.font.size = Pt(10)
    return table


def build():
    doc = Document()
    set_doc_fonts(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    add_title(doc, "MyLunarCore AI 辅助功能优化方案")
    add_subtitle(
        doc,
        f"基于现有实现代码审查 · 问题诊断与可落地方案 · {date.today().isoformat()}",
    )
    add_body(
        doc,
        "本文档基于仓库当前已落地的 AI 辅助能力（方案 A 规则教练、方案 B LLM+RAG、"
        "方案 C 旁路微服务 ai-assist-service、方案 D 战斗启发式/特征采集、方案 E 攻略包）"
        "进行代码级体检，列出仍需优化的问题，并给出分优先级的解决方案与落地步骤。"
        "原则不变：AI 只读权威状态、可降级、可热更、可审计，不得直接写货币/抽卡/战斗结算。",
    )

    # ---------- 1 现状 ----------
    add_heading(doc, "一、当前实现盘点（已具备）", 1)
    add_table(
        doc,
        ["能力", "关键组件", "成熟度"],
        [
            ["规则教练", "PlayerCoachApplicationService / CoachRuleEngine / CoachTips.json", "可用（MVP）"],
            ["自然语言助手", "AiAssistApplicationService / LlmAssistGateway / AssistSafetyFilter", "可用（可降级）"],
            ["轻量 RAG", "RagKnowledgeService（关键词+bigram，热更 rebuild）", "雏形"],
            ["旁路微服务", "ai-assist-service + AiAssistClient（失败降级）", "可用（默认关闭）"],
            ["配额与审计", "AssistQuotaLimiter / AssistAuditService", "基础"],
            ["聊天 @助手", "HallNettyService + AssistChatMentionParser", "可用"],
            ["攻略包 Local First", "GuidePackRepository + FAQ 打分匹配", "可用（粗糙）"],
            ["战斗提示", "HeuristicBattleAssistPolicy（固定技能1打首怪）", "极简"],
            ["数据飞轮", "BattleFeatureCollector / AvatarUsageStatsService", "采集偏弱"],
        ],
    )
    add_body(
        doc,
        "协议面已齐：AskCoachHint / AskAiAssist / AiHintNotify / GetGuidePack / AskBattleHint（CmdIds 900–912、211–212）。"
        "主链路已在 gameBusinessExecutorGroup 上执行，不会直接堵死 Netty EventLoop；"
        "但 LLM/远程 HTTP 仍会长时间占用业务线程，高并发时仍会影响同池其它玩法请求。",
    )

    # ---------- 2 总览 ----------
    add_heading(doc, "二、优化项总览（按优先级）", 1)
    add_table(
        doc,
        ["优先级", "优化项", "痛点", "预期收益", "建议工期"],
        [
            ["P0", "AskAiAssist 异步化 + 先回执后推送", "同步等 LLM/远程，占满业务线程", "延迟可控、游戏玩法不被拖慢", "3–5 天"],
            ["P0", "回答缓存 + 场景化配额拆分", "重复问浪费 Token；教练与 LLM 抢同一配额", "成本↓、体验↑", "2–3 天"],
            ["P0", "安全过滤配置化与出站强化", "硬编码正则易漏拦/误拦", "合规风险↓", "2–4 天"],
            ["P1", "RAG 升级（向量/混合检索）", "关键词召回不准，易幻觉", "回答相关性↑", "1–2 周"],
            ["P1", "上下文精简与按场景裁剪", "默认投影背包128/角色24，Token 膨胀", "成本↓、延迟↓", "3–5 天"],
            ["P1", "可观测性打通（Micrometer/指标）", "审计计数不全，remote-* 未统计", "可运维、可告警", "2–3 天"],
            ["P2", "FAQ 语义匹配 + 攻略包完整字段", "打分阈值粗；size=0", "Local First 命中率↑", "3–5 天"],
            ["P2", "多轮会话与引用展示", "单次问答，无历史", "体验接近真人助手", "1 周"],
            ["P2", "战斗策略可解释升级", "永远 skill1+首怪", "PVE 托管可用性↑", "1–2 周"],
            ["P3", "旁路服务去重与契约统一", "游戏服/微服务双套规则与 LLM", "维护成本↓", "与微服务节奏对齐"],
            ["P3", "反馈闭环与有用率", "无点赞/纠错", "持续优化内容质量", "1 周"],
        ],
    )

    # ---------- 3 P0 ----------
    add_heading(doc, "三、P0：延迟与稳定性（必须先做）", 1)

    add_heading(doc, "3.1 AskAiAssist 同步阻塞 → 异步回执 + AiHintNotify", 2)
    add_body(
        doc,
        "现状：AssistPacketHandlers.onAskAiAssist 在业务线程内同步调用 AiAssistApplicationService.ask()，"
        "内部可能串行：本地规则教练 → 远程 AiAssistClient.httpClient.send → 本地 LlmAssistGateway.httpClient.send。"
        "协议里已有 AiHintNotify，但同步路径几乎不用，仅聊天旁路 askAndNotify 会推送。",
    )
    add_body(doc, "解决方案：")
    add_bullet(doc, "协议约定：AskAiAssistScRsp 立即返回 retcode=0 + source=accepted（或保留同步短路径仅给 FAQ/规则）。")
    add_bullet(doc, "将 LLM/远程调用提交到独立 AssistInferenceExecutor（与战斗/登录业务池隔离，队列有界、拒绝策略可降级）。")
    add_bullet(doc, "推理完成后用 channel 事件线程 writeAndFlush AiHintNotify；channel 已断开则只写审计。")
    add_bullet(doc, "超时：超过 llmTimeoutMs/remoteTimeoutMs 推送 source=fallback 的规则答案，retcode 语义对齐 proto 注释中的「4=超时/降级」。")
    add_bullet(doc, "可选双模：配置 syncMode=true 供联调；生产默认 async。")
    add_code(
        doc,
        "Client --AskAiAssistCsReq--> AssistNettyService\n"
        "  |-- 快速路径：Guide FAQ / 规则教练（<20ms）→ AskAiAssistScRsp\n"
        "  |-- 慢路径：提交 AssistInferenceExecutor\n"
        "         |-- remote / llm\n"
        "         v\n"
        "      AiHintNotify (request_id 关联)",
    )

    add_heading(doc, "3.2 回答缓存与幂等", 2)
    add_body(
        doc,
        "现状：同一玩家短时间内重复问「今天做什么」会重复打远程/LLM；无 request_id 幂等。",
    )
    add_body(doc, "解决方案：")
    add_bullet(doc, "缓存键：uid + scene + normalize(question) 的 hash；TTL 建议 60–180s；仅缓存 source∈{guide-local,rule,llm,remote-*} 且 retcode=0。")
    add_bullet(doc, "normalize：去空白、全角半角、常见口语同义词映射（「抽卡/卡池」）。")
    add_bullet(doc, "实现可用 Caffeine 本地缓存；多节点后再上 Redis（与配额一并考虑）。")
    add_bullet(doc, "命中缓存时 source 追加 -cache，便于审计区分。")

    add_heading(doc, "3.3 配额拆分与桶生命周期", 2)
    add_body(
        doc,
        "现状：AssistQuotaLimiter 用 ConcurrentHashMap 懒建 Bucket，无淘汰；"
        "教练提示与 LLM 提问共用 perUidPerMinute=20，高频点「今日建议」会挤掉真正问答配额。",
    )
    add_body(doc, "解决方案：")
    add_bullet(doc, "拆分配额：coachPerUidPerMinute（可高，如 60）与 llmPerUidPerMinute（可低，如 10）；日配额 dailyLlmBudget 防刷 Token。")
    add_bullet(doc, "远程成功与本地 LLM 共用「贵配额」，FAQ/规则命中只扣「便宜配额」或不扣。")
    add_bullet(doc, "桶淘汰：Idle 超过 N 分钟移除；或改用弱引用/定时清理，避免长生命周期服内存膨胀。")
    add_bullet(doc, "限流响应带 retryAfterSeconds，客户端可做倒计时提示。")

    add_heading(doc, "3.4 安全过滤配置化与强化", 2)
    add_body(
        doc,
        "现状：AssistSafetyFilter / RemoteAssistEngine 硬编码少量正则；"
        "隐私探测与充值诱导覆盖面窄；出站命中则整段替换，体验生硬。",
    )
    add_body(doc, "解决方案：")
    add_bullet(doc, "将拦截词表落到 data/AssistSafetyRules.json，走 ConfigFileService + 热更，与 CoachTips 一致。")
    add_bullet(doc, "入站分级：block（隐私/违法）/ soft-block（付费诱导提问改官方文案）/ allow。")
    add_bullet(doc, "出站：对「必出」类做句子级替换或拒绝该句，保留其余有用建议；保留长度上限 800。")
    add_bullet(doc, "抽卡场景强制追加官方概率免责声明模板（与 Gacha 文案同源）。")
    add_bullet(doc, "单测扩充：边界语料库（同音、空格插入、英文混写）。")

    # ---------- 4 P1 ----------
    add_heading(doc, "四、P1：回答质量与成本", 1)

    add_heading(doc, "4.1 RAG：从关键词雏形到混合检索", 2)
    add_body(
        doc,
        "现状：RagKnowledgeService 用 tokenize + contains 打分；卡池知识块几乎无机制文案；"
        "商店仅「商品数」；热更已挂钩 HotReloadCoordinator / ActivityConfigHotReloadService（这点应保留）。",
    )
    add_body(doc, "解决方案（分两步）：")
    add_bullet(doc, "Step1（低成本）：丰富 chunk 文本——Banner 写入官方保底说明字段、Shop 写入商品名/货币类型/限购；任务写入目标场景/奖励摘要；切块长度 200–400 字。")
    add_bullet(doc, "Step1：检索加 BM25 或 TF-IDF；同义词词典（主线/任务、卡池/抽卡）；最低分阈值避免噪声召回。")
    add_bullet(doc, "Step2：引入向量索引（本地 ONNX embedding 或 pgvector/RedisSearch），接口仍暴露 search(question,scene,limit)。")
    add_bullet(doc, "混合分：0.6*向量 + 0.4*关键词；scene 加权保留。")
    add_bullet(doc, "引用约束：提示词要求回答必须带 citedConfigIds 中的 ID；否则降级规则文案。")

    add_heading(doc, "4.2 上下文白名单按场景裁剪", 2)
    add_body(
        doc,
        "现状：AssistContextBuilder 默认投影 inventory≤128、avatars≤24、talents≤16，再加 globalAvatarUsage；"
        "无论问「活动怎么玩」都会塞满上下文，Token 与延迟双高。",
    )
    add_body(doc, "解决方案：")
    add_bullet(doc, "按 scene 裁剪：quest→任务+相关材料；gacha→pity/卡池进度；growth→角色属性+天赋；activity→活动列表；meta→使用率；mail→mailSummary。")
    add_bullet(doc, "通用问题仅传 monitoringSummary + activeQuests Top3，避免整包背包。")
    add_bullet(doc, "配置化上限：contextMaxInventorySlots 可按场景覆盖（如 activity=0）。")
    add_bullet(doc, "对 LLM 请求增加 max_tokens / 回答长度提示，控制费用。")

    add_heading(doc, "4.3 审计与可观测性补齐", 2)
    add_body(
        doc,
        "现状：AssistAuditService 只把 source=llm 记入 llmHitCount；remote-llm / guide-local / blocked 等未细分；"
        "游戏服侧指标未挂到 ServerMetricsMonitor / Micrometer；旁路服务有 /ai/metrics 但游戏服未拉取。",
    )
    add_body(doc, "解决方案：")
    add_bullet(doc, "按 source 前缀分类计数：guide-local / rule / fallback / llm / remote-* / blocked / error / cache。")
    add_bullet(doc, "暴露指标：assist_ask_total、assist_latency_ms、assist_llm_tokens_est、assist_quota_reject、assist_remote_error_ratio。")
    add_bullet(doc, "接入现有 ServerMetricsMonitor 周期日志；告警：错误率>5%、P99 延迟>超时阈值、远程连续失败。")
    add_bullet(doc, "审计日志增加 requestId、latencyMs、cacheHit、citedCount；保留脱敏策略。")

    # ---------- 5 P2 ----------
    add_heading(doc, "五、P2：产品体验增强", 1)

    add_heading(doc, "5.1 GuidePack FAQ 与端侧包完整性", 2)
    add_body(
        doc,
        "现状：matchGuideFaq 用包含关系+分词加分，阈值 2；"
        "toGuidePackProto 的 size 恒为 0；downloadUrl/contentHash 依赖配置但缺少校验与增量更新策略。",
    )
    add_body(doc, "解决方案：")
    add_bullet(doc, "FAQ 增加 keywords[]、scene、priority；匹配分 = 问句相似度 + scene 加权。")
    add_bullet(doc, "服务端计算 contentHash（SHA-256）与真实 size；客户端按 hash 决定是否重新下载。")
    add_bullet(doc, "热更推送 GuidePackUpdateScNotify 时带差分版本号；弱网优先本地包。")
    add_bullet(doc, "覆盖率目标：常见 50 问 Local First 命中率 ≥ 40%，显著降 Token。")

    add_heading(doc, "5.2 多轮会话与引用展示", 2)
    add_body(doc, "解决方案：")
    add_bullet(doc, "协议扩展：AskAiAssistCsReq 增加 session_id / client_request_id；服务端保留最近 3–5 轮摘要（仅本 uid，TTL 10min）。")
    add_bullet(doc, "Rsp 强化 cited_config_ids + related_hints 的客户端 UI：可点击跳转任务/活动/商店（复用 action 字段）。")
    add_bullet(doc, "对「接着刚才」类指代，用会话摘要而非完整历史，控制 Token。")

    add_heading(doc, "5.3 战斗战术提示升级（仍限 PVE）", 2)
    add_body(
        doc,
        "现状：HeuristicBattleAssistPolicy 固定建议 skillId=1、目标=当前波次第一个存活怪；"
        "BattleFeatureCollector 仅落 battleId/playerId/endStatus 等稀疏字段，难以训练策略模型。",
    )
    add_body(doc, "解决方案：")
    add_bullet(doc, "短期：威胁评分（剩余 HP、攻击力、精英标记）选目标；按能量/冷却选技能；reason 人可读。")
    add_bullet(doc, "特征采集补齐：波次、敌我血量向量、可用技能、最终操作、胜负；写入 JSONL 保持旁路、不进结算。")
    add_bullet(doc, "接口保持 BattleAssistPolicy，后续可换 ONNX 小模型；PVP 默认禁用。")
    add_bullet(doc, "公平性：模型版本号写入战报；托管需玩家显式开关。")

    add_heading(doc, "5.4 规则教练内容运营", 2)
    add_body(doc, "解决方案：")
    add_bullet(doc, "扩充 CoachTips：体力耗尽、邮件未读、商店刷新、突破材料齐备、热门角色已拥有但未养成等。")
    add_bullet(doc, "tip 增加 cooldownHours / oncePerDay，避免同一提示刷屏。")
    add_bullet(doc, "与活动排期联动：结束前 24h 自动提高 activity tip 优先级。")

    # ---------- 6 P3 ----------
    add_heading(doc, "六、P3：架构收敛与长期能力", 1)

    add_heading(doc, "6.1 双端逻辑去重", 2)
    add_body(
        doc,
        "现状：游戏服有完整 Coach/RAG/LLM/Safety；微服务 RemoteAssistEngine 又实现一套规则与安全正则，"
        "RemoteLlmGateway 提示词更短、能力不一致。开启 remoteEnabled 时行为可能与本地路径漂移。",
    )
    add_body(doc, "解决方案：")
    add_bullet(doc, "目标态：游戏服只做鉴权、配额、上下文投影、协议；推理与 RAG 全集中在 ai-assist-service。")
    add_bullet(doc, "过渡态：抽取 common 模块共享 Safety 规则与 Prompt 模板；微服务成为唯一 LLM 出口。")
    add_bullet(doc, "契约：继续用 common-api 的 AssistAskRequest/Response；增加 schema 版本字段。")
    add_bullet(doc, "网关：api-gateway 已有 ai-assist-service-route，补齐超时/重试/熔断配置。")

    add_heading(doc, "6.2 用户反馈闭环", 2)
    add_bullet(doc, "新增协议：AssistFeedbackCsReq（request_id、useful、reason_code）。")
    add_bullet(doc, "落库或日志：用于统计有用率、坏案例回放、FAQ 补录。")
    add_bullet(doc, "运营周报：Top 未命中问题、高投诉场景（抽卡/付费）、降级率。")

    add_heading(doc, "6.3 匹配/阵容推荐（数据成熟后）", 2)
    add_bullet(doc, "复用 AvatarUsageStatsService 与 MatchCompatibilityScorer，输出「已拥有角色中的推荐阵容」。")
    add_bullet(doc, "禁止输出未拥有角色的强诱导抽卡话术；可提示「活动获取途径」类官方配置。")

    # ---------- 7 落地路线 ----------
    add_heading(doc, "七、推荐落地路线", 1)
    add_table(
        doc,
        ["阶段", "交付物", "成功标准"],
        [
            ["第1周 P0", "异步推理池 + 缓存 + 配额拆分 + 安全词表热更", "业务线程 P99 不被 LLM 拉高；重复问缓存命中≥30%"],
            ["第2–3周 P1", "RAG 文本增强 + 场景裁剪上下文 + 指标看板", "人工抽检相关性↑；单次 Token 成本↓≥40%"],
            ["第4–5周 P2", "FAQ 增强、多轮摘要、战斗启发式升级", "Local First≥40%；战斗提示可解释且仅 PVE"],
            ["后续 P3", "推理收敛到旁路服务 + 反馈闭环", "游戏服故障隔离；有用率可量化"],
        ],
    )

    # ---------- 8 代码落点 ----------
    add_heading(doc, "八、建议改动落点（与仓库对齐）", 1)
    add_table(
        doc,
        ["改动点", "文件/模块", "说明"],
        [
            ["异步调度", "AssistNettyService / AssistPacketHandlers", "快回执 + AiHintNotify"],
            ["推理线程池", "config 新增 AssistInferenceConfiguration", "与 gameBusiness 池隔离"],
            ["缓存", "新增 AssistAnswerCache", "Caffeine TTL"],
            ["配额", "AssistQuotaLimiter", "coach/llm/daily 三维"],
            ["安全", "AssistSafetyFilter + data/AssistSafetyRules.json", "热更词表"],
            ["RAG", "RagKnowledgeService", "富文本块 → 混合检索"],
            ["上下文", "AssistContextBuilder", "按 scene 裁剪"],
            ["审计指标", "AssistAuditService + ServerMetricsMonitor", "分 source 计数"],
            ["FAQ/攻略包", "AiAssistApplicationService / AssistNettyService", "hash/size/关键词"],
            ["战斗", "HeuristicBattleAssistPolicy / BattleFeatureCollector", "威胁分 + 富特征"],
            ["旁路收敛", "microservices/.../ai-assist-service", "唯一 LLM；共享契约"],
        ],
    )

    # ---------- 9 风险 ----------
    add_heading(doc, "九、风险与验收清单", 1)
    add_table(
        doc,
        ["风险", "缓解", "验收"],
        [
            ["异步乱序/重复推送", "request_id 去重；channel 绑定校验", "压测无串包"],
            ["缓存返回过期活动建议", "活动热更时 invalidate scene=activity 缓存", "热更后 1 分钟内无过期答"],
            ["向量库运维复杂", "先做 BM25 富文本，向量作可选开关", "关闭向量时行为回退"],
            ["安全误杀正常提问", "soft-block + 审计抽检", "误拦率<1%"],
            ["PVP 公平性质疑", "战斗 AI 仅 PVE；开关默认关", "PVP 无战术提示"],
            ["微服务与本地答案不一致", "统一 Prompt/Safety；灰度对比日志", "抽样一致率达标"],
        ],
    )

    # ---------- 10 结论 ----------
    add_heading(doc, "十、结论", 1)
    add_body(
        doc,
        "MyLunarCore 的 AI 辅助「骨架」已经齐全，当前最大短板不再是「有没有」，而是："
        "（1）慢路径同步占业务线程；（2）RAG/上下文导致质量与成本不稳定；"
        "（3）安全、配额、观测仍偏演示级；（4）战斗与旁路服务停留在占位实现。"
        "建议严格按 P0→P1→P2→P3 推进：先把延迟与成本打牢，再提升召回与体验，最后收敛架构并形成数据闭环。",
    )
    add_body(
        doc,
        "落地时继续坚持：只读上下文、失败可降级到规则教练、配置可热更、问答可审计。"
        "这样既提升玩家引导体验，又不会破坏战斗结算、经济与抽卡的一致性边界。",
    )

    add_subtitle(doc, "—— 文档结束 ——")
    doc.save(OUTPUT)
    print(f"Wrote: {OUTPUT}")


if __name__ == "__main__":
    build()
