# -*- coding: utf-8 -*-
"""修正 bak.docx 中 7.30 章节段落顺序。"""

import sys

sys.stdout.reconfigure(encoding="utf-8")

from pathlib import Path

from docx import Document

path = Path.home() / "Desktop" / (
    "MyLunarCore" + "\u9879\u76ee\u5404\u65b9\u9762\u8be6\u7ec6\u603b\u7ed3\u62a5\u544a" + ".bak.docx"
)
d = Document(str(path))

wanted = {}
for p in d.paragraphs:
    t = p.text.strip()
    if t.startswith("7.30 稳定性与体验深化"):
        wanted["h"] = p
    elif t.startswith("在 7.27"):
        wanted["intro"] = p
    elif t.startswith("（1）跨节点组队与私聊"):
        wanted["body"] = p
    elif t.startswith("实现入口：party"):
        wanted["impl"] = p
    elif t.startswith("测试入口：WireV4"):
        wanted["test"] = p
    elif t.startswith("客户端联调要点"):
        wanted["client"] = p

missing = [k for k in ("h", "intro", "body", "impl", "test", "client") if k not in wanted]
if missing:
    raise SystemExit(f"missing keys: {missing}")

ch8 = None
for p in d.paragraphs:
    if p.style and p.style.name == "Heading 1" and p.text.strip().startswith("八、"):
        ch8 = p
        break
if ch8 is None:
    raise SystemExit("chapter 8 not found")

order = [wanted["h"], wanted["intro"], wanted["body"], wanted["impl"], wanted["test"], wanted["client"]]
elements = [p._p for p in order]
for el in elements:
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)

ref = ch8._p
parent = ref.getparent()
idx = list(parent).index(ref)
for i, el in enumerate(elements):
    parent.insert(idx + i, el)

d.save(str(path))
print("reordered 7.30 OK ->", path)

d2 = Document(str(path))
for i, p in enumerate(d2.paragraphs):
    t = p.text
    if (
        "7.30" in t
        or t.startswith("在 7.27")
        or t.startswith("（1）跨节点")
        or "实现入口：party" in t
        or "测试入口：WireV4" in t
        or t.startswith("客户端联调")
    ):
        print(i, p.style.name, t[:70])
    if p.style and p.style.name == "Heading 1" and p.text.strip().startswith("八、"):
        print(i, ">>>", p.text)
        break
