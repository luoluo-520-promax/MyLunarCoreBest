# -*- coding: utf-8 -*-
"""Generate MyLunarCore vs mainstream open-world 二游 PC server assessment Word doc."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

DESKTOP = Path.home() / "Desktop"
OUTPUT = DESKTOP / "MyLunarCore大世界二游服务器评估与解决方案.docx"


def set_doc_fonts(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_title(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(25, 75, 140)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_subtitle(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(100, 100, 100)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_heading(doc: Document, text: str, level: int = 1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        if level == 1:
            run.font.color.rgb = RGBColor(25, 75, 140)


def add_body(doc: Document, text: str):
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_bullet(doc: Document, text: str, level: int = 0):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.5 * (level + 1))
    for run in p.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_code(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 30, 30)


def add_table(doc: Document, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.name = "微软雅黑"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for r_idx, row in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = val
            for p in row_cells[c_idx].paragraphs:
                for run in p.runs:
                    run.font.name = "微软雅黑"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()
    set_doc_fonts(doc)

    add_title(doc, "MyLunarCore 大世界二游电脑端服务器评估与解决方案")
    add_subtitle(doc, "对照主流：原神 / 鸣潮式无缝大世界 vs 崩铁式分区实例")
    add_subtitle(doc, "项目路径：c:\\Users\\ASUS\\IdeaProjects\\test\\MyLunarCore")
    add_subtitle(doc, "生成日期：2026年8月3日（基于当前代码快照）")
    doc.add_paragraph()

    # 一、总体结论
    add_heading(doc, "一、总体结论", 1)
    add_body(
        doc,
        "结论：当前项目不完全符合「原神 / 鸣潮」类无缝大世界实时动作二游的电脑端服务器主流形态；"
        "更接近「崩坏：星穹铁道」类分区探索 + 回合制战局服务端。已具备可运行的模块化单体游戏核与部分多人 Zone 雏形，"
        "但生产级多节点、无缝世界、实时战斗权威与反作弊能力仍不足。",
    )
    add_body(
        doc,
        "技术形态：Spring Boot 模块化单体（Netty TCP/KCP + Protobuf + MySQL），"
        "microservices/ 为 Spring Cloud Alibaba 脚手架，尚未替代主工程游戏核。"
        "推荐定位：短期按「崩铁式」产品打磨闭环；若目标是无缝大世界，需按下文分阶段改造，不可仅靠拆微服务。",
    )

    add_heading(doc, "1.1 与两类主流对标", 2)
    add_table(
        doc,
        ["对照维度", "原神/鸣潮式主流", "崩铁式主流", "MyLunarCore 现状"],
        [
            [
                "世界模型",
                "无缝大地图 + Cell/分片流式加载",
                "Plane/Floor 分区实例",
                "planeId+floorId→Zone，非无缝",
            ],
            [
                "战斗",
                "世界内实时动作权威",
                "遇敌进入回合/波次实例",
                "WorldEncounter→回合战，偏崩铁",
            ],
            [
                "帧率/Tick",
                "移动/战斗 20–60Hz 模拟",
                "场景低频 + 战斗独立逻辑",
                "GameServer 默认约 1s tick",
            ],
            [
                "多人可见",
                "同区域 AOI + 场景服舰队",
                "同 Plane 可见或弱社交",
                "ZoneContext+AoiGrid 九宫格已有",
            ],
            [
                "多节点",
                "Center+Scene 分片+票据+共享状态",
                "Center 路由 + 重连切服",
                "Center 雏形有，票据仅进程内存",
            ],
            [
                "缓存/在线",
                "Redis 在线表/锁/票据/聊天",
                "Redis 或等价中间件常用",
                "无 Redis，仅 MySQL+Caffeine",
            ],
            [
                "协议",
                "TCP/KCP/自研帧+Protobuf",
                "同左",
                "符合：LunarFrame+Protobuf+KCP",
            ],
        ],
    )

    add_heading(doc, "1.2 架构现状示意", 2)
    add_code(
        doc,
        "客户端 ──TCP/KCP :9000──► GameNettyServer / GameKcpServer\n"
        "                              │\n"
        "                    PacketDispatcher → *NettyService\n"
        "                              │\n"
        "              ┌───────────────┼───────────────┐\n"
        "              ▼               ▼               ▼\n"
        "         Zone/Scene      BattleManager    Player/Session\n"
        "              │               │\n"
        "              └─ 遭遇进战 ─────┘\n"
        "                              │\n"
        "                    CenterServer.planMigration\n"
        "                    (local 内存 | HTTP remote)\n"
        "\n"
        "Admin/运维 ──HTTP :8080──► Controllers + Security\n"
        "AI（可选）──HTTP──► ai-assist-service :18083（失败回退进程内）",
    )

    # 二、符合度总表
    add_heading(doc, "二、逐项符合度评估（大世界二游电脑端服务器）", 1)
    add_table(
        doc,
        ["能力域", "符合度", "现状摘要", "主流期望"],
        [
            [
                "协议与网络",
                "符合",
                "Netty TCP+KCP、Protobuf 帧、限流背压",
                "长连接二进制协议 + 弱网优化",
            ],
            [
                "会话与状态机",
                "较符合",
                "HALL/SCENE/BATTLE 等状态机 + 重连 Token",
                "会话权威、踢人、重连",
            ],
            [
                "场景/大世界",
                "部分符合",
                "共享 Zone+AOI+世界遭遇，仍是分区非无缝",
                "Cell 分片、边界移交、流式加载权威",
            ],
            [
                "战斗权威",
                "部分符合（偏崩铁）",
                "回合/波次+韧性；非世界实时战斗",
                "实时技能时间轴或明确的进战实例模型",
            ],
            [
                "中心路由/多节点",
                "雏形",
                "Local/Remote Center、迁移 API 有",
                "票据共享、容量租约、心跳摘除",
            ],
            [
                "持久化",
                "基本可用",
                "MySQL 聚合 + 周期快照 + 乐观版本",
                "在线缓存、关键写审计、分库预案",
            ],
            [
                "配置热更",
                "符合",
                "JSON/Excel 导入、分阶段热更与回滚审计",
                "配置版本化、灰度发布",
            ],
            [
                "运营后台",
                "较符合",
                "Admin 登录锁定、内部 Token、热更/导入",
                "RBAC + 审计 + 内部面隔离",
            ],
            [
                "反作弊",
                "不符合",
                "仅有包频/登录限流，无移动与伤害审计",
                "服务器权威校验 + 风控流水",
            ],
            [
                "微服务拆分",
                "脚手架",
                "gateway/auth/player/ai-assist 未替代游戏核",
                "先拆运维/AI，战斗场景勿过早拆",
            ],
            [
                "AI 助手",
                "加分项",
                "规则教练 + 可选 LLM 边车，故障回退",
                "非硬实时旁路服务（可选）",
            ],
        ],
    )

    # 三、已符合项
    add_heading(doc, "三、已符合或接近主流的部分", 1)
    add_bullet(doc, "传输层：GameNettyServer / GameKcpServer，同一套 Lunar 帧编解码，适合 PC 端弱网。")
    add_bullet(doc, "协议：Protobuf + opcode 帧；连接限流、背压、幂等处理已有工程意识。")
    add_bullet(doc, "玩法骨架：战斗、挑战、Rogue、抽卡、物品、活动、商店/皮肤、任务配置等二游标配齐全。")
    add_bullet(doc, "多人场景雏形：ZoneManager / ZoneContext / AoiGrid / SceneSyncBroadcaster / ZoneWorldService。")
    add_bullet(doc, "大世界进战钩子：WorldEncounterService 仇恨触发 FightStart，战后回锚点。")
    add_bullet(doc, "中心服接口形状：CenterServer、SceneRegistry、PlayerMigrationService、HTTP plan API。")
    add_bullet(doc, "配置运营：HotReloadCoordinator 分阶段重载与回滚；活动 Excel/JSON 导入。")
    add_bullet(doc, "AI 边车：ai-assist-service + AiAssistClient，符合「游戏核保持实时、AI 异步旁路」主流做法。")

    # 四、不符合项与解决方案
    add_heading(doc, "四、不符合项与解决方案", 1)
    add_body(
        doc,
        "以下按「是否要做无缝大世界」分两条产品线给出方案。"
        "方案 A：坚持崩铁式（推荐短期）。方案 B：转向原神/鸣潮式无缝大世界（成本高）。",
    )

    add_heading(doc, "4.1 世界模型：分区实例 vs 无缝大世界", 2)
    add_heading(doc, "差距", 3)
    add_bullet(doc, "现状用 planeId*10000+floorId 作为 zoneId，硬分区；无 Cell 流式加载、无边界物理移交。")
    add_bullet(doc, "SceneContext（玩家视图）与 ZoneContext（区域权威）双模型并存，需严格投影一致性。")

    add_heading(doc, "解决方案 A（崩铁式，推荐）", 3)
    add_bullet(doc, "明确产品语义：地图=可切换 Plane/Floor，不做无缝；强化加载门、传送门、锚点回城。")
    add_bullet(doc, "统一权威：怪物/采集物只在 ZoneWorldService 写；SceneContext 只读投影 + 差异推送。")
    add_bullet(doc, "补齐传送/切图协议与失败回滚；同 Zone 人数上限与排队。")

    add_heading(doc, "解决方案 B（无缝大世界）", 3)
    add_bullet(doc, "引入 WorldGrid：世界坐标 → cellId；按 cell 集合归属 SceneServer。")
    add_bullet(doc, "边界 handoff：玩家靠近边界时预订阅邻 cell，跨服时短冻结+状态序列化迁移。")
    add_bullet(doc, "提高世界模拟频率（建议移动 10–20Hz，战斗相关更高），GameServer 1s tick 仅保留全局调度。")
    add_bullet(doc, "独立 scene-service 舰队；Center 做 cell→节点映射与容量调度。")

    add_heading(doc, "4.2 战斗：回合进战 vs 世界实时", 2)
    add_heading(doc, "差距", 3)
    add_bullet(doc, "WorldEncounterService 遇敌后合成 FightStart，进入回合/波次 BattleContext，非世界内 hitbox。")
    add_bullet(doc, "BattleManager 以单 playerId 为主键，缺多人共战权威与回放校验。")

    add_heading(doc, "解决方案 A（保持回合制）", 3)
    add_bullet(doc, "产品上定义为「探索实时移动 + 遇敌进战」，与崩铁一致，不要强行对标鸣潮。")
    add_bullet(doc, "完善进战/结算/掉落事务；战后 Zone 刷怪与场景锚点一致性测试。")
    add_bullet(doc, "增加战斗日志与可选回放；技能 CD、伤害上下限服务器校验。")

    add_heading(doc, "解决方案 B（实时动作）", 3)
    add_bullet(doc, "新建 RealtimeCombatAuthority：技能时间轴、碰撞粗检、伤害结算队列。")
    add_bullet(doc, "客户端预测 + 服务器校正；禁止客户端直传最终伤害。")
    add_bullet(doc, "与 Zone tick 同进程或同机房低延迟；勿过早拆到异地微服务。")

    add_heading(doc, "4.3 多节点与迁移票据（关键缺口）", 2)
    add_heading(doc, "差距", 3)
    add_bullet(doc, "MigrationTicketService 票据存在进程内 ConcurrentHashMap，跨 JVM 无法核销。")
    add_bullet(doc, "SceneRegistry 首次注册胜出，无心跳、租约、摘除、排水（draining）。")
    add_bullet(doc, "InMemoryMultiNodeRoutingClient 偏演练；缺全局在线表与玩家数据跨节点交接。")

    add_heading(doc, "解决方案", 3)
    add_table(
        doc,
        ["阶段", "目标", "建议实现"],
        [
            [
                "P0",
                "票据可跨节点",
                "Redis SET ticket→payload + TTL；目标节点核销（GETDEL）",
            ],
            [
                "P1",
                "真实注册中心",
                "节点心跳、容量、zone 租约；过期剔除；planMigration 按负载选点",
            ],
            [
                "P2",
                "切服闭环",
                "源节点导出轻量会话→目标校验票据→客户端重连 advertise host/port",
            ],
            [
                "P3",
                "玩家分片",
                "uid 哈希或一致性哈希；全局 online(uid→node) 表",
            ],
        ],
    )
    add_code(
        doc,
        "建议配置项（示意）：\n"
        "lunarcore.center.mode=remote\n"
        "lunarcore.redis.url=redis://...\n"
        "lunarcore.node.advertise-host=...\n"
        "lunarcore.node.advertise-port=9000\n"
        "lunarcore.node.heartbeat-ms=3000",
    )

    add_heading(doc, "4.4 缺少 Redis（或等价中间件）", 2)
    add_heading(doc, "差距", 3)
    add_bullet(doc, "主工程无 Redis：在线状态、分布式锁、迁移票据、聊天扇出、排行榜均难做主流方案。")

    add_heading(doc, "解决方案", 3)
    add_bullet(doc, "引入 Redis：session/online、migration ticket、mailbox/chat pubsub、leaderboard ZSET。")
    add_bullet(doc, "MySQL 继续作为玩家权威落盘；Caffeine 仅作单机热点缓存。")
    add_bullet(doc, "关键货币/抽卡写路径：分布式锁或 DB 乐观锁 + 流水表，禁止无审计直写。")

    add_heading(doc, "4.5 Tick 频率与性能模型", 2)
    add_heading(doc, "差距", 3)
    add_bullet(doc, "GameServer 默认 period 约 1000ms，适合全局调度，不适合实时移动/战斗模拟。")

    add_heading(doc, "解决方案", 3)
    add_bullet(doc, "拆分时钟：GlobalTick（活动/过期清理）与 ZoneTick（10–20Hz）与 BattleTick（独立）。")
    add_bullet(doc, "Zone 按活跃度分级：无人休眠、稀疏低频、密集高频。")
    add_bullet(doc, "AOI 广播做兴趣裁剪与脏矩形/脏实体合并，避免全量刷屏。")

    add_heading(doc, "4.6 反作弊与服务器权威", 2)
    add_heading(doc, "差距", 3)
    add_bullet(doc, "缺少速度/瞬移校验、技能释放合法性、伤害区间审计、经济异常风控。")

    add_heading(doc, "解决方案", 3)
    add_bullet(doc, "移动：服务器积分位置，拒绝超速与穿墙（可用导航网格或简化碰撞）。")
    add_bullet(doc, "战斗：CD、消耗、目标合法性、伤害公式服务端重算；留存 battle_audit。")
    add_bullet(doc, "经济：Wallet/Gacha 全流程流水；运营后台可追溯。")
    add_bullet(doc, "PC 端可增加客户端完整性信号（非唯一依赖），仍以服务器权威为准。")

    add_heading(doc, "4.7 协作与匹配", 2)
    add_heading(doc, "差距", 3)
    add_bullet(doc, "存在匹配相关状态，但开放世界组队共战/副本房间权威尚未达到商业服水平。")

    add_heading(doc, "解决方案", 3)
    add_bullet(doc, "PartyService：组队、跟随、共享任务目标；同 Zone 优先同节点。")
    add_bullet(doc, "副本：MatchQueue → Room → 独立实例 Zone/Battle；结束后回原 Zone。")
    add_bullet(doc, "跨节点组队：以队长所在 Scene 为准或创建临时实例服。")

    add_heading(doc, "4.8 微服务策略（避免错误拆分）", 2)
    add_body(
        doc,
        "当前 microservices 使用 Spring Boot 3.3.x，主工程为 Boot 4.x，版本不一致。"
        "且 README 已明确：战斗/场景高一致性模块不宜先拆。",
    )
    add_table(
        doc,
        ["候选", "是否建议现在拆", "理由"],
        [
            ["Admin/配置导入/热更运维面", "是", "与 Netty tick 弱耦合，隔离故障"],
            ["AI assist 边车", "是（已具备）", "只读上下文、异步、可降级"],
            ["排行榜/邮件扇出", "可", "最终一致即可，适合 Redis"],
            ["Auth/账号", "稍后", "需与登录会话打通后再迁"],
            ["Player 聚合/钱包/抽卡", "否", "强一致写协议，拆了易双写"],
            ["Battle/Scene", "否（先做 Center）", "多节点路由未证实前拆会放大复杂度"],
        ],
    )
    add_bullet(doc, "统一依赖基线：脚手架与主工程 Spring Boot 主版本对齐，减少双栈维护成本。")

    # 五、目标架构
    add_heading(doc, "五、推荐目标架构（分产品线）", 1)

    add_heading(doc, "5.1 短期目标（崩铁式，3–6 个月可落地）", 2)
    add_code(
        doc,
        "客户端\n"
        "  ├─ TCP/KCP ──► Game Node（模块化单体：Scene/Zone/Battle/Player）\n"
        "  └─ 可选 HTTPS ──► API Gateway ──► AI Assist / Admin\n"
        "\n"
        "基础设施：MySQL（权威） + Redis（在线/票据/排行榜） + Center（可先内嵌再独立）\n"
        "运维：配置热更、发布演练、Prometheus 指标",
    )

    add_heading(doc, "5.2 中长期目标（若做无缝大世界）", 2)
    add_code(
        doc,
        "Login/Dispatch ──► Center（路由/容量/票据）\n"
        "                 ├─ SceneServer 舰队（Cell/Zone 分片 + AOI + 世界模拟）\n"
        "                 ├─ BattleServer（若实时战斗极重可旁路同机房）\n"
        "                 ├─ Hall/Social（好友/聊天/邮件）\n"
        "                 └─ Player/Economy（强一致写服务，谨慎拆分）\n"
        "Redis Cluster + MySQL + 对象存储（日志/回放）",
    )

    # 六、实施路线图
    add_heading(doc, "六、实施路线图与优先级", 1)
    add_table(
        doc,
        ["阶段", "时间建议", "交付物", "成功标准"],
        [
            [
                "Phase 0",
                "2–4 周",
                "产品定性（崩铁式 or 无缝）；补 Redis；迁移票据共享",
                "双节点切服可核销票据并重连",
            ],
            [
                "Phase 1",
                "1–2 月",
                "Zone 权威统一；AOI 压测；进战结算事务；反作弊最小集",
                "同 Zone 可见稳定；遇敌进战掉落一致",
            ],
            [
                "Phase 2",
                "2–3 月",
                "Center 心跳租约；组队/房间；经济流水；AI/Admin 边车固化",
                "节点宕机可摘除；组队副本可开",
            ],
            [
                "Phase 3",
                "按需",
                "仅当产品确认无缝：Cell 网格、边界移交、实时战斗权威、Scene 舰队",
                "跨边界连续移动与实时战斗延迟达标",
            ],
        ],
    )

    # 七、关键类索引
    add_heading(doc, "七、关键类索引（便于落地对照）", 1)
    add_table(
        doc,
        ["领域", "关键类"],
        [
            ["主循环", "GameServer, OnlinePlayer, PlayerTickRegistry"],
            ["网络", "GameNettyServer, GameKcpServer, LunarFrame*, GamePacketDispatcher"],
            ["会话", "GameSession, GameSessionManager, PlayerSessionService, PlayerSessionState"],
            [
                "场景",
                "ZoneContext, ZoneManager, AoiGrid, ZoneWorldService, WorldEncounterService, SceneSyncBroadcaster",
            ],
            [
                "中心",
                "CenterServer, LocalCenterServer, RemoteCenterServer, SceneRegistry, PlayerMigrationService, MigrationTicketService",
            ],
            ["战斗", "BattleManager, BattleContext, BattleNettyService, EntityState"],
            ["持久化", "PlayerData, PlayerDataRepository, PlayerDataPeriodicPersistenceService"],
            ["配置", "HotReloadCoordinator, ConfigImportService, ActivityImportService"],
            ["AI", "AiAssistApplicationService, AiAssistClient, AssistNettyService"],
        ],
    )

    # 八、总结
    add_heading(doc, "八、总结", 1)
    add_body(
        doc,
        "MyLunarCore 作为电脑端二游服务器，在协议层（Netty/KCP/Protobuf）、回合制战斗、运营配置热更、"
        "共享 Zone/AOI 与 AI 边车方面已贴近行业常见做法；"
        "但若以「原神/鸣潮式无缝大世界实时战斗」为标准，则在世界模型、模拟频率、多节点票据与注册、"
        "Redis 基础设施、反作弊权威等方面明显不符合主流。",
    )
    add_body(
        doc,
        "务实路径：先把产品定性为崩铁式并完成 Phase 0–2（Redis、跨节点票据、Zone 权威、反作弊最小集、Center 租约）；"
        "确认商业目标确需无缝大世界后再投入 Phase 3，避免在单体未稳时过早拆分 Battle/Scene 微服务。",
    )

    add_heading(doc, "附录：与旧评估文档的差异说明", 1)
    add_bullet(doc, "2026-07-12 旧版评估曾写「无共享 Zone / 无 CenterServer」，该结论已过时。")
    add_bullet(
        doc,
        "当前代码已存在 ZoneContext、AoiGrid、WorldEncounterService、center/* 与 PlayerMigrationService；"
        "主要缺口转为「生产级多节点与无缝/实时能力」，而非「完全没有场景多人雏形」。",
    )
    add_bullet(doc, "本文件为 2026-08-03 基于代码快照的更新版，请以代码为最终事实来源。")

    return doc


def main():
    doc = build_document()
    temp_output = DESKTOP / "_MyLunarCore_openworld_server_temp.docx"
    doc.save(str(temp_output))
    if OUTPUT.exists():
        OUTPUT.unlink()
    temp_output.rename(OUTPUT)
    print(f"已生成文档：{OUTPUT}")


if __name__ == "__main__":
    main()
