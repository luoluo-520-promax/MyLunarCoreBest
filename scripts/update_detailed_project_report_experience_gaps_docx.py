# -*- coding: utf-8 -*-
"""将近期「体验缺口」落地改动同步进桌面《MyLunarCore项目各方面详细总结报告.bak.docx》。

只原地修订该文件，不另存新文档、不生成第二份 .docx。
"""

from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DOC_PATH = Path.home() / "Desktop" / (
    "MyLunarCore"
    + "\u9879\u76ee\u5404\u65b9\u9762\u8be6\u7ec6\u603b\u7ed3\u62a5\u544a"
    + ".bak.docx"
)


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def replace_paragraph_text(paragraph, new_text):
    style = paragraph.style
    for r in list(paragraph.runs):
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        set_run_font(paragraph.runs[0])
    else:
        run = paragraph.add_run(new_text)
        set_run_font(run)
    paragraph.style = style


def insert_paragraph_after(paragraph, text, style_name="List Bullet"):
    new_para = paragraph._parent.add_paragraph(text, style=style_name)
    paragraph._p.addnext(new_para._p)
    for r in new_para.runs:
        set_run_font(
            r,
            size=13 if style_name.startswith("Heading") else 11,
            bold=style_name.startswith("Heading"),
        )
        if style_name.startswith("Heading"):
            r.font.color.rgb = RGBColor(25, 75, 140)
    return new_para


def find_para(doc, exact=None, startswith=None, contains=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if exact is not None and t == exact:
            return p
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    return None


def replace_if(doc, *, exact=None, startswith=None, contains=None, new_text=None):
    p = find_para(doc, exact=exact, startswith=startswith, contains=contains)
    if p is not None and new_text is not None:
        replace_paragraph_text(p, new_text)
        return True
    return False


def set_cell_text(cell, text):
    cell.text = text
    for para in cell.paragraphs:
        for r in para.runs:
            set_run_font(r, size=10)


def patch_table_cell(table, old_exact=None, old_contains=None, col=0, new_text=None, new_other=None):
    """按第一列（或指定列）匹配后改写单元格。"""
    for row in table.rows:
        cur = row.cells[col].text.strip()
        hit = False
        if old_exact is not None and cur == old_exact:
            hit = True
        if old_contains is not None and old_contains in cur:
            hit = True
        if not hit:
            continue
        if new_text is not None:
            set_cell_text(row.cells[col], new_text)
        if new_other is not None:
            for ci, val in new_other.items():
                set_cell_text(row.cells[ci], val)
        return True
    return False


def append_table_row(table, values):
    row = table.add_row()
    for i, val in enumerate(values):
        if i < len(row.cells):
            set_cell_text(row.cells[i], val)
    return row


def main():
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    doc = Document(str(DOC_PATH))
    already = find_para(doc, contains="ScenePreloadService") is not None
    if already:
        print(f"Already patched (ScenePreloadService present): {DOC_PATH}")
        return

    today = f"{date.today().year}年{date.today().month:02d}月{date.today().day:02d}日"
    n = 0

    # ---------- 封面 ----------
    if replace_if(
        doc,
        exact="（含技术架构、业务能力、安全运维、微服务迁移与后续建议）",
        new_text="（含技术架构、业务能力、安全运维、微服务迁移、体验缺口补齐与后续建议）",
    ):
        n += 1

    if replace_if(
        doc,
        startswith="报告生成日期：",
        new_text=(
            f"报告生成日期：{today}"
            "（已同步体验缺口补齐：场景预加载、站内 WebView 攻略、扫荡体力倍率、"
            "家园家具交互/在场、战斗打击感 FX、新手引导检查点与跳过）"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="报告依据：当前仓库 README",
        new_text=(
            "报告依据：当前仓库 README、源码包结构、配置文件与已有评估文档快照；"
            "本轮已把服务端体验缺口（预加载握手、站内攻略、扫荡、家园交互、战斗 FX、引导检查点）写入对应章节。"
        ),
    ):
        n += 1

    # ---------- 7.3 场景：预加载 / 无缝切图 ----------
    if replace_if(
        doc,
        startswith="玩家进入某个 Plane/Floor 对应的 Zone 实例后",
        new_text=(
            "玩家进入某个 Plane/Floor 对应的 Zone 实例后，可移动、与 NPC/道具交互，"
            "服务器按兴趣范围（AOI）同步周围实体。Zone 可配置人数上限与租约心跳。"
            "接近传送门时可走 ScenePreload：客户端先报目标 Plane/Floor，服务端返回资源键并在就绪后推送"
            "ScenePreloadReadyScNotify；真正切图时 MigrateSceneScRsp 带 preload_hit 与 SceneLoadMaskInfo，"
            "命中预加载则 LoadingTicket 仅作短 TTL 握手（handshake_only），缩短黑屏。"
            "跨图或跨节点时，由 Center 规划归属，必要时发放迁移票据。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        exact="入口：SceneNettyService、ZoneManager、ZoneTickService、PlayerMigrationService",
        new_text=(
            "入口：SceneNettyService、ScenePreloadService、ZoneManager、"
            "ZoneTickService、PlayerMigrationService"
        ),
    ):
        n += 1

    p_exp = find_para(doc, startswith="实验：CellBoundaryHandoffService")
    if p_exp is not None and "ScenePreload" not in p_exp.text:
        insert_paragraph_after(
            p_exp,
            "预加载协议：ScenePreloadCsReq/ScRsp、ScenePreloadReadyScNotify（CmdId 354–356）；"
            "加载完成 SceneLoadCompleteScRsp.handshake_only；黑屏遮罩 SceneLoadMaskInfo。",
        )
        n += 1

    # ---------- 7.4 战斗 FX ----------
    if replace_if(
        doc,
        startswith="场景中触发遭遇后进入战斗实例",
        new_text=(
            "场景中触发遭遇后进入战斗实例：开战、行动、结算、退出。"
            "支持遇敌配置与波次；多人共战通过 BattleContext.participantPlayerIds 等字段表达。"
            "另有启发式战斗辅助策略（偏提示，不是外挂代打）。"
            "行动经 BattleDeterministicValidator 校验后，BattleFxComposer 按暴击/弱点/击杀组装打击感元数据"
            "（camera_shake、time_scale、damage_popup），经 BattleFxScNotify（CmdId 213）推给客户端播特效，"
            "不改变结算权威。"
        ),
    ):
        n += 1

    p_battle_cfg = find_para(doc, startswith="配置：data/EncounterConfigs.json")
    if p_battle_cfg is not None and "BattleFxComposer" not in p_battle_cfg.text:
        insert_paragraph_after(
            p_battle_cfg,
            "打击感：BattleFxComposer → BattleFxScNotify；协议 battle_system.proto。",
        )
        n += 1

    # ---------- 7.5 扫荡 ----------
    if replace_if(
        doc,
        startswith="挑战玩法负责开局、结果上报、组奖励与历史",
        new_text=(
            "挑战玩法负责开局、结果上报、组奖励与历史；模拟宇宙（Rogue）则包含地图生成、移动、"
            "祝福/奇物、战斗上报与天赋等完整 Roguelike 流程。"
            "已通关关卡可走 SweepService 扫荡：以历史最低回合数解锁，跳过完整 BattleManager，"
            "按 1/2/3 倍体力消耗发放倍率掉落（单道具上限 999），适合重复刷材料。"
        ),
    ):
        n += 1

    p_ch = find_para(doc, startswith="挑战：ChallengeNettyService")
    if p_ch is not None and "SweepService" not in p_ch.text:
        insert_paragraph_after(
            p_ch,
            "扫荡：SweepService / SweepNettyService；表 player_stage_clear_best；"
            "CmdId 1010–1013；胜利结算会 recordClear。SQL：db/migration_experience_gaps.sql。",
        )
        n += 1

    # ---------- 7.14 新手引导检查点 ----------
    if replace_if(
        doc,
        startswith="成就系统记录玩家达成条件与领取状态",
        new_text=(
            "成就系统记录玩家达成条件与领取状态，适合做长期留存目标。"
            "新手引导则按配置步骤引导新玩家完成关键操作，并落库检查点（checkpoint_step_id / committed_json）："
            "断线重连从最近已提交步骤恢复；支持一键跳过（skipped=true）。"
            "两者都属于「内容运营型」能力：配置改了，行为就能跟着变。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        exact="新手引导：NewbieGuideService；配置 data/NewbieGuide.json；协议号段约 960–964",
        new_text=(
            "新手引导：NewbieGuideService（检查点 + 跳过 + 内存兜底）；"
            "配置 data/NewbieGuide.json；协议号段 960–966（含 SkipNewbieGuide 965–966）"
        ),
    ):
        n += 1

    # ---------- 7.16 家园在场 / 家具交互 ----------
    if replace_if(
        doc,
        startswith="除了战斗与经济循环，项目还补齐了「讲故事」和「家园养成」能力",
        new_text=(
            "除了战斗与经济循环，项目还补齐了「讲故事」和「家园养成」能力，方便策划做剧情与轻养成。"
            "剧情侧会记录分支与已播放状态；过场支持跳过/完成并与玩家设置联动；"
            "家园除产出、家具摆放与好友互访外，已接线家具交互与 AOI-lite 在场同步："
            "HomePresenceService 广播家园内角色位置，FurnitureInteractHandler 处理点击家具（对话/领取等）。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="家园：HomeBaseService（产出/家具/互访，DB home_state）",
        new_text=(
            "家园：HomeBaseService（产出/家具/互访，DB home_state）+ HomePresenceService、"
            "FurnitureInteractHandler；配置 data/HomeFacilityConfigs.json；"
            "协议：摆放 858–859，交互 870–871，在场同步 Notify 872、上报 873–874。"
        ),
    ):
        n += 1

    # ---------- 7.19 体力扫荡倍率 ----------
    if replace_if(
        doc,
        startswith="协议：DailyLoopPacketHandlers + daily_loop_system.proto 已接线",
        new_text=(
            "协议：DailyLoopPacketHandlers + daily_loop_system.proto 已接线"
            "（体力 996–1000、日常 1001–1005、扫荡 1010–1013）；"
            "扫荡按 1/2/3 倍消耗体力并发放倍率奖励；客户端联调可继续加深表现层。"
        ),
    ):
        n += 1

    # ---------- 7.20 切图加载态 ----------
    if replace_if(
        doc,
        startswith="好友助战允许借用好友角色快照加入战斗",
        new_text=(
            "好友助战允许借用好友角色快照加入战斗（写入 BattleContext.participantPlayerIds）；"
            "公会科技按公会等级提供全局 Buff（攻/防/血/体力恢复等）；"
            "切图时服务器会发放 LoadingTicket，加载完成前拒绝移动与开战，减少「半截进图」导致的状态错乱。"
            "若预加载已命中，完成包 handshake_only=true、TTL 更短，遮罩信息随 MigrateScene 下发。"
            "登录阶段还会交换 SupportedFeatures 能力位掩码，让客户端按服务端能力隐藏未开放入口"
            "（例如公会战入口在能力未开时返回 retcode=20）。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        exact="切图加载：PlayerLoadingStateService；CmdId 352–353（加载完成握手）。",
        new_text=(
            "切图加载：PlayerLoadingStateService（含 handshakeOnly 短 TTL）+ ScenePreloadService；"
            "CmdId 352–353（加载完成）与 354–356（预加载/就绪通知）。"
        ),
    ):
        n += 1

    # ---------- 7.22 主路径故事 ----------
    if replace_if(
        doc,
        startswith="为了让非技术同学建立整体画面",
        new_text=(
            "为了让非技术同学建立整体画面，下面用一条「普通玩家的一天」串起系统："
            "登录进服（校验协议与能力位）→ 看邮件/好友 → 消耗体力进场景走动（接近传送门可预加载、切图遮罩更短）、"
            "触发 NPC 对话或过场 → 遇敌打回合战（可带助战，客户端按 BattleFx 播打击感）→"
            "开挑战、Rogue、世界 BOSS 或深渊；已通关材料关可 1/2/3 倍体力扫荡 →"
            "回家园点点基建、摆家具并与家具交互、看到他人在场 → 抽卡与商店消费 →"
            "做每日任务/战令/版本活动拿奖 → 推主线章节解锁新地图 →"
            "进公会贡献/兑换/科技/公会战 → 看排行榜聊天 →"
            "必要时问 AI 助手（站内 WebView 打开 B站/抖音/官网攻略，正文另有 ≤200 字摘要，默认不跳系统浏览器）→"
            "新手引导按检查点续跑或一键跳过 → 改设置或提工单。"
            "运营同学则在后台改活动排期、导入配置、触发热更；运维同学盯监控与发布演练。"
        ),
    ):
        n += 1

    # ---------- 第十章 AI：站内 WebView ----------
    if replace_if(
        doc,
        startswith="主流平台外链/视频攻略（B站、抖音、官方网站）",
        new_text=(
            "主流平台外链/视频攻略（B站、抖音、官方网站）：玩家问「攻略/打法/教程/视频/官网」时，"
            "服务端从 data/ExternalGuideCatalog.json 先匹配运营精选条目（curated），不足则按三平台搜索模板生成深链；"
            "回答正文附链接摘要，协议 AskAiAssistScRsp / AiHintNotify 下发 media_links。"
            "默认 action 为 OPEN_IN_APP_WEBVIEW / OPEN_VIDEO_INLINE（render_mode=INLINE_WEBVIEW），"
            "并带 AssistInlineSummary（≤200 字、去掉裸 URL）与 forbid_external_browser=true，"
            "避免跳出游戏进系统浏览器。仅允许 https 白名单域名（bilibili / douyin / mihoyo 等），不做第三方页面抓取。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="这是本轮新增的「建议面」能力",
        new_text=(
            "这是「建议面」能力：助手不仅给文字攻略，还能把 B站、抖音、官方网站上的视频或页面交给玩家。"
            "设计原则是「白名单深链 + 站内打开」——服务端只返回可打开的 URL，内容仍在原平台播放/阅读，"
            "客户端用游戏内 WebView/内嵌播放器打开，不跳系统浏览器。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="玩家侧入口与现有助手相同",
        new_text=(
            "玩家侧入口与现有助手相同：聊天 @助手、AskAiAssist（Cmd 902）、异步 AiHintNotify。"
            "问句含「攻略/打法/通关/教程/视频/官网/B站/抖音」等意图词时触发。流程为："
            "① 安全分类与本地 FAQ/规则快路径照常执行；② ExternalGuideSearchService 读取热更目录；"
            "③ 按关键词/场景给精选条目打分，命中则返回 curated=true 的视频或文章链接；"
            "④ 精选不足时按 platforms 模板补齐 B站搜索、抖音搜索、官网首页；"
            "⑤ 域名必须在 allowedHosts 内且为 http(s)，否则丢弃；"
            "⑥ AssistInlineSummaryService 生成 ≤200 字精简摘要（剥离 URL），cited_config_ids 增加 ext-guide:{id}；"
            "⑦ 协议下发 media_links.action=OPEN_IN_APP_WEBVIEW（视频为 OPEN_VIDEO_INLINE）与"
            "forbid_external_browser；客户端在站内浮层打开，不唤起外部浏览器。"
        ),
    ):
        n += 1

    p_proto = find_para(doc, startswith="协议扩展：AskAiAssist 支持 session_id/locale")
    if p_proto is not None and "inline_summary" not in p_proto.text:
        replace_paragraph_text(
            p_proto,
            p_proto.text.strip()
            + " 另下发 inline_summary（text/max_chars/source）与 forbid_external_browser；"
            "media_links.action 默认 OPEN_IN_APP_WEBVIEW / OPEN_VIDEO_INLINE。",
        )
        n += 1

    # ---------- 15.1 已实现 ----------
    p_daily = find_para(doc, startswith="日常循环：体力、每日任务、战令")
    if p_daily is not None and "扫荡" not in p_daily.text:
        replace_paragraph_text(
            p_daily,
            p_daily.text.strip().rstrip("。")
            + "；已通关关卡扫荡（1/2/3 倍体力，CmdId 1010–1013）。",
        )
        n += 1

    p_load = find_para(doc, contains="切图 LoadingTicket 防半截进图")
    if p_load is not None and "ScenePreload" not in p_load.text:
        replace_paragraph_text(
            p_load,
            p_load.text.strip().rstrip("。")
            + "；ScenePreload 预加载命中后 handshake_only 短握手 + 切图遮罩。",
        )
        n += 1

    p_guide = find_para(doc, startswith="组队协议与多人共战/助战；成就、新手引导")
    if p_guide is not None and "检查点" not in p_guide.text:
        replace_paragraph_text(
            p_guide,
            "组队协议与多人共战/助战；成就、新手引导（检查点续跑 + 一键跳过 CmdId 960–966）、"
            "公会基础 + 公会战 + 公会科技；家园家具交互与在场同步（CmdId 858–859、870–874）。",
        )
        n += 1

    p_ai = find_para(doc, contains="AI 外链攻略（B站/抖音/官网 media_links）")
    if p_ai is not None and "WebView" not in p_ai.text:
        replace_paragraph_text(
            p_ai,
            p_ai.text.strip().rstrip("。")
            + "；默认站内 WebView/内嵌视频 + ≤200 字 inline_summary，禁止跳系统浏览器。"
            "战斗行动后推送 BattleFxScNotify（CmdId 213）。",
        )
        n += 1

    p_gap = find_para(doc, startswith="缺口补齐玩法：世界 BOSS")
    if p_gap is not None and find_para(doc, contains="ExperienceGapBusinessFlowsTest") is None:
        insert_paragraph_after(
            p_gap,
            "体验缺口回归：ExperienceGapBusinessFlowsTest 覆盖预加载握手、站内攻略摘要、扫荡倍率、"
            "家园交互/在场、战斗 FX、引导检查点；库表增量 db/migration_experience_gaps.sql。",
        )
        n += 1

    # ---------- 15.2 局限：家园已加深接线 ----------
    if replace_if(
        doc,
        exact="对话/过场/家园及缺口补齐玩法：服务端已加深；客户端表现层可继续对齐。",
        new_text=(
            "对话/过场/家园：服务端已加深（含家具交互与在场同步）；"
            "预加载遮罩、站内 WebView、扫荡入口、战斗 FX 播片仍需客户端表现层对齐。"
        ),
    ):
        n += 1

    # ---------- 16.1 短期建议 ----------
    p_short = find_para(doc, exact="继续以模块化单体打磨崩铁式体验与数值内容。")
    if p_short is not None and "站内 WebView" not in p_short.text:
        replace_paragraph_text(
            p_short,
            "继续以模块化单体打磨崩铁式体验与数值内容；"
            "客户端对齐预加载遮罩、站内攻略 WebView、扫荡入口、家园在场与战斗 FX 播片。",
        )
        n += 1

    # ---------- 结语 ----------
    if replace_if(
        doc,
        startswith="MyLunarCore 已经不是「只有空壳的演示仓库」",
        new_text=(
            "MyLunarCore 已经不是「只有空壳的演示仓库」：它在单体形态下具备较完整的崩铁式 RPG 服务端能力——"
            "联网、场景（含预加载握手与切图遮罩）、回合战（含打击感 FX 推送）、养成、经济抽卡、"
            "体力与每日循环（含 1/2/3 倍扫荡）、战令与主线章节、活动任务、对话过场与家园（含家具交互/在场）、"
            "公会/公会战/公会科技、竞技场评分、世界 BOSS/深渊/遗器词条等缺口补齐玩法、"
            "成就与带检查点/可跳过的新手引导、运维热更与回滚、业务监控，"
            "以及可选的多节点与能力增强后的 AI 助手旁路"
            "（多轮对话、个性化、主动推送、熔断灰度、合规声明、专项指标，"
            "以及主流平台攻略默认在站内 WebView 打开并附 ≤200 字摘要）。"
            "与此同时，它也很诚实地把微服务目录标成脚手架，把无缝大世界标成实验，把默认密钥标成仅限本地。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="本报告基于仓库当前代码与文档快照自动整理生成",
        new_text=(
            f"本报告基于仓库当前代码与文档快照整理，{today} 已在本文档内同步体验缺口补齐，"
            "可用于汇报、培训与决策讨论。若后续架构或配置有重大变更，建议修订对应章节后刷新目录域，"
            "以便目录页码与正文一致。"
        ),
    ):
        n += 1

    # ---------- 包一览表（表 3）----------
    t3 = doc.tables[3]
    n += patch_table_cell(
        t3,
        old_contains="回合战、遭遇、波次",
        col=2,
        new_text="回合战、遭遇、波次、战斗辅助策略、打击感 FX（BattleFxScNotify）",
    )
    n += patch_table_cell(
        t3,
        old_contains="进图、移动、AOI 同步",
        col=2,
        new_text="进图、移动、AOI 同步、Zone 管理、场景预加载握手与切图遮罩",
    )
    n += patch_table_cell(
        t3,
        old_contains="挑战开局、结算、奖励",
        col=2,
        new_text="挑战开局、结算、奖励、已通关扫荡（体力倍率）",
    )
    n += patch_table_cell(
        t3,
        old_contains="media_links",
        col=2,
        new_text=(
            "规则教练、多轮历史、画像、RAG、语义安全、远程 AI、反馈闭环；"
            "外链攻略默认站内 WebView + ≤200 字摘要（禁跳系统浏览器）"
        ),
    )
    n += patch_table_cell(
        t3,
        old_contains="引导步骤、NewbieGuide",
        col=2,
        new_text="引导步骤、检查点续跑、一键跳过、NewbieGuide 配置",
    )
    n += patch_table_cell(
        t3,
        old_contains="基建、产出、家具",
        col=2,
        new_text="基建、产出、家具摆放/交互、好友互访、家园在场同步",
    )

    # ---------- CmdId 号段表（表 5）----------
    t5 = doc.tables[5]
    n += patch_table_cell(t5, old_exact="200–212", new_text="200–213")
    n += patch_table_cell(
        t5, old_exact="352–353", new_text="352–356",
        new_other={1: "切图加载完成握手 + 场景预加载/就绪通知"},
    )
    n += patch_table_cell(
        t5, old_exact="850–857", new_text="850–859 / 870–874",
        new_other={1: "家园（含家具摆放/交互与在场同步）"},
    )
    n += patch_table_cell(
        t5, old_exact="960–964", new_text="960–966",
        new_other={1: "新手引导（含检查点跳过）"},
    )
    n += patch_table_cell(
        t5,
        old_contains="media_links 外链攻略",
        col=1,
        new_text="AI 助手（含站内 WebView 攻略 / inline_summary）",
    )
    # 扫荡号段：原表止于 1006–1008
    last_sys = t5.rows[-1].cells[0].text.strip()
    if "1010" not in last_sys and not any("1010" in r.cells[0].text for r in t5.rows):
        append_table_row(t5, ["1010–1013", "关卡扫荡（体力倍率）"])
        n += 1

    # ---------- AI 能力表（表 7）----------
    t7 = doc.tables[7]
    n += patch_table_cell(
        t7,
        old_contains="B站/抖音/官网白名单链接",
        col=2,
        new_text=(
            "B站/抖音/官网白名单；默认 OPEN_IN_APP_WEBVIEW / OPEN_VIDEO_INLINE，"
            "附 ≤200 字摘要，forbid_external_browser"
        ),
    )
    if not any("AssistInlineSummaryService" in r.cells[1].text for r in t7.rows):
        append_table_row(
            t7,
            [
                "站内打开 / 精简摘要",
                "AssistInlineSummaryService",
                "攻略默认游戏内 WebView；正文 ≤200 字，不跳系统浏览器",
            ],
        )
        n += 1

    doc.save(str(DOC_PATH))
    print(f"Updated in place: {DOC_PATH}")
    print(f"Patches applied (approx): {n}")


if __name__ == "__main__":
    main()
