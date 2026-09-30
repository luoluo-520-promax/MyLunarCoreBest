# -*- coding: utf-8 -*-
"""将 P10 运营硬化改动同步进桌面《MyMmorpg项目全方位详细介绍文档》。"""

from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DOC_PATH = Path.home() / "Desktop" / (
    "MyMmorpg" + "\u9879\u76ee\u5168\u65b9\u4f4d\u8be6\u7ec6\u4ecb\u7ecd\u6587\u6863" + ".docx"
)


def set_run_font(run, size=11, bold=False):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold


def replace_paragraph_text(paragraph, new_text):
    """清空段落 runs 后写入新文本，尽量保留段落样式。"""
    style = paragraph.style
    for r in list(paragraph.runs):
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        set_run_font(paragraph.runs[0])
    else:
        run = paragraph.add_run(new_text)
        set_run_font(run)
    paragraph.style = style


def insert_paragraph_after(paragraph, text, style_name="List Bullet"):
    """在指定段落后插入新段落。"""
    new_para = paragraph._parent.add_paragraph(text, style=style_name)
    paragraph._p.addnext(new_para._p)
    for r in new_para.runs:
        set_run_font(r)
    return new_para


def find_para(doc, exact=None, startswith=None, contains=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if exact is not None and t == exact:
            return p
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    return None


def find_para_index(doc, exact=None, contains=None):
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if exact is not None and t == exact:
            return i
        if contains is not None and contains in t:
            return i
    return -1


def main():
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    doc = Document(str(DOC_PATH))

    # ---------- 3.1 登录：资源版本封锁 ----------
    p = find_para(doc, exact="登录鉴权与会话；重复登录踢下线（KickPlayer），避免一人分身乱数据。")
    if p:
        replace_paragraph_text(
            p,
            "登录鉴权与会话；重复登录踢下线（KickPlayer），避免一人分身乱数据；"
            "登录握手校验 wire_version，并检查 client_res_version（data/ClientRequiredVersions.json），"
            "资源包过旧返回 retcode=10（ERR_CLIENT_TOO_OLD）并下发更新 URL。",
        )

    # ---------- 3.3 战斗 / 抽卡 ----------
    p = find_para(doc, exact="战斗开始 / 行动 / 结算；挑战关卡联动开战。")
    if p:
        replace_paragraph_text(
            p,
            "战斗开始 / 行动 / 结算；挑战关卡联动开战；"
            "BattleReplayService 记录 ActionTick，多人共战掉线重连可下发 ReplayPacket 快进到最新战局"
            "（不仅是最终快照）。",
        )

    p = find_para(doc, exact="抽卡：保底/十连/天井/历史；概率公示与审计摘要。")
    if p:
        replace_paragraph_text(
            p,
            "抽卡：保底/十连/天井/历史；表现层三次握手 GachaStart→DoGacha→GachaResultAck（Cmd 509–512）；"
            "配置可带 client_ui_params（特效/镜头/文案）；抽卡返利积分（GachaRebateService 星尘）。",
        )

    # ---------- 3.4 社交 ----------
    p = find_para(doc, contains="好友助战（冷却/报酬/统计）")
    if p:
        replace_paragraph_text(
            p,
            "事件临时小队；好友助战（日限：同一好友角色每天 1 次，凌晨 4 点刷新；冷却/报酬/统计）；"
            "FriendGiftService 好友体力每日赠送/领取（Redis Set 防重复领取）；家园访问与装饰同步。",
        )

    # ---------- 3.5 商业化 / 活动 ----------
    p = find_para(doc, exact="七日签到 / 月卡日领。")
    if p:
        replace_paragraph_text(
            p,
            "七日签到 / 月卡日领；MonthlyCardService 持久化月卡/小月卡状态，"
            "日切经 PeriodicResetService 发放额外货币与体力。",
        )

    p = find_para(doc, startswith="战令 / 月卡 / 首充双倍")
    if p:
        replace_paragraph_text(
            p,
            "战令 / 月卡 / 首充双倍：TopUpBonusService 记录账号首充与年度重置周期，"
            "支付回调叠加额外赠送比例；与 IAP 发货链路联动。",
        )

    p = find_para(doc, startswith="活动：条件、阶段、商店货品")
    if p:
        replace_paragraph_text(
            p,
            "活动：条件、阶段、商店货品、展示文案、导入校验；可订阅战斗结束与订单支付；"
            "ActivityResourceService 管理 CDN 资源版本，ActivityScheduleService 管开关；"
            "支持运维预加载（POST /api/admin/ops/preload）与到点原子 activate。",
        )

    p = find_para(doc, exact="热更差分、资源补丁链、CDN 预热。")
    if p:
        replace_paragraph_text(
            p,
            "热更差分、资源补丁链、CDN 预热；登录强制最低客户端资源包版本"
            "（ClientRequiredVersions.json），避免「问号图标」进服。",
        )

    # ---------- 5.3 战斗 ----------
    p = find_para(doc, startswith="战斗负责开战、行动、结算")
    if p:
        replace_paragraph_text(
            p,
            "战斗负责开战、行动、结算，并发布 BattleEnded 等事件供活动/任务订阅。"
            "状态常在 Redis；活跃战斗用 SET 索引，避免 KEYS 全库扫描。"
            "元素反应、客户端预测与 rollback、Boss BT 调试接口与 ThreatTable 仇恨表均已落地。"
            "断线重连：BattleSnapshotService 提供快照；BattleReplayService 额外记录 ActionTick 行动序列，"
            "重连时将掉线期间行动压缩为 ReplayPacket 供客户端快进到当前战局。"
            "红线：实时战斗不引入 LLM。",
        )

    # ---------- 5.4 钱包 ----------
    p = find_para(doc, startswith="发奖强调幂等")
    if p:
        replace_paragraph_text(
            p,
            "发奖强调幂等（同一凭证只成功一次）。邮件附件经 MailItemGrantPort 进背包。"
            "玩家与背包可冷热分离：Redis 热数据 + 定时刷 MySQL。"
            "钱包：Redis 分布式锁 + player.data_version 乐观锁 CAS（余额不足/版本冲突直接拒绝并重试），"
            "流水写入 wallet_ledger；高价值操作建议 UID Sticky Session 固定节点。"
            "体力 /internal/bag/resin*；装备随机词条 EquipRandomizer。",
        )

    # ---------- 5.5 商城 ----------
    p = find_para(doc, startswith="P0 经济硬化")
    if p:
        replace_paragraph_text(
            p,
            "P0 经济硬化：订单双写 shop_order、支付事件 mq_outbox、发奖/累充幂等（grant_idempotency / mq_inbox）、"
            "流水账本（wallet_ledger / item_ledger）、对账 POST /internal/shop/orders/reconcile。"
            "开发可用 MOCK；生产禁用 MOCK 并配置渠道密钥。"
            "战令/月卡/首充：MonthlyCardService（激活/剩余天数 + 日切发放）、"
            "TopUpBonusService（首充双倍与年度重置）、IapEntitlementService 权益标记。",
        )

    # ---------- 5.6 活动 ----------
    p = find_para(doc, startswith="活动配置模型较完整")
    if p:
        replace_paragraph_text(
            p,
            "活动配置模型较完整（条件、阶段、商店货品、展示、导入校验）。"
            "可消费商店已支付与战斗结束消息。SignInService 提供签到/月卡日领。"
            "世界事件提供排期、伤害与 grantPlans 结算钩子。"
            "资源与开关分离：ActivityResourceService（CDN bundle 版本/预加载缓存）+ "
            "ActivityScheduleService（排期开关）；运维可 preload 下版配置到 Redis，"
            "活动开始时刻由 ClusterJobLock 原子 activate。",
        )

    # ---------- 5.7 抽卡 ----------
    p = find_para(doc, startswith="独立 gacha-service")
    if p:
        replace_paragraph_text(
            p,
            "抽卡（单体 GachaApplicationService / GachaNettyService，亦可独立服务）："
            "保底/十连/天井/历史；表现层状态机协议 GachaStart→DoGacha→GachaResultAck（Cmd 509–512），"
            "下发 presentation_session_id 与 ClientUiParams；"
            "GachaRebateService 抽卡返利积分（星尘）可兑换限定道具；"
            "配置 data/Banners.json 支持 client_ui_params；清单见 docs/presentation-protocol-checklist.md。",
        )

    # ---------- 5.9 后台 ----------
    p = find_para(doc, startswith="Admin：配置导入")
    if p:
        replace_paragraph_text(
            p,
            "Admin：配置导入、投诉、操作日志、AI 草稿、分阶段 reload、发布审计、"
            "配置 Diff/回滚/灰度、运营控制台。"
            "新增客服时间线 GET /api/admin/ops/timeline/{uid}（聚合 wallet_ledger + 抽卡历史）；"
            "活动预加载 POST /api/admin/ops/preload 与手动 activate。"
            "RBAC 预设：OPS / PLANNER / CS / SUPERADMIN。鉴权：IP 白名单 + HMAC/API Key；生产强制开启。",
        )

    # ---------- 5.10 粘性 ----------
    p = find_para(doc, startswith="Gateway 可按 X-Account-Id")
    if p:
        replace_paragraph_text(
            p,
            "Gateway 可按 X-User-Id / X-Account-Id 粘滞到同一 player 实例（UserStickyLoadBalancer）；"
            "游戏侧 StickySessionAffinity 按 UID 哈希固定节点，降低抽卡/购买跨节点双扣窗口。"
            "GeoIP/区域头注入就近信息；FunctionNumberRoutingFilter 按功能号段路由，削弱中转瓶颈。",
        )

    # ---------- 6.1 / 6.4 账本与时间线 ----------
    p = find_para(doc, exact="MySQL 像保险箱账本；Redis 像前台便签与计分板。重要结果落库，高频读写用 Redis。")
    if p:
        replace_paragraph_text(
            p,
            "MySQL 像保险箱账本；Redis 像前台便签与计分板。重要结果落库，高频读写用 Redis。"
            "货币变动统一落 wallet_ledger（uid + created_at 索引），客服可通过 Admin 时间线一键回溯"
            "「星琼去哪了」。",
        )

    p = find_para(doc, startswith="TLogEventPublisher 默认结构化日志")
    if p:
        replace_paragraph_text(
            p,
            "TLogEventPublisher / AnalyticsEventPublisher 默认结构化日志；可选 Kafka。"
            "LogDesensitizer 脱敏；LogSamplingFilter 采样；TraceContext 跨线程 TraceId。"
            "轻量事件溯源：PlayerTimelineService 按秒排序输出收入/支出明细树，支持客服导出沟通。",
        )

    # ---------- 7.3 会话：资源版本 ----------
    p = find_para(doc, startswith="会话 RSA 生产必须固定密钥")
    if p:
        replace_paragraph_text(
            p,
            "会话 RSA 生产必须固定密钥；开发未配置可临时生成。可配置服务/网关 SSL。"
            "登录协议 player_session.proto 增加 client_res_version / required_client_res_version / "
            "client_update_url；与 PROTOCOL_WIRE_VERSION 一起构成「协议 + 资源包」双门禁。",
        )

    # ---------- 10.1 已实现 ----------
    p = find_para(doc, exact="支付对账骨架、反作弊、Feign 熔断、冷热分离、TLog、CI 与质量门禁。")
    if p:
        replace_paragraph_text(
            p,
            "支付对账骨架、反作弊、Feign 熔断、冷热分离、TLog、CI 与质量门禁；"
            "P10：表现层状态机协议、月卡/首充双倍/抽卡返利、钱包乐观锁+Sticky、"
            "活动预下载、助战日限与好友体力、玩家时间线、客户端资源版本封锁、战斗 ActionTick 回放。",
        )

    # ---------- 10.3 计划中：收窄已完成项 ----------
    p = find_para(doc, startswith="完整 RedLock、官方支付 SDK")
    if p:
        replace_paragraph_text(
            p,
            "官方支付 SDK 替换 stub、KCP 更深对齐；钱包侧已具备 Redis 锁 + data_version 乐观锁"
            "（非完整 RedLock 产品化封装）。",
        )

    # ---------- 11.1 登录故事 ----------
    p = find_para(doc, exact="登录拿到会话；玩家服记在线并规划场景节点/分线。")
    if p:
        replace_paragraph_text(
            p,
            "登录校验 wire_version 与 client_res_version；过旧则强制更新。"
            "通过后拿到会话；玩家服记在线并规划场景节点/分线。",
        )

    # ---------- 11.3 充值故事 ----------
    p = find_para(doc, exact="发货与累充幂等；活动充值进度更新。")
    if p:
        replace_paragraph_text(
            p,
            "发货与累充幂等；TopUpBonusService 判断首充/年度重置并叠加赠送；活动充值进度更新。",
        )

    # ---------- 11.4 活动上线故事 ----------
    p = find_para(doc, exact="分阶段 reload；发布审计留痕。")
    if p:
        replace_paragraph_text(
            p,
            "可提前 preload 下版 CDN/配置到 Redis；分阶段 reload；发布审计留痕；"
            "到点 ClusterJobLock 原子 activate。",
        )

    # ---------- 11.5 抽卡故事 ----------
    p = find_para(doc, exact="客户端发 14xx 抽卡命令（或 HTTP 内部接口联调）。")
    if p:
        replace_paragraph_text(
            p,
            "客户端先发 GachaStart（拿 presentation_session + client_ui），再 DoGacha 扣费发奖，"
            "动画结束后 GachaResultAck 关闭会话（或 HTTP 内部接口联调）。",
        )
    p = find_para(doc, exact="抽卡服计算保底/天井，写历史。")
    if p:
        replace_paragraph_text(
            p,
            "抽卡服计算保底/天井，写历史，并发放抽卡返利积分（星尘）。",
        )

    # ---------- 第十八章总结：追加一句 ----------
    p = find_para(doc, startswith="请记住三句话")
    if p:
        replace_paragraph_text(
            p,
            "请记住三句话：第一，客户端负责表现，服务器负责裁决与记账——表现层协议要用"
            "「开始→进行中→结束」状态机对齐动画；第二，可以先合署办公（单体），再分局扩容（微服务），"
            "高价值操作尽量 Sticky 到固定节点；第三，发奖与支付必须以幂等、验签、对账与流水时间线为底线——"
            "这是信任的基础。",
        )

    # ---------- 附录 C ----------
    p = find_para(doc, startswith="生成日期：")
    if p:
        replace_paragraph_text(p, f"生成日期：{date.today().year}年{date.today().month:02d}月{date.today().day:02d}日（已同步 P10 运营硬化）")

    # ---------- 新增小节：在 5.11 后插入 5.12 ----------
    p511 = find_para(doc, exact="5.11 反作弊与客户端完整性")
    # Find the body after 5.11
    body_511 = find_para(doc, startswith="AntiCheatService + AntiCheatRuleEngine")
    if body_511 is not None:
        # Insert heading + body after body_511
        h = insert_paragraph_after(
            body_511,
            "5.12 P10 运营硬化（表现 / 商业化 / 一致性 / 运维）",
            style_name="Heading 2",
        )
        # Heading style font color
        for r in h.runs:
            set_run_font(r, size=13, bold=True)
            r.font.color.rgb = RGBColor(25, 75, 140)

        # insert bodies after heading — each insert_after previous
        cur = h
        blocks = [
            (
                "Normal",
                "本轮（migration_p10_ops_hardening.sql + docs/feature-gap-fill.md P10）"
                "针对上线前常见短板做了工程化补齐，覆盖「有逻辑没表现、商业化分层不足、"
                "跨节点双扣窗口、活动预下载、社交粘性、客服时间线、资源版本封锁、战斗重连快进」八类问题。",
            ),
            ("List Bullet", "表现层：docs/presentation-protocol-checklist.md；抽卡三次握手；配置 client_ui_params。"),
            ("List Bullet", "商业化：MonthlyCardService、TopUpBonusService、GachaRebateService。"),
            ("List Bullet", "一致性：WalletApplicationService 乐观锁 CAS + StickySessionAffinity / 网关粘性。"),
            ("List Bullet", "活动：ActivityResourceService 预加载 + /api/admin/ops/preload|activate。"),
            ("List Bullet", "社交：SupportService 日助战次数；FriendGiftService 体力赠领。"),
            ("List Bullet", "审计：PlayerTimelineService → GET /api/admin/ops/timeline/{uid}。"),
            ("List Bullet", "版本门禁：ClientVersionGateService + data/ClientRequiredVersions.json。"),
            ("List Bullet", "战斗：BattleReplayService ActionTick → ReplayPacket。"),
            ("List Bullet", "回归测试：P10OpsHardeningFlowTest 覆盖上述业务流程。"),
        ]
        for style, text in blocks:
            cur = insert_paragraph_after(cur, text, style_name=style)

    # ---------- 文档地图：若有 feature-gap 条目则加强 ----------
    p = find_para(doc, contains="docs/feature-gap-fill.md")
    # may not exist in this doc; skip

    # ---------- 第十七章若有文档列表，尝试追加 ----------
    # Soft-add near 总结 if we find a bullet about docs
    p = find_para(doc, exact="输出路径：C:\\Users\\ASUS\\Desktop\\MyMmorpg项目全方位详细介绍文档.docx")
    if p:
        replace_paragraph_text(
            p,
            "输出路径：C:\\Users\\ASUS\\Desktop\\MyMmorpg项目全方位详细介绍文档.docx；"
            "仓库对照：docs/feature-gap-fill.md、docs/presentation-protocol-checklist.md。",
        )

    backup = DOC_PATH.with_suffix(".bak.docx")
    try:
        if not backup.exists():
            import shutil
            shutil.copyfile(DOC_PATH, backup)
    except Exception:
        pass

    doc.save(str(DOC_PATH))
    print(f"Updated: {DOC_PATH}")
    if backup.exists():
        print(f"Backup: {backup}")


if __name__ == "__main__":
    main()
