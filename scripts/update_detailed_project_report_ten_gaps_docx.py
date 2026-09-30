# -*- coding: utf-8 -*-
"""将「十大体验/生产缺口补齐」（CmdId 1130–1142）同步进桌面原报告。

目标文件：Desktop/MyLunarCore项目各方面详细总结报告.bak.docx
只原地修订该文件，不另存新文档、不生成第二份 .docx。
幂等标记：正文出现 WalletTccSagaService 或 CANCEL_WINDOW_SC_NOTIFY 即视为已写入本轮。
"""

from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DOC_PATH = Path.home() / "Desktop" / (
    "MyLunarCore"
    + "\u9879\u76ee\u5404\u65b9\u9762\u8be6\u7ec6\u603b\u7ed3\u62a5\u544a"
    + ".bak.docx"
)


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def replace_paragraph_text(paragraph, new_text):
    style = paragraph.style
    for r in list(paragraph.runs):
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        set_run_font(paragraph.runs[0])
    else:
        run = paragraph.add_run(new_text)
        set_run_font(run)
    paragraph.style = style


def insert_paragraph_after(paragraph, text, style_name="List Bullet"):
    new_para = paragraph._parent.add_paragraph(text, style=style_name)
    paragraph._p.addnext(new_para._p)
    for r in new_para.runs:
        set_run_font(
            r,
            size=13 if style_name.startswith("Heading") else 11,
            bold=style_name.startswith("Heading"),
        )
        if style_name.startswith("Heading"):
            r.font.color.rgb = RGBColor(25, 75, 140)
    return new_para


def find_para(doc, exact=None, startswith=None, contains=None, heading_only=False):
    for p in doc.paragraphs:
        t = p.text.strip()
        if heading_only and not (p.style and p.style.name.startswith("Heading")):
            continue
        if exact is not None and t == exact:
            return p
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    return None


def find_body_after_heading(doc, heading_startswith):
    """定位 Heading 标题后的首个非空正文，避开目录区 Normal 同名项。"""
    seen = False
    for p in doc.paragraphs:
        t = p.text.strip()
        if p.style and p.style.name.startswith("Heading") and t.startswith(heading_startswith):
            seen = True
            continue
        if not seen:
            continue
        if p.style and p.style.name.startswith("Heading"):
            return None
        if t:
            return p
    return None


def replace_if(doc, *, exact=None, startswith=None, contains=None, new_text=None, heading_only=False):
    p = find_para(doc, exact=exact, startswith=startswith, contains=contains, heading_only=heading_only)
    if p is not None and new_text is not None:
        replace_paragraph_text(p, new_text)
        return True
    return False


def set_cell_text(cell, text):
    cell.text = text
    for para in cell.paragraphs:
        for r in para.runs:
            set_run_font(r, size=10)


def patch_table_cell(table, old_exact=None, old_contains=None, col=0, new_text=None, new_other=None):
    for row in table.rows:
        cur = row.cells[col].text.strip()
        hit = False
        if old_exact is not None and cur == old_exact:
            hit = True
        if old_contains is not None and old_contains in cur:
            hit = True
        if not hit:
            continue
        if new_text is not None:
            set_cell_text(row.cells[col], new_text)
        if new_other is not None:
            for ci, val in new_other.items():
                set_cell_text(row.cells[ci], val)
        return True
    return False


def append_table_row(table, values):
    row = table.add_row()
    for i, val in enumerate(values):
        if i < len(row.cells):
            set_cell_text(row.cells[i], val)
    return row


def table_has(table, needle, col=0):
    for row in table.rows:
        if needle in row.cells[col].text:
            return True
    return False


def main():
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    doc = Document(str(DOC_PATH))
    if find_para(doc, contains="WalletTccSagaService") is not None:
        print(f"Already patched (WalletTccSagaService present): {DOC_PATH}")
        return
    if find_para(doc, contains="CANCEL_WINDOW_SC_NOTIFY") is not None:
        print(f"Already patched (CANCEL_WINDOW_SC_NOTIFY present): {DOC_PATH}")
        return

    today = f"{date.today().year}年{date.today().month:02d}月{date.today().day:02d}日"
    n = 0

    # ---------- 封面副标题 ----------
    if replace_if(
        doc,
        startswith="（含技术架构、业务能力、安全运维",
        new_text=(
            "（含技术架构、业务能力、安全运维、微服务迁移、体验缺口补齐、体验闭环、体验沉浸、"
            "体验优化二期/三期、AI 助手增强、战斗手感与拥堵防雪崩、性能深化、运营热更七项短板、"
            "wire v4 稳定性与体验深化，以及「十大体验/生产缺口」补齐："
            "客户端联调清单、钱包 TCC/对账、战斗取消窗与 QTE、公会语音/Ping/DPS、"
            "分级熔断与 SLA、隐私合规、配置审批与灰度回滚、快速重连、跨端 UI、版本主题赛季）"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="报告生成日期：",
        new_text=(
            f"报告生成日期：{today}"
            "（在 wire v4 稳定性深化 CmdId 1109–1125 基础上，已同步落地「十大体验/生产缺口」："
            "①客户端联调优先级清单 + PresentationMockClient + ClientProtocolIntegrationTest；"
            "②WalletTccSagaService / WalletReconciliationJob / PartyMigrateSagaDeadLetterService；"
            "③CancelWindow / BattleQte / camera_fov_impact（CmdId 1130–1133）；"
            "④ScenePing / RaidDamageStatistics / 语音房（1134–1137）；"
            "⑤GradedCircuitBreaker / HPA 指标 / SLA Grafana；"
            "⑥UserDataService / PrivacyConsent / UID 日志哈希；"
            "⑦配置审批工作流 / ConfigGrayAutoRollback / JSON diff 高亮；"
            "⑧FastReconnect 30s 快照与战斗暂停（1138–1139）；"
            "⑨UiAdaptConfig ui_scale/触摸热区（1140–1141）；"
            "⑩VersionTheme + BattlePassTheme（1142）。"
            "定向回归 ExperienceGapE2EFlowTest + ExperienceGapFillFlowTest 等 184 通过 / 0 失败。）"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="报告依据：当前仓库 README",
        new_text=(
            "报告依据：当前仓库 README、源码包结构、配置文件与已有评估文档快照；"
            "此前已写入体验闭环（CmdId 1020–1045）、体验沉浸（1046–1069）、体验优化二期（1070–1076）、"
            "三期（1078–1097）、AI 增强（1098–1104）、手感/拥堵（1105–1107）、运营热更七项、"
            "wire v4 稳定性深化（1109–1125）；"
            "本轮再把「十大体验/生产缺口」写入对应章节（CmdId 1130–1142，清单见 docs/client-integration-priority.md）。"
        ),
    ):
        n += 1

    # ---------- 目录速览追加一行提示 ----------
    p_toc_note = find_para(doc, startswith="下面先给出章节速览")
    if p_toc_note is not None and "7.31" not in p_toc_note.text:
        replace_paragraph_text(
            p_toc_note,
            p_toc_note.text.strip()
            + " 第七章已新增 7.31「十大体验/生产缺口补齐（CmdId 1130–1142）」。",
        )
        n += 1

    # ---------- 7.4 战斗：补充取消窗/QTE/镜头 ----------
    p_battle = find_para(doc, startswith="场景中触发遭遇后进入战斗实例")
    if p_battle is not None and "CancelWindowService" not in p_battle.text:
        replace_paragraph_text(
            p_battle,
            p_battle.text.strip()
            + " 本轮再补：CancelWindowService（普攻命中后默认 100ms 取消窗，CmdId 1130）、"
            "BattleQteService（破盾/击杀 QTE，1131–1133）、"
            "BattleFxMeta.camera_fov_impact / camera_shake_intensity（按技能等级动态计算）。",
        )
        n += 1

    p_buf = find_para(doc, contains="PlayerInputBufferService")
    if p_buf is not None and "CancelWindow" not in p_buf.text:
        replace_paragraph_text(
            p_buf,
            p_buf.text.strip()
            + "；取消窗 CancelWindowService（1130）；QTE BattleQteService（1131–1133）；"
            "镜头冲击字段见 battle_system.proto BattleFxMeta。",
        )
        n += 1

    # ---------- 7.7 钱包 TCC / 对账 ----------
    p_wallet = find_para(doc, startswith="商店购买会变更钱包并推送货币变化")
    if p_wallet is not None and "WalletTccSagaService" not in p_wallet.text:
        replace_paragraph_text(
            p_wallet,
            p_wallet.text.strip()
            + " 关键写路径已引入 WalletTccSagaService（Try 预留→Confirm 实扣→Cancel/Rollback 补偿，"
            "配合 LocalTxLogService），降低多节点双扣窗口；"
            "WalletReconciliationJob 可按日比对余额与 wallet_ledger，异常可管理员一键修正。",
        )
        n += 1

    p_wal_entry = find_para(doc, contains="WalletApplicationService")
    if p_wal_entry is not None and "WalletTcc" not in p_wal_entry.text:
        replace_paragraph_text(
            p_wal_entry,
            "入口：EconomyNettyService、ShopApplicationService、WalletApplicationService、"
            "WalletWalService、WalletTccSagaService、WalletReconciliationJob",
        )
        n += 1

    # ---------- 7.11 / 7.15 社交与公会 ----------
    p_guild = find_para(doc, startswith="公会已落地基础能力")
    if p_guild is not None and "RaidDamageStatisticsService" not in p_guild.text:
        replace_paragraph_text(
            p_guild,
            p_guild.text.strip()
            + " Raid 侧已补：进本自动加入 VoiceSignalingService 临时语音房；"
            "RaidDamageStatisticsService 按秒推送 DPS 排名（CmdId 1137）；"
            "ScenePingService 支持集火/集合/危险/求助标记经 AOI 广播（1134–1136）。",
        )
        n += 1

    p_guild_bullet = find_para(doc, startswith="已有：创建/加入/退出/贡献/商店")
    if p_guild_bullet is not None and "DPS" not in p_guild_bullet.text:
        replace_paragraph_text(
            p_guild_bullet,
            "已有：创建/加入/退出/贡献/商店 + 公会战匹配报分结算；"
            "Raid/贡献榜/语音房/实时 DPS/战术 Ping 可用",
        )
        n += 1

    # ---------- 7.13 设置：跨端 UI ----------
    p_settings = find_para(doc, startswith="7.13")
    # find body near settings
    p_set_body = find_para(doc, contains="键位云同步")
    if p_set_body is not None and "UiAdaptConfigService" not in p_set_body.text:
        # might be in table; try paragraph that mentions Settings
        pass
    p_set2 = find_para(doc, contains="SettingsNettyService")
    if p_set2 is not None and "UiAdapt" not in p_set2.text:
        replace_paragraph_text(
            p_set2,
            p_set2.text.strip()
            + "；跨端适配 UiAdaptConfigService（云端 ui_scale + TouchHeatmap，CmdId 1140–1141）。",
        )
        n += 1

    # ---------- 插入 7.31 新章节（必须锚在 7.30 正文末段，勿匹配目录） ----------
    p_anchor = None
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("客户端联调要点：上报 wire_version") and "PROTOCOL_CAPABILITY" in t:
            p_anchor = p
            break
    if p_anchor is not None and find_para(doc, contains="7.31 十大体验", heading_only=True) is None:
        # insert_paragraph_after 每次插在锚点后，故按「先插末段、后插标题」倒序写入
        insert_paragraph_after(
            p_anchor,
            "客户端联调要点：优先消费 P0——SceneLoadMaskInfo.transition_type（DISSOLVE/RIFT）、"
            "client_pre_fx_ms、BattleFx 镜头冲击、CancelWindow（1130）、Reconnect（1138–1139）；"
            "P1 含 QTE/Ping/语音/DPS；Mock 预览 GET /api/admin/mock/presentation-preview；"
            "清单 docs/client-integration-priority.md。",
            style_name="List Bullet",
        )
        insert_paragraph_after(
            p_anchor,
            "测试入口：ExperienceGapFillFlowTest、ExperienceGapE2EFlowTest（端到端 A–J）、"
            "ClientProtocolIntegrationTest；CI 已纳入上述用例；"
            "与 StabilityEnhancementFlowTest / WireV4EnhancementBusinessFlowsTest 等合计定向回归 "
            "184 通过 / 0 失败。",
            style_name="List Bullet",
        )
        insert_paragraph_after(
            p_anchor,
            "实现入口：tools/PresentationMockClient；economy/WalletTccSagaService、"
            "WalletReconciliationJob；party/PartyMigrateSagaDeadLetterService；"
            "battle/CancelWindowService、BattleQteService、BattleFxComposer；"
            "scene/ScenePingService；guild/RaidDamageStatisticsService；"
            "ops/GradedCircuitBreakerService；privacy/UserDataService、PrivacyConsentService；"
            "common/ConfigGrayAutoRollbackService、StructuredJsonLogger；"
            "admin/ConfigChangeApprovalController；player/FastReconnectService；"
            "settings/UiAdaptConfigService；activity/VersionThemeService；"
            "battlepass/BattlePassThemeService；deploy/grafana/dashboards/sla-health.json。",
            style_name="List Bullet",
        )
        body = (
            "在 7.30 wire v4（CmdId 1109–1125）之后，本轮按产品报告中的十大短板一次性补齐服务端能力，"
            "新增号段 1130–1142（推送与确认为主）。"
            "（1）客户端联调：docs/client-integration-priority.md 按 P0/P1/P2 分级；"
            "PresentationMockClient 生成 DISSOLVE/RIFT/BattleFx/取消窗/QTE 样本；"
            "CI 用 ClientProtocolIntegrationTest 校验 transition_type 等解析。"
            "（2）资产安全：WalletTccSagaService 做 TCC 预留/确认/回滚；"
            "PartyMigrateSagaDeadLetterService 提供 Retry+DeadLetter+混沌注入；"
            "WalletReconciliationJob 日对账并可一键修正。"
            "（3）战斗手感动作化：取消后摇窗口、破盾/击杀 QTE、镜头 FOV/震屏强度下发。"
            "（4）公会协作：Raid 语音房、战术 Ping、实时 DPS 看板。"
            "（5）运维细粒度：GradedCircuitBreaker 按活动/渠道/分片熔断；"
            "BusinessMetrics 暴露 match_queue_depth / battle_pool_usage / SLA P95；"
            "Grafana sla-health 红绿灯看板；K8s HPA 说明见 deploy/k8s/hpa-custom-metrics.md。"
            "（6）隐私合规：导出/软删、登录同意与访客模式、日志 UID 哈希（保留前缀）+ 30 天保留提示。"
            "（7）配置治理：SensitiveOp 双人审批 CONFIG_CHANGE、灰度错误率自动回滚、"
            "admin/config-versions.html JSON 差异高亮。"
            "（8）弱网重连：FastReconnectService 缓存 30s Entity 快照，断线战斗进 AWAITING_RECONNECT 最长 30s，"
            "probe 支持前台预连接。"
            "（9）跨端：按分辨率下发 ui_scale 与 TouchHeatmap（键位仍走 KeyBindCloudService）。"
            "（10）长线运营：VersionThemeService 主题生命周期 + BattlePassThemeService 赛季绑定，"
            "配置 data/VersionThemeConfigs.json。"
        )
        insert_paragraph_after(p_anchor, body, style_name="Normal")
        insert_paragraph_after(
            p_anchor,
            "7.31 十大体验/生产缺口补齐（CmdId 1130–1142）",
            style_name="Heading 2",
        )
        n += 5
    elif find_para(doc, contains="7.31 十大体验") is None:
        raise SystemExit("未找到 7.30 末段锚点，无法插入 7.31")

    # ---------- 8.4 协议版本：扩展号段说明 ----------
    p_wire = find_para(doc, startswith="为避免新旧客户端")
    if p_wire is not None and "1130" not in p_wire.text:
        replace_paragraph_text(
            p_wire,
            p_wire.text.strip()
            + " 十大缺口补齐新增 CmdId 1130–1142（取消窗/QTE/Ping/DPS/重连/UI/主题），"
            "仍归 wire v4 能力协商；wire < 4 时继续屏蔽 1109+（含本轮号段）。",
        )
        n += 1

    # ---------- 11 管理后台 ----------
    p_admin = find_para(doc, startswith="管理后台走 HTTP")
    if p_admin is not None and "ConfigChangeApproval" not in p_admin.text:
        replace_paragraph_text(
            p_admin,
            p_admin.text.strip()
            + " 配置变更可走 ConfigChangeApprovalController（申请→他人批准→apply 触发热更）；"
            "表现层 Mock 预览 GET /api/admin/mock/presentation-preview；"
            "配置版本页支持 JSON 行级差异高亮。",
        )
        n += 1

    # ---------- 12 可观测性 ----------
    p_metrics = find_para(doc, startswith="主工程通过 BusinessMetrics")
    if p_metrics is not None and "battle_pool_usage" not in p_metrics.text:
        replace_paragraph_text(
            p_metrics,
            p_metrics.text.strip()
            + " 本轮补充 HPA/SLA 指标：lunarcore.battle.pool_usage、"
            "lunarcore.sla.battle_success_rate / action_latency_p95_ms / scene_load_p95_ms；"
            "看板 deploy/grafana/dashboards/sla-health.json；"
            "分级熔断 GradedCircuitBreakerService 可按活动/渠道/分片独立止血。",
        )
        n += 1

    p_wallet_consist = find_para(doc, contains="本地事务日志")
    if p_wallet_consist is not None and "WalletTcc" not in p_wallet_consist.text:
        # find 12.2 heading body
        pass
    p_12_2 = find_para(doc, startswith="12.2")
    # body after 12.2 - search LocalTx
    for p in doc.paragraphs:
        t = p.text.strip()
        if "LocalTxLogService" in t and "WalletTcc" not in t:
            replace_paragraph_text(
                p,
                t
                + " 钱包关键路径可用 WalletTccSagaService；日对账 WalletReconciliationJob；"
                "组队迁移失败进 PartyMigrateSagaDeadLetterService DeadLetter。",
            )
            n += 1
            break

    # ---------- 13 安全 / 隐私（只改 Heading「十三」后的正文，避开目录） ----------
    p_13_intro = find_body_after_heading(doc, "十三、安全机制")
    if p_13_intro is not None and "UserDataService" not in p_13_intro.text:
        replace_paragraph_text(
            p_13_intro,
            p_13_intro.text.strip()
            + " 隐私侧已补 UserDataService（导出/软删）、PrivacyConsentService（不同意则访客不可存档）、"
            "StructuredJsonLogger UID 哈希化（保留前缀关联）。",
        )
        n += 1

    # ---------- 15.1 成熟度 ----------
    p_mature = find_para(doc, contains="稳定性与体验深化（wire v4）")
    if p_mature is not None and "十大" not in p_mature.text and "WalletTcc" not in p_mature.text:
        replace_paragraph_text(
            p_mature,
            p_mature.text.strip()
            + "；十大缺口：WalletTcc/对账/Saga DLQ、CancelWindow/QTE/镜头冲击、"
            "Raid 语音+Ping+DPS、分级熔断+SLA、隐私导出同意、配置审批灰度回滚、"
            "FastReconnect、UiAdapt、VersionTheme（CmdId 1130–1142）。",
        )
        n += 1

    # update limitation about client sync
    p_lim = find_para(doc, startswith="公会基础 + 公会战闭环已落地")
    if p_lim is not None and "联调优先级清单" not in p_lim.text:
        replace_paragraph_text(
            p_lim,
            "公会基础 + 公会战闭环已落地；Raid 语音/Ping/DPS 服务端已齐。"
            "表现层仍依赖客户端按 docs/client-integration-priority.md 消费 P0 Notify，"
            "可用 PresentationMockClient /admin mock 预览参数。",
        )
        n += 1

    p_wallet_lim = find_para(doc, startswith="钱包多节点依赖 Redis")
    if p_wallet_lim is not None and "TCC" not in p_wallet_lim.text:
        replace_paragraph_text(
            p_wallet_lim,
            "钱包多节点依赖 Redis 分布式锁；单节点洪峰可用 WalletWalService；"
            "本轮已上 WalletTccSagaService + WalletReconciliationJob 缩小双扣窗口并提供日对账修正。"
            "跨节点仍建议压测验证锁+TCC+补偿日志闭环。",
        )
        n += 1

    # ---------- 16 后续建议：把已落地项从「待做」改为「已落地」 ----------
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("完善大厅社交与匹配体验") and "已落地" not in t:
            replace_paragraph_text(
                p,
                "完善大厅社交与匹配体验：跨节点组队 Saga+DeadLetter/混沌、"
                "公会 Raid 语音房/战术 Ping/实时 DPS 已落地；"
                "继续压测 Center remote + 共享 Redis，并完善第三方 RTC AppId 生产注入。",
            )
            n += 1
            break

    # ---------- 22.2 CI ----------
    p_ci = find_para(doc, startswith="定向单测：CmdId 唯一性")
    if p_ci is not None and "ExperienceGapE2E" not in p_ci.text:
        replace_paragraph_text(
            p_ci,
            p_ci.text.strip()
            + "；十大缺口：ClientProtocolIntegrationTest、ExperienceGapFillFlowTest、"
            "ExperienceGapE2EFlowTest。",
        )
        n += 1

    # ---------- 24 结语 ----------
    p_end = find_para(doc, startswith="本报告基于仓库当前代码与文档快照整理")
    if p_end is not None and "十大体验/生产缺口" not in p_end.text:
        replace_paragraph_text(
            p_end,
            p_end.text.strip()
            + f" {today} 已在本文档内同步「十大体验/生产缺口」落地"
            "（客户端联调清单与 Mock、钱包 TCC/对账/Saga DLQ、战斗取消窗与 QTE 与镜头冲击、"
            "公会语音/Ping/DPS、分级熔断与 SLA 看板、隐私导出同意与 UID 哈希、"
            "配置审批与灰度自动回滚、快速重连、跨端 UI 适配、版本主题与战令主题；"
            "CmdId 1130–1142；测试 ExperienceGapE2EFlowTest 等）。",
        )
        n += 1

    p_conclude = find_para(doc, startswith="MyLunarCore 已经不是")
    if p_conclude is not None and "FastReconnect" not in p_conclude.text:
        replace_paragraph_text(
            p_conclude,
            p_conclude.text.strip()
            + " 近期又补齐生产向十大短板：联调可测、资产可对账、战斗更「动作」、"
            "公会可指挥、运维可细粒度熔断与扩容、隐私可导出删除、配置可审批回滚、"
            "断线可快速恢复、多端 UI 可协同、赛季主题可轮换。",
        )
        n += 1

    # ---------- 表 5 CmdId ----------
    t5 = doc.tables[5]
    extra_cmd_rows = [
        ["1130", "战斗取消后摇窗口 CancelWindowScNotify"],
        ["1131–1133", "战斗 QTE 触发/响应 BattleQte"],
        ["1134–1136", "场景战术标记 ScenePing"],
        ["1137", "Raid 实时伤害统计 RaidDamageStatsScNotify"],
        ["1138–1139", "快速重连 Reconnect"],
        ["1140", "触摸热区 TouchHeatmapScNotify"],
        ["1141", "云端 UI 缩放 UiScaleScNotify"],
        ["1142", "版本主题切换 VersionThemeScNotify"],
    ]
    if not table_has(t5, "1130"):
        for values in extra_cmd_rows:
            append_table_row(t5, values)
            n += 1

    # patch existing rows that mention related systems
    n += int(
        patch_table_cell(
            t5,
            old_contains="200–213",
            new_other={
                1: "战斗（含打击感 FX/Hit-stop/取消窗 1130/QTE 1131–1133；镜头 FOV 见 BattleFxMeta）"
            },
        )
    )
    n += int(
        patch_table_cell(
            t5,
            old_exact="970–983",
            new_other={1: "公会基础（Raid 状态见 1123；DPS 见 1137；Ping 见 1134–1136）"},
        )
    )

    # ---------- 包结构表 T3（若有 wallet/battle 行） ----------
    t3 = doc.tables[3]
    n += int(
        patch_table_cell(
            t3,
            old_contains="economy",
            col=2,
            new_text="商店/钱包/IAP；WAL 预扣；TCC Try/Confirm/Cancel；日对账修正",
        )
    )
    n += int(
        patch_table_cell(
            t3,
            old_contains="battle",
            col=0,
            new_other={
                2: "回合战、Auto、Hit-stop、预输入、取消窗、QTE、镜头冲击、重连回放"
            },
        )
    )

    # ---------- 术语表 ----------
    # find glossary table - usually near 十九
    glossary = None
    for ti, table in enumerate(doc.tables):
        if len(table.rows) > 0 and "术语" in table.rows[0].cells[0].text or (
            len(table.rows) > 1 and "KCP" in table.rows[1].cells[0].text
        ):
            # check
            head = table.rows[0].cells[0].text.strip()
            if head in ("术语", "词条", "名词") or "生活化" in table.rows[0].cells[1].text:
                glossary = table
                break
            if "KCP" in "".join(c.text for c in table.rows[0].cells) or (
                len(table.rows) > 2 and "Protobuf" in table.rows[2].cells[0].text
            ):
                glossary = table
    # table 17 from outline was 术语
    if len(doc.tables) > 17:
        t17 = doc.tables[17]
        extra_terms = [
            ["TCC 钱包", "Try 预留额度→Confirm 实扣→Cancel/Rollback 补偿，缩小跨节点双扣窗口"],
            ["取消窗", "普攻命中后短时间内允许更高优先级技能取消后摇（默认 100ms）"],
            ["战斗 QTE", "破盾/击杀时限时按键，成功给予行动点或伤害加成"],
            ["战术 Ping", "场景内标记集火/集合点，经 AOI 广播给队友"],
            ["快速重连", "保留 session 快照 30s，战斗断线暂停倒计时，无需整包重载"],
            ["版本主题", "主线/活动/卡池/商店/Boss 配置集按时钟统一切换，可预加载下一主题"],
            ["分级熔断", "按活动 ID/渠道/分片独立熔断并切保底邮件，而非整服关闭购买"],
            ["UID 哈希日志", "结构化日志写 userIdHash+前缀，不可逆且可值班关联"],
        ]
        for term, expl in extra_terms:
            if not table_has(t17, term[:4]):
                append_table_row(t17, [term, expl])
                n += 1

    # ---------- FAQ 表 ----------
    if len(doc.tables) > 18:
        t18 = doc.tables[18]
        extra_faq = [
            [
                "客户端怎么预览服务端特效参数？",
                "Admin 登录后访问 /api/admin/mock/presentation-preview，或使用 PresentationMockClient 生成 Base64 样本。",
            ],
            [
                "多节点钱包还会双扣吗？",
                "已上 TCC+对账+分布式锁；仍建议混沌/压测验证，异常由 WalletReconciliationJob 标记并可一键修正。",
            ],
            [
                "断线重连还要重新进图吗？",
                "走 ReconnectCsReq（1138）可恢复 30s 内快照与战斗暂停状态；超时需重新登录加载。",
            ],
        ]
        for q, a in extra_faq:
            if not table_has(t18, q[:8]):
                append_table_row(t18, [q, a])
                n += 1

    doc.save(str(DOC_PATH))
    print(f"Patched {n} places -> {DOC_PATH}")


if __name__ == "__main__":
    main()
