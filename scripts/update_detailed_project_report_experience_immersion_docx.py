# -*- coding: utf-8 -*-
"""将本轮「体验沉浸」8 项服务端功能同步进桌面《MyLunarCore项目各方面详细总结报告.bak.docx》。

只原地修订该文件，不另存新文档、不生成第二份 .docx。
幂等标记：正文出现 CalculateUpgradeMaterials 即视为已写入本轮。
"""

from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DOC_PATH = Path.home() / "Desktop" / (
    "MyLunarCore"
    + "\u9879\u76ee\u5404\u65b9\u9762\u8be6\u7ec6\u603b\u7ed3\u62a5\u544a"
    + ".bak.docx"
)


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def replace_paragraph_text(paragraph, new_text):
    style = paragraph.style
    for r in list(paragraph.runs):
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        set_run_font(paragraph.runs[0])
    else:
        run = paragraph.add_run(new_text)
        set_run_font(run)
    paragraph.style = style


def insert_paragraph_after(paragraph, text, style_name="List Bullet"):
    new_para = paragraph._parent.add_paragraph(text, style=style_name)
    paragraph._p.addnext(new_para._p)
    for r in new_para.runs:
        set_run_font(
            r,
            size=13 if style_name.startswith("Heading") else 11,
            bold=style_name.startswith("Heading"),
        )
        if style_name.startswith("Heading"):
            r.font.color.rgb = RGBColor(25, 75, 140)
    return new_para


def find_para(doc, exact=None, startswith=None, contains=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if exact is not None and t == exact:
            return p
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    return None


def replace_if(doc, *, exact=None, startswith=None, contains=None, new_text=None):
    p = find_para(doc, exact=exact, startswith=startswith, contains=contains)
    if p is not None and new_text is not None:
        replace_paragraph_text(p, new_text)
        return True
    return False


def set_cell_text(cell, text):
    cell.text = text
    for para in cell.paragraphs:
        for r in para.runs:
            set_run_font(r, size=10)


def patch_table_cell(table, old_exact=None, old_contains=None, col=0, new_text=None, new_other=None):
    for row in table.rows:
        cur = row.cells[col].text.strip()
        hit = False
        if old_exact is not None and cur == old_exact:
            hit = True
        if old_contains is not None and old_contains in cur:
            hit = True
        if not hit:
            continue
        if new_text is not None:
            set_cell_text(row.cells[col], new_text)
        if new_other is not None:
            for ci, val in new_other.items():
                set_cell_text(row.cells[ci], val)
        return True
    return False


def append_table_row(table, values):
    row = table.add_row()
    for i, val in enumerate(values):
        if i < len(row.cells):
            set_cell_text(row.cells[i], val)
    return row


def table_has(table, needle, col=0):
    return any(needle in row.cells[col].text for row in table.rows)


def main():
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    doc = Document(str(DOC_PATH))
    if find_para(doc, contains="CalculateUpgradeMaterials") is not None:
        print(f"Already patched (CalculateUpgradeMaterials present): {DOC_PATH}")
        return

    today = f"{date.today().year}年{date.today().month:02d}月{date.today().day:02d}日"
    n = 0

    # ---------- 封面 ----------
    if replace_if(
        doc,
        exact="（含技术架构、业务能力、安全运维、微服务迁移、体验缺口补齐、体验闭环与后续建议）",
        new_text="（含技术架构、业务能力、安全运维、微服务迁移、体验缺口补齐、体验闭环、体验沉浸与后续建议）",
    ):
        n += 1

    if replace_if(
        doc,
        startswith="报告生成日期：",
        new_text=(
            f"报告生成日期：{today}"
            "（已同步体验沉浸 8 项：养成计划计算器与一键连续扫荡、Auto 策略权重与手动大招、"
            "移动状态机与落地物理反馈、表情包仓库、键位云同步与特效降级、剧情树 impact_tags 与快照、"
            "家园好友助产与溢出推送、体力溢出提醒与合成路径；协议号段 CmdId 1046–1069）"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="报告依据：当前仓库 README",
        new_text=(
            "报告依据：当前仓库 README、源码包结构、配置文件与已有评估文档快照；"
            "此前已写入体验闭环 8 项（CmdId 1020–1045、wire v3）；"
            "本轮再把体验沉浸 8 项（CmdId 1046–1069）写入对应章节，"
            "回归测试 ExperienceImmersionFlowsTest / ExperienceImmersionProtocolFlowsTest。"
        ),
    ):
        n += 1

    # ---------- 7.2 养成计划 ----------
    if replace_if(
        doc,
        startswith="覆盖创角、成长晋阶、天赋、属性计算",
        new_text=(
            "覆盖创角、成长晋阶、天赋、属性计算，以及皮肤衣柜的拥有与装备。"
            "属性计算由 AttributeCalculator 等组件完成，皮肤配置来自 data/SkinConfigs.json。"
            "养成计划（DevelopmentPlanService）按目标角色+目标等级汇总材料总清单、缺口与体力缺口，"
            "并按掉落效率排序推荐关卡；协议 CalculateUpgradeMaterials（CmdId 1046–1047）。"
        ),
    ):
        n += 1

    p_char = find_para(doc, startswith="入口：CharacterNettyService、CharacterCreationApplicationService")
    if p_char is not None and "DevelopmentPlanService" not in p_char.text:
        insert_paragraph_after(
            p_char,
            "养成计划：DevelopmentPlanService / CharacterDevelopmentNettyService；"
            "CalculateUpgradeMaterials CsReq/ScRsp（CmdId 1046–1047）。",
        )
        n += 1

    # ---------- 7.3 移动状态机 ----------
    if replace_if(
        doc,
        startswith="玩家进入某个 Plane/Floor 对应的 Zone 实例后",
        new_text=(
            "玩家进入某个 Plane/Floor 对应的 Zone 实例后，可移动、与 NPC/道具交互，"
            "服务器按兴趣范围（AOI）同步周围实体。Zone 可配置人数上限与租约心跳。"
            "MoveCsReq 带 move_state（走/跑/跳/攀爬）；MovementPhysics 按坐标分区给出地板材质、"
            "触觉强度与脚步声效，JUMP 落地推送 SceneInteractPhysicsScNotify（CmdId 1053）。"
            "接近传送门时服务端强制推送 ScenePreloadPushScNotify（CmdId 1028），不必等客户端先请求；"
            "客户端仍可走 ScenePreloadCsReq（354–356）。真正切图时 MigrateSceneScRsp 带 preload_hit 与"
            "SceneLoadMaskInfo：命中预加载用 FADE 短握手（handshake_only），未命中用 BLACK 全黑屏。"
            "跨图或跨节点时，由 Center 规划归属，必要时发放迁移票据。"
        ),
    ):
        n += 1

    p_scene = find_para(doc, startswith="入口：SceneNettyService、ScenePreloadService")
    if p_scene is not None and "MovementPhysics" not in p_scene.text:
        replace_paragraph_text(
            p_scene,
            "入口：SceneNettyService、ScenePreloadService、MovementPhysics、"
            "ZoneManager、ZoneTickService、PlayerMigrationService",
        )
        n += 1

    p_preload = find_para(doc, startswith="预加载协议：ScenePreloadCsReq")
    if p_preload is not None and "1053" not in p_preload.text:
        insert_paragraph_after(
            p_preload,
            "移动物理反馈：MovementPhysics；SceneInteractPhysicsScNotify（CmdId 1053）"
            "下发地板材质、触觉、脚步声效与 landing 标记。",
        )
        n += 1

    # ---------- 7.4 Auto 策略 + 手动大招 + FX 降级 ----------
    if replace_if(
        doc,
        startswith="场景中触发遭遇后进入战斗实例",
        new_text=(
            "场景中触发遭遇后进入战斗实例：开战、行动、结算、退出。"
            "支持遇敌配置与波次；多人共战通过 BattleContext.participantPlayerIds 等字段表达。"
            "BattleAutoService 按启发式真正出手（不是仅提示）：基准等待 1500ms，"
            "speed_multiplier 仅允许 1/2/3；SetBattleAuto 增加 auto_strategy"
            "（PRIORITY_SKILL / PRIORITY_BASIC / SAVE_ENERGY）。"
            "Auto 中可发 BattleManualUlt（CmdId 1051–1052）插入大招：本轮暂停 AI，打完后恢复 Auto。"
            "掉线进入托管不中断战斗，回主界面才 abort。"
            "协议 SetBattleAuto / BattleAutoScNotify / SetBattleSpeed（CmdId 1020–1024）。"
            "行动经 BattleDeterministicValidator 校验后，BattleFxComposer 按暴击/弱点/击杀组装打击感元数据"
            "（camera_shake、time_scale、damage_popup），经 BattleFxScNotify（CmdId 213）推给客户端播特效；"
            "FxQualityAdvisor 按心跳 RTT/丢包把 fx_quality_level 降为高/中/低，不改变结算权威。"
        ),
    ):
        n += 1

    p_auto = find_para(doc, startswith="Auto/倍速：BattleAutoService")
    if p_auto is not None and "BattleManualUlt" not in p_auto.text:
        replace_paragraph_text(
            p_auto,
            "Auto/倍速/策略：BattleAutoService；auto_strategy 优先战技/普攻/攒能量；"
            "BattleManualUlt 手动大招覆盖（1051–1052）；CmdId 1020–1024；掉线托管、回主界面 abort。"
            "特效降级：FxQualityAdvisor + BattleFxScNotify.fx_quality_level。",
        )
        n += 1

    # ---------- 7.5 一键连续扫荡 ----------
    if replace_if(
        doc,
        startswith="挑战玩法负责开局、结果上报、组奖励与历史",
        new_text=(
            "挑战玩法负责开局、结果上报、组奖励与历史；模拟宇宙（Rogue）则包含地图生成、移动、"
            "祝福/奇物、战斗上报与天赋等完整 Roguelike 流程。"
            "已通关关卡可走 SweepService 扫荡：以历史最低回合数解锁，跳过完整 BattleManager，"
            "按 1/2/3 倍体力消耗发放倍率掉落（单道具上限 999），适合重复刷材料。"
            "养成计划确认总消耗后，可 BatchSweep（CmdId 1048–1049）按推荐关卡顺序连续扫荡，"
            "统一结算发奖，并通过 BatchSweepProgressNotify（1050）回传进度，中间不必反复点界面。"
        ),
    ):
        n += 1

    p_sweep = find_para(doc, startswith="扫荡：SweepService / SweepNettyService")
    if p_sweep is not None and "1048" not in p_sweep.text:
        replace_paragraph_text(
            p_sweep,
            "扫荡：SweepService / SweepNettyService；表 player_stage_clear_best；"
            "单次倍率扫荡 CmdId 1010–1013；一键连续扫荡 BatchSweep 1048–1050。"
            "胜利结算会 recordClear。SQL：db/migration_experience_gaps.sql。",
        )
        n += 1

    # ---------- 7.8 合成路径 ----------
    if replace_if(
        doc,
        startswith="提供背包快照以及使用、装备、强化",
        new_text=(
            "提供背包快照以及使用、装备、强化、升阶、锁定、丢弃等操作。道具表可由 CSV 等素材导入维护。"
            "材料不足时可 QueryItemSource（CmdId 1037–1038）反查来源：返回 stage_ids，"
            "并指向 FightStartCsReq / SweepStageCsReq 以便直接开战或扫荡；"
            "高级材料（如 301/401）同时下发 craft_paths 与 craft_cmd=CraftItemCsReq，给出低级合成路径。"
            "过期活动道具由 ExpiredItemRecycleService 按汇率转为信用点（币种 101），"
            "推送 ItemRecycleNotify（1027），禁止静默删除。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="入口：ItemNettyService、ItemApplicationService、ExpiredItemRecycleService",
        new_text=(
            "入口：ItemNettyService、ItemApplicationService、ExpiredItemRecycleService；"
            "材料反查 QueryItemSource（1037–1038，含合成路径 craft_paths），"
            "过期转化 ItemRecycleNotify（1027）。"
        ),
    ):
        n += 1

    # ---------- 7.10 表情包 ----------
    if replace_if(
        doc,
        startswith="大厅模块把社交能力收在一起",
        new_text=(
            "大厅模块把社交能力收在一起：好友、邮件、世界/私聊，以及历史最高分类排行榜。"
            "排行榜与世界聊天在开启 Redis 后可跨进程共享；聊天里 @助手 可转给 AI 模块。"
            "好友在线由 FriendOnlineStatusService 经 Redis online_status_channel 跨节点广播，"
            "上线/下线推送 FriendOnlineNotify（CmdId 1033，含 Plane）。"
            "表情包仓库（EmoteInventoryService）默认发放挥手/跳舞/击掌与点赞贴纸；"
            "聊天 SendChatEmote（1056–1058）与场景 SceneEmote（1059–1061）走 AOI 同步，未拥有贴纸失败。"
        ),
    ):
        n += 1

    p_hall = find_para(doc, startswith="入口：HallNettyService、ChatService")
    if p_hall is not None and "Emote" not in p_hall.text:
        replace_paragraph_text(
            p_hall,
            p_hall.text.strip().rstrip("。")
            + "、EmoteNettyService / EmoteInventoryService（表情 1054–1061）",
        )
        n += 1

    # ---------- 7.13 键位云同步 ----------
    if replace_if(
        doc,
        startswith="版本/热更相关协议可向客户端推送更新信息",
        new_text=(
            "版本/热更相关协议可向客户端推送更新信息；设置与客服工单允许玩家改偏好并提交问题单，后台可继续处理。"
            "键位支持云同步：SyncKeyBind（CmdId 1062–1063）按设备（pc/mobile/gamepad）写入"
            "KeyBindCloudService（Redis 优先，失败回退内存）；登录拉取设置时推送 PushKeyBindScNotify（1064）。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        exact="VersionNettyService；SettingsNettyService、SupportTicketApplicationService",
        new_text=(
            "VersionNettyService；SettingsNettyService、SupportTicketApplicationService、"
            "KeyBindCloudService（SyncKeyBind 1062–1064）"
        ),
    ):
        n += 1

    # ---------- 7.16 剧情树 + 家园助产 ----------
    if replace_if(
        doc,
        startswith="除了战斗与经济循环，项目还补齐了「讲故事」和「家园养成」能力",
        new_text=(
            "除了战斗与经济循环，项目还补齐了「讲故事」和「家园养成」能力，方便策划做剧情与轻养成。"
            "剧情侧在选项节点写入 SAVEPOINT（savepointNodeId），选项可带 impact_tags"
            "（如 +Affinity / UnlockRoute / -Morality）；"
            "GetStoryTreeSnapshot（CmdId 1065–1066）返回节点锁/解锁与当前节点缩略。"
            "DialogueProgressService.markPlayed / advance 必须保留存档点，"
            "ResumeFromBranchCsReq（1039–1040）可从分支续跑。"
            "过场支持跳过/完成，并可用 ReplayCutsceneCsReq（1041–1042）校验已播后发放回放 ticket。"
            "家园除产出、家具摆放与好友互访外，已接线家具交互、AOI-lite 在场同步，以及拜访日志："
            "HomeVisitorLogService 记录来访/留言（CmdId 1030–1035）。"
            "好友助产 HomeHarvestAssist（1067–1068）可为设施减少约 5% 产出 CD（每日每设施一次）；"
            "产出堆积达阈值时推送 HomeOverflowNotify（1069）。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="对话树：DialogueTriggerEngine、DialogueProgressService",
        new_text=(
            "对话树：DialogueTriggerEngine、DialogueProgressService（选项节点 SAVEPOINT，"
            "markPlayed/advance 保留 savepointNodeId；选项 impact_tags）；"
            "ResumeFromBranch 1039–1040；剧情树快照 GetStoryTreeSnapshot 1065–1066。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="家园：HomeBaseService",
        new_text=(
            "家园：HomeBaseService（产出/家具/互访/好友助产，DB home_state）+ HomePresenceService、"
            "FurnitureInteractHandler、HomeVisitorLogService；配置 data/HomeFacilityConfigs.json；"
            "协议：摆放 858–859，交互 870–871，在场 872–874，拜访日志/留言/好友在线 1030–1035，"
            "助产/溢出 1067–1069。"
        ),
    ):
        n += 1

    # ---------- 7.19 体力溢出提醒 ----------
    if replace_if(
        doc,
        startswith="这是近期补齐的「日常运营循环」能力",
        new_text=(
            "这是近期补齐的「日常运营循环」能力，让玩家每天都有明确目标，也让数值消耗可控。"
            "体力（Stamina）会自然恢复，打本消耗，也可日购补充；"
            "满体时自然恢复的 30% 进入储备体力（上限 240，StaminaOverflowService），"
            "玩家用 ClaimReserveStamina（CmdId 1025–1026）提取。"
            "体力 ≥80% 且约 10 分钟不活跃时，StaminaOverflowHintService 经 AiHintNotify 提醒尽快刷本"
            "（30 分钟冷却，避免刷屏）。"
            "每日任务按目标累计进度并领奖，与日切刷新耦合；战令提供 XP/等级以及免费轨与付费轨奖励；"
            "主线章节记录进度，并作为进入某些 Plane/Floor 的门闸。"
            "版本活动则用树形节点 + 共享 ActivityToken，把一版活动内容组织起来。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="体力：StaminaService + StaminaOverflowService",
        new_text=(
            "体力：StaminaService + StaminaOverflowService + StaminaOverflowHintService；"
            "配置 data/StaminaConfigs.json；满体 30% 进储备（上限 240）；ClaimReserveStamina 1025–1026；"
            "≥80% 且不活跃 10 分钟推送溢出提醒。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="协议：DailyLoopPacketHandlers + daily_loop_system.proto 已接线",
        new_text=(
            "协议：DailyLoopPacketHandlers + daily_loop_system.proto 已接线"
            "（体力 996–1000、储备提取 1025–1026、过期转化 1027、日常 1001–1005、"
            "扫荡 1010–1013、一键连续扫荡 1048–1050）；"
            "扫荡按 1/2/3 倍消耗体力并发放倍率奖励；客户端联调可继续加深表现层。"
        ),
    ):
        n += 1

    # ---------- 7.22 主路径故事 ----------
    if replace_if(
        doc,
        startswith="为了让非技术同学建立整体画面",
        new_text=(
            "为了让非技术同学建立整体画面，下面用一条「普通玩家的一天」串起系统："
            "登录进服（校验协议 wire ≥ 2，交换 input_methods 与推荐布局，拉取云端键位）→"
            "看邮件/好友（跨节点在线推送），在聊天或场景打表情包 →"
            "消耗体力进场景走动（跳跃落地有材质/触觉/声效；接近传送门强制预加载；命中 FADE、未命中 BLACK），"
            "触发 NPC 对话（选项带影响标签，可看剧情树快照与 SAVEPOINT，过场可回放）→"
            "遇敌打回合战（可开 Auto 并切换优先战技/普攻/攒能量，大招亮起可手动插入，弱网自动降特效档）→"
            "开挑战、Rogue、世界 BOSS 或深渊；养成计划算出材料缺口后，已通关材料关可一键连续扫荡，"
            "材料不足可反查关卡或走合成路径 →"
            "满体溢出进储备、不活跃时提醒刷本、过期道具转信用点 →"
            "回家园点点基建、摆家具、好友助产加速产出、看拜访日志与他人在场 → 抽卡与商店消费 →"
            "做每日任务/战令/版本活动拿奖，成就达成即时推送 → 推主线章节解锁新地图 →"
            "进公会贡献/兑换/科技/公会战 → 大厅排队匹配不锁界面，成局 10 秒 Ready Check →"
            "必要时问 AI 助手（站内 WebView 打开 B站/抖音/官网攻略）→"
            "新手引导按检查点续跑或一键跳过 → 改设置、同步键位或提工单。"
            "运营同学则在后台改活动排期、导入配置、触发热更；运维同学盯监控与发布演练。"
        ),
    ):
        n += 1

    # ---------- 7.23 专节 ----------
    if find_para(doc, contains="7.23 体验沉浸补齐") is None:
        anchor = find_para(doc, startswith="建议面：AI/教练异步回答")
        if anchor is None:
            raise SystemExit("未找到 7.23 插入锚点")
        h = insert_paragraph_after(anchor, "7.23 体验沉浸补齐（CmdId 1046–1069）", "Heading 2")
        p1 = insert_paragraph_after(
            h,
            "在体验闭环（1020–1045）之后，本轮按「养成路径、战斗手感、箱庭微交互、社交表情、"
            "多端键位、剧情可读、家园互助、体力消耗引导」八条线补齐服务端能力。协议成对规则仍是"
            "CS_REQ = N、SC_RSP = N+1；号段占用 1046–1069。",
            "Normal",
        )
        p2 = insert_paragraph_after(
            p1,
            "养成计划 CalculateUpgradeMaterials（1046–1047）按目标等级返回材料总清单、缺口、"
            "按掉落效率排序的推荐关卡与体力缺口；确认后 BatchSweep（1048–1050）按顺序连续扫荡并统一发奖。"
            "Auto 增加 auto_strategy，BattleManualUlt（1051–1052）允许打断 AI 插入大招。"
            "移动状态机落地推 SceneInteractPhysicsScNotify（1053）。"
            "表情仓库 1054–1061 覆盖查询、聊天与场景 AOI。"
            "键位云同步 1062–1064；战斗特效按 RTT/丢包降档。"
            "剧情树快照 1065–1066 带 impact_tags。"
            "家园助产 1067–1069 减产出 CD 并在堆积时推送溢出。"
            "体力 ≥80% 且不活跃时 AiHintNotify 提醒；QueryItemSource 对高级材料给出合成路径。",
            "Normal",
        )
        insert_paragraph_after(
            p2,
            "测试入口：ExperienceImmersionFlowsTest（领域 8 项）、"
            "ExperienceImmersionProtocolFlowsTest（协议门面与 PacketHandler 回包）、"
            "SweepServiceTest.batchSweep、CmdIdUniquenessTest / CmdIdProtoCompletenessTest。",
            "List Bullet",
        )
        n += 1

    # ---------- 8.4 协议版本 ----------
    if replace_if(
        doc,
        startswith="为避免新旧客户端「各说各话」",
        new_text=(
            "为避免新旧客户端「各说各话」，工程引入了协议线版本概念。"
            "当前 PROTOCOL_WIRE_VERSION = 3（登录握手增加 input_methods / recommended_layout_id）；"
            "ProtocolCompatService.isCompatible 仍接受 wire ≥ 2，拒绝过旧的 v1 客户端"
            "（v1 角色号段 120–129 会撞上组队）。"
            "角色相关命令号已迁到 160–173。体验闭环占用 CmdId 1020–1045，"
            "体验沉浸占用 1046–1069（docs/meta/cmdid-config-meta.yaml 范围 [1020, 1069]）。"
            "因此：改协议必须同步客户端，并尽量用自动化测试（如 CmdIdUniquenessTest）防止号段再重叠。"
            "仓库还提供 buf.yaml，作为 Protobuf lint / breaking 检查基线。"
        ),
    ):
        n += 1

    # ---------- 13.1 ----------
    p_wire = find_para(doc, startswith="确认客户端协议线版本与 CmdId 已对齐 PROTOCOL_WIRE_VERSION=3")
    if p_wire is not None and "1069" not in p_wire.text:
        replace_paragraph_text(
            p_wire,
            "确认客户端协议线版本与 CmdId 已对齐 PROTOCOL_WIRE_VERSION=3"
            "（兼容守卫仍为 wire ≥ 2；v3 字段为 input_methods / recommended_layout_id；"
            "沉浸协议 1046–1069 需与客户端同步）。",
        )
        n += 1

    # ---------- 15.1 ----------
    p_daily = find_para(doc, startswith="日常循环：体力、每日任务、战令")
    if p_daily is not None and "1048" not in p_daily.text:
        replace_paragraph_text(
            p_daily,
            "日常循环：体力、每日任务、战令、主线章节门闸、版本活动容器（DailyLoop/BattlePass 协议已接线）；"
            "已通关关卡扫荡（1/2/3 倍体力，CmdId 1010–1013）与一键连续扫荡（1048–1050）；"
            "满体 30% 溢出进储备（上限 240，ClaimReserve 1025–1026）；过期道具转信用点 101（1027）；"
            "体力 ≥80% 不活跃提醒（StaminaOverflowHintService）。",
        )
        n += 1

    p_loop = find_para(doc, startswith="体验缺口回归：ExperienceGapBusinessFlowsTest")
    if p_loop is not None and find_para(doc, contains="ExperienceImmersionFlowsTest") is None:
        insert_after = find_para(doc, contains="ExperienceLoopBusinessFlowsTest")
        target = insert_after if insert_after is not None else p_loop
        insert_paragraph_after(
            target,
            "体验沉浸回归：ExperienceImmersionFlowsTest / ExperienceImmersionProtocolFlowsTest 覆盖"
            "养成计划、BatchSweep、Auto 策略与手动大招、移动物理反馈、表情包、键位云同步、"
            "剧情树快照、家园助产、体力提醒与合成路径（CmdId 1046–1069）。",
        )
        n += 1

    p_guide = find_para(doc, startswith="组队协议与多人共战/助战")
    if p_guide is not None and "1046" not in p_guide.text:
        replace_paragraph_text(
            p_guide,
            p_guide.text.strip().rstrip("。")
            + "；养成计划与一键连续扫荡（1046–1050）、Auto 策略/手动大招（1051–1052）、"
            "移动物理反馈（1053）、表情包（1054–1061）、键位云同步（1062–1064）、"
            "剧情树快照（1065–1066）、家园助产/溢出（1067–1069）。",
        )
        n += 1

    # ---------- 15.2 ----------
    if replace_if(
        doc,
        startswith="对话/过场/家园：服务端已加深",
        new_text=(
            "对话/过场/家园：服务端已加深（SAVEPOINT/回放/剧情树快照、家具交互、在场同步、拜访日志、好友助产）；"
            "预加载 FADE/BLACK、站内 WebView、扫荡/连续扫荡入口、战斗 Auto 策略 UI 与手动大招、FX 播片与降档、"
            "Ready Check 倒计时、多端布局与云端键位、表情动作表现仍需客户端对齐。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="协议已迁到 wire v3（兼容守卫仍 ≥ 2）",
        new_text=(
            "协议已迁到 wire v3（兼容守卫仍 ≥ 2）；v1 客户端必须升级；"
            "改 CmdId（含 1020–1045 闭环与 1046–1069 沉浸）必须与客户端同步。"
        ),
    ):
        n += 1

    # ---------- 16.1 ----------
    p_short = find_para(doc, startswith="继续以模块化单体打磨崩铁式体验与数值内容")
    if p_short is not None and "养成计划" not in p_short.text:
        replace_paragraph_text(
            p_short,
            "继续以模块化单体打磨崩铁式体验与数值内容；"
            "客户端对齐养成计划/连续扫荡、Auto 策略与手动大招、跳跃落地反馈、表情包、云端键位、"
            "剧情树快照、家园助产、体力溢出提醒与合成台跳转，以及既有 Auto/倍速、储备体力、"
            "FADE/BLACK、Ready Check 与 recommended_layout_id。",
        )
        n += 1

    # ---------- 结语 ----------
    if replace_if(
        doc,
        startswith="MyLunarCore 已经不是「只有空壳的演示仓库」",
        new_text=(
            "MyLunarCore 已经不是「只有空壳的演示仓库」：它在单体形态下具备较完整的崩铁式 RPG 服务端能力——"
            "联网、场景（含强制预加载、FADE/BLACK 切图与跳跃落地物理反馈）、"
            "回合战（含 Auto 策略权重、手动大招覆盖、倍速真正出手、打击感 FX 与弱网降档）、"
            "养成（含目标导向材料计划与一键连续扫荡）、经济抽卡、"
            "体力与每日循环（含扫荡、满体 30% 储备溢出、不活跃溢出提醒、过期道具转信用点与合成路径）、"
            "战令与主线章节、活动任务、对话 SAVEPOINT / impact_tags / 剧情树快照与过场回放、"
            "家园（含家具交互/在场/拜访日志/好友助产）、公会/公会战/公会科技、竞技场评分、"
            "世界 BOSS/深渊/遗器词条等缺口补齐玩法、成就即时推送与材料反查、"
            "带检查点/可跳过的新手引导、匹配排队不锁界面与 10 秒 Ready Check、"
            "表情包、键位云同步、多端输入协商（wire v3）、运维热更与回滚、业务监控，"
            "以及可选的多节点好友在线与能力增强后的 AI 助手旁路"
            "（多轮对话、个性化、主动推送、熔断灰度、合规声明、专项指标，"
            "以及主流平台攻略默认在站内 WebView 打开并附 ≤200 字摘要）。"
            "与此同时，它也很诚实地把微服务目录标成脚手架，把无缝大世界标成实验，把默认密钥标成仅限本地。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="本报告基于仓库当前代码与文档快照整理",
        new_text=(
            f"本报告基于仓库当前代码与文档快照整理，{today} 已在本文档内同步体验沉浸 8 项"
            "（CmdId 1046–1069；此前闭环为 1020–1045、PROTOCOL_WIRE_VERSION=3），"
            "可用于汇报、培训与决策讨论。若后续架构或配置有重大变更，建议修订对应章节后刷新目录域，"
            "以便目录页码与正文一致。"
        ),
    ):
        n += 1

    # ---------- 表 3 包一览 ----------
    t3 = doc.tables[3]
    n += patch_table_cell(
        t3,
        old_contains="创角、晋阶、天赋、属性计算",
        col=2,
        new_text="创角、晋阶、天赋、属性计算、养成计划材料计算器",
    )
    n += patch_table_cell(
        t3,
        old_contains="回合战、遭遇、波次",
        col=2,
        new_text="回合战、遭遇、波次、BattleAuto 策略/倍速/手动大招、打击感 FX 与弱网降档",
    )
    n += patch_table_cell(
        t3,
        old_contains="进图、移动、AOI 同步",
        col=2,
        new_text="进图、移动状态机、落地物理反馈、AOI 同步、Zone 管理、强制 Preload、FADE/BLACK",
    )
    n += patch_table_cell(
        t3,
        old_contains="材料反查 QueryItemSource",
        col=2,
        new_text="道具使用、装备、强化、材料反查 QueryItemSource、合成路径、过期转化信用点",
    )
    n += patch_table_cell(
        t3,
        old_contains="跨节点好友在线",
        col=2,
        new_text="好友、邮件、聊天、排行榜入口、跨节点好友在线、表情包仓库",
    )
    n += patch_table_cell(
        t3,
        old_contains="选项 SAVEPOINT",
        col=2,
        new_text="NPC 对话树、impact_tags、剧情树快照、选项 SAVEPOINT、ResumeFromBranch",
    )
    n += patch_table_cell(
        t3,
        old_contains="拜访日志",
        col=2,
        new_text="基建、产出、家具摆放/交互、好友互访/助产、在场同步、拜访日志/留言、溢出推送",
    )
    n += patch_table_cell(
        t3,
        old_contains="设置同步、工单提交",
        col=2,
        new_text="设置同步、键位云同步、工单提交",
    )
    n += patch_table_cell(
        t3,
        old_contains="赠礼、公会红包、语音信令",
        col=2,
        new_text="赠礼、公会红包、语音信令、表情包播放",
    )

    # ---------- 表 5 CmdId ----------
    t5 = doc.tables[5]
    n += patch_table_cell(
        t5,
        old_exact="100–118",
        new_other={1: "大厅社交（表情包见 1054–1061）"},
    )
    n += patch_table_cell(
        t5,
        old_exact="160–173",
        new_other={1: "角色养成（养成计划见 1046–1047）"},
    )
    n += patch_table_cell(
        t5,
        old_exact="200–213",
        new_other={1: "战斗（含打击感 FX 与特效降档；Auto/倍速见 1020–1024，手动大招见 1051–1052）"},
    )
    n += patch_table_cell(
        t5,
        old_exact="300–351",
        new_other={1: "场景（移动物理反馈见 1053）"},
    )
    n += patch_table_cell(
        t5,
        old_exact="850–859 / 870–874",
        new_other={1: "家园（家具/在场/拜访日志见 1030–1035；助产见 1067–1069）"},
    )
    n += patch_table_cell(
        t5,
        old_exact="860–868",
        new_other={1: "对话 / 过场（SAVEPOINT / 回放见 1039–1042；剧情树快照见 1065–1066）"},
    )
    n += patch_table_cell(
        t5,
        old_exact="930+",
        new_other={1: "设置 / 客服（键位云同步见 1062–1064）"},
    )
    n += patch_table_cell(
        t5,
        old_exact="1010–1013",
        new_other={1: "关卡扫荡（体力倍率；一键连续扫荡见 1048–1050）"},
    )
    extra_cmd_rows = [
        ["1046–1047", "养成计划材料计算器 CalculateUpgradeMaterials"],
        ["1048–1050", "一键连续扫荡 BatchSweep + 进度推送"],
        ["1051–1052", "Auto 中手动大招 BattleManualUlt"],
        ["1053", "场景落地物理反馈 SceneInteractPhysics"],
        ["1054–1061", "表情包仓库 / 聊天表情 / 场景表情"],
        ["1062–1064", "键位云同步 SyncKeyBind"],
        ["1065–1066", "剧情树快照 GetStoryTreeSnapshot"],
        ["1067–1069", "家园助产 / 产出溢出推送"],
    ]
    if not table_has(t5, "1046"):
        for values in extra_cmd_rows:
            append_table_row(t5, values)
            n += 1

    # ---------- 表 6 配置文件 ----------
    t6 = doc.tables[6]
    n += patch_table_cell(
        t6,
        old_exact="DialogueTrees.json",
        new_other={1: "NPC 对话树（含选项 impact_tags）"},
    )

    # ---------- 表 17 术语 ----------
    t17 = doc.tables[17]
    extra_terms = [
        ["养成计划", "输入目标角色与等级，一次算出材料缺口、推荐刷本顺序和体力够不够"],
        ["一键连续扫荡", "确认总消耗后服务端按关卡顺序连刷多次，统一发奖，不用反复点扫荡"],
        ["Auto 策略 / 手动大招", "自动战斗可选优先战技、普攻或攒能量；大招亮起时可手动插入一刀再恢复 Auto"],
        ["impact_tags", "对话选项上的影响标签，如加好感、解锁路线、扣道德，供剧情树展示"],
        ["键位云同步", "按设备把按键/按钮布局存到服务器，换端登录自动推送"],
    ]
    for term, expl in extra_terms:
        if not table_has(t17, term):
            append_table_row(t17, [term, expl])
            n += 1

    # ---------- 表 18 FAQ ----------
    t18 = doc.tables[18]
    extra_faq = [
        [
            "角色升到 80 级要刷哪些本？",
            "用养成计划接口一次算出材料缺口和推荐关卡；确认体力后可一键连续扫荡，不必手动记关卡。",
        ],
        [
            "Auto 能不能手动放大招？",
            "能。开启 Auto 后点大招会发 BattleManualUlt，本轮暂停 AI，打完继续自动。",
        ],
    ]
    for q, a in extra_faq:
        if not table_has(t18, q[:8]):
            append_table_row(t18, [q, a])
            n += 1

    doc.save(str(DOC_PATH))
    print(f"Updated in place: {DOC_PATH}")
    print(f"Patches applied (approx): {n}")


if __name__ == "__main__":
    main()
