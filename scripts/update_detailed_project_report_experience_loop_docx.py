# -*- coding: utf-8 -*-
"""将本轮「体验闭环」8 项服务端功能同步进桌面《MyLunarCore项目各方面详细总结报告.bak.docx》。

只原地修订该文件，不另存新文档、不生成第二份 .docx。
幂等标记：正文出现 BattleAutoService 即视为已写入本轮。
"""

from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DOC_PATH = Path.home() / "Desktop" / (
    "MyLunarCore"
    + "\u9879\u76ee\u5404\u65b9\u9762\u8be6\u7ec6\u603b\u7ed3\u62a5\u544a"
    + ".bak.docx"
)


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def replace_paragraph_text(paragraph, new_text):
    style = paragraph.style
    for r in list(paragraph.runs):
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        set_run_font(paragraph.runs[0])
    else:
        run = paragraph.add_run(new_text)
        set_run_font(run)
    paragraph.style = style


def insert_paragraph_after(paragraph, text, style_name="List Bullet"):
    new_para = paragraph._parent.add_paragraph(text, style=style_name)
    paragraph._p.addnext(new_para._p)
    for r in new_para.runs:
        set_run_font(
            r,
            size=13 if style_name.startswith("Heading") else 11,
            bold=style_name.startswith("Heading"),
        )
        if style_name.startswith("Heading"):
            r.font.color.rgb = RGBColor(25, 75, 140)
    return new_para


def find_para(doc, exact=None, startswith=None, contains=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if exact is not None and t == exact:
            return p
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    return None


def replace_if(doc, *, exact=None, startswith=None, contains=None, new_text=None):
    p = find_para(doc, exact=exact, startswith=startswith, contains=contains)
    if p is not None and new_text is not None:
        replace_paragraph_text(p, new_text)
        return True
    return False


def set_cell_text(cell, text):
    cell.text = text
    for para in cell.paragraphs:
        for r in para.runs:
            set_run_font(r, size=10)


def patch_table_cell(table, old_exact=None, old_contains=None, col=0, new_text=None, new_other=None):
    for row in table.rows:
        cur = row.cells[col].text.strip()
        hit = False
        if old_exact is not None and cur == old_exact:
            hit = True
        if old_contains is not None and old_contains in cur:
            hit = True
        if not hit:
            continue
        if new_text is not None:
            set_cell_text(row.cells[col], new_text)
        if new_other is not None:
            for ci, val in new_other.items():
                set_cell_text(row.cells[ci], val)
        return True
    return False


def append_table_row(table, values):
    row = table.add_row()
    for i, val in enumerate(values):
        if i < len(row.cells):
            set_cell_text(row.cells[i], val)
    return row


def table_has(table, needle, col=0):
    return any(needle in row.cells[col].text for row in table.rows)


def main():
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    doc = Document(str(DOC_PATH))
    if find_para(doc, contains="BattleAutoService") is not None:
        print(f"Already patched (BattleAutoService present): {DOC_PATH}")
        return

    today = f"{date.today().year}年{date.today().month:02d}月{date.today().day:02d}日"
    n = 0

    # ---------- 封面 ----------
    if replace_if(
        doc,
        exact="（含技术架构、业务能力、安全运维、微服务迁移、体验缺口补齐与后续建议）",
        new_text="（含技术架构、业务能力、安全运维、微服务迁移、体验缺口补齐、体验闭环与后续建议）",
    ):
        n += 1

    if replace_if(
        doc,
        startswith="报告生成日期：",
        new_text=(
            f"报告生成日期：{today}"
            "（已同步体验闭环：战斗 Auto/倍速、体力溢出与过期转化、强制 Preload 与 FADE/BLACK、"
            "家园拜访日志与跨节点好友在线、成就即时推送与材料反查、对话 SAVEPOINT 与过场回放、"
            "匹配 Ready Check、多端输入协商；协议线 PROTOCOL_WIRE_VERSION=3）"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="报告依据：当前仓库 README",
        new_text=(
            "报告依据：当前仓库 README、源码包结构、配置文件与已有评估文档快照；"
            "本轮已把服务端体验闭环 8 项（CmdId 1020–1045、wire v3）写入对应章节，"
            "回归测试 ExperienceLoopBusinessFlowsTest / ExperienceLoopFillTest，"
            "库表增量 db/migration_experience_loop_fill.sql。"
        ),
    ):
        n += 1

    # ---------- 7.1 登录：多端输入协商 ----------
    if replace_if(
        doc,
        startswith="玩家用账号密码登录后，服务器建立会话",
        new_text=(
            "玩家用账号密码登录后，服务器建立会话、校验心跳、处理登出，并维护在线人数与超时。"
            "密码支持哈希校验；明文兼容开关默认关闭，避免不安全的比对方式在不知情时开启。"
            "wire v3 起 LoginCsReq 携带 input_methods（bit0 键鼠 / bit1 触屏 / bit2 手柄）与 device_id；"
            "LoginScRsp 回传 input_methods 与 recommended_layout_id"
            "（pc_km / mobile_touch / console_gamepad，见 InputCapability）。"
            "PROTOCOL_WIRE_VERSION=3，ProtocolCompatService.isCompatible 仍为客户端 wire ≥ 2。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        exact="关键类：PlayerSessionService、PlayerLoginApplicationService、AccountPasswordService",
        new_text=(
            "关键类：PlayerSessionService、PlayerLoginApplicationService、"
            "AccountPasswordService、InputCapability"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        exact="协议：player_session.proto；命令号大致在会话段 68–83",
        new_text=(
            "协议：player_session.proto；命令号大致在会话段 68–83；"
            "LoginCsReq/ScRsp 含 input_methods、device_id、recommended_layout_id（wire v3）。"
        ),
    ):
        n += 1

    # ---------- 7.3 场景：强制 Preload + FADE/BLACK ----------
    if replace_if(
        doc,
        startswith="玩家进入某个 Plane/Floor 对应的 Zone 实例后",
        new_text=(
            "玩家进入某个 Plane/Floor 对应的 Zone 实例后，可移动、与 NPC/道具交互，"
            "服务器按兴趣范围（AOI）同步周围实体。Zone 可配置人数上限与租约心跳。"
            "接近传送门时服务端强制推送 ScenePreloadPushScNotify（CmdId 1028），不必等客户端先请求；"
            "客户端仍可走 ScenePreloadCsReq（354–356）。真正切图时 MigrateSceneScRsp 带 preload_hit 与"
            "SceneLoadMaskInfo：命中预加载用 FADE 短握手（handshake_only），未命中用 BLACK 全黑屏。"
            "跨图或跨节点时，由 Center 规划归属，必要时发放迁移票据。"
        ),
    ):
        n += 1

    p_preload = find_para(doc, startswith="预加载协议：ScenePreloadCsReq")
    if p_preload is not None and "1028" not in p_preload.text:
        replace_paragraph_text(
            p_preload,
            "预加载协议：ScenePreloadCsReq/ScRsp、ScenePreloadReadyScNotify（CmdId 354–356）；"
            "靠近传送门强制推送 ScenePreloadPushScNotify（1028）；"
            "加载完成 SceneLoadCompleteScRsp.handshake_only；"
            "过渡分级 SceneLoadMaskInfo.transition_type = FADE（命中）/ BLACK（未命中）。",
        )
        n += 1

    # ---------- 7.4 战斗 Auto + 倍速 ----------
    if replace_if(
        doc,
        startswith="场景中触发遭遇后进入战斗实例",
        new_text=(
            "场景中触发遭遇后进入战斗实例：开战、行动、结算、退出。"
            "支持遇敌配置与波次；多人共战通过 BattleContext.participantPlayerIds 等字段表达。"
            "BattleAutoService 按启发式真正出手（不是仅提示）：基准等待 1500ms，"
            "speed_multiplier 仅允许 1/2/3，压缩等待后由服务端代打；"
            "掉线进入托管不中断战斗，回主界面才 abort。"
            "协议 SetBattleAuto / BattleAutoScNotify / SetBattleSpeed（CmdId 1020–1024）。"
            "行动经 BattleDeterministicValidator 校验后，BattleFxComposer 按暴击/弱点/击杀组装打击感元数据"
            "（camera_shake、time_scale、damage_popup），经 BattleFxScNotify（CmdId 213）推给客户端播特效，"
            "不改变结算权威。"
        ),
    ):
        n += 1

    p_fx = find_para(doc, startswith="打击感：BattleFxComposer")
    if p_fx is not None and "BattleAutoService" not in p_fx.text:
        insert_paragraph_after(
            p_fx,
            "Auto/倍速：BattleAutoService；基准等待 1500ms；CmdId 1020–1024；掉线托管、回主界面 abort。",
        )
        n += 1
    elif p_fx is None:
        p_battle_cfg = find_para(doc, startswith="配置：data/EncounterConfigs.json")
        if p_battle_cfg is not None:
            insert_paragraph_after(
                p_battle_cfg,
                "Auto/倍速：BattleAutoService；基准等待 1500ms；CmdId 1020–1024；掉线托管、回主界面 abort。",
            )
            n += 1

    # ---------- 7.8 背包：材料反查 + 过期转化 ----------
    if replace_if(
        doc,
        startswith="提供背包快照以及使用、装备、强化",
        new_text=(
            "提供背包快照以及使用、装备、强化、升阶、锁定、丢弃等操作。道具表可由 CSV 等素材导入维护。"
            "材料不足时可 QueryItemSource（CmdId 1037–1038）反查来源：返回 stage_ids，"
            "并指向 FightStartCsReq / SweepStageCsReq 以便直接开战或扫荡。"
            "过期活动道具由 ExpiredItemRecycleService 按汇率转为信用点（币种 101），"
            "推送 ItemRecycleNotify（1027），禁止静默删除。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        exact="入口：ItemNettyService、ItemApplicationService",
        new_text=(
            "入口：ItemNettyService、ItemApplicationService、ExpiredItemRecycleService；"
            "材料反查 QueryItemSource（1037–1038），过期转化 ItemRecycleNotify（1027）。"
        ),
    ):
        n += 1

    # ---------- 7.10 大厅：跨节点好友在线 ----------
    if replace_if(
        doc,
        startswith="大厅模块把社交能力收在一起",
        new_text=(
            "大厅模块把社交能力收在一起：好友、邮件、世界/私聊，以及历史最高分类排行榜。"
            "排行榜与世界聊天在开启 Redis 后可跨进程共享；聊天里 @助手 可转给 AI 模块。"
            "好友在线由 FriendOnlineStatusService 经 Redis online_status_channel 跨节点广播，"
            "上线/下线推送 FriendOnlineNotify（CmdId 1033，含 Plane）。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        exact="入口：HallNettyService、ChatService、FriendApplicationService、MailApplicationService",
        new_text=(
            "入口：HallNettyService、ChatService、FriendApplicationService、"
            "MailApplicationService、FriendOnlineStatusService"
        ),
    ):
        n += 1

    # ---------- 7.11 匹配：不锁界面 + Ready Check ----------
    if replace_if(
        doc,
        startswith="Party 是大世界组队雏形",
        new_text=(
            "Party 是大世界组队雏形（同进程权威，人数上限约 4 人）。"
            "Match/Room 面向副本匹配：排队、开房、兼容性打分，并与挑战匹配协调器联动。"
            "大厅排队保持会话 HALL、MatchPhase.QUEUED，不锁界面；"
            "成局推送 MatchSuccessScNotify（CmdId 1043，10 秒 Ready Check），"
            "双方 MatchReadyCheckCsReq（1044–1045）确认后才进入 LOCKED/MATCHING。"
            "跨节点组队与完整私聊跨服仍有局限，报告后文会再次提醒。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        exact="匹配：MatchNettyService、RoomService、MatchQueue、ChallengeMatchCoordinator",
        new_text=(
            "匹配：MatchNettyService、MatchmakingService、RoomService、MatchQueue、"
            "ChallengeMatchCoordinator；Ready Check 超时 10000ms（CmdId 1043–1045）。"
        ),
    ):
        n += 1

    # ---------- 7.14 成就即时推送 ----------
    if replace_if(
        doc,
        startswith="成就系统记录玩家达成条件与领取状态",
        new_text=(
            "成就系统记录玩家达成条件与领取状态，适合做长期留存目标。"
            "进度达标立即推送 AchievementUnlockPushScNotify（CmdId 1036），不必等客户端轮询领取列表。"
            "新手引导则按配置步骤引导新玩家完成关键操作，并落库检查点（checkpoint_step_id / committed_json）："
            "断线重连从最近已提交步骤恢复；支持一键跳过（skipped=true）。"
            "两者都属于「内容运营型」能力：配置改了，行为就能跟着变。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="成就：AchievementService",
        new_text=(
            "成就：AchievementService；表 player_achievement；协议号段约 950–954，"
            "即时推送 AchievementUnlockPushScNotify（1036）。"
        ),
    ):
        n += 1

    # ---------- 7.16 对话 SAVEPOINT / 过场回放 / 家园拜访日志 ----------
    if replace_if(
        doc,
        startswith="除了战斗与经济循环，项目还补齐了「讲故事」和「家园养成」能力",
        new_text=(
            "除了战斗与经济循环，项目还补齐了「讲故事」和「家园养成」能力，方便策划做剧情与轻养成。"
            "剧情侧在选项节点写入 SAVEPOINT（savepointNodeId）；"
            "DialogueProgressService.markPlayed / advance 必须保留存档点，"
            "ResumeFromBranchCsReq（1039–1040）可从分支续跑。"
            "过场支持跳过/完成，并可用 ReplayCutsceneCsReq（1041–1042）校验已播后发放回放 ticket。"
            "家园除产出、家具摆放与好友互访外，已接线家具交互、AOI-lite 在场同步，以及拜访日志："
            "HomeVisitorLogService 记录来访/留言（CmdId 1030–1035）。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="对话树：DialogueTriggerEngine、DialogueProgressService",
        new_text=(
            "对话树：DialogueTriggerEngine、DialogueProgressService（选项节点 SAVEPOINT，"
            "markPlayed/advance 保留 savepointNodeId）；ResumeFromBranch CmdId 1039–1040。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="过场：CutsceneTriggerService",
        new_text=(
            "过场：CutsceneTriggerService；配置 data/CutsceneConfigs.json；"
            "ReplayCutsceneCsReq 校验已播后发 ticket（CmdId 1041–1042）。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="家园：HomeBaseService",
        new_text=(
            "家园：HomeBaseService（产出/家具/互访，DB home_state）+ HomePresenceService、"
            "FurnitureInteractHandler、HomeVisitorLogService；配置 data/HomeFacilityConfigs.json；"
            "协议：摆放 858–859，交互 870–871，在场 872–874，拜访日志/留言/好友在线 1030–1035。"
        ),
    ):
        n += 1

    # ---------- 7.19 体力溢出 ----------
    if replace_if(
        doc,
        startswith="这是近期补齐的「日常运营循环」能力",
        new_text=(
            "这是近期补齐的「日常运营循环」能力，让玩家每天都有明确目标，也让数值消耗可控。"
            "体力（Stamina）会自然恢复，打本消耗，也可日购补充；"
            "满体时自然恢复的 30% 进入储备体力（上限 240，StaminaOverflowService），"
            "玩家用 ClaimReserveStamina（CmdId 1025–1026）提取。"
            "每日任务按目标累计进度并领奖，与日切刷新耦合；战令提供 XP/等级以及免费轨与付费轨奖励；"
            "主线章节记录进度，并作为进入某些 Plane/Floor 的门闸。"
            "版本活动则用树形节点 + 共享 ActivityToken，把一版活动内容组织起来。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="体力：StaminaService",
        new_text=(
            "体力：StaminaService + StaminaOverflowService；配置 data/StaminaConfigs.json；"
            "满体 30% 进储备（上限 240）；ClaimReserveStamina 1025–1026。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="协议：DailyLoopPacketHandlers + daily_loop_system.proto 已接线",
        new_text=(
            "协议：DailyLoopPacketHandlers + daily_loop_system.proto 已接线"
            "（体力 996–1000、储备提取 1025–1026、过期转化 1027、日常 1001–1005、扫荡 1010–1013）；"
            "扫荡按 1/2/3 倍消耗体力并发放倍率奖励；客户端联调可继续加深表现层。"
        ),
    ):
        n += 1

    # ---------- 7.20 切图过渡分级 ----------
    if replace_if(
        doc,
        startswith="好友助战允许借用好友角色快照加入战斗",
        new_text=(
            "好友助战允许借用好友角色快照加入战斗（写入 BattleContext.participantPlayerIds）；"
            "公会科技按公会等级提供全局 Buff（攻/防/血/体力恢复等）；"
            "切图时服务器会发放 LoadingTicket，加载完成前拒绝移动与开战，减少「半截进图」导致的状态错乱。"
            "靠近传送门强制推送 ScenePreloadPushScNotify（1028）；命中预加载则 handshake_only=true、"
            "过渡 FADE，未命中 BLACK。登录阶段还会交换 SupportedFeatures 与 recommended_layout_id，"
            "让客户端按输入能力切换 UI 布局，并按服务端能力隐藏未开放入口"
            "（例如公会战入口在能力未开时返回 retcode=20）。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="切图加载：PlayerLoadingStateService",
        new_text=(
            "切图加载：PlayerLoadingStateService（含 handshakeOnly 短 TTL）+ ScenePreloadService；"
            "CmdId 352–353（加载完成）、354–356（客户端预加载）与 1028（服务端强制推送）；"
            "过渡分级 FADE / BLACK。"
        ),
    ):
        n += 1

    p_feat = find_para(doc, startswith="能力握手：登录响应含 supported_features")
    if p_feat is not None and "recommended_layout_id" not in p_feat.text:
        replace_paragraph_text(
            p_feat,
            "能力握手：登录响应含 supported_features / enabled_features、session_crypto_key，"
            "以及 input_methods / recommended_layout_id（wire v3）。",
        )
        n += 1

    # ---------- 7.22 主路径故事 ----------
    if replace_if(
        doc,
        startswith="为了让非技术同学建立整体画面",
        new_text=(
            "为了让非技术同学建立整体画面，下面用一条「普通玩家的一天」串起系统："
            "登录进服（校验协议 wire ≥ 2，交换 input_methods 与推荐布局）→ 看邮件/好友（跨节点在线推送）→"
            "消耗体力进场景走动（接近传送门强制预加载；命中 FADE、未命中 BLACK），"
            "触发 NPC 对话（选项节点可 SAVEPOINT，过场可回放）→"
            "遇敌打回合战（可开 Auto/1–3 倍速代打，可带助战，客户端按 BattleFx 播打击感）→"
            "开挑战、Rogue、世界 BOSS 或深渊；已通关材料关可 1/2/3 倍体力扫荡，材料不足可反查关卡来源 →"
            "满体溢出进储备、过期道具转信用点 →"
            "回家园点点基建、摆家具、看拜访日志与他人在场 → 抽卡与商店消费 →"
            "做每日任务/战令/版本活动拿奖，成就达成即时推送 → 推主线章节解锁新地图 →"
            "进公会贡献/兑换/科技/公会战 → 大厅排队匹配不锁界面，成局 10 秒 Ready Check →"
            "必要时问 AI 助手（站内 WebView 打开 B站/抖音/官网攻略）→"
            "新手引导按检查点续跑或一键跳过 → 改设置或提工单。"
            "运营同学则在后台改活动排期、导入配置、触发热更；运维同学盯监控与发布演练。"
        ),
    ):
        n += 1

    # ---------- 8.4 协议版本 ----------
    if replace_if(
        doc,
        startswith="为避免新旧客户端「各说各话」",
        new_text=(
            "为避免新旧客户端「各说各话」，工程引入了协议线版本概念。"
            "当前 PROTOCOL_WIRE_VERSION = 3（登录握手增加 input_methods / recommended_layout_id）；"
            "ProtocolCompatService.isCompatible 仍接受 wire ≥ 2，拒绝过旧的 v1 客户端"
            "（v1 角色号段 120–129 会撞上组队）。"
            "角色相关命令号已迁到 160–173。本轮体验闭环占用 CmdId 1020–1045。"
            "因此：改协议必须同步客户端，并尽量用自动化测试（如 CmdIdUniquenessTest）防止号段再重叠。"
            "仓库还提供 buf.yaml，作为 Protobuf lint / breaking 检查基线。"
        ),
    ):
        n += 1

    # ---------- 13.1 上线清单 ----------
    if replace_if(
        doc,
        exact="确认客户端协议线版本与 CmdId 已对齐 PROTOCOL_WIRE_VERSION=2。",
        new_text=(
            "确认客户端协议线版本与 CmdId 已对齐 PROTOCOL_WIRE_VERSION=3"
            "（兼容守卫仍为 wire ≥ 2；v3 字段为 input_methods / recommended_layout_id）。"
        ),
    ):
        n += 1

    # ---------- 15.1 已实现 ----------
    p_daily = find_para(doc, startswith="日常循环：体力、每日任务、战令")
    if p_daily is not None and "储备" not in p_daily.text:
        replace_paragraph_text(
            p_daily,
            "日常循环：体力、每日任务、战令、主线章节门闸、版本活动容器（DailyLoop/BattlePass 协议已接线）；"
            "已通关关卡扫荡（1/2/3 倍体力，CmdId 1010–1013）；"
            "满体 30% 溢出进储备（上限 240，ClaimReserve 1025–1026）；过期道具转信用点 101（1027）。",
        )
        n += 1

    p_load = find_para(doc, contains="切图 LoadingTicket 防半截进图")
    if p_load is not None and "1028" not in p_load.text:
        replace_paragraph_text(
            p_load,
            p_load.text.strip().rstrip("。")
            + "；靠近传送门强制 ScenePreloadPushScNotify（1028），命中 FADE、未命中 BLACK。",
        )
        n += 1

    p_guide = find_para(doc, startswith="组队协议与多人共战/助战；成就")
    if p_guide is not None and "1036" not in p_guide.text:
        replace_paragraph_text(
            p_guide,
            "组队协议与多人共战/助战；成就即时推送（1036）与材料反查（1037–1038）、"
            "新手引导（检查点续跑 + 一键跳过 CmdId 960–966）、"
            "公会基础 + 公会战 + 公会科技；家园家具交互/在场/拜访日志（858–859、870–874、1030–1035）；"
            "对话 SAVEPOINT 与过场回放（1039–1042）；匹配排队不锁界面 + 10 秒 Ready Check（1043–1045）。",
        )
        n += 1

    p_ai = find_para(doc, contains="AI 外链攻略（B站/抖音/官网 media_links）")
    if p_ai is not None and "BattleAutoService" not in p_ai.text:
        replace_paragraph_text(
            p_ai,
            p_ai.text.strip().rstrip("。")
            + " 战斗 Auto/倍速由 BattleAutoService 真正出手（CmdId 1020–1024，基准等待 1500ms）；"
            "登录协商 input_methods 与 recommended_layout_id（wire v3）。",
        )
        n += 1

    p_gap = find_para(doc, startswith="体验缺口回归：ExperienceGapBusinessFlowsTest")
    if p_gap is not None and find_para(doc, contains="ExperienceLoopBusinessFlowsTest") is None:
        insert_paragraph_after(
            p_gap,
            "体验闭环回归：ExperienceLoopBusinessFlowsTest / ExperienceLoopFillTest 覆盖 Auto 代打、"
            "储备体力、强制 Preload 与 FADE/BLACK、好友在线、成就推送、材料反查、对话 SAVEPOINT、"
            "过场回放、匹配 Ready Check、多端布局；库表增量 db/migration_experience_loop_fill.sql。",
        )
        n += 1

    # ---------- 15.2 局限 ----------
    if replace_if(
        doc,
        startswith="对话/过场/家园：服务端已加深",
        new_text=(
            "对话/过场/家园：服务端已加深（SAVEPOINT/回放、家具交互、在场同步、拜访日志）；"
            "预加载 FADE/BLACK、站内 WebView、扫荡入口、战斗 Auto UI 与 FX 播片、"
            "Ready Check 倒计时、多端布局皮肤仍需客户端表现层对齐。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        exact="协议已迁到 wire v2，旧客户端必须升级；改 CmdId 必须与客户端同步。",
        new_text=(
            "协议已迁到 wire v3（兼容守卫仍 ≥ 2）；v1 客户端必须升级；"
            "改 CmdId（含 1020–1045）必须与客户端同步。"
        ),
    ):
        n += 1

    # ---------- 16.1 短期建议 ----------
    p_short = find_para(doc, startswith="继续以模块化单体打磨崩铁式体验与数值内容")
    if p_short is not None and "Ready Check" not in p_short.text:
        replace_paragraph_text(
            p_short,
            "继续以模块化单体打磨崩铁式体验与数值内容；"
            "客户端对齐 Auto/倍速、储备体力领取、FADE/BLACK 过渡、Ready Check 倒计时、"
            "材料反查跳转、对话续跑/过场回放与多端 recommended_layout_id。",
        )
        n += 1

    # ---------- 结语 ----------
    if replace_if(
        doc,
        startswith="MyLunarCore 已经不是「只有空壳的演示仓库」",
        new_text=(
            "MyLunarCore 已经不是「只有空壳的演示仓库」：它在单体形态下具备较完整的崩铁式 RPG 服务端能力——"
            "联网、场景（含强制预加载与 FADE/BLACK 切图）、回合战（含 Auto/倍速真正出手与打击感 FX）、"
            "养成、经济抽卡、体力与每日循环（含扫荡、满体 30% 储备溢出、过期道具转信用点）、"
            "战令与主线章节、活动任务、对话 SAVEPOINT 与过场回放、家园（含家具交互/在场/拜访日志）、"
            "公会/公会战/公会科技、竞技场评分、世界 BOSS/深渊/遗器词条等缺口补齐玩法、"
            "成就即时推送与材料反查、带检查点/可跳过的新手引导、匹配排队不锁界面与 10 秒 Ready Check、"
            "多端输入协商（wire v3）、运维热更与回滚、业务监控，"
            "以及可选的多节点好友在线与能力增强后的 AI 助手旁路"
            "（多轮对话、个性化、主动推送、熔断灰度、合规声明、专项指标，"
            "以及主流平台攻略默认在站内 WebView 打开并附 ≤200 字摘要）。"
            "与此同时，它也很诚实地把微服务目录标成脚手架，把无缝大世界标成实验，把默认密钥标成仅限本地。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="本报告基于仓库当前代码与文档快照整理",
        new_text=(
            f"本报告基于仓库当前代码与文档快照整理，{today} 已在本文档内同步体验闭环 8 项"
            "（CmdId 1020–1045、PROTOCOL_WIRE_VERSION=3），"
            "可用于汇报、培训与决策讨论。若后续架构或配置有重大变更，建议修订对应章节后刷新目录域，"
            "以便目录页码与正文一致。"
        ),
    ):
        n += 1

    # ---------- 包一览表（表 3）----------
    t3 = doc.tables[3]
    n += patch_table_cell(
        t3,
        old_contains="登录会话、在线状态",
        col=2,
        new_text="登录会话、在线状态、多端 input_methods / 推荐布局、玩家数据加载与落库",
    )
    n += patch_table_cell(
        t3,
        old_contains="回合战、遭遇、波次",
        col=2,
        new_text="回合战、遭遇、波次、BattleAuto 真正出手/倍速、打击感 FX（BattleFxScNotify）",
    )
    n += patch_table_cell(
        t3,
        old_contains="进图、移动、AOI 同步",
        col=2,
        new_text="进图、移动、AOI 同步、Zone 管理、强制 Preload 推送、FADE/BLACK 切图遮罩",
    )
    n += patch_table_cell(
        t3,
        old_contains="道具使用、装备、强化",
        col=2,
        new_text="道具使用、装备、强化、材料反查 QueryItemSource、过期转化信用点",
    )
    n += patch_table_cell(
        t3,
        old_contains="好友、邮件、聊天、排行榜入口",
        col=2,
        new_text="好友、邮件、聊天、排行榜入口、跨节点好友在线（FriendOnlineNotify）",
    )
    n += patch_table_cell(
        t3,
        old_contains="队列、房间、挑战匹配",
        col=2,
        new_text="队列（HALL/QUEUED 不锁界面）、房间、Ready Check、挑战匹配",
    )
    n += patch_table_cell(
        t3,
        old_contains="成就进度与领取",
        col=2,
        new_text="成就进度与领取、达标即时推送 AchievementUnlockPush",
    )
    n += patch_table_cell(
        t3,
        old_contains="NPC 对话树触发、分支进度",
        col=2,
        new_text="NPC 对话树触发、选项 SAVEPOINT、ResumeFromBranch",
    )
    n += patch_table_cell(
        t3,
        old_contains="剧情过场触发、跳过/完成联动",
        col=2,
        new_text="剧情过场触发、跳过/完成联动、已播校验后回放 ticket",
    )
    n += patch_table_cell(
        t3,
        old_contains="基建、产出、家具摆放",
        col=2,
        new_text="基建、产出、家具摆放/交互、好友互访、在场同步、拜访日志/留言",
    )
    n += patch_table_cell(
        t3,
        old_contains="时钟、热更协调、活动、指标、优雅停机、日周刷新、体力",
        col=2,
        new_text="时钟、热更协调、活动、指标、优雅停机、日周刷新、体力与储备溢出",
    )

    # ---------- CmdId 号段表（表 5）----------
    t5 = doc.tables[5]
    n += patch_table_cell(
        t5,
        old_exact="68–83",
        new_other={1: "登录会话（含 input_methods / recommended_layout_id，wire v3）"},
    )
    n += patch_table_cell(
        t5,
        old_exact="200–213",
        new_other={1: "战斗（含打击感 FX；Auto/倍速见 1020–1024）"},
    )
    n += patch_table_cell(
        t5,
        old_exact="352–356",
        new_other={1: "切图加载完成握手 + 场景预加载/就绪通知（强制推送见 1028）"},
    )
    n += patch_table_cell(
        t5,
        old_exact="800–806",
        new_other={1: "匹配（排队不锁界面；Ready Check 见 1043–1045）"},
    )
    n += patch_table_cell(
        t5,
        old_exact="850–859 / 870–874",
        new_other={1: "家园（含家具摆放/交互、在场同步；拜访日志见 1030–1035）"},
    )
    n += patch_table_cell(
        t5,
        old_exact="860–868",
        new_other={1: "对话 / 过场（SAVEPOINT / 回放见 1039–1042）"},
    )
    n += patch_table_cell(
        t5,
        old_exact="950–954",
        new_other={1: "成就（即时推送见 1036）"},
    )
    n += patch_table_cell(
        t5,
        old_exact="996–1000",
        new_other={1: "体力（储备提取见 1025–1026）"},
    )

    extra_cmd_rows = [
        ["1020–1024", "战斗 Auto / 倍速（基准等待 1500ms）"],
        ["1025–1027", "体力溢出提取 / 过期道具转化信用点"],
        ["1028", "场景强制预加载推送（FADE/BLACK）"],
        ["1030–1035", "家园拜访日志 / 留言 / 好友在线"],
        ["1036–1038", "成就即时推送 / 材料反查"],
        ["1039–1042", "对话 SAVEPOINT / 过场回放"],
        ["1043–1045", "匹配 Ready Check（超时 10s）"],
    ]
    if not table_has(t5, "1020"):
        for values in extra_cmd_rows:
            append_table_row(t5, values)
            n += 1

    # ---------- 术语表（表 17）----------
    t17 = doc.tables[17]
    if not table_has(t17, "Ready Check"):
        append_table_row(
            t17,
            [
                "Ready Check",
                "匹配成局后 10 秒确认窗；确认前保持大厅可操作，确认后才 LOCKED/MATCHING",
            ],
        )
        n += 1
    if not table_has(t17, "Battle Auto"):
        append_table_row(
            t17,
            [
                "Battle Auto / 倍速",
                "服务端按启发式真正出手；speed_multiplier 1/2/3 压缩 1500ms 等待；掉线可托管",
            ],
        )
        n += 1
    if not table_has(t17, "SAVEPOINT"):
        append_table_row(
            t17,
            [
                "对话 SAVEPOINT",
                "选项节点存档点；ResumeFromBranch 从分支续跑，markPlayed 不得清掉 savepointNodeId",
            ],
        )
        n += 1

    doc.save(str(DOC_PATH))
    print(f"Updated in place: {DOC_PATH}")
    print(f"Patches applied (approx): {n}")


if __name__ == "__main__":
    main()
