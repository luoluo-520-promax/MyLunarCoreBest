# -*- coding: utf-8 -*-
"""将本轮「体验优化二期」8 项服务端功能同步进桌面《MyLunarCore项目各方面详细总结报告.bak.docx》。

只原地修订该文件，不另存新文档、不生成第二份 .docx。
幂等标记：正文出现 CalculateOptimalSchedule 或 CmdId 1070 即视为已写入本轮。
"""

from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DOC_PATH = Path.home() / "Desktop" / (
    "MyLunarCore"
    + "\u9879\u76ee\u5404\u65b9\u9762\u8be6\u7ec6\u603b\u7ed3\u62a5\u544a"
    + ".bak.docx"
)


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def replace_paragraph_text(paragraph, new_text):
    style = paragraph.style
    for r in list(paragraph.runs):
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        set_run_font(paragraph.runs[0])
    else:
        run = paragraph.add_run(new_text)
        set_run_font(run)
    paragraph.style = style


def insert_paragraph_after(paragraph, text, style_name="List Bullet"):
    new_para = paragraph._parent.add_paragraph(text, style=style_name)
    paragraph._p.addnext(new_para._p)
    for r in new_para.runs:
        set_run_font(
            r,
            size=13 if style_name.startswith("Heading") else 11,
            bold=style_name.startswith("Heading"),
        )
        if style_name.startswith("Heading"):
            r.font.color.rgb = RGBColor(25, 75, 140)
    return new_para


def find_para(doc, exact=None, startswith=None, contains=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if exact is not None and t == exact:
            return p
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    return None


def replace_if(doc, *, exact=None, startswith=None, contains=None, new_text=None):
    p = find_para(doc, exact=exact, startswith=startswith, contains=contains)
    if p is not None and new_text is not None:
        replace_paragraph_text(p, new_text)
        return True
    return False


def set_cell_text(cell, text):
    cell.text = text
    for para in cell.paragraphs:
        for r in para.runs:
            set_run_font(r, size=10)


def patch_table_cell(table, old_exact=None, old_contains=None, col=0, new_text=None, new_other=None):
    for row in table.rows:
        cur = row.cells[col].text.strip()
        hit = False
        if old_exact is not None and cur == old_exact:
            hit = True
        if old_contains is not None and old_contains in cur:
            hit = True
        if not hit:
            continue
        if new_text is not None:
            set_cell_text(row.cells[col], new_text)
        if new_other is not None:
            for ci, val in new_other.items():
                set_cell_text(row.cells[ci], val)
        return True
    return False


def append_table_row(table, values):
    row = table.add_row()
    for i, val in enumerate(values):
        if i < len(row.cells):
            set_cell_text(row.cells[i], val)
    return row


def table_has(table, needle, col=0):
    return any(needle in row.cells[col].text for row in table.rows)


def main():
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    doc = Document(str(DOC_PATH))
    if find_para(doc, contains="CalculateOptimalSchedule") is not None:
        print(f"Already patched (CalculateOptimalSchedule present): {DOC_PATH}")
        return
    if find_para(doc, contains="CmdId 1070") is not None or find_para(doc, contains="1070–1076") is not None:
        print(f"Already patched (1070–1076 present): {DOC_PATH}")
        return

    today = f"{date.today().year}年{date.today().month:02d}月{date.today().day:02d}日"
    n = 0

    # ---------- 封面 ----------
    if replace_if(
        doc,
        exact="（含技术架构、业务能力、安全运维、微服务迁移、体验缺口补齐、体验闭环、体验沉浸与后续建议）",
        new_text=(
            "（含技术架构、业务能力、安全运维、微服务迁移、体验缺口补齐、体验闭环、"
            "体验沉浸、体验优化二期与后续建议）"
        ),
    ):
        n += 1
    elif replace_if(
        doc,
        contains="体验沉浸与后续建议",
        new_text=(
            "（含技术架构、业务能力、安全运维、微服务迁移、体验缺口补齐、体验闭环、"
            "体验沉浸、体验优化二期与后续建议）"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="报告生成日期：",
        new_text=(
            f"报告生成日期：{today}"
            "（已同步体验优化二期 8 项：场景 DISSOLVE/RIFT + 方向预判预加载、"
            "Auto 大招抢占与 estimated_cast_time、跨节点组队跟随迁移与私聊邮箱化、"
            "养成日程表 CalculateOptimalSchedule 与双倍关 is_bonus_today、"
            "剧情 impact_tags 环境联动 SceneEnvironmentModify、"
            "家园助产视觉正反馈与虚影在场、体力溢出提醒场景感知与列车长叙事、"
            "匹配 Ready Check 迷你交互；协议号段 CmdId 1070–1076）"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="报告依据：当前仓库 README",
        new_text=(
            "报告依据：当前仓库 README、源码包结构、配置文件与已有评估文档快照；"
            "此前已写入体验闭环（CmdId 1020–1045）与体验沉浸（1046–1069）；"
            "本轮再把体验优化二期 8 项（CmdId 1070–1076）写入对应章节，"
            "回归测试 ExperienceOptV2FlowsTest / ExperienceImmersionFlowsTest / "
            "ExperienceLoopBusinessFlowsTest。"
        ),
    ):
        n += 1

    # ---------- 7.2 养成日程表 ----------
    if replace_if(
        doc,
        startswith="覆盖创角、成长晋阶、天赋、属性计算",
        new_text=(
            "覆盖创角、成长晋阶、天赋、属性计算，以及皮肤衣柜的拥有与装备。"
            "属性计算由 AttributeCalculator 等组件完成，皮肤配置来自 data/SkinConfigs.json。"
            "养成计划（DevelopmentPlanService）按目标角色+目标等级汇总材料总清单、缺口与体力缺口，"
            "并按掉落效率排序推荐关卡；协议 CalculateUpgradeMaterials（CmdId 1046–1047）。"
            "进一步提供养成日程表 CalculateOptimalSchedule（CmdId 1071–1072）："
            "结合当前体力、储备体力（上限 240）、日购剩余次数与自然恢复，输出今日可完成进度 %"
            "与预计自然回体达标时间，并推送 DevelopmentScheduleScNotify（1073）。"
        ),
    ):
        n += 1

    p_char = find_para(doc, startswith="养成计划：DevelopmentPlanService")
    if p_char is not None and "CalculateOptimalSchedule" not in p_char.text:
        replace_paragraph_text(
            p_char,
            "养成计划：DevelopmentPlanService / CharacterDevelopmentNettyService；"
            "CalculateUpgradeMaterials（1046–1047）；"
            "养成日程表 CalculateOptimalSchedule（1071–1072）+ DevelopmentScheduleScNotify（1073）。",
        )
        n += 1

    # ---------- 7.3 场景过渡升级 ----------
    if replace_if(
        doc,
        startswith="玩家进入某个 Plane/Floor 对应的 Zone 实例后",
        new_text=(
            "玩家进入某个 Plane/Floor 对应的 Zone 实例后，可移动、与 NPC/道具交互，"
            "服务器按兴趣范围（AOI）同步周围实体。Zone 可配置人数上限与租约心跳。"
            "MoveCsReq 带 move_state（走/跑/跳/攀爬）；MovementPhysics 按坐标分区给出地板材质、"
            "触觉强度与脚步声效，JUMP 落地推送 SceneInteractPhysicsScNotify（CmdId 1053）。"
            "接近传送门时服务端强制推送 ScenePreloadPushScNotify（CmdId 1028）；"
            "并基于移动方向提前 2~3 身位触发预判预加载（reason=3），将未命中 BLACK 比例压低。"
            "真正切图时 MigrateSceneScRsp 带 preload_hit 与 SceneLoadMaskInfo："
            "预加载命中用 DISSOLVE（元素溶解）或 RIFT（裂缝穿梭）软过渡并下发 duration_ms，"
            "由客户端着色器在特效窗口内完成加载；未命中仍用 BLACK。"
            "PlayerLoadingStateService 将 FADE/DISSOLVE/RIFT 统一识别为软过渡握手。"
            "跨图或跨节点时，由 Center 规划归属，必要时发放迁移票据。"
        ),
    ):
        n += 1

    p_preload = find_para(doc, startswith="预加载协议：ScenePreloadCsReq")
    if p_preload is not None:
        replace_paragraph_text(
            p_preload,
            "预加载协议：ScenePreloadCsReq/ScRsp、ScenePreloadReadyScNotify（CmdId 354–356）；"
            "服务端强制/方向预判推送 ScenePreloadPushScNotify（1028，reason=1/2/3）；"
            "切图完成 SceneLoadCompleteScRsp.handshake_only；遮罩分级"
            "SceneLoadMaskInfo.transition_type = DISSOLVE | RIFT | FADE | BLACK，含 duration_ms。",
        )
        n += 1

    p_loading = find_para(doc, startswith="切图加载：PlayerLoadingStateService")
    if p_loading is not None and "DISSOLVE" not in p_loading.text:
        replace_paragraph_text(
            p_loading,
            "切图加载：PlayerLoadingStateService（含 handshakeOnly 短 TTL）+ ScenePreloadService"
            "（CmdId 352–353 兼容完成；354–356 客户端预加载；1028 服务端强制/方向预判；"
            "遮罩分级 DISSOLVE / RIFT / FADE / BLACK）。",
        )
        n += 1

    # ---------- 7.4 大招抢占 ----------
    if replace_if(
        doc,
        startswith="场景中触发遭遇后进入战斗实例",
        new_text=(
            "场景中触发遭遇后进入战斗实例：开战、行动、结算、退出。"
            "支持遇敌配置与波次；多人共战通过 BattleContext.participantPlayerIds 等字段表达。"
            "BattleAutoService 按启发式真正出手：基准等待 1500ms，speed_multiplier 仅允许 1/2/3；"
            "SetBattleAuto 增加 auto_strategy（PRIORITY_SKILL / PRIORITY_BASIC / SAVE_ENERGY）。"
            "Auto 中可发 BattleManualUlt（CmdId 1051–1052）插入大招："
            "抢占式指令槽立即暂停当前 1500ms 倒计时并结算；"
            "ScRsp 携带 estimated_cast_time_ms 与 client_pre_fx_ms（默认 200），"
            "客户端可先播 0.2 秒高亮闪屏并提前蓄力，掩盖网络抖动。"
            "掉线进入托管不中断战斗，回主界面才 abort。"
            "协议 SetBattleAuto / BattleAutoScNotify / SetBattleSpeed（CmdId 1020–1024）。"
            "行动经 BattleDeterministicValidator 校验后，BattleFxComposer 按暴击/弱点/击杀组装打击感元数据"
            "（camera_shake、time_scale、damage_popup），经 BattleFxScNotify（CmdId 213）推给客户端播特效；"
            "FxQualityAdvisor 按心跳 RTT/丢包把 fx_quality_level 降为高/中/低，不改变结算权威。"
        ),
    ):
        n += 1

    p_auto = find_para(doc, startswith="Auto/倍速/策略：BattleAutoService")
    if p_auto is not None:
        replace_paragraph_text(
            p_auto,
            "Auto/倍速/策略：BattleAutoService；auto_strategy 优先战技/普攻/攒能量；"
            "BattleManualUlt 抢占式手动大招（1051–1052，estimated_cast_time_ms / client_pre_fx_ms）；"
            "CmdId 1020–1024；掉线托管、回主界面 abort。"
            "特效降级：FxQualityAdvisor + BattleFxScNotify.fx_quality_level。",
        )
        n += 1

    # ---------- 7.8 双倍关置顶 ----------
    if replace_if(
        doc,
        startswith="提供背包快照以及使用、装备、强化",
        new_text=(
            "提供背包快照以及使用、装备、强化、升阶、锁定、丢弃等操作。道具表可由 CSV 等素材导入维护。"
            "材料不足时可 QueryItemSource（CmdId 1037–1038）反查来源：返回 stage_ids，"
            "并指向 FightStartCsReq / SweepStageCsReq 以便直接开战或扫荡；"
            "若掉落关卡恰为当日双倍/加成活动关，置顶并标记 is_bonus_today=true。"
            "高级材料（如 301/401）同时下发 craft_paths 与 craft_cmd=CraftItemCsReq，给出低级合成路径。"
            "过期活动道具由 ExpiredItemRecycleService 按汇率转为信用点（币种 101），"
            "推送 ItemRecycleNotify（1027），禁止静默删除。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="入口：ItemNettyService、ItemApplicationService、ExpiredItemRecycleService",
        new_text=(
            "入口：ItemNettyService、ItemApplicationService、ExpiredItemRecycleService；"
            "材料反查 QueryItemSource（1037–1038，含合成路径 craft_paths 与 is_bonus_today），"
            "过期转化 ItemRecycleNotify（1027）。"
        ),
    ):
        n += 1

    # ---------- 7.10 / 7.11 社交 ----------
    if replace_if(
        doc,
        startswith="大厅模块把社交能力收在一起",
        new_text=(
            "大厅模块把社交能力收在一起：好友、邮件、世界/私聊，以及历史最高分类排行榜。"
            "排行榜与世界聊天在开启 Redis 后可跨进程共享；聊天里 @助手 可转给 AI 模块。"
            "好友在线由 FriendOnlineStatusService 经 Redis online_status_channel 跨节点广播，"
            "上线/下线推送 FriendOnlineNotify（CmdId 1033，含 Plane）。"
            "私聊采用「离线消息邮箱化」+「在线 Push 透传」双重保障："
            "在线走 Redis Pub/Sub，离线写入 offline_chat_message，上线 flushOffline / mailboxUnreadCount 投递。"
            "表情包仓库（EmoteInventoryService）默认发放挥手/跳舞/击掌与点赞贴纸；"
            "聊天 SendChatEmote（1056–1058）与场景 SceneEmote（1059–1061）走 AOI 同步，未拥有贴纸失败。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="Party 是开放世界组队雏形",
        new_text=(
            "Party 是开放世界组队雏形（同节点权威，Redis 租约约 4 人）；"
            "Match/Room 负责副本匹配：排队、成局、可选打分/超时扩展与挑战匹配协调。"
            "跨节点组队已提升为生产级能力：队长邀请成功后，若队员不在队长权威节点，"
            "PartyFollowMigrationService 将队员临时场景实体序列化为迁移票据，"
            "推送 PartyFollowMigrateScNotify（CmdId 1076），实现「跟随迁移」而非要求重新登录。"
            "大厅排队保持会话 HALL（MatchPhase.QUEUED）不锁界面；"
            "成局推送 MatchSuccessScNotify（CmdId 1043）附 10 秒 Ready Check，"
            "并额外推送 MatchCountdownInteractiveScNotify（1075）："
            "允许待机动作点击（踢石子/耍武器/挥手）并展示队友养成名片，填充等待无聊感；"
            "双方 MatchReadyCheckCsReq（1044–1045）确认后才 LOCKED/MATCHING。"
        ),
    ):
        n += 1

    p_party = find_para(doc, startswith="组队：party.PartyService")
    if p_party is not None and "Follow" not in p_party.text:
        replace_paragraph_text(
            p_party,
            "组队：party.PartyService + RedisPartyStore + PartyFollowMigrationService"
            "（跟随迁移 PartyFollowMigrateScNotify 1076）。",
        )
        n += 1

    p_match = find_para(doc, startswith="匹配：MatchNettyService")
    if p_match is not None and "1075" not in p_match.text:
        replace_paragraph_text(
            p_match,
            "匹配：MatchNettyService、MatchmakingService、RoomService、MatchQueue、"
            "ChallengeMatchCoordinator；Ready Check 超时 10000ms（CmdId 1043–1045）；"
            "倒计时迷你交互 MatchCountdownInteractiveScNotify（1075）。",
        )
        n += 1

    # ---------- 7.16 剧情环境 + 家园视觉 ----------
    if replace_if(
        doc,
        startswith="除了战斗与经济循环，项目还补齐了「讲故事」和「家园养成」能力",
        new_text=(
            "除了战斗与经济循环，项目还补齐了「讲故事」和「家园养成」能力，方便策划做剧情与轻养成。"
            "剧情侧在选项节点写入 SAVEPOINT（savepointNodeId），选项可带 impact_tags"
            "（如 +Affinity / UnlockRoute / -Morality）；"
            "GetStoryTreeSnapshot（CmdId 1065–1066）返回节点锁/解锁与当前节点缩略。"
            "当玩家选择含 +Affinity 的选项时，DialogueTriggerEngine 主动下发"
            "SceneEnvironmentModifyScNotify（CmdId 1070），触发 NPC 微表情、爱心粒子与 BGM 变奏，"
            "让选择「肉眼可见」地影响世界。"
            "DialogueProgressService.markPlayed / advance 必须保留存档点，"
            "ResumeFromBranchCsReq（1039–1040）可从分支续跑。"
            "家园除产出、家具摆放与好友互访外，已接线家具交互、AOI-lite 在场同步，以及拜访日志："
            "HomeVisitorLogService 记录来访/留言/助产（CmdId 1030–1035，含设施名称与图标）。"
            "好友助产 HomeHarvestAssist（1067–1068）可为设施减少约 5% 产出 CD（每日每设施一次），"
            "并广播 HomeAssistEffectScNotify（1074）触发设施金光一闪与飘字「xxx 帮你加速了产量！」；"
            "非好友访客以 silhouette 虚影同步，保留「有人在家」的温暖感并控制开销。"
            "产出堆积达阈值时推送 HomeOverflowNotify（1069）。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="对话树：DialogueTriggerEngine、DialogueProgressService",
        new_text=(
            "对话树：DialogueTriggerEngine、DialogueProgressService（选项节点 SAVEPOINT，"
            "markPlayed/advance 保留 savepointNodeId；选项 impact_tags；"
            "+Affinity 触发 SceneEnvironmentModifyScNotify 1070）；"
            "ResumeFromBranch 1039–1040；剧情树快照 GetStoryTreeSnapshot 1065–1066。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="家园：HomeBaseService",
        new_text=(
            "家园：HomeBaseService（产出/家具/互访/好友助产，DB home_state）+ HomePresenceService"
            "（含 silhouette 虚影）、FurnitureInteractHandler、HomeVisitorLogService；"
            "配置 data/HomeFacilityConfigs.json；"
            "协议：摆放 858–859，交互 870–871，在场 872–874，拜访日志/留言/好友在线 1030–1035，"
            "助产/溢出 1067–1069，助产视觉 HomeAssistEffectScNotify 1074。"
        ),
    ):
        n += 1

    # ---------- 7.19 体力场景感知 ----------
    if replace_if(
        doc,
        startswith="这是近期补齐的「日常运营循环」能力",
        new_text=(
            "这是近期补齐的「日常运营循环」能力，让玩家每天都有明确目标，也让数值消耗可控。"
            "体力（Stamina）会自然恢复，打本消耗，也可日购补充；"
            "满体时自然恢复的 30% 进入储备体力（上限 240，StaminaOverflowService），"
            "玩家用 ClaimReserveStamina（CmdId 1025–1026）提取。"
            "体力 ≥80% 且约 10 分钟不活跃时，StaminaOverflowHintService 经 AiHintNotify 提醒；"
            "若玩家处于 STORY / DIALOGUE / BATTLE 状态则推迟推送，回到 HALL 或 SCENE 后再发，"
            "文案改为列车长/帕姆类叙事口吻（source=conductor），降低系统弹窗冰冷感（30 分钟冷却）。"
            "每日任务按目标累计进度并领奖，与日切刷新耦合；战令提供 XP/等级以及免费轨与付费轨奖励；"
            "主线章节记录进度，并作为进入某些 Plane/Floor 的门闸。"
            "版本活动则用树形节点 + 共享 ActivityToken，把一版活动内容组织起来。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="体力：StaminaService + StaminaOverflowService + StaminaOverflowHintService",
        new_text=(
            "体力：StaminaService + StaminaOverflowService + StaminaOverflowHintService；"
            "配置 data/StaminaConfigs.json；满体 30% 进储备（上限 240）；ClaimReserveStamina 1025–1026；"
            "≥80% 且不活跃 10 分钟推送溢出提醒；剧情/对话/战斗中推迟，叙事文案由列车长口吻下发。"
        ),
    ):
        n += 1

    # ---------- 7.22 主路径 ----------
    if replace_if(
        doc,
        startswith="为了让非技术同学建立整体画面",
        new_text=(
            "为了让非技术同学建立整体画面，下面用一条「普通玩家的一天」串起系统："
            "登录进服（校验协议 wire ≥ 2，交换 input_methods 与推荐布局，拉取云端键位）→"
            "看邮件/好友（跨节点在线推送与离线私聊邮箱），在聊天或场景打表情包 →"
            "消耗体力进场景走动（跳跃落地有材质/触觉/声效；接近传送门强制/方向预判预加载；"
            "命中 DISSOLVE/RIFT、未命中 BLACK），"
            "触发 NPC 对话（选项带影响标签，选 +Affinity 时 NPC 微笑/爱心粒子/BGM 变奏，"
            "可看剧情树快照与 SAVEPOINT，过场可回放）→"
            "遇敌打回合战（可开 Auto 并切换优先战技/普攻/攒能量；点大招立即抢拍，"
            "回包带蓄力时长且客户端可先闪屏 0.2 秒；弱网自动降特效档）→"
            "开挑战、Rogue、世界 BOSS 或深渊；养成计划算出材料缺口后可再看日程表"
            "（今日能完成几成、自然回体何时达标），已通关材料关可一键连续扫荡，"
            "材料不足可反查关卡（当日双倍关置顶高亮）或走合成路径 →"
            "满体溢出进储备；看风景/看剧情时不打扰，回到大厅再由列车长提醒刷本；"
            "过期道具转信用点 →"
            "回家园点点基建、摆家具、好友助产时设施金光一闪与飘字、非好友虚影串门、看拜访日志 →"
            "抽卡与商店消费 → 做每日任务/战令/版本活动拿奖，成就达成即时推送 → 推主线章节解锁新地图 →"
            "进公会贡献/兑换/科技/公会战 → 大厅排队匹配不锁界面，成局 10 秒 Ready Check 期间"
            "可点待机动作、看队友名片破冰；跨节点好友邀请可跟随迁移到队长节点 →"
            "必要时问 AI 助手（站内 WebView 打开 B站/抖音/官网攻略）→"
            "新手引导按检查点续跑或一键跳过 → 改设置、同步键位或提工单。"
            "运营同学则在后台改活动排期、导入配置、触发热更；运维同学盯监控与发布演练。"
        ),
    ):
        n += 1

    # ---------- 7.24 专节 ----------
    if find_para(doc, contains="7.24 体验优化二期") is None:
        anchor = find_para(doc, contains="7.23 体验沉浸补齐")
        # 插在 7.23 测试入口段落后、第八章前
        insert_at = find_para(doc, startswith="测试入口：ExperienceImmersionFlowsTest")
        if insert_at is None:
            insert_at = find_para(doc, contains="ExperienceImmersionProtocolFlowsTest")
        if insert_at is None and anchor is not None:
            insert_at = anchor
        if insert_at is None:
            raise SystemExit("未找到 7.24 插入锚点")
        h = insert_paragraph_after(
            insert_at, "7.24 体验优化二期（CmdId 1070–1076）", "Heading 2"
        )
        p1 = insert_paragraph_after(
            h,
            "在体验沉浸（1046–1069）之后，本轮针对箱庭割裂感、大招手感、跨节点社交、养成闭环深度、"
            "剧情即时反馈、家园烟火气、体力打扰感与匹配等待无聊感做了 8 项优化。"
            "协议仍遵循 CS_REQ = N、SC_RSP = N+1；新增推送号段占用 1070–1076。",
            "Normal",
        )
        p2 = insert_paragraph_after(
            p1,
            "场景：预加载命中升级为 DISSOLVE/RIFT（带 duration_ms），方向预判提前 2~3 身位推送 1028。"
            "战斗：BattleManualUlt 抢占 1500ms 倒计时，ScRsp 带 estimated_cast_time_ms / client_pre_fx_ms。"
            "社交：PartyFollowMigrateScNotify（1076）跟随迁移；私聊离线邮箱 + 在线 Push。"
            "养成：CalculateOptimalSchedule（1071–1073）输出今日进度与回体 ETA；"
            "QueryItemSource 增加 is_bonus_today。"
            "剧情：+Affinity 推 SceneEnvironmentModifyScNotify（1070）。"
            "家园：HomeAssistEffectScNotify（1074）金光飘字；非好友 silhouette 虚影。"
            "体力：剧情/对话中推迟 AiHint，列车长叙事文案。"
            "匹配：MatchCountdownInteractiveScNotify（1075）填充待机动作与队友名片。",
            "Normal",
        )
        insert_paragraph_after(
            p2,
            "测试入口：ExperienceOptV2FlowsTest（二期 8 项业务流程全覆盖）、"
            "ExperienceImmersionFlowsTest / ExperienceLoopBusinessFlowsTest（回归）、"
            "ScenePreloadServiceTest（方向预判与 DISSOLVE）。",
            "List Bullet",
        )
        n += 1

    # ---------- 8.4 协议版本 ----------
    if replace_if(
        doc,
        startswith="为避免新旧客户端「各说各话」",
        new_text=(
            "为避免新旧客户端「各说各话」，工程引入了协议线版本概念。"
            "当前 PROTOCOL_WIRE_VERSION = 3（登录握手增加 input_methods / recommended_layout_id）；"
            "ProtocolCompatService.isCompatible 仍接受 wire ≥ 2，拒绝过旧的 v1 客户端"
            "（v1 角色号段 120–129 会撞上组队）。"
            "角色相关命令号已迁到 160–173。体验闭环占用 CmdId 1020–1045，"
            "体验沉浸占用 1046–1069，体验优化二期占用 1070–1076"
            "（docs/meta/cmdid-config-meta.yaml 体验号段可持续扩展）。"
            "因此：改协议必须同步客户端，并尽量用自动化测试（如 CmdIdUniquenessTest）防止号段再重叠。"
            "仓库还提供 buf.yaml，作为 Protobuf lint / breaking 检查基线。"
        ),
    ):
        n += 1

    # ---------- 15.x 清单 ----------
    p_wire = find_para(doc, startswith="确认客户端协议线版本与 CmdId 已对齐 PROTOCOL_WIRE_VERSION=3")
    if p_wire is not None and "1076" not in p_wire.text:
        replace_paragraph_text(
            p_wire,
            "确认客户端协议线版本与 CmdId 已对齐 PROTOCOL_WIRE_VERSION=3"
            "（兼容守卫仍为 wire ≥ 2；v3 字段为 input_methods / recommended_layout_id；"
            "沉浸协议 1046–1069 与优化二期 1070–1076 需与客户端同步）。",
        )
        n += 1

    p_daily = find_para(doc, startswith="日常循环：体力、每日任务、战令")
    if p_daily is not None and "1071" not in p_daily.text:
        replace_paragraph_text(
            p_daily,
            "日常循环：体力、每日任务、战令、主线章节门闸、版本活动容器（DailyLoop/BattlePass 协议已接线）；"
            "已通关关卡扫荡（1/2/3 倍体力，CmdId 1010–1013）与一键连续扫荡（1048–1050）；"
            "养成日程表 CalculateOptimalSchedule（1071–1073）；"
            "满体 30% 溢出进储备（上限 240，ClaimReserve 1025–1026）；过期道具转信用点 101（1027）；"
            "体力 ≥80% 不活跃提醒（场景感知推迟 + 列车长叙事）。",
        )
        n += 1

    p_immersion = find_para(doc, startswith="体验沉浸回归：ExperienceImmersionFlowsTest")
    if p_immersion is not None and find_para(doc, contains="ExperienceOptV2FlowsTest") is None:
        insert_paragraph_after(
            p_immersion,
            "体验优化二期回归：ExperienceOptV2FlowsTest 覆盖 DISSOLVE/RIFT 与方向预判、"
            "大招抢占与 estimated_cast_time、跟随迁移、养成日程与双倍关、剧情环境联动、"
            "家园助产特效与虚影、体力场景感知、匹配迷你交互（CmdId 1070–1076）。",
        )
        n += 1

    p_guide = find_para(doc, startswith="组队协议与多人共战/助战")
    if p_guide is not None and "1070" not in p_guide.text:
        replace_paragraph_text(
            p_guide,
            p_guide.text.strip().rstrip("。")
            + "；体验优化二期：场景环境联动（1070）、养成日程（1071–1073）、"
            "家园助产特效（1074）、匹配迷你交互（1075）、组队跟随迁移（1076）。",
        )
        n += 1

    # ---------- 15.2 ----------
    if replace_if(
        doc,
        startswith="对话/过场/家园：服务端已加深",
        new_text=(
            "对话/过场/家园：服务端已加深（SAVEPOINT/回放/剧情树快照/环境联动、家具交互、在场虚影、"
            "拜访日志、好友助产视觉）；"
            "预加载 DISSOLVE/RIFT/BLACK、站内 WebView、扫荡/连续扫荡/日程表入口、"
            "战斗 Auto 策略 UI 与大招预表现、FX 播片与降档、"
            "Ready Check 迷你交互、多端布局与云端键位、表情动作表现仍需客户端对齐。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="协议已迁到 wire v3（兼容守卫仍 ≥ 2）",
        new_text=(
            "协议已迁到 wire v3（兼容守卫仍 ≥ 2）；v1 客户端必须升级；"
            "改 CmdId（含 1020–1045 闭环、1046–1069 沉浸与 1070–1076 优化二期）必须与客户端同步。"
        ),
    ):
        n += 1

    # ---------- 16.1 ----------
    p_short = find_para(doc, startswith="继续以模块化单体打磨崩铁式体验与数值内容")
    if p_short is not None and "DISSOLVE" not in p_short.text:
        replace_paragraph_text(
            p_short,
            "继续以模块化单体打磨崩铁式体验与数值内容；"
            "客户端对齐 DISSOLVE/RIFT 切图、大招预表现、跟随迁移、养成日程表、双倍关高亮、"
            "剧情环境反馈、家园助产飘字与虚影、列车长体力提醒、Ready Check 迷你交互，"
            "以及既有养成计划/连续扫荡、Auto 策略、云端键位与 recommended_layout_id。",
        )
        n += 1

    # ---------- 结语 ----------
    if replace_if(
        doc,
        startswith="MyLunarCore 已经不是「只有空壳的演示仓库」",
        new_text=(
            "MyLunarCore 已经不是「只有空壳的演示仓库」：它在单体形态下具备较完整的崩铁式 RPG 服务端能力——"
            "联网、场景（含强制/方向预判预加载、DISSOLVE/RIFT/BLACK 切图与跳跃落地物理反馈）、"
            "回合战（含 Auto 策略权重、抢占式手动大招与蓄力预估、倍速真正出手、打击感 FX 与弱网降档）、"
            "养成（含目标导向材料计划、日程表与一键连续扫荡）、经济抽卡、"
            "体力与每日循环（含扫荡、满体 30% 储备溢出、场景感知溢出提醒、过期道具转信用点与合成路径/"
            "双倍关置顶）、"
            "战令与主线章节、活动任务、对话 SAVEPOINT / impact_tags / 环境联动与过场回放、"
            "家园（含家具交互/虚影在场/拜访日志/助产视觉）、公会/公会战/公会科技、竞技场评分、"
            "世界 BOSS/深渊/遗器词条等缺口补齐玩法、成就即时推送与材料反查、"
            "带检查点/可跳过的新手引导、匹配排队不锁界面与 Ready Check 迷你交互、"
            "跨节点组队跟随迁移与私聊邮箱、表情包、键位云同步、多端输入协商（wire v3）、"
            "运维热更与回滚、业务监控，"
            "以及可选的多节点好友在线与能力增强后的 AI 助手旁路"
            "（多轮对话、个性化、主动推送、熔断灰度、合规声明、专项指标，"
            "以及主流平台攻略默认在站内 WebView 打开并附 ≤200 字摘要）。"
            "与此同时，它也很诚实地把微服务目录标成脚手架，把无缝大世界标成实验，把默认密钥标成仅限本地。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="本报告基于仓库当前代码与文档快照整理",
        new_text=(
            f"本报告基于仓库当前代码与文档快照整理，{today} 已在本文档内同步体验优化二期 8 项"
            "（CmdId 1070–1076；此前沉浸为 1046–1069、闭环为 1020–1045、PROTOCOL_WIRE_VERSION=3），"
            "可用于汇报、培训与决策讨论。若后续架构或配置有重大变更，建议修订对应章节后刷新目录域，"
            "以便目录页码与正文一致。"
        ),
    ):
        n += 1

    # ---------- 表 3 包一览 ----------
    t3 = doc.tables[3]
    n += int(
        patch_table_cell(
            t3,
            old_contains="养成计划材料计算器",
            col=2,
            new_text="创角、晋阶、天赋、属性计算、养成计划材料计算器、养成日程表",
        )
    )
    n += int(
        patch_table_cell(
            t3,
            old_contains="BattleAuto 策略/倍速/手动大招",
            col=2,
            new_text="回合战、遭遇、波次、BattleAuto 策略/倍速/抢占式手动大招、打击感 FX 与弱网降档",
        )
    )
    n += int(
        patch_table_cell(
            t3,
            old_contains="强制 Preload、FADE/BLACK",
            col=2,
            new_text="进图、移动状态机、落地物理反馈、AOI 同步、Zone 管理、强制/方向预判 Preload、DISSOLVE/RIFT/BLACK",
        )
    )
    n += int(
        patch_table_cell(
            t3,
            old_contains="合成路径、过期转化信用点",
            col=2,
            new_text="道具使用、装备、强化、材料反查 QueryItemSource（含 is_bonus_today）、合成路径、过期转化信用点",
        )
    )
    n += int(
        patch_table_cell(
            t3,
            old_contains="表情包仓库",
            col=2,
            new_text="好友、邮件、聊天（离线邮箱+在线 Push）、排行榜入口、跨节点好友在线、表情包仓库",
        )
    )
    n += int(
        patch_table_cell(
            t3,
            old_contains="impact_tags、剧情树快照",
            col=2,
            new_text="NPC 对话树、impact_tags、环境联动、剧情树快照、选项 SAVEPOINT、ResumeFromBranch",
        )
    )
    n += int(
        patch_table_cell(
            t3,
            old_contains="好友互访/助产",
            col=2,
            new_text="基建、产出、家具摆放/交互、好友互访/助产视觉、虚影在场、拜访日志/留言、溢出推送",
        )
    )

    # ---------- 表 5 CmdId ----------
    t5 = doc.tables[5]
    n += int(
        patch_table_cell(
            t5,
            old_exact="1046–1047",
            new_other={1: "养成计划材料计算器 CalculateUpgradeMaterials（日程表见 1071–1073）"},
        )
    )
    n += int(
        patch_table_cell(
            t5,
            old_exact="1051–1052",
            new_other={1: "Auto 中抢占式手动大招 BattleManualUlt（含 estimated_cast_time）"},
        )
    )
    n += int(
        patch_table_cell(
            t5,
            old_exact="1067–1069",
            new_other={1: "家园助产 / 产出溢出推送（助产视觉见 1074）"},
        )
    )
    extra_cmd_rows = [
        ["1070", "剧情环境联动 SceneEnvironmentModifyScNotify"],
        ["1071–1072", "养成日程表 CalculateOptimalSchedule"],
        ["1073", "养成日程推送 DevelopmentScheduleScNotify"],
        ["1074", "家园助产视觉 HomeAssistEffectScNotify"],
        ["1075", "匹配倒计时迷你交互 MatchCountdownInteractiveScNotify"],
        ["1076", "跨节点组队跟随迁移 PartyFollowMigrateScNotify"],
    ]
    if not table_has(t5, "1070"):
        for values in extra_cmd_rows:
            append_table_row(t5, values)
            n += 1

    # ---------- 表 17 术语 ----------
    t17 = doc.tables[17]
    extra_terms = [
        ["DISSOLVE / RIFT", "切图软过渡：场景元素溶解或动态裂缝穿梭，加载在特效窗口内完成，替代生硬黑屏"],
        ["方向预判预加载", "沿玩家移动方向提前 2~3 身位触发目标 Plane 预加载，降低 BLACK 比例"],
        ["大招抢占", "Auto 倒计时未结束时点大招立即打断等待并结算，回包带蓄力时长供客户端预表现"],
        ["跟随迁移", "跨节点组队时把队员临时实体序列化迁到队长节点，无需重新登录"],
        ["养成日程表", "结合体力/储备/日购/自然恢复，告诉玩家今日能完成几成、何时自然回体达标"],
        ["is_bonus_today", "材料反查时标记当日双倍关并置顶高亮"],
    ]
    for term, expl in extra_terms:
        if not table_has(t17, term[:6]):
            append_table_row(t17, [term, expl])
            n += 1

    # ---------- 表 18 FAQ ----------
    t18 = doc.tables[18]
    extra_faq = [
        [
            "切图为什么还有黑屏？",
            "只有预加载未命中才走 BLACK；命中时用 DISSOLVE/RIFT。靠近传送门并朝向移动会提前预加载。",
        ],
        [
            "点大招为什么偶尔没反应？",
            "已改为抢占式：收到 BattleManualUlt 立即打断 Auto 等待；客户端可用 0.2 秒闪屏预表现掩盖延迟。",
        ],
        [
            "好友不在同一节点能组队吗？",
            "能。邀请成功后走跟随迁移（1076），队员实体迁到队长节点，不必重新登录。",
        ],
        [
            "材料缺口多久能刷完？",
            "用 CalculateOptimalSchedule 看今日进度 % 与自然回体达标时间；双倍关会在反查里置顶。",
        ],
    ]
    for q, a in extra_faq:
        if not table_has(t18, q[:6]):
            append_table_row(t18, [q, a])
            n += 1

    doc.save(str(DOC_PATH))
    print(f"Updated in place: {DOC_PATH}")
    print(f"Patches applied (approx): {n}")


if __name__ == "__main__":
    main()
