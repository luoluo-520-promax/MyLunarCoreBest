# -*- coding: utf-8 -*-
"""修复十大缺口报告补丁：将误插到目录区的 7.31 移到正文 7.30 之后，并清理目录污染。

仍只原地修订 Desktop 原 .bak.docx，不另存新文件。
"""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DOC_PATH = Path.home() / "Desktop" / (
    "MyLunarCore"
    + "\u9879\u76ee\u5404\u65b9\u9762\u8be6\u7ec6\u603b\u7ed3\u62a5\u544a"
    + ".bak.docx"
)


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def replace_paragraph_text(paragraph, new_text):
    style = paragraph.style
    for r in list(paragraph.runs):
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        set_run_font(paragraph.runs[0])
    else:
        run = paragraph.add_run(new_text)
        set_run_font(run)
    paragraph.style = style


def delete_paragraph(paragraph):
    el = paragraph._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def insert_paragraph_after(paragraph, text, style_name="List Bullet"):
    new_para = paragraph._parent.add_paragraph(text, style=style_name)
    paragraph._p.addnext(new_para._p)
    for r in new_para.runs:
        set_run_font(
            r,
            size=13 if style_name.startswith("Heading") else 11,
            bold=style_name.startswith("Heading"),
        )
        if style_name.startswith("Heading"):
            r.font.color.rgb = RGBColor(25, 75, 140)
    return new_para


def main():
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    doc = Document(str(DOC_PATH))
    n = 0

    # 1) 清理目录区误插的 7.31（位于 TOC「八、网络」与「九、数据」之间）
    to_delete = []
    in_bad_block = False
    for p in doc.paragraphs:
        t = p.text.strip()
        style = p.style.name if p.style else ""
        if t == "八、网络与通信协议" and style == "Normal":
            # 下一组若出现误插 Heading 2 7.31，开始标记
            in_bad_block = "pending"
            continue
        if in_bad_block == "pending" and style.startswith("Heading") and "7.31" in t:
            in_bad_block = True
            to_delete.append(p)
            continue
        if in_bad_block is True:
            if t.startswith("九、数据存储") and style == "Normal":
                in_bad_block = False
                continue
            if "7.31" in t or t.startswith("在 7.30 wire v4") or t.startswith("实现入口：tools/PresentationMockClient") \
                    or t.startswith("测试入口：ExperienceGapFillFlowTest") \
                    or t.startswith("客户端联调要点：优先消费 P0"):
                to_delete.append(p)
                continue
            # 其他则结束
            in_bad_block = False

    for p in to_delete:
        delete_paragraph(p)
        n += 1

    # 重新加载段落视图（删除后需新 Document 更稳妥）
    doc.save(str(DOC_PATH))
    doc = Document(str(DOC_PATH))

    # 2) 修复被污染的目录项「十四、微服务…」
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("十四、微服务脚手架现状") and "UserDataService" in t:
            replace_paragraph_text(p, "十四、微服务脚手架现状")
            n += 1
            break

    # 3) 在真正的 Heading「十三、安全机制」后写隐私说明（若正文尚未写）
    heading13 = None
    for p in doc.paragraphs:
        if p.style and p.style.name.startswith("Heading") and p.text.strip().startswith("十三、安全机制"):
            heading13 = p
            break
    if heading13 is not None:
        # 找紧随其后的正文
        body13 = None
        seen = False
        for p in doc.paragraphs:
            if p._element is heading13._element:
                seen = True
                continue
            if not seen:
                continue
            if p.style and p.style.name.startswith("Heading"):
                break
            if p.text.strip():
                body13 = p
                break
        if body13 is not None and "UserDataService" not in body13.text:
            replace_paragraph_text(
                body13,
                body13.text.strip()
                + " 隐私侧已补 UserDataService（导出/软删）、PrivacyConsentService（不同意则访客不可存档）、"
                "StructuredJsonLogger UID 哈希化（保留前缀关联）。",
            )
            n += 1

    # 4) 若正文尚无正确位置的 7.31，插到 7.30 末段之后、Heading「八、网络」之前
    has_body_731 = False
    for p in doc.paragraphs:
        if p.style and p.style.name.startswith("Heading") and "7.31" in p.text:
            # 确认不在目录附近：后面应紧跟 wire v4 正文或实现入口，且前方有 7.30
            has_body_731 = True
            break

    # 更严：7.31 Heading 必须出现在 7.30 Heading 之后
    idx_730 = idx_731 = None
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if p.style and p.style.name.startswith("Heading") and t.startswith("7.30"):
            idx_730 = i
        if p.style and p.style.name.startswith("Heading") and "7.31" in t:
            idx_731 = i
    if idx_731 is not None and idx_730 is not None and idx_731 > idx_730:
        has_body_731 = True
    else:
        has_body_731 = False
        # 删掉仍残留在错误位置的 7.31（若清理未净）
        if idx_731 is not None and (idx_730 is None or idx_731 < idx_730):
            # 删除从该 Heading 起连续误插段
            kill = []
            started = False
            for p in doc.paragraphs:
                t = p.text.strip()
                if p._element is doc.paragraphs[idx_731]._element or (
                    started is False and p.style and p.style.name.startswith("Heading") and "7.31" in t
                ):
                    started = True
                    kill.append(p)
                    continue
                if started:
                    if (p.style and p.style.name.startswith("Heading") and not t.startswith("7.31")) \
                            or (t.startswith("九、") and p.style and p.style.name == "Normal") \
                            or (t.startswith("八、") and p.style and p.style.name.startswith("Heading")):
                        break
                    if t.startswith("在 7.30") or t.startswith("实现入口：tools/") \
                            or t.startswith("测试入口：ExperienceGap") \
                            or t.startswith("客户端联调要点：优先消费"):
                        kill.append(p)
                    else:
                        break
            for p in kill:
                delete_paragraph(p)
                n += 1
            doc.save(str(DOC_PATH))
            doc = Document(str(DOC_PATH))

    if not has_body_731:
        anchor = None
        for p in doc.paragraphs:
            t = p.text.strip()
            if t.startswith("客户端联调要点：上报 wire_version") and "PROTOCOL_CAPABILITY" in t:
                anchor = p
                break
        if anchor is None:
            raise SystemExit("未找到 7.30 客户端联调锚点")

        # 倒序插入，使最终顺序为：标题 → 正文 → 实现 → 测试 → 联调要点
        insert_paragraph_after(
            anchor,
            "客户端联调要点：优先消费 P0——SceneLoadMaskInfo.transition_type（DISSOLVE/RIFT）、"
            "client_pre_fx_ms、BattleFx 镜头冲击、CancelWindow（1130）、Reconnect（1138–1139）；"
            "P1 含 QTE/Ping/语音/DPS；Mock 预览 GET /api/admin/mock/presentation-preview；"
            "清单 docs/client-integration-priority.md。",
            style_name="List Bullet",
        )
        insert_paragraph_after(
            anchor,
            "测试入口：ExperienceGapFillFlowTest、ExperienceGapE2EFlowTest（端到端 A–J）、"
            "ClientProtocolIntegrationTest；CI 已纳入上述用例；"
            "与 StabilityEnhancementFlowTest / WireV4EnhancementBusinessFlowsTest 等合计定向回归 "
            "184 通过 / 0 失败。",
            style_name="List Bullet",
        )
        insert_paragraph_after(
            anchor,
            "实现入口：tools/PresentationMockClient；economy/WalletTccSagaService、"
            "WalletReconciliationJob；party/PartyMigrateSagaDeadLetterService；"
            "battle/CancelWindowService、BattleQteService、BattleFxComposer；"
            "scene/ScenePingService；guild/RaidDamageStatisticsService；"
            "ops/GradedCircuitBreakerService；privacy/UserDataService、PrivacyConsentService；"
            "common/ConfigGrayAutoRollbackService、StructuredJsonLogger；"
            "admin/ConfigChangeApprovalController；player/FastReconnectService；"
            "settings/UiAdaptConfigService；activity/VersionThemeService；"
            "battlepass/BattlePassThemeService；deploy/grafana/dashboards/sla-health.json。",
            style_name="List Bullet",
        )
        body = (
            "在 7.30 wire v4（CmdId 1109–1125）之后，本轮按产品报告中的十大短板一次性补齐服务端能力，"
            "新增号段 1130–1142（推送与确认为主）。"
            "（1）客户端联调：docs/client-integration-priority.md 按 P0/P1/P2 分级；"
            "PresentationMockClient 生成 DISSOLVE/RIFT/BattleFx/取消窗/QTE 样本；"
            "CI 用 ClientProtocolIntegrationTest 校验 transition_type 等解析。"
            "（2）资产安全：WalletTccSagaService 做 TCC 预留/确认/回滚；"
            "PartyMigrateSagaDeadLetterService 提供 Retry+DeadLetter+混沌注入；"
            "WalletReconciliationJob 日对账并可一键修正。"
            "（3）战斗手感动作化：取消后摇窗口、破盾/击杀 QTE、镜头 FOV/震屏强度下发。"
            "（4）公会协作：Raid 语音房、战术 Ping、实时 DPS 看板。"
            "（5）运维细粒度：GradedCircuitBreaker 按活动/渠道/分片熔断；"
            "BusinessMetrics 暴露 match_queue_depth / battle_pool_usage / SLA P95；"
            "Grafana sla-health 红绿灯看板；K8s HPA 说明见 deploy/k8s/hpa-custom-metrics.md。"
            "（6）隐私合规：导出/软删、登录同意与访客模式、日志 UID 哈希（保留前缀）+ 30 天保留提示。"
            "（7）配置治理：SensitiveOp 双人审批 CONFIG_CHANGE、灰度错误率自动回滚、"
            "admin/config-versions.html JSON 差异高亮。"
            "（8）弱网重连：FastReconnectService 缓存 30s Entity 快照，断线战斗进 AWAITING_RECONNECT 最长 30s，"
            "probe 支持前台预连接。"
            "（9）跨端：按分辨率下发 ui_scale 与 TouchHeatmap（键位仍走 KeyBindCloudService）。"
            "（10）长线运营：VersionThemeService 主题生命周期 + BattlePassThemeService 赛季绑定，"
            "配置 data/VersionThemeConfigs.json。"
        )
        insert_paragraph_after(anchor, body, style_name="Normal")
        insert_paragraph_after(
            anchor,
            "7.31 十大体验/生产缺口补齐（CmdId 1130–1142）",
            style_name="Heading 2",
        )
        n += 5

    doc.save(str(DOC_PATH))
    print(f"Repaired {n} places -> {DOC_PATH}")


if __name__ == "__main__":
    main()
