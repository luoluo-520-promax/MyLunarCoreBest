# -*- coding: utf-8 -*-
"""
将「二游运营热更七项短板」落地能力同步进桌面原 bak 总结报告，覆盖保存原文件（不另存新文档）。

目标文件：~/Desktop/MyLunarCore项目各方面详细总结报告.bak.docx

覆盖：
1. 活动 Groovy 脚本热更
2. 配置生效窗 + Time Travel
3. 资源 Manifest 版本锁（RESOURCE_OUTDATED / Cmd 语义）
4. Override 热补丁 + 活动全局熔断
5. Server Freeze + 优雅停机强制落盘
6. QA 可见性掩码
7. 负向发放 + 钱包回滚快照
"""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DOC_PATH = Path.home() / "Desktop" / "MyLunarCore项目各方面详细总结报告.bak.docx"

MARKER = "CommercialOpsBusinessFlowTest"
SECTION_TITLE = "7.29 二游运营热更七项短板补齐（脚本/生效窗/资源锁/热补丁/Freeze/QA/负向发放）"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def replace_paragraph_text(paragraph, new_text, bold=False):
    style = paragraph.style
    for r in list(paragraph.runs):
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        set_run_font(paragraph.runs[0], bold=bold)
    else:
        run = paragraph.add_run(new_text)
        set_run_font(run, bold=bold)
    paragraph.style = style


def insert_paragraph_after(paragraph, text, style_name="List Bullet"):
    new_para = paragraph._parent.add_paragraph(text, style=style_name)
    paragraph._p.addnext(new_para._p)
    for r in new_para.runs:
        set_run_font(
            r,
            size=13 if style_name.startswith("Heading") else 11,
            bold=style_name.startswith("Heading"),
            color=RGBColor(25, 75, 140) if style_name.startswith("Heading") else None,
        )
    return new_para


def find_para(doc, exact=None, startswith=None, contains=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if exact is not None and t == exact:
            return p
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    return None


def already_has(doc, needle):
    return any(needle in (p.text or "") for p in doc.paragraphs)


def main():
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    doc = Document(str(DOC_PATH))
    n = 0

    # ---------- 副标题 ----------
    p = find_para(doc, startswith="（含技术架构、业务能力")
    if p and "运营热更七项" not in p.text:
        replace_paragraph_text(
            p,
            "（含技术架构、业务能力、安全运维、微服务迁移、体验缺口补齐、体验闭环、体验沉浸、"
            "体验优化二期/三期、AI 助手增强、战斗手感与拥堵防雪崩、性能深化，"
            "以及运营热更七项短板补齐与后续建议）",
        )
        n += 1

    # ---------- 报告生成日期 ----------
    p = find_para(doc, startswith="报告生成日期")
    if p and "运营热更七项" not in p.text:
        replace_paragraph_text(
            p,
            "报告生成日期：2026年09月02日（在性能深化 CmdId 1107 与手感/拥堵 1105–1106 基础上，"
            "已同步「二游运营热更七项短板」落地："
            "①活动 Groovy 脚本热更（data/scripts/activity + ActivityScriptEngine，/reload 生效）；"
            "②配置生效窗 display/effect + ServerClock Time Travel；"
            "③资源 Manifest 锁，登录 retcode=RESOURCE_OUTDATED(11)；"
            "④Override 热补丁 /api/admin/ops/hotfix/override + 活动全局熔断 basic_mail；"
            "⑤Server Freeze（SERVER_MAINTENANCE_PUSH=1108）+ 优雅停机强制同步落盘；"
            "⑥QA 可见性掩码（QaTesterUids / is_qa_tester，不进正式排行榜）；"
            "⑦负向发放 + 钱包版本回滚快照。"
            "回归 CommercialOpsBusinessFlowTest / CommercialOpsEnhancementTest）",
        )
        n += 1

    # ---------- 报告依据 ----------
    p = find_para(doc, startswith="报告依据：当前仓库 README")
    if p and MARKER not in p.text:
        base = p.text.rstrip("。").rstrip()
        replace_paragraph_text(
            p,
            base
            + "；本轮再将运营热更七项写入 7.1/7.9/7.29、第九/十一/十二/十五/十六章与附录，"
            "业务流回归 CommercialOpsBusinessFlowTest（嵌套 7+1 流程）及 "
            "CommercialOpsEnhancementTest、ActivityTemplate/Query、ClientVersionGate、"
            "HotReloadCoordinator、PublishDrill 等相关用例。",
        )
        n += 1

    # ---------- 一句话定位 ----------
    p = find_para(doc, startswith="MyLunarCore 是一套基于 Spring Boot 4")
    if p and "活动脚本热更" not in p.text:
        replace_paragraph_text(
            p,
            "MyLunarCore 是一套基于 Spring Boot 4 + Netty/KCP + Protobuf + MySQL 的"
            "「崩坏：星穹铁道式」游戏私服 / 自研服务端：玩家在分区场景里探索，遇敌后进入回合制战斗；"
            "同时提供抽卡、商店支付、活动任务、组队匹配、排行榜聊天，以及管理后台热更与可选多节点切服。"
            "真正可运行的主业务在根目录单体工程；microservices/ 是迁移脚手架。"
            "运营侧已补齐「头部二游」常见热更能力：活动独有逻辑可用 Groovy 脚本热更、"
            "未来版本配置可 Time Travel 预演、登录强校验客户端资源 Manifest、"
            "点路径 Override 热补丁与活动熔断、发版 Server Freeze、QA 隐身可见性与负向追缴/钱包回滚点。",
        )
        n += 1

    # ---------- 目录速览：追加 7.29 ----------
    p = find_para(doc, exact="七、核心业务能力详解")
    # 在自动目录前的章节速览里不一定逐条列 7.x；正文新增 Heading 即可，Word 更新域后出现
    # 若有「7.28」相关速览句则追加
    p = find_para(doc, contains="7.28 AI 沉浸增强")
    if p and "7.29" not in p.text and p.style and p.style.name == "Normal":
        # don't touch long body; section insert handles it
        pass

    # ---------- 7.1 登录 ----------
    p = find_para(doc, startswith="玩家用账号密码登录后，服务器建立会话")
    if p and "RESOURCE_OUTDATED" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "。登录门禁在协议线版本与 ClientRequiredVersions 之外，"
            "增加 ClientResourceManifestService 资源版本锁："
            "客户端资源哈希/版本落后于服务端 Manifest 时返回 retcode=11（RESOURCE_OUTDATED），"
            "并下发 required_client_res_version / client_update_url，强制跳转更新页，"
            "避免下发无法渲染的新角色/碰撞网格导致闪退。",
        )
        n += 1

    # ---------- 7.9 活动 ----------
    p = find_para(doc, startswith="活动侧有排期与配置查询")
    if p and "ActivityScriptEngine" not in p.text:
        replace_paragraph_text(
            p,
            "活动侧有排期与配置查询，支持热更推送，也支持 Excel/JSON 导入。"
            "独有计算逻辑（如签到第 7 天多倍、限时挑战特殊积分）可抽到 "
            "data/scripts/activity/*.groovy，由 ActivityScriptEngine（Groovy JSR-223）编译执行；"
            "ActivityTemplateService.invokeScript 接入，/reload 与 HotReloadCoordinator 会重载脚本，"
            "无需为新活动机制发版 Java。"
            "配置增加 displayStart/effectStart/displayEnd/effectEnd 生效窗："
            "可提前展示入口但未到 effect 不可玩；ServerClockService 支持 Time Travel 预演未来时间。"
            "ActivityVisibilityService + data/QaTesterUids.json：QA 隐身账号可看未正式开启活动，"
            "操作不记入正式排行榜；普通玩家 GetActivityInfo 对未开始活动返回 RET_NOT_STARTED。"
            "任务侧支持接取、提交、放弃，以及进度与触发引擎，方便做引导与运营目标。",
        )
        n += 1

    p = find_para(doc, startswith="活动：ActivityNettyService")
    if p and "scriptId" not in p.text:
        replace_paragraph_text(
            p,
            "活动：ActivityNettyService、ActivityScheduleService、ActivityConfigService、"
            "ActivityScriptEngine、ActivityVisibilityService、ActivityCircuitBreakerService；"
            "data/Activity*.json、data/activities/、data/scripts/activity/、data/QaTesterUids.json",
        )
        n += 1

    p = find_para(doc, startswith="活动模板：ActivityTemplateService")
    if p and "invokeScript" not in p.text:
        replace_paragraph_text(
            p,
            "活动模板：ActivityTemplateService + ActivityTypeCatalog（转盘/拼图/积分兑换）+ "
            "invokeScript（脚本扩展机制）；熔断时限时玩法返回 503，奖励可降级 basic_mail。",
        )
        n += 1

    # ---------- 新增 7.29 章节（插在 7.28 测试 bullet 之后、第八章之前）----------
    if not already_has(doc, SECTION_TITLE):
        anchor = find_para(doc, startswith="测试：AssistImmersionIntegrationTest")
        if anchor is None:
            anchor = find_para(doc, startswith="在 7.26 截图 VLM")
        if anchor is None:
            anchor = find_para(doc, exact="八、网络与通信协议")
            # insert before chapter 8 by inserting after previous para is harder; use contains
        if anchor is not None:
            h = insert_paragraph_after(anchor, SECTION_TITLE, "Heading 2")
            cur = h
            cur = insert_paragraph_after(
                cur,
                "本轮针对报告中「活动机制硬编码、缺未来版本沙盒、客户端资源与服耦合、"
                "紧急屏蔽粒度粗、版本更新断线落盘风险、QA 灰度不彻底、刷产物回溯弱」七项短板，"
                "在单体工程落地可运营闭环（能力摘要如下）。",
                "Normal",
            )
            items = [
                "（1）活动脚本热更：策划 Excel + 几行 Groovy 放入 data/scripts/activity/，"
                "HotReloadCoordinator.reloadAll 重编译；示例 sign_in_bonus / challenge_score。",
                "（2）生效窗 + Time Travel：display_* 控制展示，effect_* 控制可玩；"
                "Admin /api/admin/ops/time-travel 可将内部机逻辑时钟拨到未来版本日做零延迟切换预演。",
                "（3）资源版本锁：登录握手校验客户端资源版本/哈希相对服务端 Manifest；"
                "落后返回 RESOURCE_OUTDATED(11)+更新 URL，阻断无法渲染的数据下发。",
                "（4）Override 热补丁 + 熔断：/api/admin/ops/hotfix/override 按 configId+JSON Path "
                "覆写（如 gacha.base_probability=0.006），落盘 hotfix-overrides.json，避免整文件回滚二次故障；"
                "/circuit/trip 一键关闭限时玩法并切 basic_mail 奖励保服。",
                "（5）Server Freeze：发版脚本先 freeze（推送 1108、拒新战斗、同步战斗快照与 WAL、清 session_dirty），"
                "再 escalate 到维护拒登录；GracefulShutdownCoordinator 停机前强制同步落盘，"
                "杜绝「仅 saveAsync + 快速杀进程」丢最后一击。",
                "（6）QA 可见性掩码：Session.isQaTester；未开启活动对普通玩家隐藏/未开始，"
                "QA 返回完整配置并可交互预演；excludeFromLeaderboard 保证操作不污染正式榜。",
                "（7）负向发放与钱包回滚点：Admin POST /api/admin/ops/grant/negative 支持 uid+item_id+负数 count，"
                "强制扣道具/货币并推送 ItemChange / ItemRecycle（ItemSubtract 语义）与货币变更；"
                "WalletRollbackSnapshotService 在大版本前打 data/wallet-snapshots/{label}.json，"
                "事故时可 dryRun 对比或 confirm=ROLLBACK_WALLET 回退货币状态（需配套公告补偿）。",
                "实现入口：ActivityScriptEngine、ConfigActiveWindow、ServerClockService、"
                "ClientResourceManifestService.checkResourceLock、ConfigOverrideHotfixService、"
                "ActivityCircuitBreakerService、ServerFreezeService、ActivityVisibilityService、"
                "NegativeGrantService、WalletRollbackSnapshotService、CommercialOpsController；"
                "CmdIds.SERVER_MAINTENANCE_PUSH=1108；发布脚本 scripts/release/server_freeze_publish.ps1。",
                "测试：CommercialOpsBusinessFlowTest（7 项嵌套流程 + 发版日组合流水线）、"
                "CommercialOpsEnhancementTest；相关回归 ActivityTemplateServiceTest / "
                "ActivityQueryServiceTest / ClientVersionGateServiceTest / "
                "HotReloadCoordinatorTest / PublishDrillTest 等。",
            ]
            for text in items:
                style = "List Bullet"
                cur = insert_paragraph_after(cur, text, style)
            n += 1

    # ---------- 9.2 data 目录 ----------
    p = find_para(doc, startswith="很多玩法数值不以硬编码写死在 Java 里")
    if p and "scripts/activity" not in p.text:
        replace_paragraph_text(
            p,
            "很多玩法数值不以硬编码写死在 Java 里，而是放在 data/ 下的 JSON、CSV、Excel。"
            "这样改活动、改卡池、改商店，往往不必重新发整包程序，只要走热更/导入流程。"
            "本轮进一步把「活动独有计算」也配置化：data/scripts/activity/*.groovy 可随 /reload 生效；"
            "hotfix-overrides.json 保存点路径紧急覆写；QaTesterUids.json 配置 QA 白名单；"
            "wallet-snapshots/ 保存大版本货币回滚基准点。",
        )
        n += 1

    # ---------- 9.4 热更 ----------
    p = find_para(doc, startswith="可以把热更想成「换一本新说明书")
    if p and "Override" not in p.text:
        replace_paragraph_text(
            p,
            "可以把热更想成「换一本新说明书，但不要把正在营业的店关掉」。"
            "协调器分阶段加载新配置；若中途失败，回滚到旧的内存指针；成功则广播通知并留下审计。"
            "控制台 stdin 输入 /reload 之类的能力在生产默认关闭，避免误操作。"
            "更进一步，配置发布可写入 config_release 指纹表（ConfigReleaseService），"
            "支持版本追溯与回滚思路，详见 docs/config-versioning.md。"
            "紧急场景可用 Override 热补丁只改错误字段（不必回滚整份 Banners.json）；"
            "活动脚本一并纳入 HotReloadCoordinator；内部机可用 Time Travel 预加载下版本生效窗。",
        )
        n += 1

    p = find_para(doc, startswith="核心类：HotReloadCoordinator")
    if p and "ActivityScriptEngine" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；活动脚本 ActivityScriptEngine；覆盖补丁 ConfigOverrideHotfixService；"
            "逻辑时钟 ServerClockService。",
        )
        n += 1

    p = find_para(doc, exact="后台接口示例：POST /api/admin/ops/reload")
    if p and "hotfix/override" not in p.text:
        replace_paragraph_text(
            p,
            "后台接口示例：POST /api/admin/ops/reload；"
            "POST /api/admin/ops/hotfix/override；POST /api/admin/ops/circuit/trip；"
            "POST /api/admin/ops/time-travel；POST /api/admin/ops/freeze；"
            "POST /api/admin/ops/grant/negative；POST /api/admin/ops/wallet/snapshot|rollback；"
            "GET/POST /api/admin/ops/visibility/qa；POST /api/admin/ops/scripts/reload",
        )
        n += 1

    # ---------- 第十一章 管理后台 ----------
    p = find_para(doc, startswith="管理后台走 HTTP，与玩家游戏口隔离")
    if p and "CommercialOpsController" not in p.text:
        replace_paragraph_text(
            p,
            "管理后台走 HTTP，与玩家游戏口隔离。登录使用 Session，并有 CSRF 与 RBAC 权限控制。"
            "常见能力包括：热更触发、配置导入、中心路由计划查询、客服工单处理、登录失败锁定等。"
            "CommercialOpsController 聚合运营热更增强："
            "Override 热补丁、活动熔断、Time Travel、Server Freeze、负向发放、钱包快照回滚、"
            "脚本重载与 QA 可见性管理（均需 admin:ops:write）。",
        )
        n += 1

    p = find_para(doc, startswith="发布前建议跑通发布演练脚本")
    if p and "server_freeze_publish" not in p.text:
        replace_paragraph_text(
            p,
            "发布前建议跑通发布演练脚本 scripts/publish_drill.ps1，"
            "大版本停机可先跑 scripts/release/server_freeze_publish.ps1"
            "（钱包快照 → Freeze 推送 → 可选 escalate 维护），"
            "或运行 PublishDrillTest / CommercialOpsBusinessFlowTest 等关键测试，"
            "确认编译、热更、Freeze、Center、匹配等核心路径仍健康。"
            "管理操作（热更、导入、工单、回滚、Override、负向发放等）会写入 admin_audit_log，"
            "便于事后追责与回放；归档任务可按策略清理过期审计日志。",
        )
        n += 1

    # ---------- 12.4 优雅停机 ----------
    p = find_para(doc, startswith="想象商场要打烊")
    if p and "ServerFreeze" not in p.text and "1108" not in p.text:
        replace_paragraph_text(
            p,
            "想象商场要打烊：先向在场顾客广播预计打烊时间（ServerMaintenancePush / CmdId 1108）→"
            "停止接待新战斗（ServerFreezeService.acceptNewBattles=false）→"
            "强制同步落盘进行中战斗快照与钱包 WAL、清 session_dirty →"
            "再停止新客人进匹配队列 → 中止尚未结算完的战斗并妥善收尾 → 停止网络监听 → "
            "把该落盘的数据落盘（GracefulShutdownCoordinator）。"
            "这样比重启时直接杀进程更安全，也减少玩家「打到一半掉线且状态诡异」的投诉。",
        )
        n += 1

    # ---------- 十五 成熟度 ----------
    p = find_para(doc, startswith="Admin 热更/导入/审计")
    if p and "Server Freeze" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；运营热更七项：活动脚本、生效窗/Time Travel、资源 Manifest 锁、"
            "Override+熔断、Server Freeze(1108)、QA 可见性、负向发放/钱包快照。",
        )
        n += 1

    # 更新「已知局限」里关于活动硬编码 / 灰度的表述
    p = find_para(doc, startswith="配置灰度已有 ConfigGrayRelease")
    if p and "Visibility Mask" not in p.text:
        replace_paragraph_text(
            p,
            "配置灰度已有 ConfigGrayRelease；卡池等业务读路径需显式按 inGray 分支。"
            "活动侧已升级为 Visibility Mask（QA 隐身账号），可与灰度策略叠加使用。",
        )
        n += 1

    # ---------- 十六 短期建议：把已落地项改写 ----------
    p = find_para(doc, startswith="把 Admin/热更、排行榜、AI 旁路做成更稳定的可独立部署件")
    if p and "Freeze/Override" not in p.text:
        replace_paragraph_text(
            p,
            "把 Admin/热更、排行榜、AI 旁路做成更稳定的可独立部署件；"
            "单体侧 Freeze/Override/脚本热更/资源锁已可用，后续可把 CommercialOps 接口纳入网关鉴权统一面。",
        )
        n += 1

    p = find_para(doc, startswith="坚持发布演练、热更审计")
    if p and "server_freeze_publish" not in p.text:
        replace_paragraph_text(
            p,
            "坚持发布演练、热更审计，大版本走 server_freeze_publish → 优雅停机；"
            "并把 Grafana 告警阈值用压测校准。",
        )
        n += 1

    # ---------- 附录速览 / 一分钟介绍 ----------
    p = find_para(doc, startswith="如果你只有一分钟向领导")
    if p and "活动脚本" not in p.text:
        replace_paragraph_text(
            p,
            "如果你只有一分钟向领导、同事或客户介绍 MyLunarCore，可以按下面四句说："
            "（1）这是一套崩坏星穹铁道风格的游戏服务端：分区场景探索，遇敌进回合战，并带抽卡、商店、活动。"
            "（2）真正能开服的是根目录「一个主程序」；旁边的 microservices 只是拆分试验，不能单独当生产账号系统。"
            "（3）运营改活动、卡池、任务主要靠 data 配置和后台热更；活动独有规则可用 Groovy 脚本热更，"
            "紧急字段可 Override，发版可 Server Freeze；AI 助手是可选旁路，挂了游戏照样玩。"
            "（4）要上生产必须换掉演示密钥、关掉 IAP Mock、对齐客户端协议与资源 Manifest，并跑监控与发布演练。",
        )
        n += 1

    p = find_para(doc, startswith="运维：8080 管后台")
    if p and "freeze" not in p.text:
        replace_paragraph_text(
            p,
            "运维：8080 管后台，9000 跑游戏；Prometheus/Grafana 看业务健康；"
            "发布前跑 publish_drill；大版本先 freeze（1108）再停机。",
        )
        n += 1

    # ---------- 文末总结 ----------
    p = find_para(doc, startswith="本报告基于仓库当前代码与文档快照整理")
    if p and "运营热更七项" not in p.text:
        replace_paragraph_text(
            p,
            "本报告基于仓库当前代码与文档快照整理，2026年09月02日 已在本文档内同步"
            "「二游运营热更七项短板」落地（活动脚本 / 生效窗+Time Travel / 资源 Manifest 锁 / "
            "Override+熔断 / Server Freeze 1108 / QA 可见性 / 负向发放+钱包快照；"
            "此前性能深化 1107、手感/拥堵 1105–1106、AI 增强 1098–1104、体验优化三期 1078–1097、"
            "二期 1070–1076、沉浸 1046–1069、闭环 1020–1045、PROTOCOL_WIRE_VERSION=3），"
            "可用于汇报、培训与开源说明；打开 Word 后请右键目录「更新域」以刷新页码。",
        )
        n += 1

    # ---------- Cmd 表（若有）追加 1108 ----------
    for table in doc.tables:
        if not table.rows:
            continue
        header = " ".join(c.text.strip() for c in table.rows[0].cells)
        if ("Cmd" in header or "命令" in header) and ("说明" in header or "含义" in header or "名称" in header):
            exists = any("1108" in c.text for row in table.rows for c in row.cells)
            if not exists and len(table.rows[0].cells) >= 2:
                row = table.add_row()
                cells = row.cells
                texts = [
                    "1108",
                    "SERVER_MAINTENANCE_PUSH",
                    "服务端维护预告推送（Server Freeze / 预计停机分钟数）",
                ]
                for i, val in enumerate(texts):
                    if i < len(cells):
                        cell = cells[i]
                        if cell.paragraphs:
                            for r in list(cell.paragraphs[0].runs):
                                r.text = ""
                            if cell.paragraphs[0].runs:
                                cell.paragraphs[0].runs[0].text = val
                                set_run_font(cell.paragraphs[0].runs[0])
                            else:
                                run = cell.paragraphs[0].add_run(val)
                                set_run_font(run)
                        else:
                            cell.text = val
                n += 1
                break

    doc.save(str(DOC_PATH))
    print(f"updated {DOC_PATH} patches={n}")


if __name__ == "__main__":
    main()
