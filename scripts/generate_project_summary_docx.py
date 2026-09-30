# -*- coding: utf-8 -*-
"""Generate MyLunarCore project summary Word document to Desktop."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

DESKTOP = Path.home() / "Desktop"
# Windows 控制台编码下中文文件名易乱码：先写 ASCII 名，再尝试重命名为中文名
OUTPUT_ASCII = DESKTOP / "MyLunarCore_Project_Summary.docx"
OUTPUT_CN = DESKTOP / "MyLunarCore项目总结.docx"
OUTPUT = OUTPUT_CN


def set_doc_fonts(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_title(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(25, 75, 140)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_subtitle(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(100, 100, 100)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_heading(doc: Document, text: str, level: int = 1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        if level == 1:
            run.font.color.rgb = RGBColor(25, 75, 140)


def add_body(doc: Document, text: str):
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_bullet(doc: Document, text: str, level: int = 0):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.5 * (level + 1))
    for run in p.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_code(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 30, 30)


def add_table(doc: Document, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.name = "微软雅黑"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for r_idx, row in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = val
            for p in row_cells[c_idx].paragraphs:
                for run in p.runs:
                    run.font.name = "微软雅黑"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()
    set_doc_fonts(doc)

    add_title(doc, "MyLunarCore 项目总结")
    add_subtitle(doc, "项目路径：c:\\Users\\ASUS\\IdeaProjects\\test\\MyLunarCore")
    add_subtitle(doc, "生成日期：2026年7月31日")
    doc.add_paragraph()

    # 一、项目概述
    add_heading(doc, "一、项目概述", 1)
    add_body(
        doc,
        "MyLunarCore 是一款基于 Spring Boot 的模块化单体游戏服务器，定位为星穹铁道风格的"
        "实例化 RPG 私服/自研服务端。核心采用 Netty TCP + KCP 双通道承载游戏协议，"
        "Protobuf 序列化，MySQL 持久化，并提供 HTTP 管理运维平面。",
    )
    add_body(
        doc,
        "当前形态以单体游戏核心为主；microservices/ 目录为 Spring Cloud Alibaba 迁移脚手架，"
        "尚未替代主工程。项目更接近单人/实例化玩法服务端，而非传统分区 MMO。",
    )

    add_heading(doc, "一句话定位", 2)
    add_body(
        doc,
        "以 Netty 游戏循环为主体的星铁风格单体私服 + HTTP 运维平面；"
        "微服务目录仅作迁移脚手架，AI 助手等少量能力可旁路到 sidecar。",
    )

    # 二、技术栈
    add_heading(doc, "二、技术栈", 1)
    add_table(
        doc,
        ["类别", "技术选型", "说明"],
        [
            ["语言 / JDK", "Java 17", "编译与运行目标"],
            ["框架", "Spring Boot 4.0.4", "主工程依赖与自动配置"],
            ["游戏网络", "Netty 4.2.9 + KCP", "TCP/KCP 同端口，默认 9000"],
            ["协议", "Protobuf 4.34", "约 58 个 .proto，自研 Lunar 帧"],
            ["持久化", "JDBC + MySQL 8", "玩家、账号、道具等"],
            ["运维 HTTP", "Spring MVC + Security", "管理后台，默认 8080"],
            ["缓存 / 限流", "Caffeine、Bucket4j", "会话缓存与接口限流"],
            ["可观测性", "Actuator + Prometheus", "指标与健康检查"],
            ["微服务脚手架", "Spring Cloud Alibaba", "Boot 3.3 / Nacos / Gateway"],
        ],
    )

    # 三、总体架构
    add_heading(doc, "三、总体架构", 1)
    add_code(
        doc,
        "客户端\n"
        "  ├─ TCP/KCP :9000  ──►  GameNettyServer / GameKcpServer\n"
        "  │                      └─ PacketDispatcher → 各业务 *NettyService\n"
        "  │                      └─ GameServer Tick（游戏主循环）\n"
        "  └─（可选）AI / 运营\n"
        "\n"
        "Admin / 运维 HTTP :8080  ──►  配置导入、热更、登录 RBAC、内部 API\n"
        "\n"
        "MySQL ◄── player / repo / 周期落库\n"
        "data/*.json ◄── 热更 / 活动 / 商店 / 任务等静态配置\n"
        "\n"
        "（可选旁路）ai-assist-service :18083\n"
        "（脚手架）api-gateway :18443 → auth / player / ai-assist",
    )

    add_heading(doc, "架构原则", 2)
    add_bullet(doc, "游戏实时面（战斗、场景、玩家写路径）保留在单体 Netty 核心。")
    add_bullet(doc, "运维面（Admin、配置导入、热更）可优先外拆，与 Tick 解耦。")
    add_bullet(doc, "AI 助手为只读建议面，适合 sidecar，失败回退进程内规则教练。")
    add_bullet(doc, "Center 路由支持 local / remote，为多节点场景迁移做准备。")

    # 四、业务模块
    add_heading(doc, "四、核心业务模块", 1)
    add_table(
        doc,
        ["模块", "包名", "能力摘要"],
        [
            ["网络层", "net", "TCP/KCP、帧编解码、分发、限流、背压、TLS"],
            ["玩家会话", "player", "登录状态机、聚合缓存、异步加载、周期落库"],
            ["战斗", "battle", "回合制战斗、遭遇配置、波次、战斗辅助策略"],
            ["场景", "scene", "场景/区域、AOI、遭遇、同步广播"],
            ["中心路由", "center", "多节点场景路由、迁移票据、SceneRegistry"],
            ["角色养成", "character", "创角、成长、天赋、属性计算"],
            ["道具", "item", "道具业务与配置解析"],
            ["经济", "economy", "钱包、商店、IAP（默认 Mock 验签）"],
            ["抽卡", "gacha", "卡池热更与抽卡引擎"],
            ["挑战", "challenge", "挑战玩法运行时"],
            ["肉鸽", "rogue", "Rogue 管理与运行时"],
            ["任务", "quest", "任务配置、进度、触发引擎"],
            ["大厅社交", "hall", "好友、邮件、聊天、排行榜"],
            ["匹配", "matchmaking", "匹配队列、房间、兼容性打分"],
            ["皮肤", "skin", "皮肤配置、拥有与装备"],
            ["AI 助手", "assist", "规则教练、RAG、安全过滤、远程 sidecar"],
            ["管理后台", "admin", "登录 RBAC、配置导入、热更、工单"],
            ["横切公共", "common", "活动、热更协调、指标、优雅停机"],
        ],
    )

    # 五、网络与运维
    add_heading(doc, "五、网络与运维平面", 1)
    add_heading(doc, "游戏网络", 2)
    add_bullet(doc, "TCP：GameNettyServer，默认端口 lunarcore.netty-port=9000。")
    add_bullet(doc, "KCP：GameKcpServer，与 TCP 同绑 9000；interval=20，mtu=1400。")
    add_bullet(doc, "协议：Lunar 帧 + Protobuf；GamePacketDispatcher 路由至各 PacketHandlers。")
    add_bullet(doc, "游戏循环：lunarcore.game-loop.period-ms 默认 1000ms。")

    add_heading(doc, "运维与配置", 2)
    add_bullet(doc, "HTTP 管理端口默认 8080，Spring Security 保护 /admin。")
    add_bullet(doc, "数据目录 lunarcore.data-dir=data，含活动、商店、任务、皮肤、遭遇等 JSON/CSV。")
    add_bullet(doc, "HotReloadCoordinator 分阶段热更并广播；支持配置导入、回滚审计。")
    add_bullet(doc, "内部接口使用 X-Internal-Token（与 AI 旁路共用）。")

    # 六、AI 助手
    add_heading(doc, "六、AI 助手能力", 1)
    add_body(
        doc,
        "AI 助手默认启用规则教练（rule-coach），LLM 与远程 sidecar 默认关闭。"
        "处理链路：本地规则快速判定 → 上下文拼装 → 可选远程/本地 LLM → 安全过滤、缓存、审计、配额。",
    )
    add_bullet(doc, "能力面：任务/探索引导、阵容推荐、RAG、聊天 @助手、战斗特征采集等。")
    add_bullet(doc, "Sidecar：ai-assist-service（默认 18083），提供 /ai/health 与 POST /internal/ai/ask。")
    add_bullet(doc, "失败时回退到进程内规则教练 / 本地 LLM，保证游戏服可用性。")

    # 七、微服务迁移
    add_heading(doc, "七、微服务迁移现状", 1)
    add_body(
        doc,
        "microservices/ 为迁移脚手架，使用 Spring Boot 3.3 + Spring Cloud Alibaba，"
        "配套 docker-compose（Nacos、MySQL、Redis）。当前模块包括：",
    )
    add_bullet(doc, "common-api：共享 API 模型（含 Assist 请求/响应约定）。")
    add_bullet(doc, "api-gateway：HTTPS 网关（默认 18443），鉴权、限流、AI 降级回退。")
    add_bullet(doc, "auth-service / player-service：演示级认证与玩家查询。")
    add_bullet(doc, "ai-assist-service：AI 旁路服务，可被游戏服按需调用。")

    add_heading(doc, "拆分建议（与 README 一致）", 2)
    add_table(
        doc,
        ["候选", "与游戏循环一致性", "是否近期拆分", "理由"],
        [
            ["Admin / 配置热更", "低", "优先", "运维平面与 Tick 隔离"],
            ["排行榜", "低–中", "其次", "大厅 Netty 可作薄客户端"],
            ["AI 助手", "低", "Sidecar", "只读建议，故障隔离"],
            ["Auth / 账号", "中", "稍后", "脚手架已有 auth-service"],
            ["玩家聚合 / 钱包 / 抽卡", "高", "暂不", "需共享写协议"],
            ["战斗 / 场景 / 匹配", "高", "暂不", "先验证 Center 多节点路由"],
        ],
    )

    # 八、实现成熟度
    add_heading(doc, "八、实现成熟度概览", 1)
    add_heading(doc, "偏已落地", 2)
    add_bullet(doc, "Netty 业务面：登录、战斗、场景 AOI、挑战、肉鸽、抽卡、商店/IAP、任务、匹配、皮肤、活动、AI。")
    add_bullet(doc, "中心路由 local/remote、热更协调、配置导入审计、Prometheus、限流背压。")

    add_heading(doc, "偏脚手架 / 可开关", 2)
    add_bullet(doc, "microservices/ 整体体量偏演示；auth/player 源码较少。")
    add_bullet(doc, "LLM 与远程 AI 默认关闭；IAP 默认 mock-verify；战斗辅助提示默认关闭。")

    # 九、目录结构
    add_heading(doc, "九、关键目录结构", 1)
    add_code(
        doc,
        "MyLunarCore/\n"
        "├── src/main/java/.../mylunarcore/   # 单体游戏核心业务\n"
        "├── src/main/proto/                  # Protobuf 协议定义\n"
        "├── data/                            # 运行时 JSON/CSV 配置\n"
        "├── microservices/                   # Cloud Alibaba 迁移脚手架\n"
        "│   ├── common/common-api\n"
        "│   └── services/{api-gateway,auth,player,ai-assist}\n"
        "├── scripts/                         # 文档生成、导入校验、发布演练\n"
        "└── pom.xml                          # 主工程 Maven 描述",
    )

    # 十、总结与展望
    add_heading(doc, "十、总结与后续方向", 1)
    add_body(
        doc,
        "MyLunarCore 已具备较完整的实例化 RPG 服务端能力：网络层、战斗、场景、养成、"
        "经济、活动与运维热更形成闭环。当前最优路径仍是「模块化单体 + 选择性旁路」，"
        "而非一次性微服务化。",
    )
    add_bullet(doc, "短期：巩固 Admin/配置热更外拆、AI sidecar 稳定性、发布演练与热更审计。")
    add_bullet(doc, "中期：完善大厅社交与匹配体验，验证 Center 多节点路由与玩家迁移。")
    add_bullet(doc, "长期：在写路径边界清晰后再拆分玩家聚合、战斗/场景等高一致性域。")

    add_body(
        doc,
        "本文档基于仓库当前代码与配置自动整理，用于项目汇报、新人上手与架构决策参考。",
    )

    return doc


def main():
    doc = build_document()
    DESKTOP.mkdir(parents=True, exist_ok=True)
    # 先落盘 ASCII 文件名，再重命名为中文，避免 Windows 控制台代码页导致乱码文件名
    doc.save(OUTPUT_ASCII)
    try:
        if OUTPUT_CN.exists():
            OUTPUT_CN.unlink()
        OUTPUT_ASCII.rename(OUTPUT_CN)
        print(f"Saved: {OUTPUT_CN}")
    except OSError:
        print(f"Saved: {OUTPUT_ASCII}")


if __name__ == "__main__":
    main()
