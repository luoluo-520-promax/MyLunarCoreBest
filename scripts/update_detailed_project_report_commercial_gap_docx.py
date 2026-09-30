# -*- coding: utf-8 -*-
"""将「商业箱庭缺口」落地改动同步进桌面《MyLunarCore项目各方面详细总结报告》。"""

from datetime import date
from pathlib import Path
import shutil

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DOC_PATH = Path.home() / "Desktop" / (
    "MyLunarCore"
    + "\u9879\u76ee\u5404\u65b9\u9762\u8be6\u7ec6\u603b\u7ed3\u62a5\u544a"
    + ".docx"
)


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def replace_paragraph_text(paragraph, new_text):
    style = paragraph.style
    for r in list(paragraph.runs):
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        set_run_font(paragraph.runs[0])
    else:
        run = paragraph.add_run(new_text)
        set_run_font(run)
    paragraph.style = style


def insert_paragraph_after(paragraph, text, style_name="List Bullet"):
    new_para = paragraph._parent.add_paragraph(text, style=style_name)
    paragraph._p.addnext(new_para._p)
    for r in new_para.runs:
        set_run_font(r, size=13 if style_name.startswith("Heading") else 11,
                     bold=style_name.startswith("Heading"))
        if style_name.startswith("Heading"):
            r.font.color.rgb = RGBColor(25, 75, 140)
    return new_para


def find_para(doc, exact=None, startswith=None, contains=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if exact is not None and t == exact:
            return p
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    return None


def replace_if(doc, *, exact=None, startswith=None, contains=None, new_text=None):
    p = find_para(doc, exact=exact, startswith=startswith, contains=contains)
    if p is not None and new_text is not None:
        replace_paragraph_text(p, new_text)
        return True
    return False


def main():
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    backup = DOC_PATH.with_suffix(".bak.docx")
    if not backup.exists():
        shutil.copyfile(DOC_PATH, backup)

    doc = Document(str(DOC_PATH))
    today = f"{date.today().year}年{date.today().month:02d}月{date.today().day:02d}日"
    n = 0

    # ---------- 副标题 / 一句话定位 ----------
    if replace_if(
        doc,
        exact="（含技术架构、业务能力、安全运维、微服务迁移与后续建议）",
        new_text="（含技术架构、业务能力、安全运维、微服务迁移、商业箱庭缺口补齐与后续建议）",
    ):
        n += 1

    if replace_if(
        doc,
        startswith="MyLunarCore 是一套基于 Spring Boot 4 + Netty/KCP + Protobuf + MySQL",
        new_text=(
            "MyLunarCore 是一套基于 Spring Boot 4 + Netty/KCP + Protobuf + MySQL 的"
            "「崩坏：星穹铁道式」游戏私服 / 自研服务端：玩家在分区场景里探索，遇敌后进入回合制战斗；"
            "同时提供抽卡、商店支付、活动任务、组队匹配、排行榜聊天，以及管理后台热更与可选多节点切服。"
            "真正可运行的主业务在根目录单体工程；microservices/ 是迁移脚手架，但本轮已将"
            "Player Profile/Session、Chat/Match 挂入 api-gateway，并补齐动态分线、配置 Delta、"
            "IAP 退款冻结、维护模式等商业箱庭能力（见 docs/commercial-gap-solutions.md）。"
        ),
    ):
        n += 1

    # ---------- 目录速览：追加商业缺口 ----------
    if replace_if(
        doc,
        exact="十四、微服务脚手架现状",
        new_text="十四、微服务脚手架现状（含 Chat/Match 入网关与 Player 中心化）",
    ):
        n += 1

    # ---------- 5.1 架构鸟瞰 ----------
    if replace_if(
        doc,
        startswith="脚手架（未替代主工程）",
        new_text=(
            "脚手架（未替代主工程，但已可联调）\n"
            "  ├─ api-gateway :18443 → auth / player / ai-assist / chat / match（网关已挂）\n"
            "  ├─ chat-service :18084（世界削峰 backlog + 离线私聊 7 天）\n"
            "  ├─ match-service :18085（全局评分队列，撮合后返回 zoneId）\n"
            "  ├─ player-service :18082（Profile/Session 中心化 API，非单纯演示桩）\n"
            "  └─ auth :18081（JWT 演示登录，生产仍以单体账号为准）"
        ),
    ):
        n += 1

    # 架构图里旧的两行可能是独立段落
    replace_if(
        doc,
        exact="  ├─ api-gateway :18443 → auth / player / ai-assist（网关已挂）",
        new_text="  ├─ api-gateway :18443 → auth / player / ai-assist / chat / match（网关已挂）",
    )
    replace_if(
        doc,
        exact="  ├─ chat-service :18084、match-service :18085（独立 HTTP，未进网关）",
        new_text="  ├─ chat-service :18084、match-service :18085（评分撮合 / 削峰离线；已进网关）",
    )
    replace_if(
        doc,
        exact="  └─ auth :18081 / player :18082（演示桩）",
        new_text="  └─ auth :18081（演示登录）/ player :18082（Profile/Session 中心化 API）",
    )

    # ---------- 7.3 场景 ----------
    if replace_if(
        doc,
        startswith="玩家进入某个 Plane/Floor 对应的 Zone 实例后",
        new_text=(
            "玩家进入某个 Plane/Floor 对应的 Zone 实例后，可移动、与 NPC/道具交互，"
            "服务器按兴趣范围（AOI）同步周围实体。单线默认容量已抬升至约 1000 人，"
            "并支持动态分线（DynamicZoneLineAllocator）、轻量 Zone Actor 邮箱串行坐标更新，"
            "以及 AOI 广播异步队列（AoiAsyncProcessor），避免同屏广播拖垮战斗时钟。"
            "跨图或跨节点时，由 Center 规划归属，必要时发放迁移票据。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        exact="入口：SceneNettyService、ZoneManager、ZoneTickService、PlayerMigrationService",
        new_text=(
            "入口：SceneNettyService、ZoneManager、DynamicZoneLineAllocator、"
            "AoiAsyncProcessor、ZoneTickService、PlayerMigrationService"
        ),
    ):
        n += 1

    # ---------- 7.4 战斗 ----------
    if replace_if(
        doc,
        startswith="场景中触发遭遇后进入战斗实例",
        new_text=(
            "场景中触发遭遇后进入战斗实例：开战、行动、结算、退出。"
            "支持遇敌配置与波次；多人共战通过 BattleContext.participantPlayerIds 等字段表达。"
            "启发式战斗辅助（HeuristicBattleAssistPolicy）之外，可开启 AI 策略代理实验开关"
            "（AiBattleStrategyProxy）：输出可一键采纳的行动指令 JSON（confirmRequired=true），"
            "仍需客户端确认后发包，避免外挂直连。"
        ),
    ):
        n += 1

    # ---------- 7.7 IAP ----------
    if replace_if(
        doc,
        contains="IAP：IapVerifyGateway",
        new_text=(
            "IAP：IapVerifyGateway、各渠道 Verifier、IapProductionGuard；"
            "退款/撤销：IapRefundWebhookController（Google RTDN / Apple ASN）+ "
            "NegativeBalanceFreezeService（负资产冻结，抽卡 retcode=9 禁消费）"
        ),
    ):
        n += 1

    # ---------- 7.11 匹配 ----------
    if replace_if(
        doc,
        startswith="Party 是大世界组队雏形",
        new_text=(
            "Party 是大世界组队雏形（同进程权威，人数上限约 4 人）。"
            "Match/Room 面向副本匹配：排队、开房、兼容性打分，并与挑战匹配协调器联动。"
            "微服务 match-service 已升级为全局评分队列，撮合成功返回 zoneId 供客户端跳转，"
            "并已挂入 api-gateway（/v1/match/**）；单体 MatchNettyService 仍是权威实现主路径。"
        ),
    ):
        n += 1

    # ---------- 7.10 聊天 ----------
    if replace_if(
        doc,
        startswith="大厅模块把社交能力收在一起",
        new_text=(
            "大厅模块把社交能力收在一起：好友、邮件、世界/私聊，以及历史最高分类排行榜。"
            "排行榜与世界聊天在开启 Redis 后可跨进程共享；聊天里 @助手 可转给 AI 模块。"
            "微服务 chat-service 提供世界频道削峰 backlog 与离线私聊缓存（7 天语义），"
            "已挂入 api-gateway（/v1/chat/**）；生产可无缝替换为 Kafka/Redis Stream。"
        ),
    ):
        n += 1

    # ---------- 7.14 成就 ----------
    if replace_if(
        doc,
        startswith="成就系统记录玩家达成条件与领取状态",
        new_text=(
            "成就系统记录玩家达成条件与领取状态，适合做长期留存目标。"
            "登录成功路径会调用 OfflineAchievementCompensator："
            "按离线窗口内战报/抽卡/归档摘要异步补发成就进度，减少「断网达成却未触发」投诉。"
            "新手引导则按配置步骤引导新玩家完成关键操作，降低上手门槛。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        exact="成就：AchievementService；表 player_achievement；协议号段约 950–954",
        new_text=(
            "成就：AchievementService + OfflineAchievementCompensator；"
            "表 player_achievement；协议号段约 950–954"
        ),
    ):
        n += 1

    # ---------- 7.18 战斗安全 ----------
    if replace_if(
        doc,
        startswith="回合战不能「客户端报多少伤害服务器就信多少」",
        new_text=(
            "回合战不能「客户端报多少伤害服务器就信多少」。"
            "BattleDeterministicValidator 与 BattleDamageFormula 用服务端公式做确定性校验；"
            "BattleActionIntegrityGuard 进一步对关键动作签发 Salt，客户端须回传 Hash，"
            "校验失败立即中断战斗并可踢线（相对纯事后审计）。"
            "BattleAuditService 留下审计线索；BattleSnapshotService 将中间态写入 Redis/内存，"
            "支持断线重连与维护后 hydrate。"
        ),
    ):
        n += 1

    # ---------- 8.5 连接保护：包优先级 ----------
    if replace_if(
        doc,
        exact="KCP 拥塞与重传：KcpRetransmitAlgo（fixed/adaptive）+ KcpRttMonitor，弱网更稳。",
        new_text=(
            "KCP 拥塞与重传：KcpRetransmitAlgo（fixed/adaptive）+ KcpRttMonitor，弱网更稳；"
            "出站包优先级 PacketPriorityOutboundHandler："
            "战斗 Cmd 200–299 走 HIGH 队列，场景 300–399 走 NORMAL，弱网下优先保操作延迟。"
        ),
    ):
        n += 1

    # ---------- 9.1 / 9.4 存储与热更 ----------
    p_mysql = find_para(doc, startswith="玩家的账号、角色、背包、抽卡历史")
    if p_mysql is None:
        p_mysql = find_para(doc, contains="MySQL 是最终记账本")
    if p_mysql is not None:
        t = p_mysql.text.strip()
        if "读写分离" not in t:
            replace_paragraph_text(
                p_mysql,
                t + " 可选开启读写分离与影子库路由（RoutingDataSourceConfig / PtLoadRoutingFilter，"
                "请求头 x-pt-load=true 打到影子库）；钱包等热数据可按 UID Hash 分 16 逻辑库"
                "（UidShardRouter）。",
            )
            n += 1

    if replace_if(
        doc,
        startswith="可以把热更想成「换一本新说明书，但不要把正在营业的店关掉」",
        new_text=(
            "可以把热更想成「换一本新说明书，但不要把正在营业的店关掉」。"
            "协调器分阶段加载新配置；若中途失败，回滚到旧的内存指针；成功则广播通知并留下审计。"
            "本轮新增 ConfigDeltaPatchService：按 JSON Path 做指针级增量替换，"
            "并生成带校验头的二进制快照，降低大版本全量反序列化造成的 CPU/GC 毛刺。"
            "客户端资源侧 ClientResourceManifestService 维护 Manifest Diff，"
            "Admin 可上传/对比（/api/admin/resources/manifest/**），登录侧可下发 CDN 差异列表。"
            "控制台 stdin /reload 在生产默认关闭；配置发布指纹见 docs/config-versioning.md。"
        ),
    ):
        n += 1

    # ---------- 10 AI ----------
    if replace_if(
        doc,
        exact="战斗特征采集后的辅助建议（可开关）；多语言 locale / Accept-Language",
        new_text=(
            "战斗特征采集后的辅助建议（可开关）；"
            "AI 策略代理实验开关 battle-strategy-proxy-enabled："
            "输出可一键采纳行动 JSON（仍需人工确认）；多语言 locale / Accept-Language"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="入站安全在关键词/正则",
        new_text=(
            "入站安全在关键词/正则（AssistSafetyRules）之外增加轻量语义意图分类，"
            "降低绕过词表变体的风险；出站仍做敏感句清洗。"
            "送入 LLM 前对 UID/手机号/邮箱等做 PII 脱敏（AssistPiiRedactor）；"
            "审计日志使用 uidHash，默认不落完整隐私原文。"
            "配额按教练/LLM 分桶，并支持等级/VIP 差异化日预算与「同一问题」短时限流。"
            "内容素材来自 GuidePack、CoachTips、AssistFeatureContent 等配置，热更时可增量刷新 RAG。"
            "记住：AI 给出的是建议；策略代理也只产出指令草案，真正扣费、开战、发奖仍由游戏权威逻辑决定。"
        ),
    ):
        n += 1

    # ---------- 12 可观测 ----------
    if replace_if(
        doc,
        exact="Zone 默认人数上限约 300（可按 Tick 耗时自适应下调）；匹配队列超时默认约 60 秒。",
        new_text=(
            "Zone 单线默认人数上限约 1000（可按 Tick 耗时自适应下调），并支持动态分线与 AOI 异步；"
            "匹配队列超时默认约 60 秒；压测可用 x-pt-load 路由影子库；"
            "战斗 DEBUG 日志按 UID 白名单采样（BattleLogSampler）；"
            "告警可推飞书/钉钉 Webhook（AlertWebhookNotifier）。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="想象商场要打烊",
        new_text=(
            "想象商场要打烊：先停止新客人进匹配队列 → 中止尚未结算完的战斗并妥善收尾 →"
            "停止网络监听 → 把该落盘的数据落盘。"
            "本轮增加 MaintenanceModeService：维护窗口拒绝新登录，停机前 abort 战斗并依赖"
            "已有 Redis BattleSnapshot，维护结束后玩家可 hydrate/重连回战场。"
            "Admin：POST /api/admin/ops/maintenance/enable|disable。"
        ),
    ):
        n += 1

    # ---------- 14 微服务 ----------
    if replace_if(
        doc,
        startswith="microservices/ 目录的官方定位非常明确",
        new_text=(
            "microservices/ 目录仍是迁移脚手架，不是完整生产 IdP，也不应「关掉单体只启脚手架就开服」。"
            "但本轮已显著推进：api-gateway 挂载 /v1/chat/**、/v1/match/**；"
            "player-service 提供可独立部署的 Profile/Session 查询接口；"
            "match-service 实现评分撮合并回传 zoneId；chat-service 实现世界削峰与离线缓存。"
            "auth-service 仍为 JWT 演示登录；生产账号与权威写路径仍在单体。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="docker compose -f microservices/docker-compose.yml up -d",
        new_text=(
            "docker compose -f microservices/docker-compose.yml up -d 可拉起 Nacos、MySQL、Redis 等依赖；"
            "也可再拉起 Prometheus/Grafana。ClickHouse 分析仓用 profile analytics。"
            "推荐联调启动顺序：auth → player → chat → match → ai-assist → api-gateway。"
            "健康检查示例：网关上的 /auth/health、/players/{uid}、/v1/chat/health、/v1/match/health、/ai/health；"
            "版本化路径 /v1/** 与旧路径并存。"
        ),
    ):
        n += 1

    # ---------- 15 已实现 / 未完成 ----------
    if replace_if(
        doc,
        exact="Zone 租约心跳与人数上限（默认约 300，可按 Tick 耗时自适应下调）；Global / Zone / Battle 分时钟。",
        new_text=(
            "Zone 租约心跳与人数上限（默认约 1000 + 动态分线 + AOI 异步 + Actor 邮箱）；"
            "Global / Zone / Battle 分时钟。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        exact="可运行主业务在单体；微服务 auth/player 是演示桩；chat/match 仅为独立 HTTP 骨架且未挂网关。",
        new_text=(
            "可运行主业务在单体；auth 仍为演示登录；player 已提供中心化 Profile/Session API；"
            "chat/match 已挂入 api-gateway 并可联调（权威写路径仍以单体为准）。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        exact="真 IAP 需配齐渠道密钥；battle-service 未列入脚手架模块，勿当作可用服务。",
        new_text=(
            "真 IAP 需配齐渠道密钥；已支持退款 Webhook 与负资产冻结（仍需渠道侧签名校验加固）；"
            "battle-service 未列入脚手架模块，勿当作可用服务。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        exact="压测参考见 docs/load-test-baseline.md；目标硬件须用 k6 回填，文档初值不可直接当 SLA。",
        new_text=(
            "压测参考见 docs/load-test-baseline.md；建议开启影子库 + 请求头 x-pt-load=true；"
            "目标硬件须用 k6 回填，文档初值不可直接当 SLA。"
            "商业缺口落地说明：docs/commercial-gap-solutions.md；回归：CommercialGapFlowTest。"
        ),
    ):
        n += 1

    # ---------- 18 docker 注释 ----------
    if replace_if(
        doc,
        exact="# 可选另启：chat-service :18084、match-service :18085（不经网关）",
        new_text="# 建议同启：chat-service :18084、match-service :18085（经 api-gateway /v1/chat|/v1/match）",
    ):
        n += 1

    # ---------- 一分钟介绍 ----------
    if replace_if(
        doc,
        startswith="如果你只有一分钟向领导、同事或客户介绍 MyLunarCore",
        new_text=(
            "如果你只有一分钟向领导、同事或客户介绍 MyLunarCore，可以按下面四句说："
            "（1）这是一套崩坏星穹铁道风格的游戏服务端：分区场景探索，遇敌进回合战，并带抽卡、商店、活动。"
            "（2）真正能开服的是根目录「一个主程序」；microservices 做拆分试验，本轮已把 Chat/Match 挂上网关、"
            "Player Profile 可独立查询。（3）运营改活动、卡池主要靠 data 配置和后台热更；"
            "大配置可用 Delta 增量，客户端资源有 Manifest Diff；AI 是可选旁路，还可实验「一键采纳」策略指令。"
            "（4）要上线必须切 prod、换密钥、关假支付；开服热门区用动态分线+AOI 异步；"
            "压测走影子库；维护窗口可拒绝登录并靠战斗快照重连。"
        ),
    ):
        n += 1

    # ---------- 结尾三段结论 ----------
    if replace_if(
        doc,
        startswith="MyLunarCore 已经不是「只有空壳的演示仓库」",
        new_text=(
            "MyLunarCore 已经不是「只有空壳的演示仓库」：它在单体形态下具备较完整的崩铁式 RPG 服务端能力——"
            "联网、场景、回合战、养成、经济抽卡、体力与每日循环、战令与主线章节、活动任务、对话过场与家园、"
            "公会/公会战/公会科技、竞技场评分、世界 BOSS/深渊/遗器词条等缺口补齐玩法、成就引导、运维热更与回滚、"
            "业务监控，以及可选的多节点与 AI 助手旁路。"
            "本轮（商业箱庭缺口）进一步补齐了：Zone 千人级+动态分线+AOI 异步、读写分离/影子库、配置 Delta、"
            "KCP 包优先级、战斗 Salt 即时校验、IAP 退款冻结、资源 Manifest、维护模式、Chat/Match 入网关等。"
            "与此同时，它也很诚实地把 auth 演示密钥标成仅限本地，把无缝大世界标成实验。"
        ),
    ):
        n += 1

    if replace_if(
        doc,
        startswith="对普通人最重要的结论只有三句",
        new_text=(
            "对普通人最重要的结论只有三句：第一，想体验或继续开发，请先跑主工程；"
            "第二，想上线，请走 prod 并认真对待密钥、支付验签、监控告警与维护窗口演练；"
            "第三，想变「更大、更微服务、更无缝」，请先用压测校准 Zone/DB/网关边界，再决定拆什么——"
            "而不是为了拆而拆。细节见 docs/commercial-gap-solutions.md。"
        ),
    ):
        n += 1

    # ---------- 生成日期（若有） ----------
    p_date = find_para(doc, startswith="生成日期")
    if p_date is None:
        p_date = find_para(doc, contains="生成日期：")
    if p_date is not None:
        replace_paragraph_text(
            p_date,
            f"生成日期：{today}（已同步商业箱庭缺口落地：性能 / 玩法 / 运维）",
        )
        n += 1
    else:
        # 尝试更新封面附近「报告用途」后追加日期备注
        p_use = find_para(doc, startswith="报告用途：")
        if p_use is not None:
            insert_paragraph_after(
                p_use,
                f"文档修订：{today} 已同步商业箱庭缺口解决方案（详见第十五章增补 / docs/commercial-gap-solutions.md）。",
                style_name="Normal",
            )
            n += 1

    # ---------- 插入新章节：商业箱庭缺口补齐（放在「十四、微服务」正文段落后） ----------
    if find_para(doc, exact="十五、商业箱庭缺口补齐（本轮落地）") is None:
        anchor = find_para(doc, startswith="箱庭回合制战斗虽是 CPU 密集")
        if anchor is None:
            anchor = find_para(doc, contains="开服活动洪峰最大且允许最终一致的")
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "十五、商业箱庭缺口补齐（本轮落地）",
                style_name="Heading 1",
            )
            cur = h
            blocks = [
                ("Normal",
                 "对应「高并发性能 / 玩法闭环 / 生产运维」三块商业箱庭评审缺口，本轮已在单体与微服务落地可运行骨架。"
                 "默认开关偏保守，生产按需开启。仓库说明：docs/commercial-gap-solutions.md；"
                 "回归测试：CommercialGapFlowTest 及 match/chat/player 微服务流程测试。"),
                ("Heading 2", "15.1 性能与基础设施"),
                ("List Bullet",
                 "Zone：单线默认约 1000；动态分线 DynamicZoneLineAllocator；Actor 邮箱；AOI 异步 AoiAsyncProcessor。"),
                ("List Bullet",
                 "数据库：ReadWriteRouter + RoutingDataSourceConfig；UID Hash 分库 UidShardRouter；"
                 "压测头 x-pt-load → 影子库 PtLoadRoutingFilter。"),
                ("List Bullet",
                 "配置：ConfigDeltaPatchService（JSON Path Delta + 二进制快照）；"
                 "Admin POST /api/admin/resources/config/delta。"),
                ("List Bullet",
                 "弱网：PacketPriorityOutboundHandler，战斗包优先于场景移动包。"),
                ("Heading 2", "15.2 功能深度与玩法闭环"),
                ("List Bullet",
                 "Player：/players/{uid}/profile|session 中心化接口；Chat/Match 已入 api-gateway。"),
                ("List Bullet",
                 "战斗：BattleActionIntegrityGuard（Salt+Hash 即时校验）；"
                 "AI：AiBattleStrategyProxy（一键采纳草案）。"),
                ("List Bullet",
                 "资源：ClientResourceManifestService Manifest Diff；"
                 "成就：OfflineAchievementCompensator 登录补偿。"),
                ("List Bullet",
                 "IAP：退款 Webhook + NegativeBalanceFreezeService 负资产冻结禁抽卡。"),
                ("Heading 2", "15.3 运维与可观测"),
                ("List Bullet",
                 "维护模式 MaintenanceModeService + Redis BattleSnapshot 重连；"
                 "Admin /api/admin/ops/maintenance/*。"),
                ("List Bullet",
                 "BattleLogSampler 白名单采样；AlertWebhookNotifier 飞书/钉钉；"
                 "压测隔离影子库。"),
                ("Heading 2", "15.4 推荐启用顺序"),
                ("List Bullet", "开服前：zone.aoi-async + dynamic-line + max-players=1000（已默认）。"),
                ("List Bullet", "压测：routing-enabled + shadow URL + 请求头 x-pt-load=true。"),
                ("List Bullet", "灰度：battle-strategy-proxy；战斗 Hash 接到 FightAction 协议字段。"),
                ("List Bullet", "生产：replica/shadow URL、告警 Webhook、维护窗口演练。"),
            ]
            for style, text in blocks:
                cur = insert_paragraph_after(cur, text, style_name=style)
            n += 1

    # ---------- 文档地图 ----------
    p_map = find_para(doc, contains="docs/production-hardening.md")
    if p_map is not None and "commercial-gap" not in p_map.text:
        replace_paragraph_text(
            p_map,
            p_map.text.strip() + " 商业箱庭缺口：docs/commercial-gap-solutions.md。",
        )
        n += 1

    doc.save(str(DOC_PATH))
    print(f"Updated: {DOC_PATH}")
    print(f"Backup: {backup}")
    print(f"Patches applied (approx): {n}")


if __name__ == "__main__":
    main()
