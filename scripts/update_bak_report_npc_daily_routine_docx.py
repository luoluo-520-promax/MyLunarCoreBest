# -*- coding: utf-8 -*-
"""
将「环境生态假死感」优化——NPC/生物独立日程作息（EntityBehaviorService / DailyRoutineBehavior /
NPCScheduleConfig / SceneNpcBehaviorScNotify CmdId 1143）同步进桌面原 bak 总结报告。

只覆盖保存原文件，不另存新文档。
幂等标记：正文出现 EntityBehaviorService 或 7.33 NPC 或 SCENE_NPC_BEHAVIOR 即视为已写入。
"""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DOC_PATH = Path.home() / "Desktop" / "MyLunarCore项目各方面详细总结报告.bak.docx"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def replace_paragraph_text(paragraph, new_text, bold=False):
    style = paragraph.style
    for r in list(paragraph.runs):
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        set_run_font(paragraph.runs[0], bold=bold)
    else:
        run = paragraph.add_run(new_text)
        set_run_font(run, bold=bold)
    paragraph.style = style


def insert_paragraph_after(paragraph, text, style_name="List Bullet"):
    new_para = paragraph._parent.add_paragraph(text, style=style_name)
    paragraph._p.addnext(new_para._p)
    for r in new_para.runs:
        set_run_font(
            r,
            size=13 if style_name.startswith("Heading") else 11,
            bold=style_name.startswith("Heading"),
        )
        if style_name.startswith("Heading"):
            r.font.color.rgb = RGBColor(25, 75, 140)
    return new_para


def find_para(doc, exact=None, startswith=None, contains=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if exact is not None and t == exact:
            return p
        if startswith is not None and t.startswith(startswith):
            return p
        if contains is not None and contains in t:
            return p
    return None


def set_cell_text(cell, text):
    if not cell.paragraphs:
        cell.text = text
        return
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r.text = ""
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0], size=10)
    else:
        run = p.add_run(text)
        set_run_font(run, size=10)
    for extra in cell.paragraphs[1:]:
        extra.clear()


def add_table_row(table, values):
    row = table.add_row()
    for i, value in enumerate(values):
        if i < len(row.cells):
            set_cell_text(row.cells[i], value)
    return row


def table_has(table, needle, col=None):
    for row in table.rows:
        cells = row.cells
        if col is not None:
            if needle in cells[col].text:
                return True
        else:
            for cell in cells:
                if needle in cell.text:
                    return True
    return False


def find_table_by_header(doc, *needles):
    for table in doc.tables:
        if not table.rows:
            continue
        header = " ".join(c.text.strip() for c in table.rows[0].cells)
        if all(n in header for n in needles):
            return table
    return None


def patch_table_cell(table, old_contains, col, new_text):
    for row in table.rows:
        if old_contains in row.cells[col].text:
            set_cell_text(row.cells[col], new_text)
            return True
    return False


def already_patched(doc):
    blob = "\n".join(p.text for p in doc.paragraphs)
    return (
        "EntityBehaviorService" in blob
        or "7.33 NPC" in blob
        or "SCENE_NPC_BEHAVIOR" in blob
        or "SceneNpcBehaviorScNotify" in blob
    )


def main():
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    doc = Document(str(DOC_PATH))
    if already_patched(doc):
        print(f"文档已含 NPC 日程作息描述，跳过重复写入: {DOC_PATH}")
        return

    n = 0

    # ---------- 封面副标题 ----------
    p = find_para(doc, startswith="（含技术架构、业务能力")
    if p is not None and "NPC 日程作息" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("）")
            + "；以及箱庭生态作息：EntityBehaviorService + DailyRoutineBehavior +"
            " NPCScheduleConfig + SceneNpcBehaviorScNotify（CmdId 1143），"
            "解决 NPC/生物仅休眠降频造成的「假死感」）",
        )
        n += 1

    # ---------- 报告头日期 ----------
    p = find_para(doc, startswith="报告生成日期")
    if p is not None:
        replace_paragraph_text(
            p,
            "报告生成日期：2026年09月04日（在「十大生产/体验优化」与 CmdId 1130–1142 基础上，"
            "已同步落地「环境生态假死感」优化："
            "EntityBehaviorService 为场景 NPC/生物挂载 DailyRoutineBehavior"
            "（WaypointMove 巡逻 / IdleAnimation 待机 / InteractionSpot 摆摊）；"
            "data/NPCScheduleConfig.json 提供时段→行为映射"
            "（dawn/noon/dusk/night，兼容 morning→dawn；如 MORNING:PATROL、NIGHT:DESPAWN_OR_SLEEP）；"
            "WorldTimeService 增加 WorldTimePeriodListener，切时段时主动推送"
            " SceneNpcBehaviorScNotify（CmdId 1143）驱动客户端动画状态机，"
            "而非仅靠 EntitySleepService 降心跳；ZoneTick 推进巡逻并与距离休眠正交；"
            "回归 NpcDailyRoutineFlowTest + EntityBehaviorServiceTest）",
        )
        n += 1

    # ---------- 目录速览 ----------
    p = find_para(doc, startswith="下面先给出章节速览")
    if p is not None and "7.33" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + " 第七章已新增 7.33「NPC/生物日程作息与生态循环（CmdId 1143）」。",
        )
        n += 1

    # ---------- 报告依据 ----------
    p = find_para(doc, startswith="报告依据：当前仓库 README")
    if p is not None and "NpcDailyRoutineFlowTest" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；本轮再将 NPC/生物独立日程作息写入 7.3/7.25/7.33、第八章 CmdId 表与成熟度清单"
            "（SceneNpcBehaviorScNotify=1143；测试 NpcDailyRoutineFlowTest、"
            "EntityBehaviorServiceTest）。",
        )
        n += 1

    # ---------- 6.2 scene 包职责 ----------
    pkg = find_table_by_header(doc, "包名", "普通人理解")
    if pkg is not None:
        for row in pkg.rows:
            if row.cells[0].text.strip() == "scene":
                cur = row.cells[2].text.strip()
                if "日程作息" not in cur and "EntityBehavior" not in cur:
                    set_cell_text(
                        row.cells[2],
                        cur.rstrip("。")
                        + "、NPC/生物日程作息（DailyRoutine + SceneNpcBehavior 1143）",
                    )
                    n += 1
                break

    # ---------- 7.3 场景正文 ----------
    p = find_para(doc, startswith="玩家进入某个 Plane/Floor 对应的 Zone 实例后")
    if p is not None and "DailyRoutineBehavior" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "。"
            "生态作息：播种 NPC 后 EntityBehaviorService.bindNpc 按当前 WorldTime 时段套用"
            " DailyRoutineBehavior；ZoneTickService 在 tickRefresh 之后推进 PATROL 路点移动"
            "（仍受 EntitySleepService 距离门控，与「玩法睡觉/消失」正交）；"
            "切时段经 WorldTimePeriodListener 批量切换行为，AOI 推送 SceneNpcBehaviorScNotify"
            "（1143，含 behavior/anim_hint/target_pos/visible），"
            "客户端切动画状态机；夜间 DESPAWN 从 SceneContext 投影移除，黎明再投影。",
        )
        n += 1

    if find_para(doc, startswith="NPC 日程作息：EntityBehaviorService") is None:
        anchor = find_para(doc, startswith="实体休眠 EntitySleepService")
        if anchor is not None:
            insert_paragraph_after(
                anchor,
                "NPC 日程作息：EntityBehaviorService + DailyRoutineBehavior"
                "（PATROL/IDLE/WORK/OPEN_SHOP/SLEEP/DESPAWN）；"
                "配置 data/NPCScheduleConfig.json；时段监听 WorldTimePeriodListener；"
                "推送 SceneNpcBehaviorScNotify（CmdId 1143）。",
                "List Bullet",
            )
            n += 1

    # ---------- 7.25 三期 WorldTime 句 ----------
    p = find_para(doc, startswith="移动：EnvInteractDetector 微交互池推送 1078")
    if p is not None and "DailyRoutineBehavior" not in p.text:
        replace_paragraph_text(
            p,
            p.text.replace(
                "时空：WorldTimeService 黎明/正午/黄昏/深夜（1083–1085）；DailyMission timePeriod 挂钩夜间任务。",
                "时空：WorldTimeService 黎明/正午/黄昏/深夜（1083–1085）并广播 WorldTimeScNotify；"
                "DailyMission timePeriod 挂钩夜间任务；"
                "另经 WorldTimePeriodListener 驱动 EntityBehaviorService，"
                "切时段推送 SceneNpcBehaviorScNotify（1143）切换 NPC 动画状态机"
                "（配置见 NPCScheduleConfig.json）。",
            ),
        )
        n += 1

    # ---------- 新增 7.33（插在 7.32 测试条后、第八章前） ----------
    if find_para(doc, contains="7.33 NPC/生物日程作息") is None:
        anchor = find_para(doc, startswith="测试入口：TenGapsOptimizationFlowTest")
        if anchor is None:
            anchor = find_para(doc, contains="7.32 十大生产/体验优化落地")
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "7.33 NPC/生物日程作息与生态循环（CmdId 1143）",
                "Heading 2",
            )
            p1 = insert_paragraph_after(
                h,
                "针对箱庭世界「静物画/假死感」：此前虽有 EntitySleepService（距离休眠降频）与"
                " WorldTimeService（昼夜渲染切换），但 NPC/生物缺少独立日程行为树。"
                "本轮升级 EntityBehaviorService，为场景 NPC（及配置化动物）挂载"
                " DailyRoutineBehavior，包含 WaypointMove 巡逻、IdleAnimation 待机、"
                "InteractionSpot/OPEN_SHOP 摆摊工作；夜间 DESPAWN_OR_SLEEP"
                "（有 homeSpot 则回屋 SLEEP，否则 DESPAWN 隐藏）。"
                "配置文件 data/NPCScheduleConfig.json 以「时段→状态」映射描述作息"
                "（dawn/noon/dusk/night；兼容 morning→dawn）。"
                "WorldTimeService.tick/forcePeriod 切换时段时，除全服 WorldTimeScNotify（1085）外，"
                "通过 WorldTimePeriodListener 回调主动推送 SceneNpcBehaviorScNotify（1143），"
                "客户端切换动画状态机，而非仅靠休眠降低心跳。"
                "ZoneTickService 在有玩家的 Zone 内推进 PATROL；"
                "位置同步进各玩家 SceneContext；invisible 实体不参与 projectAliveToScene。",
                "Normal",
            )
            insert_paragraph_after(
                p1,
                "实现入口：EntityBehaviorService、DailyRoutineBehavior、NpcScheduleConfigRepository、"
                "WorldTimePeriodListener、WorldTimeService.forcePeriod、ZoneContext.ZoneNpc"
                "（behavior/waypoint/visible）、SceneSyncBroadcaster.broadcastNpcBehavior、"
                "ZoneTickService、ZoneWorldService.bindNpc；"
                "协议 scene_system.proto → SceneNpcBehaviorScNotify / NpcBehaviorType；"
                "CmdIds.SCENE_NPC_BEHAVIOR_SC_NOTIFY=1143；"
                "配置 data/NPCScheduleConfig.json；"
                "测试 NpcDailyRoutineFlowTest（配置/绑定/一日循环/巡逻/ZoneTick/AOI 1143）、"
                "EntityBehaviorServiceTest。",
                "List Bullet",
            )
            n += 1

    # ---------- 9.2 data 目录 ----------
    p = find_para(doc, startswith="很多玩法数值不以硬编码写死在 Java 里")
    if p is not None and "NPCScheduleConfig" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；场景生态另增 data/NPCScheduleConfig.json（NPC/生物时段作息与路点组），"
            "由 NpcScheduleConfigRepository 启动加载。",
        )
        n += 1

    # ---------- 玩家一天串场 ----------
    p = find_para(doc, startswith="为了让非技术同学建立整体画面")
    if p is not None and "NPC 作息" not in p.text and "1143" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；场景里摊贩白天巡逻/摆摊、夜里回屋或消失，时段切换会收到 1143 行为推送"
            "（与天空盒 1085 配套），世界不再像静物画",
        )
        n += 1

    # ---------- 成熟度 ----------
    p = find_para(doc, startswith="日常循环：体力、每日任务、战令")
    if p is not None and "DailyRoutineBehavior" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；箱庭 NPC/生物日程作息 DailyRoutineBehavior + NPCScheduleConfig +"
            " SceneNpcBehaviorScNotify（1143）。",
        )
        n += 1

    p = find_para(doc, startswith="组队协议与多人共战/助战")
    if p is not None and "1143" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；NPC 日程作息 EntityBehaviorService（1143）。",
        )
        n += 1

    p = find_para(doc, startswith="协议已迁到 wire v4")
    if p is not None and "1143" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；新增 SceneNpcBehaviorScNotify（1143）需客户端动画状态机对接。",
        )
        n += 1

    # ---------- 短期建议 ----------
    p = find_para(doc, startswith="继续以模块化单体打磨崩铁式体验")
    if p is not None and "1143" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；客户端对齐 SceneNpcBehaviorScNotify（1143）切换 NPC 巡逻/摆摊/睡觉/消失"
            " 状态机，并把真实 npcId 写入 NPCScheduleConfig.json",
        )
        n += 1

    # ---------- CmdId 表 ----------
    cmd_table = find_table_by_header(doc, "号段", "系统")
    if cmd_table is not None and not table_has(cmd_table, "1143"):
        add_table_row(
            cmd_table,
            ["1143", "NPC/生物日程行为推送 SceneNpcBehaviorScNotify"],
        )
        n += 1

    # ---------- 术语表 ----------
    glossary = find_table_by_header(doc, "术语", "人话解释")
    if glossary is not None:
        if patch_table_cell(
            glossary,
            "WorldTime 时段",
            1,
            "黎明/正午/黄昏/深夜预设切换；可与每日任务 timePeriod 挂钩；"
            "并经监听器驱动 NPC 日程，推送 1143 切动画状态机",
        ):
            n += 1
        if not table_has(glossary, "NPC 日程作息"):
            add_table_row(
                glossary,
                [
                    "NPC 日程作息",
                    "按昼夜让 NPC/生物巡逻、摆摊、回屋睡觉或消失；"
                    "配置 NPCScheduleConfig，推送 SceneNpcBehaviorScNotify（1143）",
                ],
            )
            n += 1
        if not table_has(glossary, "EntitySleep"):
            add_table_row(
                glossary,
                [
                    "EntitySleep（距离休眠）",
                    "远离玩家时降频 AI 心跳的性能手段；"
                    "与玩法上的「夜间睡觉/消失」是两回事，可同时存在",
                ],
            )
            n += 1

    # ---------- 结语 ----------
    p = find_para(doc, startswith="MyLunarCore 已经不是「只有空壳的演示仓库」")
    if p is not None and "日程作息" not in p.text and "1143" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；箱庭侧已具备 NPC/生物独立日程作息（巡逻/摆摊/回屋/消失，CmdId 1143），"
            "与世界时段、距离休眠正交，减轻「静物画/假死感」",
        )
        n += 1

    p = find_para(doc, startswith="本报告基于仓库当前代码与文档快照整理")
    if p is not None:
        # 在现有长文末追加本轮同步说明（保留历史日期叙述）
        if "2026年09月04日" not in p.text:
            replace_paragraph_text(
                p,
                p.text.rstrip("。")
                + " 2026年09月04日 已在本文档内同步「环境生态假死感」优化"
                "（EntityBehaviorService / DailyRoutineBehavior / NPCScheduleConfig.json /"
                " WorldTimePeriodListener / SceneNpcBehaviorScNotify CmdId 1143；"
                "测试 NpcDailyRoutineFlowTest）；打开 Word 后请右键目录「更新域」以刷新页码。",
            )
            n += 1

    # ---------- CI 测试列表（若有） ----------
    p = find_para(doc, startswith="定向单测：CmdId 唯一性")
    if p is not None and "NpcDailyRoutineFlowTest" not in p.text:
        replace_paragraph_text(
            p,
            p.text.rstrip("。")
            + "；NpcDailyRoutineFlowTest / EntityBehaviorServiceTest（NPC 日程作息）",
        )
        n += 1

    doc.save(str(DOC_PATH))
    print(f"已覆盖更新原文档: {DOC_PATH}")
    print(f"本轮变更点数: {n}")


if __name__ == "__main__":
    main()
